import sys, types, os, json, importlib.util, tempfile

# ---------- Simulador mínimo de tkinter ----------
class Var:
    def __init__(self, value=None, **k): self.v = value; self.cbs = []
    def get(self): return self.v
    def set(self, v):
        self.v = v
        for cb in self.cbs: cb()
    def trace_add(self, mode, cb): self.cbs.append(cb)

class Widget:
    def __init__(self, parent=None, **kw):
        self.cfg = dict(kw); self.parent = parent; self.mapped = True
        self.bindings = {}; self.text = kw.get("text", ""); self._focus = False
    def config(self, **kw): self.cfg.update(kw); self.text = self.cfg.get("text", self.text)
    configure = config
    def cget(self, k): return self.cfg.get(k)
    def pack(self, **k): self.mapped = True
    def pack_forget(self): self.mapped = False
    def grid(self, **k): pass
    def grid_rowconfigure(self, *a, **k): pass
    def grid_columnconfigure(self, *a, **k): pass
    def bind(self, seq, fn=None, add=None): self.bindings[seq] = fn
    def focus_set(self): Widget.focused = self
    def winfo_ismapped(self): return self.mapped
    def winfo_width(self): return 700
    def winfo_height(self): return 300
    def destroy(self): pass
    def after(self, ms, fn=None, *a): return "job"
    def after_cancel(self, j): pass
    def title(self, *a): pass
    def geometry(self, *a): pass
    def minsize(self, *a): pass
    def iconbitmap(self, *a): pass
    def protocol(self, *a): pass
    def mainloop(self): pass
    def tag_configure(self, *a, **k): pass
    def __getattr__(self, name):
        # Cualquier otro método de tkinter (place, lift, create_polygon...)
        if name.startswith("__"): raise AttributeError(name)
        return lambda *a, **k: 1

class Entry(Widget):
    def __init__(self, parent=None, **kw):
        super().__init__(parent, **kw); self.s = ""; self.tv = kw.get("textvariable")
    def get(self): return self.s
    def delete(self, a, b=None): self.s = ""
    def insert(self, i, t): self.s = t
    def index(self, i): return len(self.s)
    def icursor(self, i): pass
    def select_range(self, a, b): pass

class Combobox(Widget):
    def __init__(self, parent=None, **kw):
        super().__init__(parent, **kw); self.val = ""; self.idx = 0
        self.values = list(kw.get("values", []))
        tv = kw.get("textvariable"); self.tv = tv
    def configure(self, **kw):
        if "values" in kw: self.values = list(kw["values"])
        super().config(**kw)
    config = configure
    def current(self, i=None):
        if i is None: return self.idx
        self.idx = i
        if i < len(self.values): self.val = self.values[i]
    def set(self, v): self.val = v
    def get(self): return self.val

class Canvas(Widget):
    def __init__(self, *a, **k): super().__init__(*a, **k); self.n = 0
    def _id(self): self.n += 1; return self.n
    def create_line(self, *a, **k): return self._id()
    def create_text(self, *a, **k): return self._id()
    def create_rectangle(self, *a, **k): return self._id()
    def create_window(self, *a, **k): return self._id()
    def delete(self, *a): pass
    def coords(self, *a): pass
    def itemconfigure(self, *a, **k): pass
    def tag_lower(self, *a): pass
    def tag_raise(self, *a): pass
    def canvasx(self, x): return x
    def canvasy(self, y): return y
    def xview(self, *a): pass
    def yview(self, *a): pass
    def xview_moveto(self, *a): pass
    def yview_moveto(self, *a): pass
    def xview_scroll(self, *a): pass
    def yview_scroll(self, *a): pass

class Treeview(Widget):
    def __init__(self, *a, **k): super().__init__(*a, **k); self.rows = {}
    def insert(self, p, i, iid=None, values=(), **k): self.rows[iid] = values
    def delete(self, iid): self.rows.pop(iid, None)
    def get_children(self): return list(self.rows)
    def exists(self, iid): return iid in self.rows
    def column(self, *a, **k): pass
    def heading(self, *a, **k): pass
    def yview(self, *a): pass

class Style:
    def theme_use(self, n): pass
    def configure(self, *a, **k): pass
    def map(self, *a, **k): pass
    def layout(self, *a, **k): pass

tk = types.ModuleType("tkinter")
tk.TclError = Exception; tk.END = "end"; tk.INSERT = "insert"
for n in ("Frame", "Label", "Button", "Checkbutton", "Tk", "Text", "Toplevel", "Menu"): setattr(tk, n, type(n, (Widget,), {}))
tk.Entry = Entry; tk.Canvas = Canvas
tk.StringVar = Var; tk.BooleanVar = Var
tk.Misc = type("Misc", (), {"tkraise": staticmethod(lambda w: None)})
ttk = types.ModuleType("tkinter.ttk")
ttk.Combobox = Combobox; ttk.Treeview = Treeview; ttk.Scrollbar = type("Scrollbar", (Widget,), {"set": lambda s, a, b: None})
ttk.Style = Style
fd = types.ModuleType("tkinter.filedialog")
tk.ttk = ttk; tk.filedialog = fd
sys.modules.update({"tkinter": tk, "tkinter.ttk": ttk, "tkinter.filedialog": fd})
w = types.ModuleType("winsound"); w.MB_ICONASTERISK = 1; w.MB_ICONHAND = 2; w.MB_ICONEXCLAMATION = 3
beeps = []; w.MessageBeep = lambda x: beeps.append(x); sys.modules["winsound"] = w

spec = importlib.util.spec_from_file_location("snc", os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "serial_number_checker.py"))
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
import openpyxl

TMP = tempfile.mkdtemp()
m.os.path.expanduser = lambda p: TMP  # sesión aislada

def new_app():
    root = tk.Tk()
    return m.SerialNumberChecker(root)

def scan(app, text):
    app.entry.s = text
    app.on_key_release()
    if app.mode == "create" or len(text) != 17:
        app.on_enter()

fails = []
def check(name, cond, extra=""):
    print(("OK   " if cond else "FAIL ") + name, extra)
    if not cond: fails.append(name)

# ---------- 1. Cabeceras ----------
for h in ["S/N", "s.n.", "Nº de serie", "Número de serie", "N/S", "Serial Number", "Serial Nº", "Serie"]:
    check(f"cabecera {h!r}", m.normalize_header(h) in m.SN_HEADERS or h == "Serial Nº")
check("Serial Nº se reconoce como SERIALNO", m.normalize_header("Serial Nº") in m.SN_HEADERS)

# ---------- 2. Modo lista ----------
app = new_app()
app.mode = "create"; app.update_mode_view()
for s in ["A1", "A2", "A3"]: scan(app, s)
check("pallet 1 creado sin diálogos", len(app.pallets) == 1 and app.pallets[0]["mode"] == "column")
app.next_pallet()
for s in ["B1", "B2"]: scan(app, s)
lay = m.compute_layout(app.create_items, app.pallets)
cells = lay["sheets"]["Pallet 1"]["cells"]
check("col A empieza en A1 (no en B)", cells[(1, 1)]["text"] == "Pallet 1" and cells[(2, 1)]["text"] == "A1")
check("col B pallet 2", cells[(1, 2)]["text"] == "Pallet 2" and cells[(3, 2)]["text"] == "B2")
check("contador = únicos", app.counter.text == "5")

# duplicado
scan(app, "A2")
check("duplicado entra marcado", app.create_items[-1]["dup"] is True and app.dup_count() == 1)
check("contador sin cambio con duplicado", app.counter.text == "5")
check("contador amarillo", app.counter.cfg["bg"] == m.YELLOW_BG)
check("banner visible", app.banner_visible and "A2" in app.banner_label.text)
check("sonido de duplicado", m.winsound.MB_ICONEXCLAMATION in beeps)
check("celda azul en visor", any(c["kind"] == "dup" for c in app.layout["sheets"]["Pallet 1"]["cells"].values()))

# seguir escaneando mantiene amarillo
scan(app, "Z9")
check("sigue amarillo tras otro escaneo", app.counter.cfg["bg"] == m.YELLOW_BG)

# texto libre, espacios, minúsculas
scan(app, "ab 12 cd/x")
check("texto tal cual (sin mayúsculas ni recorte interno)", app.create_items[-1]["serial"] == "ab 12 cd/x")
scan(app, "X" * 40)
check("longitud arbitraria", len(app.create_items[-1]["serial"]) == 40)

# deshacer
n = len(app.create_items)
app.undo_last()
check("deshacer añade", len(app.create_items) == n - 1)

# eliminar duplicados + deshacer
app.remove_duplicates()
check("eliminar duplicados", app.dup_count() == 0)
check("contador vuelve a normal", app.counter.cfg["bg"] == m.PANEL)
app.undo_last()
check("deshacer eliminar duplicados", app.dup_count() == 1)

# eliminar seleccionado
app.sheet_grid.selected = lay["positions"][0][1:]
app.refresh_create_view()
app.sheet_grid.selected = (2, 1)
before = len(app.create_items)
app.delete_selected_item()
check("eliminar seleccionado", len(app.create_items) == before - 1)
app.undo_last()
check("deshacer eliminar", len(app.create_items) == before)

# edición de celda
app.on_grid_edit(2, 1, "EDIT1")
check("edición celda", app.create_items[0]["serial"] == "EDIT1")
app.undo_last()
check("deshacer edición", app.create_items[0]["serial"] == "A1")

# destino hoja nueva
app.dest_index = 1; app.pallet_name_var.set("Mi hoja")
app.next_pallet(); scan(app, "H1"); scan(app, "H2")
lay = m.compute_layout(app.create_items, app.pallets)
check("hoja nueva", "Mi hoja" in lay["sheets"] and lay["sheets"]["Mi hoja"]["cells"][(2, 1)]["text"] == "H1")
# destino celda: sin selección
app.dest_index = 2; app.sheet_grid.selected = None
app.next_pallet()
check("celda sin selección avisa", app.status_key == "select_cell_first")
app.sheet_grid.selected = (4, 3); app.view_sheet = "Pallet 1"
app.next_pallet(); scan(app, "C1"); scan(app, "C2")
lay = m.compute_layout(app.create_items, app.pallets)
check("desde celda C4", lay["sheets"]["Pallet 1"]["cells"][(4, 3)]["text"] == "C1" and lay["sheets"]["Pallet 1"]["cells"][(5, 3)]["text"] == "C2")

# exportar XLSX
out = os.path.join(TMP, "lista.xlsx")
fd.asksaveasfilename = lambda **k: out
app.export_create_list()
wb = openpyxl.load_workbook(out)
check("XLSX exportado con hojas", set(wb.sheetnames) == {"Pallet 1", "Mi hoja"}, wb.sheetnames)
ws = wb["Pallet 1"]
check("export = visor (A1/A2)", ws["A1"].value == "Pallet 1" and ws["A2"].value == "A1")
check("dup en azul en export", any(c.fill.fill_type == "solid" and c.fill.fgColor.rgb.endswith("BDD7EE") for r in ws.iter_rows() for c in r))
check("status export ok", app.status_key == "list_exported")
# fórmula como texto
app.create_items.append({"serial": "=1+1", "pallet": app.pallets[-1]["number"], "dup": False})
app.refresh_create_view(); app.export_create_list()
wb = openpyxl.load_workbook(out)
vals = [c.value for ws in wb.worksheets for r in ws.iter_rows() for c in r]
check("=1+1 queda como texto", "=1+1" in vals and 2 not in vals)
# CSV
outc = os.path.join(TMP, "lista.csv"); fd.asksaveasfilename = lambda **k: outc
app.export_create_list()
check("CSV exportado", os.path.exists(outc) and "Duplicado" in open(outc, encoding="utf-8-sig").read())

# reinicio con doble clic
app.reset_app()
check("reset pide confirmación", app.reset_armed and app.create_items)
app.reset_app()
check("reset borra solo lista", not app.create_items and app.counter.text == "0")

# sesión
app2 = new_app()
app2.mode = "create"; app2.update_mode_view(); scan(app2, "S1"); scan(app2, "S2"); scan(app2, "S1")
app2.save_sessions()
app3 = new_app()
check("sesión recuperada (modo y datos)", app3.mode == "create" and len(app3.create_items) == 3 and app3.dup_count() == 1)

# ---------- 3. Checker ----------
app4 = new_app(); app4.mode = "checker"; app4.update_mode_view()
wbk = openpyxl.Workbook(); ws = wbk.active; ws.title = "H1"
ws.append(["S/N", "x"])
for v in ["ABCDEF1GHJKLMNPQR", "ABCDEF1GHJKLMNPQS", "ABCDEF1GHJKLMNPQR", "ABCDEFAGHJKLMNPQT"]: ws.append([v, "d"])
ws2 = wbk.create_sheet("H2"); ws2.append(["Nº de serie"]); ws2.append(["ZZZ"])
src = os.path.join(TMP, "orig.xlsx"); wbk.save(src)
fd.askopenfilename = lambda **k: src
app4.load_file()
check("carga: 4 únicos (acepta 'Nº de serie')", len(app4.serials) == 4 and app4.loaded_count == 5, (len(app4.serials), app4.loaded_count))
scan(app4, "ABCDEF1GHJKLMNPQR")
check("exacto verde", app4.status_key == "found_exact" and len(app4.found) == 1)
scan(app4, "ABCDEF1GHJKLMNPQR")
check("repetido -> amarillo + otro sonido, sin subir", app4.status_key == "already_scanned" and len(app4.found) == 1 and app4.counter.cfg["bg"] == m.YELLOW_BG and app4.standby)
scan(app4, "ABCDEF9GHJKLMNPQT")
check("variante A/B/U", app4.status_key == "found_variant" and not app4.standby and len(app4.found) == 2, app4.status_key)
app4.alt_matching_var.set(False)
scan(app4, "ABCDEF9GHJKLMNPQX")
check("no encontrado", app4.status_key == "not_found")
app4.alt_matching_var.set(False)
scan(app4, "ABCDEF9GHJKLMNPQS")
check("A/B/U desactivado no usa variante", app4.status_key == "not_found")
app4.alt_matching_var.set(True)
check("pendientes", app4.pending_list() == ["ABCDEF1GHJKLMNPQS", "ZZZ"], app4.pending_list())

# exportar y reanudar
out2 = os.path.join(TMP, "orig_resultado.xlsx"); fd.asksaveasfilename = lambda **k: out2
app4.export_report()
wb = openpyxl.load_workbook(out2)
fills = [(c.value, c.fill.fgColor.rgb[2:] if c.fill.fill_type else None) for r in wb["H1"].iter_rows(min_row=2, max_col=1) for c in r]
check("colores export", fills == [("ABCDEF1GHJKLMNPQR", "C6EFCE"), ("ABCDEF1GHJKLMNPQS", "FFC7CE"), ("ABCDEF1GHJKLMNPQR", "BDD7EE"), ("ABCDEFAGHJKLMNPQT", "C6EFCE")], fills)
app5 = new_app(); app5.mode = "checker"; app5.update_mode_view()
fd.askopenfilename = lambda **k: out2
app5.load_file()
check("aviso de resume integrado", app5.banner_visible and len(app5.pending_resume) == 2)
app5.resume_continue()
check("resume recupera 2", len(app5.found) == 2 and app5.resumed and not app5.banner_visible)
app5.reset_app(); app5.reset_app()
check("reset checker", not app5.serials and app5.counter.text == "0" and app5.counter_total.text == "/ 0")

# CSV con ; y encabezado S/N
csvp = os.path.join(TMP, "p.csv"); open(csvp, "w", encoding="utf-8").write("Equipo;S/N\nA;QQQ1\nB;QQQ2\n")
fd.askopenfilename = lambda **k: csvp
app5.load_file()
check("CSV con ;", app5.serials == {"QQQ1", "QQQ2"}, app5.serials)
check("CSV: fila del S/N", app5.locations.get("QQQ2") == (None, 3), app5.locations.get("QQQ2"))

# ---------- 4. Funciones nuevas del Checker ----------
app6 = new_app(); app6.mode = "checker"; app6.update_mode_view()
fd.askopenfilename = lambda **k: src
app6.load_file()
check("ubicación hoja/fila", app6.locations.get("ABCDEF1GHJKLMNPQR") == ("H1", 2) and app6.locations.get("ZZZ") == ("H2", 2), app6.locations)
scan(app6, "ABCDEF1GHJKLMNPQR")
check("mensaje con hoja y fila", ("detail_location", {"sheet": "H1", "row": 2}) in app6.status_details, app6.status_details)
check("detalle visible", "H1" in app6.details.text, app6.details.text)

scan(app6, "QQQQQQQQQQQQQQQQQ"); scan(app6, "QQQQQQQQQQQQQQQQQ")
check("sobrante guardado y contado", app6.extras.get("QQQQQQQQQQQQQQQQQ", {}).get("count") == 2 and app6.status_key == "not_found")
check("pestaña sobrantes", "(1)" in app6.tab_buttons["extras"].text, app6.tab_buttons["extras"].text)

scan(app6, "7501234567890")
check("formato distinto avisa y no es sobrante", app6.status_key == "bad_format" and "7501234567890" not in app6.extras)
n_log = len(app6.scan_log)
app6.undo_scan()
check("deshacer lectura de formato", len(app6.scan_log) == n_log - 1 and app6.status_key == "scan_undone")
app6.undo_scan()
check("deshacer sobrante resta", app6.extras["QQQQQQQQQQQQQQQQQ"]["count"] == 1)
app6.on_undo_key()
check("Ctrl+Z en Checker deshace", "QQQQQQQQQQQQQQQQQ" not in app6.extras)
app6.undo_scan()
check("deshacer encontrado", not app6.found and app6.counter.text == "0")
app6.undo_scan()
check("nada que deshacer", app6.status_key == "nothing_to_undo")

scan(app6, "ABCDEF1GHJKLMNPQR"); scan(app6, "ABCDEF1GHJKLMNPQR")
check("repetido no se cuenta como sobrante", app6.standby and not app6.extras)
app6.undo_scan()
check("deshacer repetido quita el amarillo", not app6.standby and app6.found == {"ABCDEF1GHJKLMNPQR"})

# pegar varios
app6.root.clipboard_get = lambda: "ABCDEF1GHJKLMNPQS\r\nZZZ\nNOPE1234567890123; ABCDEF1GHJKLMNPQS"
before = len(beeps)
check("pegar varios devuelve break", app6.on_paste() == "break")
check("pegar: resumen", app6.status_key == "paste_summary" and app6.status_params == {"total": 4, "found": 2, "missing": 1, "repeated": 1}, app6.status_params)
check("pegar: un solo sonido", len(beeps) - before == 1, len(beeps) - before)
check("pegar: sobrante", "NOPE1234567890123" in app6.extras)
app6.root.clipboard_get = lambda: "UNSOLO"
check("pegar un valor: pegado normal", app6.on_paste() is None)

# buscador y pestañas
app6.set_list_tab("pending")
check("pendientes en lista", [v[0] for v in app6.pending_tree.rows.values()] == ["ABCDEFAGHJKLMNPQT"], app6.pending_tree.rows)
app6.set_list_tab("extras")
check("lista de sobrantes", [v[0] for v in app6.pending_tree.rows.values()] == ["NOPE1234567890123"])
app6.search_var.set("XYZ")
check("buscador filtra", [v[0] for v in app6.pending_tree.rows.values()] == [app6.t("extras_none")])
app6.search_var.set(""); app6.set_list_tab("pending")

# vista grande
app6.toggle_big_view()
check("vista grande activa", app6.big_view)
app6.on_escape()
check("Esc sale de la vista grande", not app6.big_view)

# exportar con hojas nuevas
out3 = os.path.join(TMP, "orig2_resultado.xlsx"); fd.asksaveasfilename = lambda **k: out3
app6.export_report()
wb = openpyxl.load_workbook(out3)
check("hojas de reporte", {"Resumen", "Escaneos", "Sobrantes"} <= set(wb.sheetnames), wb.sheetnames)
check("escaneos registrados", wb["Escaneos"].max_row == len(app6.scan_log) + 1, (wb["Escaneos"].max_row, len(app6.scan_log)))
check("sobrantes exportados", wb["Sobrantes"]["A2"].value == "NOPE1234567890123")
check("resumen", [r[1].value for r in wb["Resumen"].iter_rows()][1:3] == [4, 3], [r[1].value for r in wb["Resumen"].iter_rows()])

# reanudar un reporte con hojas nuevas: se ignoran
app7 = new_app(); app7.mode = "checker"; app7.update_mode_view()
fd.askopenfilename = lambda **k: out3
app7.load_file()
check("reanudar ignora hojas del reporte", len(app7.serials) == 4 and len(app7.pending_resume) == 3, (len(app7.serials), len(app7.pending_resume)))
app7.resume_continue()
fd.asksaveasfilename = lambda **k: out3
app7.export_report()
wb = openpyxl.load_workbook(out3)
check("re-exportar no duplica hojas", sorted(wb.sheetnames) == sorted(["H1", "H2", "Resumen", "Escaneos"]), wb.sheetnames)

# sesión con sobrantes, registro y ubicaciones
app6.save_sessions()
app8 = new_app()
check("sesión: sobrantes y registro", "NOPE1234567890123" in app8.extras and len(app8.scan_log) == len(app6.scan_log) and app8.locations.get("ZZZ") == ("H2", 2))

# idiomas
for lang in ("English", "Português", "Español"):
    app8.language_var.set(lang); app8.change_language()
    app8.set_mode("create"); app8.set_mode("checker")
check("cambio de idioma y modo sin errores", app8.lang == "es")
keys = [set(v) for v in m.TEXTS.values()]
check("mismas claves en es/en/pt", keys[0] == keys[1] == keys[2])

# ---------- 5. Pegar en Crear lista ----------
app9 = new_app(); app9.set_mode("create"); app9.reset_app(); app9.reset_app()
app9.root.clipboard_get = lambda: "a, b\nc\n\nc"
app9.on_paste()
check("pegar en lista (comas se conservan)", [i["serial"] for i in app9.create_items] == ["a, b", "c", "c"] and app9.status_key == "paste_added" and app9.status_params["duplicates"] == 1, [i["serial"] for i in app9.create_items])

# ---------- 6. Nueva hoja, Retroceso y anchos de columna ----------
app10 = new_app(); app10.set_mode("create"); app10.reset_app(); app10.reset_app()
scan(app10, "N1"); app10.pallet_name_var.set("Extra"); app10.new_sheet()
check("Nueva hoja salta a ella", app10.view_sheet == "Extra", app10.view_sheet)
scan(app10, "N2")
lay = m.compute_layout(app10.create_items, app10.pallets)
check("lo escaneado va a la hoja nueva", lay["sheets"]["Extra"]["cells"][(2, 1)]["text"] == "N2")
app10.show_sheet("Pallet 1"); app10.sheet_grid.selected = (2, 1)
app10.entry.s = ""; app10._last_backspace = 0
app10.on_backspace()
check("Retroceso con campo vacío elimina", [i["serial"] for i in app10.create_items] == ["N2"])
app10.undo_last()
app10.sheet_grid.selected = (2, 1); app10.on_backspace()
check("Retroceso mantenido no borra (repetición)", len(app10.create_items) == 2)
app10.entry.s = "abc"; app10._last_backspace = 0
check("Retroceso con texto borra letras, no registros", app10.on_backspace() is None and len(app10.create_items) == 2)
app10.entry.s = ""
app10.col_widths.setdefault("Pallet 1", {})[1] = 300
app10.save_sessions(); app11 = new_app()
check("anchos de columna en la sesión", app11.col_widths.get("Pallet 1", {}).get(1) == 300, app11.col_widths)
out5 = os.path.join(TMP, "anchos.xlsx"); fd.asksaveasfilename = lambda **k: out5
app11.export_create_list()
wb = openpyxl.load_workbook(out5)
check("Excel usa el ancho del visor", abs(wb["Pallet 1"].column_dimensions["A"].width - 300 / 7) < 0.01)

# ---------- 7. Modo Rack ----------
app12 = new_app(); app12.set_mode("create"); app12.reset_app(); app12.reset_app()
app12.rack_mode_var.set(True); app12.on_rack_toggle()
app12.add_rack()
check("rack pide ubicación", app12.status_key == "rack_need_location" and not app12.rack_pallets())
app12.rack_location_var.set("DC05"); app12.rack_rows_var.set("2"); app12.rack_cols_var.set("3")
app12.add_rack()
check("rack DC05-1 creado", [p["name"] for p in app12.rack_pallets()] == ["DC05-1"] and app12.view_sheet == "DC05-1")
check("guía en (0,0)", app12.counter.text == "(0,0)", app12.counter.text)
for v in ["R0", "R1", "R2"]: scan(app12, v)
slots = [tuple(i["slot"]) for i in app12.create_items]
check("orden horizontal (x,y)", slots == [(0, 0), (1, 0), (2, 0)], slots)
check("mensaje con ubicación", app12.status_key == "rack_added" and app12.status_params["x"] == 2 and app12.status_params["y"] == 0)
app12.on_grid_select(2, 3, None)
check("clic en ubicación vacía mueve la guía", app12.counter.text == "(2,1)", app12.counter.text)
scan(app12, "R3"); scan(app12, "R4")
check("salta a la siguiente libre", [tuple(i["slot"]) for i in app12.create_items][-2:] == [(2, 1), (0, 1)])
scan(app12, "R0")
check("duplicado en rack", app12.create_items[-1]["dup"] and app12.status_key == "duplicate_warning")
check("rack completo avisa", app12.status_key == "duplicate_warning" and app12.counter.text == "✓" and app12.banner_visible and "DC05-1" in app12.banner_label.text, (app12.counter.text, app12.banner_label.text))
n = len(app12.create_items)
scan(app12, "R9")
check("rack lleno no registra", len(app12.create_items) == n and app12.status_key == "rack_full_scan")
app12.undo_last()
check("deshacer vuelve a esa ubicación", app12.counter.text == "(1,1)", app12.counter.text)
scan(app12, "R5")
app12.set_rack_order("v"); app12.add_rack()
check("DC05-2 creado y orden vertical", app12.rack_pallets()[-1]["name"] == "DC05-2" and app12.rack_pallets()[-1]["order"] == "v")
for v in ["V0", "V1", "V2"]: scan(app12, v)
check("orden vertical (x,y)", [tuple(i["slot"]) for i in app12.create_items][-3:] == [(0, 0), (0, 1), (1, 0)])
out6 = os.path.join(TMP, "racks.xlsx"); fd.asksaveasfilename = lambda **k: out6
app12.export_create_list()
wb = openpyxl.load_workbook(out6)
check("export racks: hojas", wb.sheetnames == ["DC05-1", "DC05-2", "Lista"], wb.sheetnames)
ws = wb["DC05-1"]
check("rack dibujado (x arriba, y a la izquierda)", ws["B1"].value == 0 and ws["D1"].value == 2 and ws["A3"].value == 1 and ws["B2"].value == "R0" and ws["D3"].value == "R3", (ws["B2"].value, ws["D3"].value))
check("rack sin encabezados de Excel", ws.sheet_view.showRowColHeaders is False and not ws.print_options.headings)
rows = [[c.value for c in r] for r in wb["Lista"].iter_rows()]
check("lista de ubicaciones", rows[0][:3] == ["Ubicación", "Rack", "Posición (x,y)"] and rows[1][:6] == ["DC05", "DC05-1", "(0,0)", 0, 0, "R0"], rows[:2])
app12.save_sessions(); app13 = new_app()
check("sesión rack", app13.rack_mode_var.get() and len(app13.rack_pallets()) == 2 and app13.rack_location_var.get() == "DC05" and app13.create_items[0]["slot"] == [0, 0])

# ---------- 8. Imprimir ----------
printed = []
m.os.startfile = lambda path, verb=None: printed.append((os.path.splitext(path)[1], verb))
app12.ask_print()
check("imprimir pregunta en el aviso", app12.print_prompt == "create" and app12.banner_visible and app12.t("print_ask") in app12.banner_label.text)
app12.print_create(True)
check("imprime hoja actual", printed[-1] == (".xlsx", "print") and app12.status_key == "print_sent" and app12.print_prompt is None)
app6.ask_print()
check("checker: casillas en el aviso", app6.print_prompt == "checker")
for v in app6.print_vars.values(): v.set(False)
app6.print_checker()
check("checker: sin opciones avisa", app6.status_key == "print_nothing")
app6.print_vars["pending"].set(True); app6.print_vars["extras"].set(True)
wbp = app6.build_checker_print(["summary", "pending", "extras"])
check("libro de impresión", wbp.sheetnames == ["Resumen", "Pendientes", "Sobrantes"], wbp.sheetnames)
app6.print_checker()
check("checker imprime", app6.status_key == "print_sent" and len(printed) == 2)
def fail_xlsx(path, verb=None):
    if path.endswith(".xlsx"): raise OSError("sin Excel")
    printed.append((".txt", verb))
m.os.startfile = fail_xlsx
app6.ask_print(); app6.print_checker()
check("sin Excel imprime con el Bloc de notas", printed[-1] == (".txt", "print"))
app6.ask_print(); app6.on_escape()
check("Esc cancela imprimir", app6.print_prompt is None)
del m.os.startfile

# ---------- 9. README ----------
for lang, name in (("es", "README.txt"), ("en", "README_EN.txt"), ("pt", "README_PT.txt")):
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", name)
    check(f"README {lang} incrustado = {name}", m.README_TEXTS[lang] == open(path, encoding="utf-8").read())
blocks = m.SerialNumberChecker.readme_blocks(m.README_TEXTS["es"])
check("README: título y secciones", blocks[2] == ("title", "SERIAL NUMBER CHECKER") and ("h2", "12. VALIDAR LISTA") in blocks and ("h2", "13. MODO RACK") in blocks and ("h2", "19. GitHub") in blocks and ("h3", "VERSIÓN 2.1A - Validar lista y nombres de hoja") in blocks, blocks[:4])
check("README: creador en su línea", blocks[0] == ("body", 'Creador: Desblack "Dess"'), blocks[0])
check("README: tablas intactas", any(s == "mono" and "Ctrl+P" in v for s, v in blocks))
app6.toggle_readme()
check("README abre", app6.readme_open)
app6.on_escape()
check("Esc cierra el README", not app6.readme_open)

# ---------- 10. Versión 2.1A ----------
# Validar lista
lista = os.path.join(TMP, "validar.xlsx")
wbv = openpyxl.Workbook(); wsv = wbv.active; wsv.append(["S/N"])
for v in ["VAL0000001AAAAAA1", "VAL0000002AAAAAA2", "VAL0000003AAAAAA3", "VAL000A004AAAAAA4"]: wsv.append([v])
wbv.save(lista)
app14 = new_app(); app14.set_mode("create"); app14.reset_app(); app14.reset_app()
app14.rack_mode_var.set(False); app14.on_rack_toggle()
app14.validate_var.set(True); app14.on_validate_toggle()
check("validar: aviso sin lista", app14.banner_visible and app14.t("validate_banner") in app14.banner_label.text)
scan(app14, "VAL0000001AAAAAA1")
check("validar: sin lista no registra", not app14.create_items and app14.status_key == "validate_need_list")
fd.askopenfilename = lambda **k: lista
app14.load_validation_list()
check("validar: lista cargada", len(app14.valid_set) == 4 and app14.status_key == "validate_loaded" and not app14.banner_visible)
scan(app14, "val0000001aaaaaa1")
check("validar: válido entra (como en la lista)", [i["serial"] for i in app14.create_items] == ["VAL0000001AAAAAA1"])
scan(app14, "NOESTA00000000000")
check("validar: no válido no entra y avisa en rojo", len(app14.create_items) == 1 and app14.status_key == "validate_invalid" and m.STATUS_KINDS["validate_invalid"] == "error")
scan(app14, "VAL0000001AAAAAA1")
check("validar: duplicado entra en azul", app14.create_items[-1]["dup"] and app14.status_key == "duplicate_warning")
scan(app14, "VAL0009004AAAAAA4")
check("validar: alternativa A/B/U", app14.create_items[-1]["serial"] == "VAL000A004AAAAAA4", app14.create_items[-1]["serial"])
app14.valid_alt_var.set(False); scan(app14, "VAL0009004AAAAAA4")
check("validar: sin A/B/U no acepta variante", app14.status_key == "validate_invalid")
check("validar: faltan", "Validados 2 de 4 · faltan 2" in app14.counter_info.text, app14.counter_info.text)
app14.root.clipboard_get = lambda: "VAL0000002AAAAAA2\nMALO\nVAL0000003AAAAAA3"
app14.on_paste()
check("validar: pegar resume", app14.status_key == "paste_validated" and app14.status_params["added"] == 2 and app14.status_params["invalid"] == 1, app14.status_params)
app14.rack_mode_var.set(True); app14.on_rack_toggle()
check("Modo Rack desmarca y bloquea Validar lista", not app14.validate_var.get() and not app14.validate_check.enabled and app14.valid_set)
app14.validate_check._toggle()
check("Validar lista bloqueada no se marca", not app14.validate_var.get())
app14.rack_mode_var.set(False); app14.on_rack_toggle()
check("sin Modo Rack se puede volver a marcar", app14.validate_check.enabled)
app14.validate_check._toggle()
check("la lista cargada se conserva", app14.validate_var.get() and len(app14.valid_set) == 4)
check("con Validar lista, Modo Rack bloqueado", not app14.rack_check.enabled)
app14.rack_check._toggle()
check("Modo Rack bloqueado no se marca", not app14.rack_mode_var.get())
app14.save_sessions(); app15 = new_app()
check("sesión: validar lista", app15.validate_var.get() and len(app15.valid_set) == 4 and app15.valid_path == lista)

# Hojas: Pallet 1 y renombrar
app16 = new_app(); app16.set_mode("create"); app16.reset_app(); app16.reset_app(); app16.validate_var.set(False); app16.on_validate_toggle()
scan(app16, "H1")
check("la primera hoja se llama Pallet 1", app16.view_sheet == "Pallet 1" and "Pallet 1" in app16.layout["sheets"])
app16.col_widths["Pallet 1"] = {1: 222}
check("renombrar hoja principal", app16.rename_sheet("Pallet 1", "Zona A") and app16.view_sheet == "Zona A" and app16.col_widths.get("Zona A") == {1: 222})
app16.pallet_name_var.set("B"); app16.new_sheet(); scan(app16, "H2")
check("renombrar hoja de pallet", app16.rename_sheet("B", "Pasillo 2") and "Pasillo 2" in app16.layout["sheets"])
check("nombre repetido se rechaza", not app16.rename_sheet("Pasillo 2", "zona a") and app16.status_key == "sheet_name_taken")
check("caracteres no válidos se quitan", app16.rename_sheet("Pasillo 2", "P/2:*") and "P2" in app16.layout["sheets"])
out7 = os.path.join(TMP, "renombrada.xlsx"); fd.asksaveasfilename = lambda **k: out7
app16.export_create_list()
check("export con nombres nuevos", openpyxl.load_workbook(out7).sheetnames == ["Zona A", "P2"])
app16.save_sessions(); app17 = new_app()
check("sesión: nombres de hoja", app17.default_sheet == "Zona A" and "P2" in app17.layout["sheets"])

# Sesión de la 2.0 (hoja "Lista")
old = json.load(open(app17.session_file, encoding="utf-8"))
del old["create"]["default_sheet"]; old["view_sheet"] = "Lista"; old["create"]["col_widths"] = {"Lista": {"1": 180}}
json.dump(old, open(app17.session_file, "w", encoding="utf-8"))
app18 = new_app()
check("sesión 2.0: Lista pasa a Pallet 1", app18.default_sheet == "Pallet 1" and app18.col_widths.get("Pallet 1") == {1: 180})

# Anchos justos al exportar
wsr = openpyxl.load_workbook(out3)["H1"]
check("reporte: columna S/N al ancho justo", wsr.column_dimensions["A"].width == len("ABCDEF1GHJKLMNPQR") + 2, wsr.column_dimensions["A"].width)
wsl = openpyxl.load_workbook(out7)["P2"]
check("lista: ancho justo", wsl.column_dimensions["A"].width == 6, wsl.column_dimensions["A"].width)

# Copiar S/N pendiente
app6.set_list_tab("pending")
copied = []
app6.root.clipboard_clear = lambda: copied.clear()
app6.root.clipboard_append = lambda v: copied.append(v)
first = next(iter(app6.pending_tree.rows))
app6.pending_tree.selection = lambda: (first,)
app6.pending_tree.item = lambda iid, option=None: app6.pending_tree.rows[iid]
app6.entry.selection_present = lambda: False
check("Ctrl+C copia el S/N pendiente", app6.on_copy() == "break" and copied == [app6.pending_tree.rows[first][0]] and app6.status_key == "copied")
app6.entry.selection_present = lambda: True
check("Ctrl+C con texto seleccionado copia el texto", app6.on_copy() is None)

print("\nFALLOS:", fails if fails else "ninguno")
sys.exit(1 if fails else 0)
