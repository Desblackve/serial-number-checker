import csv
import json
import os
import re
import sys
import tempfile
import time
import unicodedata
from collections import OrderedDict
import tkinter as tk
from tkinter import filedialog, ttk

try:
    from tkinter import font as tkfont
except ImportError:  # Simulador de pruebas: sin medición de fuentes.
    tkfont = None

try:
    import winsound
except ImportError:  # Fuera de Windows la app funciona, solo sin sonidos.
    winsound = None

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import PatternFill, Font
except ImportError:
    load_workbook = None
    Workbook = None
    PatternFill = None
    Font = None


APP_TITLE = "Serial Number Checker"
FLASH_MS = 2000
SAVE_DELAY_MS = 800
RESET_ARM_MS = 3000
HISTORY_SIZE = 5
PENDING_LIMIT = 500
CHECKER_LENGTH = 17
SESSION_VERSION = 2
CHECKER_UNDO_LIMIT = 50
SCAN_LOG_LIMIT = 50000
RATE_WINDOW_S = 600
BACKSPACE_REPEAT_S = 0.5

DEFAULT_SHEET = "Pallet 1"
OLD_DEFAULT_SHEET = "Lista"  # nombre en la versión 2.0 (sesiones antiguas)
RACK_ROWS = 6
RACK_COLS = 8
RACK_MAX = 99
SN_HEADER_TEXT = "S/N"

# Tema oscuro: tarjetas sobre un fondo negro con patrón de líneas.
BG = "#050505"
PANEL = "#14161b"
FIELD = "#0d0f13"
CONTROL = "#1b1e25"
WHITE = "#f4f4f5"
MUTED = "#a1a1aa"
DIM = "#8b8b94"
BORDER = "#262a33"
BORDER_STRONG = "#343a46"
BLUE = "#2563eb"
BLUE_DARK = "#1d4ed8"
BLUE_LIGHT = "#3b82f6"
BLUE_TEXT = "#93c5fd"
GREEN = "#22c55e"
GREEN_BG = "#0f2417"
GREEN_BORDER = "#166534"
RED = "#ef4444"
RED_BG = "#2a0d0d"
RED_BORDER = "#991b1b"
YELLOW = "#f59e0b"
YELLOW_BG = "#3f3108"
YELLOW_BORDER = "#a16207"
YELLOW_TEXT = "#fbbf24"
BUTTON_BG = "#1f232b"
BUTTON_ACTIVE = "#2a2f39"
DUP_BG = "#1d4ed8"
LIME = "#b4ed3d"
HEADER_BG = "#15171c"
HEADER_TINT = "#141d33"
EMPTY_BG = "#0d0f13"
GRID_LINE = "#1f232b"

# Fuentes: se eligen al abrir según lo instalado (ver pick_fonts).
UI_FONT = "Segoe UI"
MONO_FONT = "Consolas"
EMPTY_BG = "#0e0e0e"

# Encabezados reconocidos como columna de números de serie.
# Se comparan ya normalizados (sin acentos, signos ni espacios).
SN_HEADERS = {
    "SN",
    "NS",
    "SERIAL",
    "SERIALNO",
    "SERIALNR",
    "SERIALNBR",
    "SERIALNUM",
    "SERIALNUMBER",
    "SERIE",
    "NSERIE",
    "NDESERIE",
    "NOSERIE",
    "NODESERIE",
    "NUMSERIE",
    "NUMDESERIE",
    "NUMEROSERIE",
    "NUMERODESERIE",
    "NSERIAL",
    "NUMEROSERIAL",
}

# Tipo de estado de la línea de mensajes.
STATUS_KINDS = {
    "found_exact": "ok",
    "found_variant": "ok",
    "not_found": "error",
    "already_scanned": "warn",
    "bad_format": "error",
    "scan_undone": "info",
    "paste_summary": "info",
    "paste_added": "info",
    "rack_ready": "info",
    "rack_added": "ok",
    "rack_completed": "warn",
    "rack_full_scan": "warn",
    "rack_need_location": "warn",
    "rack_bad_size": "warn",
    "rack_target": "info",
    "validate_need_list": "warn",
    "validate_invalid": "error",
    "validate_loaded": "ok",
    "paste_validated": "info",
    "sheet_renamed": "info",
    "sheet_name_taken": "warn",
    "copied": "info",
    "print_sent": "ok",
    "print_error": "error",
    "print_unavailable": "warn",
    "print_nothing": "warn",
    "file_loaded": "info",
    "resumed": "ok",
    "load_first": "warn",
    "unsupported": "warn",
    "no_serials": "warn",
    "read_error": "error",
    "export_dependency": "error",
    "export_success": "ok",
    "export_error": "error",
    "same_file": "warn",
    "create_added": "ok",
    "duplicate_warning": "warn",
    "dups_removed": "info",
    "pallet_ready": "info",
    "select_cell_first": "warn",
    "deleted_msg": "info",
    "undone": "info",
    "nothing_to_undo": "warn",
    "nothing_to_select": "warn",
    "nothing_to_export": "warn",
    "list_exported": "ok",
    "reset_done": "info",
    "checker_session_restored": "info",
    "create_session_restored": "info",
}


# ============================================================
# TRADUCCIONES
# ============================================================

LANGUAGES = {
    "Español": "es",
    "English": "en",
    "Português": "pt",
}

TEXTS = {
    "es": {
        "mode_checker": "Checker",
        "mode_create": "Crear lista S/N",
        "alt_match": "Alternativas A/B/U",
        "big_view": "Vista grande",
        "big_hint": "F11 o Esc para volver",
        "no_file": "No hay ningún archivo CSV o Excel cargado",
        "load": "Cargar CSV / Excel",
        "export": "Exportar reporte",
        "input_label": "(S/N) o secuencia a buscar",
        "input_label_create": "S/N o texto a registrar (Enter)",
        "details": (
            "Carga un CSV o Excel y escanea. Un XLSX permite múltiples hojas; "
            "un CSV contiene una sola tabla. También puedes pegar una lista."
        ),
        "create_hint": (
            "Escanea cualquier texto; Enter lo registra. "
            "También puedes pegar una lista (Ctrl+V)."
        ),
        "footer": "Para tutorial, leer README · Att: Dess",
        "select_file": "Seleccionar archivo",
        "file_types": "Archivos CSV y Excel",
        "csv_files": "Archivos CSV",
        "excel_files": "Archivos Excel",
        "all_files": "Todos los archivos",
        "unsupported": "Selecciona un archivo CSV (.csv) o Excel (.xlsx).",
        "no_serials": (
            "No encontré números de serie en las columnas SN. "
            "Verifica los encabezados de tu archivo."
        ),
        "file_loaded": "Archivo cargado. Listo para comprobar.",
        "file_summary": (
            "{name}  ·  {count} números leídos  ·  {unique} S/N únicos"
        ),
        "read_error": "No se pudo leer el archivo: {error}",
        "excel_dependency": (
            "Para abrir archivos Excel instala openpyxl "
            "(pip install openpyxl)."
        ),
        "csv_encoding": "No se pudo determinar la codificación del CSV.",
        "load_first": "Primero carga un archivo CSV o Excel.",
        "save_report": "Guardar reporte coloreado",
        "output_file": "Archivo Excel",
        "export_dependency": (
            "Para exportar instala openpyxl (pip install openpyxl)."
        ),
        "csv_read_error": "No se pudo leer la codificación del CSV.",
        "invalid_format": "Formato de archivo no compatible.",
        "same_file": "El reporte debe guardarse como un archivo nuevo.",
        "export_success": (
            "Reporte exportado: {path}  ·  Coloreadas: {cells}  ·  "
            "Duplicados (azul): {duplicates}  ·  Encontrados: {found}  ·  "
            "Pendientes: {remaining}  ·  Sobrantes: {extras}"
        ),
        "export_error": "No se pudo exportar: {error}",
        "found_exact": "ENCONTRADO — coincidencia exacta",
        "found_variant": "ENCONTRADO — variante: {serial}",
        "not_found": "NO ENCONTRADO — {serial}",
        "already_scanned": "YA ESCANEADO — {serial}",
        "bad_format": "NO PARECE UN S/N — {serial}",
        "detail_location": "Hoja {sheet} · fila {row}",
        "detail_location_csv": "Fila {row} del archivo",
        "detail_variant": "Leído como {original}",
        "detail_repeat": "El contador no sube. Esc para quitar el aviso.",
        "detail_extra": (
            "Guardado en Sobrantes (se exporta en una hoja aparte)."
        ),
        "detail_bad_format": (
            "Tiene {length} caracteres y los S/N del archivo tienen "
            "{expected}. ¿Se escaneó otro código de la caja?"
        ),
        "legend_blue": "Azul: SN duplicado (al exportar)",
        "legend_red": "Rojo: no escaneado",
        "legend_green": "Verde: escaneado",
        "legend_yellow": "Amarillo: ya escaneado",
        "legend_dup": "Azul: duplicado (se puede eliminar)",
        "legend_warn": "Amarillo: aviso activo",
        "reset_arm": "¿Seguro?",
        "reset_done": "Sesión reiniciada.",
        "resume_prompt": (
            "Este archivo tiene {done} de {total} S/N marcados en verde "
            "(reporte exportado antes). ¿Continuar donde te quedaste?"
        ),
        "resume_continue": "Continuar",
        "resume_fresh": "Empezar de cero",
        "resumed": "Progreso recuperado: {done} de {total} S/N ya escaneados.",
        "history_title": "Últimos escaneos",
        "history_empty": "Todavía no hay escaneos",
        "hist_exact": "exacto",
        "hist_variant": "variante",
        "hist_not_found": "no encontrado",
        "hist_repeated": "repetido",
        "hist_bad_format": "formato",
        "undo": "Deshacer",
        "scan_undone": "Escaneo deshecho: {serial}",
        "tab_pending": "Pendientes ({count})",
        "tab_extras": "Sobrantes ({count})",
        "search": "Buscar",
        "pending_col": "S/N pendiente",
        "pending_more": "… y {count} más",
        "pending_none": "No quedan S/N pendientes",
        "extras_col": "S/N sobrante (no está en el archivo)",
        "count_col": "Veces",
        "time_col": "Hora",
        "extras_none": "No hay sobrantes",
        "caption_scanned": "Escaneados",
        "caption_unique": "Total único",
        "progress_info": "{pct}% completado · {left} pendientes",
        "rate_info": "{rate}/min · faltan ~{eta} min",
        "destination": "Destino",
        "dest_column": "Columna nueva",
        "dest_sheet": "Hoja nueva",
        "dest_cell": "Desde la celda seleccionada",
        "pallet_name": "Nombre (opcional)",
        "next_pallet": "Siguiente pallet",
        "delete_selected": "Eliminar",
        "export_list": "Exportar lista",
        "create_count": (
            "Pallet {pallet}: {pallet_count} S/N · Duplicados: {duplicates}"
        ),
        "create_added": "REGISTRADO — {serial}",
        "duplicate_warning": "DUPLICADO — ya existe: {serial}",
        "dup_banner": "Duplicados ({count}): {list}",
        "dup_remove": "Eliminar duplicados",
        "dup_hide": "Ocultar aviso",
        "dups_removed": "{count} duplicados eliminados · Ctrl+Z para deshacer",
        "pallet_ready": "Pallet {pallet} listo. Escanea para añadir números.",
        "select_cell_first": "Selecciona primero una celda en el visor.",
        "deleted_msg": "Eliminado: {serial} · Ctrl+Z para deshacer",
        "undone": "Acción deshecha.",
        "nothing_to_undo": "No hay nada que deshacer.",
        "nothing_to_select": "Selecciona un registro en el visor.",
        "nothing_to_export": "Todavía no hay registros para exportar.",
        "export_list_title": "Guardar lista de números de serie",
        "list_exported": (
            "Lista exportada: {path}  ·  {count} registros  ·  "
            "{duplicates} duplicados en azul"
        ),
        "csv_headers": "SN|Pallet|Hoja|Celda|Duplicado",
        "yes": "Sí",
        "no": "No",
        "session_saved": "Sesión guardada",
        "session_save_error": "No se pudo guardar la sesión: {error}",
        "checker_session_restored": "Sesión de Checker recuperada.",
        "create_session_restored": "Sesión de creación de lista recuperada.",
        "paste_summary": (
            "Pegados {total}: {found} encontrados · {missing} no encontrados "
            "· {repeated} repetidos"
        ),
        "paste_added": "Pegados {count} registros · {duplicates} duplicados",
        "sheet_summary": "Resumen",
        "sheet_log": "Escaneos",
        "sheet_extras": "Sobrantes",
        "log_headers": "Hora|Leído|Resultado|Coincide con",
        "extras_headers": "Sobrante|Veces|Primera lectura",
        "sum_file": "Archivo",
        "sum_total": "S/N únicos en el archivo",
        "sum_found": "Encontrados",
        "sum_pending": "Pendientes",
        "sum_extras": "Sobrantes (no están en el archivo)",
        "sum_scans": "Lecturas registradas",
        "sum_first": "Primera lectura",
        "sum_last": "Última lectura",
        "sum_exported": "Exportado",
        "rack_mode": "Modo Rack",
        "rack_location": "Ubicación",
        "rack_rows": "Filas",
        "rack_cols": "Columnas",
        "rack_order": "Orden",
        "rack_order_h": "→ Horizontal",
        "rack_order_v": "↓ Vertical",
        "add_rack": "Agregar nuevo rack",
        "new_sheet": "Nueva hoja",
        "caption_next": "Siguiente ubicación",
        "rack_none": "Escribe la ubicación (por ejemplo, DC05) y pulsa «Agregar nuevo rack».",
        "rack_count": "{rack} · {used}/{total} ubicaciones\nDuplicados: {duplicates}",
        "rack_ready": "Rack {rack} listo ({rows} × {cols}). Empieza en (0,0).",
        "rack_added": "REGISTRADO en {rack} ({x},{y}) — {serial}",
        "rack_completed": "RACK {rack} COMPLETO — {total} ubicaciones llenas",
        "rack_full_scan": "El rack {rack} está completo. Pulsa «Agregar nuevo rack».",
        "rack_need_location": "Escribe primero la ubicación (por ejemplo, DC05).",
        "rack_bad_size": "Filas y columnas deben ser números entre 1 y {maximum}.",
        "rack_target": "El siguiente escaneo irá a {rack} ({x},{y}).",
        "rack_where": "{rack} ({x},{y})",
        "rack_full_banner": "Se completó el rack {rack} ({total}/{total}).",
        "rack_list_sheet": "Lista",
        "rack_list_headers": "Ubicación|Rack|Posición (x,y)|X|Y|S/N|Duplicado",
        "print": "Imprimir",
        "print_ask": "¿Qué quieres imprimir?",
        "print_ask_checker": "¿Qué quieres imprimir? Marca una o varias opciones:",
        "print_current": "Hoja actual",
        "print_all": "Todas las hojas",
        "cancel": "Cancelar",
        "print_summary": "Resumen",
        "print_pending": "Pendientes",
        "print_extras": "Sobrantes",
        "print_sent": "Enviado a la impresora predeterminada.",
        "print_error": "No se pudo imprimir: {error}",
        "print_unavailable": "Imprimir solo funciona en Windows. Archivo preparado: {path}",
        "print_nothing": "Marca al menos una opción para imprimir.",
        "readme": "README",
        "reset_label": "Reiniciar",
        "keys_checker": (
            "F11|vista grande;Ctrl+Z|deshacer;Ctrl+E|exportar;Esc|quitar aviso"
        ),
        "validate_mode": "Validar lista",
        "validate_load": "Cargar lista",
        "validate_off": "Desactivar",
        "validate_select": "Seleccionar la lista de S/N esperados",
        "validate_banner": "Validar lista está activo: carga el CSV o Excel con los S/N esperados para empezar a escanear.",
        "validate_need_list": "Primero carga la lista de validación.",
        "validate_invalid": "S/N NO VÁLIDO PARA EL PALLET — {serial}",
        "validate_invalid_detail": "No está en la lista de validación: no se registró.",
        "validate_loaded": "Lista de validación cargada: {name} · {count} S/N",
        "validate_count": "Validados {done} de {total} · faltan {left}",
        "validate_no_list": "Sin lista de validación",
        "validate_file": "Lista: {name}",
        "paste_validated": "Pegados {total}: {added} registrados · {invalid} no válidos · {duplicates} duplicados",
        "sheet_renamed": "Hoja renombrada: {old} → {new}",
        "sheet_name_taken": "Ya existe una hoja llamada «{name}».",
        "tip_rack": "Escaneo guiado por ubicaciones (x,y) de un rack. Cada S/N se guarda en la siguiente posición libre y la app avisa cuando el rack se llena.",
        "tip_validate": "Carga un CSV o Excel con los S/N esperados. Los que no estén en la lista no se registran y se avisa en rojo; los repetidos entran en azul. Sin lista cargada no se puede escanear.",
        "tip_rack_locked": "Modo Rack no disponible durante la validación de S/N en pallets.",
        "tip_validate_locked": "Validar lista no disponible mientras esté activo Modo Rack.",
        "tip_tab": "Doble clic para cambiar el nombre de la hoja.",
        "copy_sn": "Copiar S/N",
        "copied": "S/N copiado: {serial}",
        "keys_rack": (
            "F2|nuevo rack;Retroceso|eliminar;Ctrl+Z|deshacer;Ctrl+E|exportar;"
            "Esc|quitar aviso"
        ),
        "keys_create": (
            "F2|pallet;Retroceso|eliminar;Ctrl+Z|deshacer;Ctrl+E|exportar;"
            "Esc|quitar aviso"
        ),
    },

    "en": {
        "mode_checker": "Checker",
        "mode_create": "Create S/N List",
        "alt_match": "A/B/U alternatives",
        "big_view": "Big view",
        "big_hint": "F11 or Esc to go back",
        "no_file": "No CSV or Excel file loaded",
        "load": "Load CSV / Excel",
        "export": "Export report",
        "input_label": "(S/N) or serial sequence to search",
        "input_label_create": "S/N or text to record (Enter)",
        "details": (
            "Load a CSV or Excel file and scan. An XLSX file can contain "
            "multiple sheets; a CSV file contains one table. You can also "
            "paste a list."
        ),
        "create_hint": (
            "Scan any text; Enter records it. "
            "You can also paste a list (Ctrl+V)."
        ),
        "footer": "For a tutorial, read README · By: Dess",
        "select_file": "Select file",
        "file_types": "CSV and Excel files",
        "csv_files": "CSV files",
        "excel_files": "Excel files",
        "all_files": "All files",
        "unsupported": "Select a CSV (.csv) or Excel (.xlsx) file.",
        "no_serials": (
            "No serial numbers were found in the SN columns. "
            "Check the headers in your file."
        ),
        "file_loaded": "File loaded. Ready to check.",
        "file_summary": (
            "{name}  ·  {count} values read  ·  {unique} unique S/Ns"
        ),
        "read_error": "Could not read the file: {error}",
        "excel_dependency": (
            "Install openpyxl to open Excel files (pip install openpyxl)."
        ),
        "csv_encoding": "Could not determine the CSV encoding.",
        "load_first": "Load a CSV or Excel file first.",
        "save_report": "Save colored report",
        "output_file": "Excel file",
        "export_dependency": (
            "Install openpyxl to export (pip install openpyxl)."
        ),
        "csv_read_error": "Could not read the CSV encoding.",
        "invalid_format": "Unsupported file format.",
        "same_file": "The report must be saved as a new file.",
        "export_success": (
            "Report exported: {path}  ·  Colored: {cells}  ·  "
            "Duplicates (blue): {duplicates}  ·  Found: {found}  ·  "
            "Pending: {remaining}  ·  Extras: {extras}"
        ),
        "export_error": "Could not export: {error}",
        "found_exact": "FOUND — exact match",
        "found_variant": "FOUND — variant: {serial}",
        "not_found": "NOT FOUND — {serial}",
        "already_scanned": "ALREADY SCANNED — {serial}",
        "bad_format": "DOES NOT LOOK LIKE AN S/N — {serial}",
        "detail_location": "Sheet {sheet} · row {row}",
        "detail_location_csv": "Row {row} of the file",
        "detail_variant": "Read as {original}",
        "detail_repeat": "The counter does not go up. Esc dismisses the alert.",
        "detail_extra": "Saved in Extras (exported on a separate sheet).",
        "detail_bad_format": (
            "It has {length} characters and the file's S/Ns have "
            "{expected}. Was another barcode on the box scanned?"
        ),
        "legend_blue": "Blue: duplicate S/N (on export)",
        "legend_red": "Red: not scanned",
        "legend_green": "Green: scanned",
        "legend_yellow": "Yellow: already scanned",
        "legend_dup": "Blue: duplicate (can be deleted)",
        "legend_warn": "Yellow: active alert",
        "reset_arm": "Sure?",
        "reset_done": "Session reset.",
        "resume_prompt": (
            "This file has {done} of {total} S/Ns marked green "
            "(previously exported report). Resume where you left off?"
        ),
        "resume_continue": "Resume",
        "resume_fresh": "Start from scratch",
        "resumed": "Progress restored: {done} of {total} S/Ns already scanned.",
        "history_title": "Latest scans",
        "history_empty": "No scans yet",
        "hist_exact": "exact",
        "hist_variant": "variant",
        "hist_not_found": "not found",
        "hist_repeated": "repeated",
        "hist_bad_format": "format",
        "undo": "Undo",
        "scan_undone": "Scan undone: {serial}",
        "tab_pending": "Pending ({count})",
        "tab_extras": "Extras ({count})",
        "search": "Search",
        "pending_col": "Pending S/N",
        "pending_more": "… and {count} more",
        "pending_none": "No pending S/Ns left",
        "extras_col": "Extra S/N (not in the file)",
        "count_col": "Times",
        "time_col": "Time",
        "extras_none": "No extras",
        "caption_scanned": "Scanned",
        "caption_unique": "Unique total",
        "progress_info": "{pct}% done · {left} pending",
        "rate_info": "{rate}/min · ~{eta} min left",
        "destination": "Destination",
        "dest_column": "New column",
        "dest_sheet": "New sheet",
        "dest_cell": "From selected cell",
        "pallet_name": "Name (optional)",
        "next_pallet": "Next pallet",
        "delete_selected": "Delete",
        "export_list": "Export list",
        "create_count": (
            "Pallet {pallet}: {pallet_count} S/N · Duplicates: {duplicates}"
        ),
        "create_added": "RECORDED — {serial}",
        "duplicate_warning": "DUPLICATE — already exists: {serial}",
        "dup_banner": "Duplicates ({count}): {list}",
        "dup_remove": "Delete duplicates",
        "dup_hide": "Hide alert",
        "dups_removed": "{count} duplicates deleted · Ctrl+Z to undo",
        "pallet_ready": "Pallet {pallet} ready. Scan to add numbers.",
        "select_cell_first": "Select a cell in the viewer first.",
        "deleted_msg": "Deleted: {serial} · Ctrl+Z to undo",
        "undone": "Action undone.",
        "nothing_to_undo": "Nothing to undo.",
        "nothing_to_select": "Select a record in the viewer.",
        "nothing_to_export": "There are no records to export yet.",
        "export_list_title": "Save serial number list",
        "list_exported": (
            "List exported: {path}  ·  {count} records  ·  "
            "{duplicates} duplicates in blue"
        ),
        "csv_headers": "SN|Pallet|Sheet|Cell|Duplicate",
        "yes": "Yes",
        "no": "No",
        "session_saved": "Session saved",
        "session_save_error": "Could not save the session: {error}",
        "checker_session_restored": "Checker session restored.",
        "create_session_restored": "List creation session restored.",
        "paste_summary": (
            "Pasted {total}: {found} found · {missing} not found "
            "· {repeated} repeated"
        ),
        "paste_added": "Pasted {count} records · {duplicates} duplicates",
        "sheet_summary": "Summary",
        "sheet_log": "Scans",
        "sheet_extras": "Extras",
        "log_headers": "Time|Read|Result|Matched",
        "extras_headers": "Extra|Times|First read",
        "sum_file": "File",
        "sum_total": "Unique S/Ns in the file",
        "sum_found": "Found",
        "sum_pending": "Pending",
        "sum_extras": "Extras (not in the file)",
        "sum_scans": "Recorded reads",
        "sum_first": "First read",
        "sum_last": "Last read",
        "sum_exported": "Exported",
        "rack_mode": "Rack mode",
        "rack_location": "Location",
        "rack_rows": "Rows",
        "rack_cols": "Columns",
        "rack_order": "Order",
        "rack_order_h": "→ Horizontal",
        "rack_order_v": "↓ Vertical",
        "add_rack": "Add new rack",
        "new_sheet": "New sheet",
        "caption_next": "Next location",
        "rack_none": "Type the location (for example, DC05) and press “Add new rack”.",
        "rack_count": "{rack} · {used}/{total} locations\nDuplicates: {duplicates}",
        "rack_ready": "Rack {rack} ready ({rows} × {cols}). Starts at (0,0).",
        "rack_added": "RECORDED at {rack} ({x},{y}) — {serial}",
        "rack_completed": "RACK {rack} COMPLETE — {total} locations filled",
        "rack_full_scan": "Rack {rack} is complete. Press “Add new rack”.",
        "rack_need_location": "Type the location first (for example, DC05).",
        "rack_bad_size": "Rows and columns must be numbers between 1 and {maximum}.",
        "rack_target": "The next scan will go to {rack} ({x},{y}).",
        "rack_where": "{rack} ({x},{y})",
        "rack_full_banner": "Rack {rack} is complete ({total}/{total}).",
        "rack_list_sheet": "List",
        "rack_list_headers": "Location|Rack|Position (x,y)|X|Y|S/N|Duplicate",
        "print": "Print",
        "print_ask": "What do you want to print?",
        "print_ask_checker": "What do you want to print? Tick one or more options:",
        "print_current": "Current sheet",
        "print_all": "All sheets",
        "cancel": "Cancel",
        "print_summary": "Summary",
        "print_pending": "Pending",
        "print_extras": "Extras",
        "print_sent": "Sent to the default printer.",
        "print_error": "Could not print: {error}",
        "print_unavailable": "Printing only works on Windows. File prepared: {path}",
        "print_nothing": "Tick at least one option to print.",
        "readme": "README",
        "reset_label": "Reset",
        "keys_checker": (
            "F11|big view;Ctrl+Z|undo;Ctrl+E|export;Esc|dismiss alert"
        ),
        "validate_mode": "Validate list",
        "validate_load": "Load list",
        "validate_off": "Turn off",
        "validate_select": "Select the list of expected S/Ns",
        "validate_banner": "Validate list is on: load the CSV or Excel file with the expected S/Ns to start scanning.",
        "validate_need_list": "Load the validation list first.",
        "validate_invalid": "S/N NOT VALID FOR THE PALLET — {serial}",
        "validate_invalid_detail": "It is not in the validation list: it was not recorded.",
        "validate_loaded": "Validation list loaded: {name} · {count} S/Ns",
        "validate_count": "Validated {done} of {total} · {left} left",
        "validate_no_list": "No validation list",
        "validate_file": "List: {name}",
        "paste_validated": "Pasted {total}: {added} recorded · {invalid} not valid · {duplicates} duplicates",
        "sheet_renamed": "Sheet renamed: {old} → {new}",
        "sheet_name_taken": "There is already a sheet called “{name}”.",
        "tip_rack": "Scanning guided by the (x,y) locations of a rack. Each S/N is saved in the next free position and the app warns when the rack is full.",
        "tip_validate": "Load a CSV or Excel file with the expected S/Ns. Those not in the list are not recorded and a red alert is shown; repeated ones come in blue. Without a loaded list you cannot scan.",
        "tip_rack_locked": "Rack mode is not available while validating S/Ns on pallets.",
        "tip_validate_locked": "Validate list is not available while Rack mode is on.",
        "tip_tab": "Double-click to rename the sheet.",
        "copy_sn": "Copy S/N",
        "copied": "S/N copied: {serial}",
        "keys_rack": (
            "F2|new rack;Backspace|delete;Ctrl+Z|undo;Ctrl+E|export;"
            "Esc|dismiss alert"
        ),
        "keys_create": (
            "F2|pallet;Backspace|delete;Ctrl+Z|undo;Ctrl+E|export;"
            "Esc|dismiss alert"
        ),
    },

    "pt": {
        "mode_checker": "Checker",
        "mode_create": "Criar lista S/N",
        "alt_match": "Alternativas A/B/U",
        "big_view": "Visão ampliada",
        "big_hint": "F11 ou Esc para voltar",
        "no_file": "Nenhum arquivo CSV ou Excel carregado",
        "load": "Carregar CSV / Excel",
        "export": "Exportar relatório",
        "input_label": "(S/N) ou sequência para pesquisar",
        "input_label_create": "S/N ou texto a registrar (Enter)",
        "details": (
            "Carregue um CSV ou Excel e escaneie. Um XLSX pode ter várias "
            "planilhas; um CSV contém uma única tabela. Também é possível "
            "colar uma lista."
        ),
        "create_hint": (
            "Escaneie qualquer texto; Enter registra. "
            "Também é possível colar uma lista (Ctrl+V)."
        ),
        "footer": "Para o tutorial, leia README · Att: Dess",
        "select_file": "Selecionar arquivo",
        "file_types": "Arquivos CSV e Excel",
        "csv_files": "Arquivos CSV",
        "excel_files": "Arquivos Excel",
        "all_files": "Todos os arquivos",
        "unsupported": "Selecione um arquivo CSV (.csv) ou Excel (.xlsx).",
        "no_serials": (
            "Nenhum número de série foi encontrado nas colunas SN. "
            "Verifique os cabeçalhos do arquivo."
        ),
        "file_loaded": "Arquivo carregado. Pronto para verificar.",
        "file_summary": (
            "{name}  ·  {count} valores lidos  ·  {unique} S/N únicos"
        ),
        "read_error": "Não foi possível ler o arquivo: {error}",
        "excel_dependency": (
            "Instale openpyxl para abrir arquivos Excel "
            "(pip install openpyxl)."
        ),
        "csv_encoding": "Não foi possível identificar a codificação do CSV.",
        "load_first": "Carregue primeiro um arquivo CSV ou Excel.",
        "save_report": "Salvar relatório colorido",
        "output_file": "Arquivo Excel",
        "export_dependency": (
            "Instale openpyxl para exportar (pip install openpyxl)."
        ),
        "csv_read_error": "Não foi possível ler a codificação do CSV.",
        "invalid_format": "Formato de arquivo não compatível.",
        "same_file": "O relatório deve ser salvo como um novo arquivo.",
        "export_success": (
            "Relatório exportado: {path}  ·  Coloridas: {cells}  ·  "
            "Duplicados (azul): {duplicates}  ·  Encontrados: {found}  ·  "
            "Pendentes: {remaining}  ·  Sobras: {extras}"
        ),
        "export_error": "Não foi possível exportar: {error}",
        "found_exact": "ENCONTRADO — correspondência exata",
        "found_variant": "ENCONTRADO — variante: {serial}",
        "not_found": "NÃO ENCONTRADO — {serial}",
        "already_scanned": "JÁ ESCANEADO — {serial}",
        "bad_format": "NÃO PARECE UM S/N — {serial}",
        "detail_location": "Planilha {sheet} · linha {row}",
        "detail_location_csv": "Linha {row} do arquivo",
        "detail_variant": "Lido como {original}",
        "detail_repeat": "O contador não sobe. Esc para ocultar o aviso.",
        "detail_extra": "Salvo em Sobras (exportado em uma planilha separada).",
        "detail_bad_format": (
            "Tem {length} caracteres e os S/N do arquivo têm "
            "{expected}. Foi escaneado outro código da caixa?"
        ),
        "legend_blue": "Azul: S/N duplicado (ao exportar)",
        "legend_red": "Vermelho: não escaneado",
        "legend_green": "Verde: escaneado",
        "legend_yellow": "Amarelo: já escaneado",
        "legend_dup": "Azul: duplicado (pode excluir)",
        "legend_warn": "Amarelo: aviso ativo",
        "reset_arm": "Certeza?",
        "reset_done": "Sessão reiniciada.",
        "resume_prompt": (
            "Este arquivo tem {done} de {total} S/N marcados em verde "
            "(relatório exportado antes). Continuar de onde parou?"
        ),
        "resume_continue": "Continuar",
        "resume_fresh": "Começar do zero",
        "resumed": "Progresso recuperado: {done} de {total} S/N já escaneados.",
        "history_title": "Últimas leituras",
        "history_empty": "Ainda não há leituras",
        "hist_exact": "exato",
        "hist_variant": "variante",
        "hist_not_found": "não encontrado",
        "hist_repeated": "repetido",
        "hist_bad_format": "formato",
        "undo": "Desfazer",
        "scan_undone": "Leitura desfeita: {serial}",
        "tab_pending": "Pendentes ({count})",
        "tab_extras": "Sobras ({count})",
        "search": "Buscar",
        "pending_col": "S/N pendente",
        "pending_more": "… e mais {count}",
        "pending_none": "Não há S/N pendentes",
        "extras_col": "S/N sobrando (não está no arquivo)",
        "count_col": "Vezes",
        "time_col": "Hora",
        "extras_none": "Não há sobras",
        "caption_scanned": "Escaneados",
        "caption_unique": "Total único",
        "progress_info": "{pct}% concluído · {left} pendentes",
        "rate_info": "{rate}/min · faltam ~{eta} min",
        "destination": "Destino",
        "dest_column": "Nova coluna",
        "dest_sheet": "Nova planilha",
        "dest_cell": "A partir da célula selecionada",
        "pallet_name": "Nome (opcional)",
        "next_pallet": "Próximo pallet",
        "delete_selected": "Excluir",
        "export_list": "Exportar lista",
        "create_count": (
            "Pallet {pallet}: {pallet_count} S/N · Duplicados: {duplicates}"
        ),
        "create_added": "REGISTRADO — {serial}",
        "duplicate_warning": "DUPLICADO — já existe: {serial}",
        "dup_banner": "Duplicados ({count}): {list}",
        "dup_remove": "Excluir duplicados",
        "dup_hide": "Ocultar aviso",
        "dups_removed": "{count} duplicados excluídos · Ctrl+Z para desfazer",
        "pallet_ready": "Pallet {pallet} pronto. Escaneie para adicionar números.",
        "select_cell_first": "Selecione primeiro uma célula no visor.",
        "deleted_msg": "Excluído: {serial} · Ctrl+Z para desfazer",
        "undone": "Ação desfeita.",
        "nothing_to_undo": "Não há nada para desfazer.",
        "nothing_to_select": "Selecione um registro no visor.",
        "nothing_to_export": "Ainda não há registros para exportar.",
        "export_list_title": "Salvar lista de números de série",
        "list_exported": (
            "Lista exportada: {path}  ·  {count} registros  ·  "
            "{duplicates} duplicados em azul"
        ),
        "csv_headers": "SN|Pallet|Planilha|Célula|Duplicado",
        "yes": "Sim",
        "no": "Não",
        "session_saved": "Sessão salva",
        "session_save_error": "Não foi possível salvar a sessão: {error}",
        "checker_session_restored": "Sessão do Checker recuperada.",
        "create_session_restored": "Sessão de criação de lista recuperada.",
        "paste_summary": (
            "Colados {total}: {found} encontrados · {missing} não encontrados "
            "· {repeated} repetidos"
        ),
        "paste_added": "Colados {count} registros · {duplicates} duplicados",
        "sheet_summary": "Resumo",
        "sheet_log": "Leituras",
        "sheet_extras": "Sobras",
        "log_headers": "Hora|Lido|Resultado|Corresponde a",
        "extras_headers": "Sobra|Vezes|Primeira leitura",
        "sum_file": "Arquivo",
        "sum_total": "S/N únicos no arquivo",
        "sum_found": "Encontrados",
        "sum_pending": "Pendentes",
        "sum_extras": "Sobras (não estão no arquivo)",
        "sum_scans": "Leituras registradas",
        "sum_first": "Primeira leitura",
        "sum_last": "Última leitura",
        "sum_exported": "Exportado",
        "rack_mode": "Modo Rack",
        "rack_location": "Localização",
        "rack_rows": "Linhas",
        "rack_cols": "Colunas",
        "rack_order": "Ordem",
        "rack_order_h": "→ Horizontal",
        "rack_order_v": "↓ Vertical",
        "add_rack": "Adicionar novo rack",
        "new_sheet": "Nova planilha",
        "caption_next": "Próxima posição",
        "rack_none": "Digite a localização (por exemplo, DC05) e clique em «Adicionar novo rack».",
        "rack_count": "{rack} · {used}/{total} posições\nDuplicados: {duplicates}",
        "rack_ready": "Rack {rack} pronto ({rows} × {cols}). Começa em (0,0).",
        "rack_added": "REGISTRADO em {rack} ({x},{y}) — {serial}",
        "rack_completed": "RACK {rack} COMPLETO — {total} posições preenchidas",
        "rack_full_scan": "O rack {rack} está completo. Clique em «Adicionar novo rack».",
        "rack_need_location": "Digite primeiro a localização (por exemplo, DC05).",
        "rack_bad_size": "Linhas e colunas devem ser números entre 1 e {maximum}.",
        "rack_target": "A próxima leitura irá para {rack} ({x},{y}).",
        "rack_where": "{rack} ({x},{y})",
        "rack_full_banner": "O rack {rack} foi completado ({total}/{total}).",
        "rack_list_sheet": "Lista",
        "rack_list_headers": "Localização|Rack|Posição (x,y)|X|Y|S/N|Duplicado",
        "print": "Imprimir",
        "print_ask": "O que você quer imprimir?",
        "print_ask_checker": "O que você quer imprimir? Marque uma ou mais opções:",
        "print_current": "Planilha atual",
        "print_all": "Todas as planilhas",
        "cancel": "Cancelar",
        "print_summary": "Resumo",
        "print_pending": "Pendentes",
        "print_extras": "Sobras",
        "print_sent": "Enviado para a impressora padrão.",
        "print_error": "Não foi possível imprimir: {error}",
        "print_unavailable": "Imprimir só funciona no Windows. Arquivo preparado: {path}",
        "print_nothing": "Marque pelo menos uma opção para imprimir.",
        "readme": "README",
        "reset_label": "Reiniciar",
        "keys_checker": (
            "F11|visão ampliada;Ctrl+Z|desfazer;Ctrl+E|exportar;"
            "Esc|ocultar aviso"
        ),
        "validate_mode": "Validar lista",
        "validate_load": "Carregar lista",
        "validate_off": "Desativar",
        "validate_select": "Selecionar a lista de S/N esperados",
        "validate_banner": "Validar lista está ativo: carregue o CSV ou Excel com os S/N esperados para começar a escanear.",
        "validate_need_list": "Carregue primeiro a lista de validação.",
        "validate_invalid": "S/N NÃO VÁLIDO PARA O PALLET — {serial}",
        "validate_invalid_detail": "Não está na lista de validação: não foi registrado.",
        "validate_loaded": "Lista de validação carregada: {name} · {count} S/N",
        "validate_count": "Validados {done} de {total} · faltam {left}",
        "validate_no_list": "Sem lista de validação",
        "validate_file": "Lista: {name}",
        "paste_validated": "Colados {total}: {added} registrados · {invalid} não válidos · {duplicates} duplicados",
        "sheet_renamed": "Planilha renomeada: {old} → {new}",
        "sheet_name_taken": "Já existe uma planilha chamada «{name}».",
        "tip_rack": "Leitura guiada pelas posições (x,y) de um rack. Cada S/N é salvo na próxima posição livre e o aplicativo avisa quando o rack fica cheio.",
        "tip_validate": "Carregue um CSV ou Excel com os S/N esperados. Os que não estiverem na lista não são registrados e aparece um aviso em vermelho; os repetidos entram em azul. Sem lista carregada não é possível escanear.",
        "tip_rack_locked": "Modo Rack indisponível durante a validação de S/N nos pallets.",
        "tip_validate_locked": "Validar lista indisponível enquanto o Modo Rack estiver ativo.",
        "tip_tab": "Clique duas vezes para renomear a planilha.",
        "copy_sn": "Copiar S/N",
        "copied": "S/N copiado: {serial}",
        "keys_rack": (
            "F2|novo rack;Backspace|excluir;Ctrl+Z|desfazer;Ctrl+E|exportar;"
            "Esc|ocultar aviso"
        ),
        "keys_create": (
            "F2|pallet;Backspace|excluir;Ctrl+Z|desfazer;Ctrl+E|exportar;"
            "Esc|ocultar aviso"
        ),
    },
}

# Hojas que el reporte añade (en cualquier idioma). Al leer o volver a
# exportar un reporte se ignoran para no confundirlas con datos.
REPORT_SHEETS = {
    TEXTS[lang][key].lower()
    for lang in TEXTS
    for key in ("sheet_summary", "sheet_log", "sheet_extras")
}


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def normalize_header(value):
    """Normaliza un encabezado: sin acentos, signos, espacios ni mayúsculas.

    "S/N", "s.n.", "Nº de serie" o "Serial Number" se comparan igual
    que SN, NODESERIE o SERIALNUMBER.
    """
    text = unicodedata.normalize("NFKD", str(value).strip()).upper()

    return "".join(
        ch for ch in text
        if ch.isalnum() and not unicodedata.combining(ch)
    )


def column_letter(index):
    """1 -> A, 2 -> B, 27 -> AA."""
    letters = ""

    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters

    return letters


def parse_cell(reference):
    """'B5' -> (fila 5, columna 2). Devuelve None si no es válida."""
    match = re.fullmatch(
        r"\s*([A-Za-z]{1,3})(\d{1,7})\s*",
        str(reference),
    )

    if not match:
        return None

    letters, row = match.groups()
    column = 0

    for ch in letters.upper():
        column = column * 26 + (ord(ch) - 64)

    row = int(row)

    if row < 1 or column < 1:
        return None

    return row, column


def safe_sheet_name(name, used):
    """Nombre de hoja válido para Excel y sin repetir (sin distinguir
    mayúsculas)."""
    cleaned = "".join(
        ch for ch in str(name) if ch not in "[]:*?/\\"
    ).strip()[:31] or DEFAULT_SHEET

    candidate = cleaned
    number = 2

    while candidate.lower() in used:
        suffix = f"_{number}"
        candidate = cleaned[:31 - len(suffix)] + suffix
        number += 1

    used.add(candidate.lower())

    return candidate


def autofit_columns(worksheet, columns=None, minimum=4, padding=2, maximum=90):
    """Ancho de cada columna justo para su texto más largo (más un margen).

    Así el archivo se ve completo en Excel o Google Sheets sin ajustar las
    columnas a mano. columns: solo esas columnas (números desde 1).
    """
    lengths = {}

    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None or (columns and cell.column not in columns):
                continue

            length = max(len(line) for line in str(cell.value).split("\n"))
            lengths[cell.column] = max(lengths.get(cell.column, 0), length)

    for column, length in lengths.items():
        worksheet.column_dimensions[column_letter(column)].width = min(
            max(length + padding, minimum), maximum,
        )


def compute_layout(items, pallets, default_sheet=DEFAULT_SHEET):
    """Calcula en qué hoja, fila y columna queda cada registro.

    Es la misma función para el visor y para exportar, así lo que se ve
    es exactamente lo que se guarda. default_sheet es el nombre de la hoja
    principal (el usuario la puede renombrar).

    items:   [{"serial", "pallet", "dup", "slot"?}]
    pallets: [{"number", "mode", "name", "sheet", "cell",
               "rows"?, "cols"?, "order"?}]
             mode: "column" | "sheet" | "cell" | "rack"

    En un rack cada registro tiene su ubicación "slot" = [x, y] (desde 0):
    la celda es la fila y+1, columna x+1 de la hoja del rack.

    Devuelve {"sheets": {nombre: {"cells", "max_row", "max_col", "rack"?}},
              "positions": {indice_item: (hoja, fila, columna)},
              "pallet_sheets": {número_pallet: hoja}}.
    """
    sheets = OrderedDict()
    positions = {}
    pallet_sheets = {}
    used = {default_sheet.lower()}
    next_column = {}

    def get_sheet(name):
        if name not in sheets:
            sheets[name] = {"cells": {}, "max_row": 0, "max_col": 0}

        return sheets[name]

    def put(sheet_name, row, col, text, kind, index=None):
        sheet = get_sheet(sheet_name)
        sheet["cells"][(row, col)] = {
            "text": text,
            "kind": kind,
            "index": index,
        }
        sheet["max_row"] = max(sheet["max_row"], row)
        sheet["max_col"] = max(sheet["max_col"], col)

        if index is not None:
            positions[index] = (sheet_name, row, col)

    by_pallet = {}

    for index, item in enumerate(items):
        by_pallet.setdefault(item["pallet"], []).append(index)

    for pallet in pallets:
        number = pallet["number"]
        indexes = by_pallet.get(number, [])
        mode = pallet.get("mode", "column")
        label = pallet.get("name") or f"Pallet {number}"

        def kind_of(index):
            return "dup" if items[index].get("dup") else "item"

        if mode == "rack":
            sheet_name = safe_sheet_name(label, used)
            sheet = get_sheet(sheet_name)
            rows = pallet.get("rows") or RACK_ROWS
            cols = pallet.get("cols") or RACK_COLS
            sheet["rack"] = {"rows": rows, "cols": cols, "pallet": number}
            sheet["max_row"] = max(sheet["max_row"], rows)
            sheet["max_col"] = max(sheet["max_col"], cols)

            for index in indexes:
                x, y = items[index].get("slot") or (0, 0)
                put(sheet_name, y + 1, x + 1, items[index]["serial"],
                    kind_of(index), index)

        elif mode == "sheet":
            sheet_name = safe_sheet_name(label, used)
            put(sheet_name, 1, 1, SN_HEADER_TEXT, "header")
            row = 2

            for index in indexes:
                put(sheet_name, row, 1, items[index]["serial"],
                    kind_of(index), index)
                row += 1

        elif mode == "cell":
            sheet_name = pallet.get("sheet") or default_sheet
            start = parse_cell(pallet.get("cell", "A1")) or (1, 1)
            row, col = start
            sheet = get_sheet(sheet_name)

            for index in indexes:
                while (row, col) in sheet["cells"]:
                    row += 1

                put(sheet_name, row, col, items[index]["serial"],
                    kind_of(index), index)
                row += 1

        else:
            sheet_name = default_sheet
            sheet = get_sheet(sheet_name)
            col = next_column.get(sheet_name, 1)

            while (1, col) in sheet["cells"]:
                col += 1

            next_column[sheet_name] = col + 1
            put(sheet_name, 1, col, label, "header")
            row = 2

            for index in indexes:
                while (row, col) in sheet["cells"]:
                    row += 1

                put(sheet_name, row, col, items[index]["serial"],
                    kind_of(index), index)
                row += 1

        pallet_sheets[number] = sheet_name

    return {
        "sheets": sheets,
        "positions": positions,
        "pallet_sheets": pallet_sheets,
    }


# ============================================================
# FUENTES, FORMAS E ICONOS
# ============================================================

def ui(size, weight="normal"):
    return (UI_FONT, size, weight)


def mono(size, weight="normal"):
    return (MONO_FONT, size, weight)


def pick_fonts(root):
    """Elige las fuentes instaladas: Segoe UI Variable en Windows 11,
    Segoe UI en Windows 10 y una alternativa en otros sistemas."""
    global UI_FONT, MONO_FONT

    if tkfont is None:
        return

    try:
        families = set(tkfont.families(root))
    except Exception:
        return

    for name in ("Segoe UI Variable Text", "Segoe UI", "Inter", "DejaVu Sans"):
        if name in families:
            UI_FONT = name
            break

    for name in ("Cascadia Mono", "Consolas", "JetBrains Mono",
                 "DejaVu Sans Mono"):
        if name in families:
            MONO_FONT = name
            break


_FONT_CACHE = {}


def text_width(font, text):
    """Ancho en píxeles de un texto (aproximado si no hay tkinter.font)."""
    if tkfont is None:
        return 7 * len(text)

    font_obj = _FONT_CACHE.get(font)

    if font_obj is None:
        try:
            font_obj = tkfont.Font(
                family=font[0], size=font[1],
                weight="bold" if "bold" in font[2:] else "normal",
            )
        except Exception:
            return 7 * len(text)

        _FONT_CACHE[font] = font_obj

    return font_obj.measure(text)


def round_rect_points(x0, y0, x1, y1, radius):
    """Puntos de un rectángulo con esquinas redondeadas (polígono suave)."""
    r = max(0, min(radius, (x1 - x0) / 2, (y1 - y0) / 2))

    return [
        x0 + r, y0, x0 + r, y0, x1 - r, y0, x1 - r, y0,
        x1, y0, x1, y0 + r, x1, y0 + r, x1, y1 - r, x1, y1 - r,
        x1, y1, x1 - r, y1, x1 - r, y1, x0 + r, y1, x0 + r, y1,
        x0, y1, x0, y1 - r, x0, y1 - r, x0, y0 + r, x0, y0 + r, x0, y0,
    ]


# Iconos de trazo en una caja de 24×24 (se escalan al dibujar).
# Cada trazo: ("line", puntos) / ("poly", puntos) / ("curve", puntos)
# / ("oval", caja) / ("arc", caja, inicio, extensión) / ("dot", x, y).
ICONS = {
    "folder": [("poly", (3, 6, 9, 6, 11, 8, 21, 8, 21, 19, 3, 19))],
    "download": [
        ("line", (12, 4, 12, 15)),
        ("line", (7, 10, 12, 15, 17, 10)),
        ("line", (5, 20, 19, 20)),
    ],
    "plus": [("line", (12, 5, 12, 19)), ("line", (5, 12, 19, 12))],
    "trash": [
        ("line", (4, 7, 20, 7)),
        ("line", (6, 7, 7, 20, 17, 20, 18, 7)),
        ("line", (9, 7, 9, 4, 15, 4, 15, 7)),
    ],
    "undo": [
        ("line", (9, 14, 4, 9, 9, 4)),
        ("curve", (4, 9, 15, 9, 20, 14, 15, 19, 11, 19)),
    ],
    "reset": [
        ("arc", (4, 4, 20, 20), 70, 290),
        ("line", (20, 3, 20, 9, 14, 9)),
    ],
    "expand": [
        ("line", (4, 9, 4, 4, 9, 4)),
        ("line", (15, 4, 20, 4, 20, 9)),
        ("line", (20, 15, 20, 20, 15, 20)),
        ("line", (9, 20, 4, 20, 4, 15)),
    ],
    "check": [("line", (5, 12, 10, 17, 19, 7))],
    "x": [("line", (6, 6, 18, 18)), ("line", (18, 6, 6, 18))],
    "warn": [
        ("poly", (12, 3, 22, 21, 2, 21)),
        ("line", (12, 10, 12, 15)),
        ("dot", 12, 18),
    ],
    "info": [
        ("oval", (3, 3, 21, 21)),
        ("line", (12, 11, 12, 17)),
        ("dot", 12, 7.5),
    ],
    "scan": [
        ("line", (3, 8, 3, 3, 8, 3)),
        ("line", (16, 3, 21, 3, 21, 8)),
        ("line", (21, 16, 21, 21, 16, 21)),
        ("line", (8, 21, 3, 21, 3, 16)),
        ("line", (8, 8, 8, 16)),
        ("line", (11, 8, 11, 16)),
        ("line", (14, 8, 14, 16)),
        ("line", (17, 8, 17, 16)),
    ],
    "globe": [
        ("oval", (3, 3, 21, 21)),
        ("line", (3, 12, 21, 12)),
        ("oval", (8, 3, 16, 21)),
    ],
    "file": [
        ("poly", (6, 3, 14, 3, 19, 8, 19, 21, 6, 21)),
        ("line", (14, 3, 14, 8, 19, 8)),
        ("line", (9, 13, 16, 13)),
        ("line", (9, 17, 16, 17)),
    ],
    "printer": [
        ("line", (7, 8, 7, 3, 17, 3, 17, 8)),
        ("poly", (4, 8, 20, 8, 20, 16, 4, 16)),
        ("poly", (7, 13, 17, 13, 17, 21, 7, 21)),
    ],
    "book": [
        ("poly", (4, 4, 10, 4, 12, 6, 12, 20, 10, 18, 4, 18)),
        ("poly", (20, 4, 14, 4, 12, 6, 12, 20, 14, 18, 20, 18)),
    ],
    "search": [
        ("oval", (4, 4, 16, 16)),
        ("line", (14.5, 14.5, 20, 20)),
    ],
}


def draw_icon(canvas, name, cx, cy, color, size=16, width=2, tags=None):
    """Dibuja un icono de trazo centrado en (cx, cy)."""
    scale = size / 24
    ox = cx - size / 2
    oy = cy - size / 2

    def pts(values):
        return [
            (ox if i % 2 == 0 else oy) + v * scale
            for i, v in enumerate(values)
        ]

    common = {"tags": tags} if tags else {}

    for stroke in ICONS.get(name, []):
        kind = stroke[0]

        if kind == "line":
            canvas.create_line(
                *pts(stroke[1]), fill=color, width=width,
                capstyle="round", joinstyle="round", **common,
            )
        elif kind == "curve":
            canvas.create_line(
                *pts(stroke[1]), fill=color, width=width, smooth=True,
                capstyle="round", **common,
            )
        elif kind == "poly":
            canvas.create_polygon(
                *pts(stroke[1]), fill="", outline=color, width=width,
                joinstyle="round", **common,
            )
        elif kind == "oval":
            canvas.create_oval(
                *pts(stroke[1]), outline=color, width=width, **common,
            )
        elif kind == "arc":
            canvas.create_arc(
                *pts(stroke[1]), start=stroke[2], extent=stroke[3],
                style="arc", outline=color, width=width, **common,
            )
        elif kind == "dot":
            x, y = pts((stroke[1], stroke[2]))
            r = max(width * 0.7, 1)
            canvas.create_oval(
                x - r, y - r, x + r, y + r, fill=color, outline="", **common,
            )


# ============================================================
# COMPONENTES DE LA INTERFAZ
# ============================================================

class Card(tk.Canvas):
    """Tarjeta con esquinas redondeadas. El contenido va en self.body.

    fill=False: la tarjeta toma la altura de su contenido.
    fill=True:  el contenido ocupa el alto que le dé la ventana.
    """

    def __init__(self, parent, bg=PANEL, border=BORDER, outer=BG, radius=12,
                 padx=18, pady=14, fill=False, border_width=1, height=10):
        # Altura inicial amplia: el contenido solo se dibuja (y avisa de su
        # tamaño) si cabe en el lienzo; luego la tarjeta se ajusta a él.
        super().__init__(
            parent, bg=outer, highlightthickness=0, bd=0,
            width=10, height=height if fill else 300,
        )
        self.card_bg = bg
        self.card_border = border
        self.radius = radius
        self.padx = padx
        self.pady = pady
        self.fill = fill
        self.border_width = border_width

        self.shape = self.create_polygon(
            0, 0, 0, 0, smooth=True, fill=bg, outline=border,
            width=border_width,
        )
        self.body = tk.Frame(self, bg=bg)
        self.window = self.create_window(
            padx, pady, anchor="nw", window=self.body,
        )
        self.bind("<Configure>", self._on_resize)

        if not fill:
            self.body.bind("<Configure>", self._fit_height)

    def _on_resize(self, event):
        inset = self.border_width / 2 + 0.5
        self.coords(
            self.shape,
            *round_rect_points(
                inset, inset, event.width - inset - 1,
                event.height - inset - 1, self.radius,
            ),
        )
        self.itemconfigure(self.window, width=max(event.width - 2 * self.padx, 1))

        if self.fill:
            self.itemconfigure(
                self.window, height=max(event.height - 2 * self.pady, 1),
            )

    def _fit_height(self, event):
        wanted = event.height + 2 * self.pady

        if str(self.cget("height")) != str(wanted):
            self.configure(height=wanted)

    def set_colors(self, bg=None, border=None):
        if bg is not None:
            self.card_bg = bg
            self.body.configure(bg=bg)

        if border is not None:
            self.card_border = border

        self.itemconfigure(
            self.shape, fill=self.card_bg, outline=self.card_border,
        )


class RoundButton(tk.Canvas):
    """Botón redondeado con estados al pasar el ratón y al pulsar.

    variant: primary | secondary | ghost | danger | warn | warn_ghost | tab
    """

    def __init__(self, parent, text="", command=None, variant="secondary",
                 icon=None, size=10, bold=True, height=36, padx=14,
                 radius=8, bg=PANEL, width=None):
        super().__init__(
            parent, height=height, width=width or 10, bg=bg,
            highlightthickness=0, bd=0, cursor="hand2",
        )
        self.text = text
        self._command = command
        self._variant = variant
        self._icon = icon
        self._size = size
        self._bold = bold
        self._height = height
        self._padx = padx
        self._radius = radius
        self._fixed_width = width
        self._parent_bg = bg
        self._hover = False
        self._pressed = False

        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<ButtonPress-1>", self._on_press)
        self.bind("<ButtonRelease-1>", self._on_release)
        self.bind("<Configure>", lambda e: self._draw())
        self._draw()

    # ----- configuración compatible con tk.Button -----

    def configure(self, cnf=None, **kw):
        changed = False

        for key, attr in (
            ("text", "text"), ("command", "_command"),
            ("variant", "_variant"), ("icon", "_icon"),
        ):
            if key in kw:
                setattr(self, attr, kw.pop(key))
                changed = True

        if "bg" in kw:
            self._parent_bg = kw["bg"]
            changed = True

        if kw:
            super().configure(**kw)

        if changed:
            self._draw()

    config = configure

    def invoke(self):
        if self._command:
            self._command()

    # ----- dibujo -----

    def _font(self):
        return ui(self._size, "bold" if self._bold else "normal")

    def _natural_width(self):
        width = 2 * self._padx

        if self._icon:
            width += 16 + (8 if self.text else 0)

        if self.text:
            width += text_width(self._font(), self.text)

        return int(max(width, self._height))

    def _colors(self):
        parent = self._parent_bg
        hover = self._hover or self._pressed

        table = {
            "primary": (BLUE_DARK if hover else BLUE, "#ffffff", BLUE_LIGHT),
            "secondary": (
                BUTTON_ACTIVE if hover else BUTTON_BG, WHITE, BORDER_STRONG,
            ),
            "ghost": (BUTTON_BG if hover else parent, MUTED, parent),
            "tab": (BUTTON_BG if hover else parent, MUTED, parent),
            "danger": ("#dc2626" if hover else RED, "#ffffff", RED),
            "warn": ("#d97706" if hover else YELLOW, "#1c1300", YELLOW),
            "warn_ghost": (
                "#4d3b0a" if hover else YELLOW_BG, YELLOW_TEXT, YELLOW_BORDER,
            ),
        }

        return table.get(self._variant, table["secondary"])

    def _draw(self):
        natural = self._natural_width()

        if self._fixed_width is None and str(self.cget("width")) != str(natural):
            super().configure(width=natural)

        self.delete("all")
        width = self.winfo_width()

        if not isinstance(width, int) or width <= 1:
            width = self._fixed_width or natural

        height = self._height
        bg, fg, border = self._colors()

        self.create_polygon(
            *round_rect_points(0.5, 0.5, width - 1.5, height - 1.5, self._radius),
            smooth=True, fill=bg, outline=border,
        )

        content = 0
        font = self._font()

        if self._icon:
            content += 16

        if self.text:
            content += text_width(font, self.text) + (8 if self._icon else 0)

        x = (width - content) / 2
        y = height / 2

        if self._icon:
            draw_icon(self, self._icon, x + 8, y, fg, size=16, width=2)
            x += 24

        if self.text:
            self.create_text(x, y, text=self.text, anchor="w", fill=fg, font=font)

    # ----- ratón -----

    def _on_enter(self, event=None):
        self._hover = True
        self._draw()

    def _on_leave(self, event=None):
        self._hover = False
        self._pressed = False
        self._draw()

    def _on_press(self, event=None):
        self._pressed = True
        self._draw()

    def _on_release(self, event=None):
        inside = (
            event is None
            or (0 <= event.x <= self.winfo_width()
                and 0 <= event.y <= self.winfo_height())
        )
        self._pressed = False
        self._draw()

        if inside and self._command:
            self._command()


class CheckBox(tk.Frame):
    """Casilla de verificación dibujada (encaja con el tema oscuro)."""

    def __init__(self, parent, variable, command=None, bg=PANEL):
        super().__init__(parent, bg=bg, cursor="hand2")
        self.variable = variable
        self.command = command
        self.box = tk.Canvas(
            self, width=18, height=18, bg=bg, highlightthickness=0, bd=0,
        )
        self.box.pack(side="left")
        self.label_color = WHITE
        self.label = tk.Label(self, text="", font=ui(10), bg=bg, fg=WHITE)
        self.label.pack(side="left", padx=(8, 0))

        self.parts = (self.box, self.label)

        for widget in (self, self.box, self.label):
            widget.bind("<Button-1>", self._toggle)

        self.redraw()

    enabled = True

    def _toggle(self, event=None):
        if not self.enabled:
            return

        self.variable.set(not self.variable.get())
        self.redraw()

        if self.command:
            self.command()

    def set_enabled(self, enabled):
        """Bloqueada: se ve en gris y no se puede marcar."""
        self.enabled = enabled
        self.config(cursor="hand2" if enabled else "arrow")
        self.label.config(fg=self.label_color if enabled else DIM)
        self.redraw()

    def redraw(self):
        on = bool(self.variable.get()) and self.enabled
        self.box.delete("all")

        if not self.enabled:
            fill, outline = CONTROL, BORDER
        else:
            fill = BLUE if on else FIELD
            outline = BLUE if on else BORDER_STRONG

        self.box.create_polygon(
            *round_rect_points(1, 1, 17, 17, 5), smooth=True,
            fill=fill, outline=outline,
        )

        if on:
            draw_icon(self.box, "check", 9, 9, "#ffffff", size=13, width=2.5)

    def set_text(self, text):
        self.label.config(text=text)


class Tooltip:
    """Leyenda corta al dejar el ratón sobre un control.

    Aparece tras medio segundo y desaparece al alejar el ratón. No toma el
    foco: el escáner sigue escribiendo en el campo de escaneo.
    """

    DELAY_MS = 500

    def __init__(self, widget, text_func):
        self.widget = widget
        self.text_func = text_func
        self.job = None
        self.tip = None

        parts = getattr(widget, "parts", ())
        parts = list(parts) if isinstance(parts, (list, tuple)) else []

        for part in [widget] + parts:
            part.bind("<Enter>", self._schedule, add="+")
            part.bind("<Leave>", self._leave, add="+")
            part.bind("<ButtonPress>", self._hide, add="+")

    def _schedule(self, event=None):
        self._cancel()
        self.job = self.widget.after(self.DELAY_MS, self._show)

    def _cancel(self):
        if self.job is not None:
            try:
                self.widget.after_cancel(self.job)
            except Exception:
                pass

            self.job = None

    def _leave(self, event=None):
        # Pasar de la casilla a su texto también es "salir": solo se oculta
        # si el ratón ya no está encima del control.
        try:
            x, y = self.widget.winfo_pointerxy()
            left, top = self.widget.winfo_rootx(), self.widget.winfo_rooty()
            inside = (
                left <= x < left + self.widget.winfo_width()
                and top <= y < top + self.widget.winfo_height()
            )
        except Exception:
            inside = False

        if not inside:
            self._hide()

    def _show(self):
        self.job = None

        if self.tip is not None:
            return

        text = self.text_func()

        if not text:
            return

        tip = tk.Toplevel(self.widget)
        tip.wm_overrideredirect(True)

        try:
            tip.attributes("-topmost", True)
        except Exception:
            pass

        frame = tk.Frame(tip, bg=BORDER_STRONG, padx=1, pady=1)
        frame.pack()
        tk.Label(
            frame, text=text, font=ui(9), bg=CONTROL, fg=WHITE,
            justify="left", wraplength=320, padx=10, pady=7,
        ).pack()
        x = self.widget.winfo_rootx()
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 6
        tip.wm_geometry(f"+{x}+{y}")
        self.tip = tip

    def _hide(self, event=None):
        self._cancel()

        if self.tip is not None:
            try:
                self.tip.destroy()
            except Exception:
                pass

            self.tip = None


class ProgressBar(tk.Canvas):
    def __init__(self, parent, bg=PANEL):
        super().__init__(
            parent, height=8, width=10, bg=bg, highlightthickness=0, bd=0,
        )
        self.value = 0.0
        self.color = GREEN
        self.bind("<Configure>", lambda e: self._draw())

    def set(self, value, color=None):
        self.value = max(0.0, min(1.0, value))

        if color:
            self.color = color

        self._draw()

    def _draw(self):
        self.delete("all")
        width = self.winfo_width()

        if not isinstance(width, int) or width <= 1:
            return

        self.create_polygon(
            *round_rect_points(0, 0, width - 1, 7, 4), smooth=True,
            fill="#23262e", outline="",
        )

        if self.value > 0:
            filled = max(8, (width - 1) * self.value)
            self.create_polygon(
                *round_rect_points(0, 0, filled, 7, 4), smooth=True,
                fill=self.color, outline="",
            )


# ============================================================
# IMAGEN DE FONDO (patrón de líneas en mosaico)
# ============================================================

def load_background_tile():
    """Devuelve el PhotoImage del patrón o None si no se puede cargar."""
    try:
        return tk.PhotoImage(data="".join(BACKGROUND_TILE))
    except Exception:
        return None


# ============================================================
# VISOR TIPO HOJA DE CÁLCULO
# ============================================================

class SheetGrid(tk.Frame):
    """Cuadrícula con columnas A, B, C... y filas numeradas.

    Muestra la hoja que se está generando, permite seleccionar una celda
    con un clic y editar su contenido con doble clic. Las columnas se
    estiran arrastrando su borde en la fila de letras (doble clic en el
    borde: ancho automático).
    """

    # Medidas mínimas; las reales se calculan con la fuente (ver __init__).
    CELL_W = 150
    CELL_H = 26
    HDR_W = 46
    HDR_H = 26
    CELL_CHARS = 21
    CELL_PAD = 8
    MIN_COL_W = 50
    MAX_COL_W = 900
    GRAB = 5
    MIN_ROWS = 14
    MIN_COLS = 6

    def __init__(self, parent, on_select=None, on_edit=None, on_done=None,
                 on_resize=None):
        super().__init__(parent, bg=PANEL)

        self.FONT = mono(10)
        self.FONT_BOLD = mono(10, "bold")

        # La celda se mide con la fuente real para que quepan 21 letras
        # con cualquier escala de Windows (125 %, 150 %...).
        self.CELL_W = max(
            self.CELL_W,
            text_width(self.FONT_BOLD, "W" * self.CELL_CHARS)
            + 2 * self.CELL_PAD + 4,
        )
        self.CELL_H = max(self.CELL_H, self._line_height(self.FONT) + 10)
        self.HDR_W = max(self.HDR_W, text_width(ui(9), "00000") + 12)
        self.HDR_H = self.CELL_H

        self.on_select = on_select
        self.on_edit = on_edit
        self.on_done = on_done
        self.on_resize = on_resize

        self.col_canvas = tk.Canvas(
            self, height=self.HDR_H, bg=HEADER_BG, highlightthickness=0,
        )
        self.row_canvas = tk.Canvas(
            self, width=self.HDR_W, bg=HEADER_BG, highlightthickness=0,
        )
        self.canvas = tk.Canvas(
            self,
            height=120,
            bg=EMPTY_BG,
            highlightthickness=0,
            yscrollincrement=self.CELL_H,
            xscrollincrement=30,
        )
        self.vbar = ttk.Scrollbar(
            self, orient="vertical", command=self._yview,
            style="Thin.Vertical.TScrollbar",
        )
        self.hbar = ttk.Scrollbar(
            self, orient="horizontal", command=self._xview,
            style="Thin.Horizontal.TScrollbar",
        )
        self.canvas.configure(
            xscrollcommand=self._on_xscroll,
            yscrollcommand=self._on_yscroll,
        )

        corner = tk.Frame(
            self, bg=HEADER_BG, width=self.HDR_W, height=self.HDR_H,
        )
        corner.grid(row=0, column=0)
        self.col_canvas.grid(row=0, column=1, sticky="ew")
        self.row_canvas.grid(row=1, column=0, sticky="ns")
        self.canvas.grid(row=1, column=1, sticky="nsew")
        self.vbar.grid(row=1, column=2, sticky="ns")
        self.hbar.grid(row=2, column=1, sticky="ew")
        self.grid_rowconfigure(1, weight=1)
        self.grid_columnconfigure(1, weight=1)

        self.rows = 0
        self.cols = 0
        self.widths = {}
        self.lefts = [0]
        self.cells = {}
        self.info = {}
        self.cell_items = {}
        self.selected = None
        self.sel_id = None
        self.editor = None
        self.editor_window = None
        self.drag = None
        self.zero_based = False
        self.guide = None
        self.guide_id = None

        self._resize(self.MIN_ROWS, self.MIN_COLS)

        self.canvas.bind("<Button-1>", self._click)
        self.canvas.bind("<Double-Button-1>", self._double)
        self.canvas.bind("<MouseWheel>", self._wheel)
        self.row_canvas.bind("<MouseWheel>", self._wheel)
        self.canvas.bind("<Shift-MouseWheel>", self._shift_wheel)
        self.canvas.bind("<Button-4>", lambda e: self._scroll_y(-1))
        self.canvas.bind("<Button-5>", lambda e: self._scroll_y(1))

        self.col_canvas.bind("<Motion>", self._header_motion)
        self.col_canvas.bind("<Leave>", lambda e: self.col_canvas.config(cursor=""))
        self.col_canvas.bind("<ButtonPress-1>", self._header_press)
        self.col_canvas.bind("<B1-Motion>", self._header_drag)
        self.col_canvas.bind("<ButtonRelease-1>", self._header_release)
        self.col_canvas.bind("<Double-Button-1>", self._header_double)

    # ----- desplazamiento -----

    def _on_xscroll(self, first, last):
        self.hbar.set(first, last)
        self.col_canvas.xview_moveto(first)

    def _on_yscroll(self, first, last):
        self.vbar.set(first, last)
        self.row_canvas.yview_moveto(first)

    def _xview(self, *args):
        self.canvas.xview(*args)

    def _yview(self, *args):
        self.canvas.yview(*args)

    def _scroll_y(self, step):
        self.canvas.yview_scroll(step * 3, "units")
        return "break"

    def _wheel(self, event):
        return self._scroll_y(-1 if event.delta > 0 else 1)

    def _shift_wheel(self, event):
        step = -1 if event.delta > 0 else 1
        self.canvas.xview_scroll(step * 3, "units")
        return "break"

    # ----- anchos de columna -----

    def col_width(self, col):
        return self.widths.get(col, self.CELL_W)

    def col_left(self, col):
        return self.lefts[col - 1]

    def configure_sheet(self, widths, zero_based=False):
        """Prepara el visor para otra hoja.

        widths: anchos de la hoja ({columna: píxeles}). Es el mismo
        diccionario de la app, así los cambios quedan en la sesión.
        zero_based: coordenadas desde 0 (racks) en vez de A, B, C / 1, 2, 3.
        """
        self.widths = widths
        self.zero_based = zero_based
        self.rows = 0
        self.cols = 0
        self._relayout()

    def set_guide(self, position):
        """Resalta la ubicación del siguiente escaneo (Modo Rack)."""
        self.guide = position
        self._draw_guide()

    def _draw_guide(self):
        if self.guide_id is not None:
            self.canvas.delete(self.guide_id)
            self.guide_id = None

        if not self.guide:
            return

        row, col = self.guide

        if row > self.rows or col > self.cols:
            return

        x0 = self.col_left(col)
        y0 = (row - 1) * self.CELL_H
        self.guide_id = self.canvas.create_rectangle(
            x0 + 2, y0 + 2, x0 + self.col_width(col) - 2, y0 + self.CELL_H - 2,
            outline=LIME, width=2, dash=(5, 3), fill="", tags="guide",
        )
        self.canvas.tag_raise(self.guide_id)

    def _border_at(self, event):
        """Columna cuyo borde derecho está bajo el ratón (o None)."""
        x = self.col_canvas.canvasx(event.x)

        for col in range(1, self.cols + 1):
            if abs(x - self.lefts[col]) <= self.GRAB:
                return col

        return None

    def _header_motion(self, event):
        if self.drag is None:
            cursor = "sb_h_double_arrow" if self._border_at(event) else ""
            self.col_canvas.config(cursor=cursor)

    def _header_press(self, event):
        col = self._border_at(event)

        if col:
            self.drag = (col, event.x, self.col_width(col))

    def _header_drag(self, event):
        if self.drag is None:
            return

        col, start_x, start_w = self.drag
        width = int(min(max(start_w + event.x - start_x, self.MIN_COL_W),
                        self.MAX_COL_W))

        if width != self.col_width(col):
            self.widths[col] = width
            self._relayout()

    def _header_release(self, event):
        if self.drag is not None:
            self.drag = None

            if self.on_resize:
                self.on_resize()

            if self.on_done:
                self.on_done()

    def _header_double(self, event):
        """Doble clic en el borde: la columna se ajusta al texto más largo."""
        col = self._border_at(event)

        if not col:
            return

        longest = 0

        for (row, column), info in self.info.items():
            if column == col:
                font = self.FONT_BOLD if info["kind"] == "header" else self.FONT
                longest = max(longest, text_width(font, info["text"]))

        if longest:
            self.widths[col] = int(min(
                max(longest + 2 * self.CELL_PAD + 4, self.MIN_COL_W),
                self.MAX_COL_W,
            ))
        else:
            self.widths.pop(col, None)

        self.drag = None
        self._relayout()

        if self.on_resize:
            self.on_resize()

        if self.on_done:
            self.on_done()

    # ----- tamaño y cuadrícula -----

    def _resize(self, rows, cols):
        self.rows, self.cols = rows, cols
        self._relayout()

    def _relayout(self):
        """Redibuja líneas, letras, números y celdas con los anchos actuales."""
        self.lefts = [0]

        for col in range(1, self.cols + 1):
            self.lefts.append(self.lefts[-1] + self.col_width(col))

        width = self.lefts[-1]
        height = self.rows * self.CELL_H

        self.canvas.delete("gridline")
        self.col_canvas.delete("all")
        self.row_canvas.delete("all")

        for row in range(1, self.rows + 1):
            self.row_canvas.create_text(
                self.HDR_W / 2, (row - 0.5) * self.CELL_H,
                text=str(row - 1 if self.zero_based else row),
                fill=DIM, font=ui(9),
            )

        for col in range(1, self.cols + 1):
            left = self.lefts[col - 1]
            self.col_canvas.create_text(
                left + self.col_width(col) / 2, self.HDR_H / 2,
                text=str(col - 1) if self.zero_based else column_letter(col),
                fill=DIM, font=ui(9, "bold"),
            )
            # Separador visible entre letras: indica que se puede arrastrar.
            self.col_canvas.create_line(
                self.lefts[col], 6, self.lefts[col], self.HDR_H - 6,
                fill=BORDER_STRONG,
            )

        for row in range(self.rows + 1):
            y = row * self.CELL_H
            self.canvas.create_line(0, y, width, y, fill=GRID_LINE, tags="gridline")

        for x in self.lefts:
            self.canvas.create_line(x, 0, x, height, fill=GRID_LINE, tags="gridline")

        self.canvas.tag_lower("gridline")
        self.canvas.configure(scrollregion=(0, 0, width, height))
        self.col_canvas.configure(scrollregion=(0, 0, width, self.HDR_H))
        self.row_canvas.configure(scrollregion=(0, 0, self.HDR_W, height))

        # Las celdas cambian de sitio: se vuelven a dibujar todas.
        for rect, text in self.cell_items.values():
            self.canvas.delete(rect)
            self.canvas.delete(text)

        self.cell_items = {}
        self.cells = {}

        for key, info in self.info.items():
            self._draw_cell(key, info)
            self.cells[key] = (info["text"], info["kind"])

        if self.editor is not None:
            self._close_editor()

        self.guide_id = None
        self._draw_selection()
        self._draw_guide()

    # ----- dibujo de celdas -----

    def clear(self):
        """Borra todas las celdas (por ejemplo, al cambiar de hoja)."""
        self._close_editor()

        for rect, text in self.cell_items.values():
            self.canvas.delete(rect)
            self.canvas.delete(text)

        self.cell_items = {}
        self.cells = {}
        self.info = {}
        self.selected = None
        self.guide = None

        if self.guide_id is not None:
            self.canvas.delete(self.guide_id)
            self.guide_id = None

        if self.sel_id is not None:
            self.canvas.delete(self.sel_id)
            self.sel_id = None

    def sync(self, cells, max_row, max_col, exact=False):
        """Actualiza el visor con las celdas indicadas.

        Solo se dibuja lo que cambió, para que el escaneo rápido no se
        ralentice con listas grandes. exact=True: tamaño fijo (un rack).
        """
        if exact:
            rows, cols = max_row, max_col
        else:
            rows = max(self.MIN_ROWS, -(-(max_row + 3) // 20) * 20)
            cols = max(self.MIN_COLS, max_col + 2)
            rows = max(rows, self.rows)
            cols = max(cols, self.cols)

        self.info = cells

        if (rows, cols) != (self.rows, self.cols):
            self._resize(rows, cols)
            return

        for key in [k for k in self.cell_items if k not in cells]:
            rect, text = self.cell_items.pop(key)
            self.canvas.delete(rect)
            self.canvas.delete(text)
            self.cells.pop(key, None)

        for key, info in cells.items():
            signature = (info["text"], info["kind"])

            if self.cells.get(key) != signature:
                self._draw_cell(key, info)
                self.cells[key] = signature

        self._draw_selection()

    def _draw_cell(self, key, info):
        row, col = key
        x0 = self.col_left(col)
        y0 = (row - 1) * self.CELL_H
        x1 = x0 + self.col_width(col)
        y1 = y0 + self.CELL_H

        kind = info["kind"]

        if kind == "header":
            bg, fg, font = HEADER_TINT, WHITE, self.FONT_BOLD
        elif kind == "dup":
            bg, fg, font = DUP_BG, "#ffffff", self.FONT
        elif row % 2:
            bg, fg, font = EMPTY_BG, "#d4d4d8", self.FONT
        else:
            bg, fg, font = "#111318", "#d4d4d8", self.FONT

        text = self._fit(info["text"], font, self.col_width(col))

        if key in self.cell_items:
            rect, text_id = self.cell_items[key]
            self.canvas.itemconfigure(rect, fill=bg)
            self.canvas.itemconfigure(text_id, text=text, fill=fg, font=font)
        else:
            rect = self.canvas.create_rectangle(
                x0, y0, x1, y1, fill=bg, outline=GRID_LINE, tags="cell",
            )
            text_id = self.canvas.create_text(
                x0 + self.CELL_PAD, (y0 + y1) / 2, text=text, anchor="w",
                fill=fg, font=font, tags="cell",
            )
            self.cell_items[key] = (rect, text_id)

    @staticmethod
    def _line_height(font):
        if tkfont is None:
            return 16

        try:
            return tkfont.Font(
                family=font[0], size=font[1],
            ).metrics("linespace")
        except Exception:
            return 16

    def _fit(self, text, font, width):
        """Recorta con "…" lo que no cabe, para no invadir la celda vecina."""
        limit = width - 2 * self.CELL_PAD

        if text_width(font, text) <= limit:
            return text

        while text and text_width(font, text + "…") > limit:
            text = text[:-1]

        return text + "…"

    def _draw_selection(self):
        if not self.selected:
            if self.sel_id is not None:
                self.canvas.delete(self.sel_id)
                self.sel_id = None
            return

        row, col = self.selected

        if col > self.cols or row > self.rows:
            return

        x0 = self.col_left(col)
        y0 = (row - 1) * self.CELL_H
        coords = (
            x0 + 1, y0 + 1,
            x0 + self.col_width(col) - 1, y0 + self.CELL_H - 1,
        )

        if self.sel_id is None:
            self.sel_id = self.canvas.create_rectangle(
                *coords, outline=BLUE_LIGHT, width=2, fill="",
                tags="selection",
            )
        else:
            self.canvas.coords(self.sel_id, *coords)

        self.canvas.tag_raise(self.sel_id)

    # ----- selección y edición -----

    def _cell_at(self, event):
        x = self.canvas.canvasx(event.x)
        y = self.canvas.canvasy(event.y)
        row = int(y // self.CELL_H) + 1
        col = next(
            (c for c in range(1, self.cols + 1) if x < self.lefts[c]), None,
        )

        if col and 1 <= row <= self.rows:
            return row, col

        return None

    def _click(self, event):
        self._close_editor()
        position = self._cell_at(event)

        if position:
            self.select_cell(*position)

    def _double(self, event):
        position = self._cell_at(event)

        if position:
            self.select_cell(*position)
            self.begin_edit(*position)

    def select_cell(self, row, col):
        self.selected = (row, col)
        self._draw_selection()

        if self.on_select:
            self.on_select(row, col, self.info.get((row, col)))

    def ensure_visible(self, row, col):
        if col > self.cols:
            return

        total_w = self.lefts[-1]
        total_h = self.rows * self.CELL_H
        x0 = self.col_left(col)
        y0 = (row - 1) * self.CELL_H
        view_w = self.canvas.winfo_width()
        view_h = self.canvas.winfo_height()

        if view_w <= 1 or view_h <= 1:
            return  # Todavía sin dibujar en pantalla.

        left = self.canvas.canvasx(0)
        top = self.canvas.canvasy(0)

        if y0 < top:
            self.canvas.yview_moveto(y0 / total_h)
        elif y0 + self.CELL_H > top + view_h:
            self.canvas.yview_moveto((y0 + self.CELL_H - view_h) / total_h)

        if x0 < left:
            self.canvas.xview_moveto(x0 / total_w)
        elif x0 + self.col_width(col) > left + view_w:
            self.canvas.xview_moveto(
                (x0 + self.col_width(col) - view_w) / total_w
            )

    def begin_edit(self, row, col):
        info = self.info.get((row, col))

        if not info or info["kind"] not in ("item", "dup"):
            return

        self._close_editor()

        editor = tk.Entry(
            self.canvas,
            font=self.FONT,
            bg=CONTROL,
            fg=WHITE,
            insertbackground=WHITE,
            relief="flat",
            highlightthickness=2,
            highlightbackground=BLUE_LIGHT,
            highlightcolor=BLUE_LIGHT,
        )
        editor.insert(0, info["text"])
        editor.select_range(0, tk.END)
        self.editor = editor
        self.editor_window = self.canvas.create_window(
            self.col_left(col), (row - 1) * self.CELL_H, anchor="nw",
            window=editor, width=self.col_width(col), height=self.CELL_H,
        )
        editor.bind("<Return>", lambda e: self._commit(row, col))
        editor.bind("<Escape>", lambda e: self._cancel_edit())
        editor.bind("<FocusOut>", lambda e: self._close_editor())
        editor.focus_set()

    def _commit(self, row, col):
        if self.editor is None:
            return "break"

        text = self.editor.get().strip()
        self._close_editor()

        if text and self.on_edit:
            self.on_edit(row, col, text)

        if self.on_done:
            self.on_done()

        return "break"

    def _cancel_edit(self):
        self._close_editor()

        if self.on_done:
            self.on_done()

        return "break"

    def _close_editor(self):
        if self.editor is None:
            return

        editor = self.editor
        window = self.editor_window
        self.editor = None
        self.editor_window = None

        try:
            self.canvas.delete(window)
            editor.destroy()
        except tk.TclError:
            pass


# ============================================================
# APLICACIÓN
# ============================================================

class SerialNumberChecker:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        height = min(980, max(self.root.winfo_screenheight() - 90, 800))
        self.root.geometry(f"1120x{height}")
        self.root.minsize(940, 800)
        self.root.configure(bg=BG)
        pick_fonts(root)

        self.lang = "es"
        self.language_var = tk.StringVar(value="Español")
        self.alt_matching_var = tk.BooleanVar(value=True)
        self.pallet_name_var = tk.StringVar(value="")
        self.search_var = tk.StringVar(value="")
        self.rack_mode_var = tk.BooleanVar(value=False)
        self.validate_var = tk.BooleanVar(value=False)
        self.valid_alt_var = tk.BooleanVar(value=True)
        self.valid_order = []
        self.valid_set = set()
        self.valid_path = None
        self._rename_entry = None
        self.rack_location_var = tk.StringVar(value="")
        self.rack_rows_var = tk.StringVar(value=str(RACK_ROWS))
        self.rack_cols_var = tk.StringVar(value=str(RACK_COLS))
        self.rack_order = "h"
        self.rack_cursor = None
        self.rack_dismissed = set()
        self.col_widths = {}

        self.mode = "checker"

        # --- Estado del Checker ---
        self.serials = set()
        self.serial_order = []
        self.serial_lengths = set()
        self.found = set()
        self.locations = {}
        self.extras = OrderedDict()
        self.scan_log = []
        self.checker_undo = []
        self.csv_path = None
        self.loaded_count = 0
        self.resumed = False
        self.pending_resume = set()
        self.history = []
        self.standby = False
        self.list_tab = "pending"
        self.pending_iids = {}

        # --- Estado de la creación de listas ---
        self.create_items = []
        self.create_seen = set()
        self.pallets = []
        self.pallet_counter = 0
        self.dest_index = 0
        self.undo_stack = []
        self.default_sheet = DEFAULT_SHEET
        self.view_sheet = DEFAULT_SHEET
        self.drawn_sheet = None
        self.sheet_tab_names = None
        self.layout = {"sheets": OrderedDict(), "positions": {}}
        self.dup_dismissed_at = 0

        # --- Mensajes, colores y temporizadores ---
        self.status_key = ""
        self.status_params = {}
        self.status_details = []
        self.banner_visible = False
        self.big_view = False
        self._last_backspace = 0.0
        self.print_prompt = None
        self.readme_open = False
        self.batch = False
        self.flash_job = None
        self.save_job = None
        self.save_label_job = None
        self.reset_armed = False
        self.reset_job = None

        self.session_file = os.path.join(
            os.path.expanduser("~"),
            ".serial_number_checker_sessions.json",
        )

        self._set_window_icon()

        self._setup_styles()
        self._build_ui()
        self._bind_keys()
        self.load_sessions()
        self.update_mode_view()
        self.root.protocol("WM_DELETE_WINDOW", self.close_app)
        use_dark_title_bar(self.root)

    def _set_window_icon(self):
        """Icono de la ventana y de la barra de tareas.

        icon.ico es arthur.ico con todos los tamaños (16-96 px): la barra
        de tareas de Windows necesita los pequeños para mostrarlo.
        """
        folders = [os.path.dirname(os.path.abspath(__file__))]

        if getattr(sys, "frozen", False):
            # .exe de PyInstaller: el icono va incluido en sys._MEIPASS.
            folders.insert(0, getattr(sys, "_MEIPASS", ""))

        for name in ("icon.ico", "arthur.ico"):
            for folder in folders:
                path = os.path.join(folder, name)

                if os.path.exists(path):
                    try:
                        self.root.iconbitmap(default=path)
                        return
                    except tk.TclError:
                        pass

    def t(self, key, **kwargs):
        """Devuelve el texto del idioma activo."""
        template = TEXTS[self.lang].get(key, key)
        return template.format(**kwargs) if kwargs else template

    # ========================================================
    # ESTILOS E INTERFAZ
    # ========================================================

    def _setup_styles(self):
        self.style = ttk.Style()

        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass

        self.style.configure(
            "Dark.TCombobox",
            fieldbackground=CONTROL,
            background=CONTROL,
            foreground=WHITE,
            arrowcolor=MUTED,
            bordercolor=BORDER_STRONG,
            lightcolor=CONTROL,
            darkcolor=CONTROL,
            padding=(8, 5),
        )
        self.style.map(
            "Dark.TCombobox",
            fieldbackground=[("readonly", CONTROL)],
            foreground=[("readonly", WHITE)],
            background=[("active", BUTTON_ACTIVE), ("readonly", CONTROL)],
            selectbackground=[("readonly", CONTROL)],
            selectforeground=[("readonly", WHITE)],
            bordercolor=[("focus", BLUE_LIGHT)],
            arrowcolor=[("active", WHITE)],
        )

        # Lista desplegable de los combos.
        self.root.option_add("*TCombobox*Listbox.background", CONTROL)
        self.root.option_add("*TCombobox*Listbox.foreground", WHITE)
        self.root.option_add("*TCombobox*Listbox.selectBackground", BLUE)
        self.root.option_add("*TCombobox*Listbox.selectForeground", "#ffffff")
        self.root.option_add("*TCombobox*Listbox.font", ui(10))
        self.root.option_add("*TCombobox*Listbox.borderWidth", 0)

        self.style.configure(
            "Dark.Treeview",
            background=PANEL,
            fieldbackground=PANEL,
            foreground=WHITE,
            bordercolor=PANEL,
            lightcolor=PANEL,
            darkcolor=PANEL,
            borderwidth=0,
            rowheight=26,
            font=mono(10),
        )
        self.style.layout(
            "Dark.Treeview", [("Treeview.treearea", {"sticky": "nswe"})],
        )
        self.style.map(
            "Dark.Treeview",
            background=[("selected", BUTTON_ACTIVE)],
            foreground=[("selected", WHITE)],
        )
        self.style.configure(
            "Dark.Treeview.Heading",
            background=PANEL,
            foreground=DIM,
            bordercolor=PANEL,
            lightcolor=PANEL,
            darkcolor=PANEL,
            relief="flat",
            font=ui(9, "bold"),
            padding=(4, 4),
        )
        self.style.map("Dark.Treeview.Heading", background=[("active", PANEL)])

        # Barras de desplazamiento finas, sin flechas.
        for orient in ("Vertical", "Horizontal"):
            name = f"Thin.{orient}.TScrollbar"
            self.style.layout(name, [
                (f"{orient}.Scrollbar.trough", {
                    "sticky": "nswe",
                    "children": [
                        (f"{orient}.Scrollbar.thumb",
                         {"expand": "1", "sticky": "nswe"}),
                    ],
                }),
            ])
            self.style.configure(
                name,
                troughcolor=FIELD,
                background=BORDER_STRONG,
                bordercolor=FIELD,
                lightcolor=BORDER_STRONG,
                darkcolor=BORDER_STRONG,
                gripcount=0,
                arrowsize=9,
            )
            self.style.map(name, background=[("active", "#4b5263")])

    def _combo(self, parent, width):
        return ttk.Combobox(
            parent,
            state="readonly",
            width=width,
            justify="left",
            style="Dark.TCombobox",
            font=ui(10),
        )

    def _label(self, parent, size=10, weight="normal", fg=WHITE, bg=PANEL,
               font=None, **kw):
        return tk.Label(
            parent, text="", font=font or ui(size, weight), bg=bg, fg=fg, **kw,
        )

    def _icon(self, parent, name, color, size=18, box=22, bg=PANEL, width=2):
        canvas = tk.Canvas(
            parent, width=box, height=box, bg=bg, highlightthickness=0, bd=0,
        )
        draw_icon(canvas, name, box / 2, box / 2, color, size=size, width=width)
        return canvas

    def _put(self, widget, row, col, span=1, last_row=False):
        """Coloca una tarjeta en la cuadrícula de la ventana.

        Los huecos entre tarjetas dejan ver el patrón del fondo.
        """
        widget.grid(
            row=row, column=col, columnspan=span, sticky="nsew",
            padx=(24 if col == 0 else 6, 24 if col + span == 3 else 6),
            pady=(14 if row == 0 else 0, 14 if last_row else 10),
        )

    def _build_legend(self, parent, definitions):
        """Leyenda de colores con muestras pequeñas."""
        frame = tk.Frame(parent, bg=PANEL)

        for key, color in definitions:
            swatch = tk.Canvas(
                frame, width=12, height=12, bg=PANEL,
                highlightthickness=0, bd=0,
            )
            swatch.create_polygon(
                *round_rect_points(1, 1, 11, 11, 3), smooth=True,
                fill=color, outline="",
            )
            swatch.pack(side="left", padx=(0, 6))

            label = self._label(frame, 9, fg=MUTED)
            label.pack(side="left", padx=(0, 18))
            self.legend_labels.append((label, key))

        return frame

    def _build_ui(self):
        root = self.root
        self.legend_labels = []
        self.key_widgets = []

        # Fondo con el patrón de líneas: se ve entre las tarjetas.
        self.bg_canvas = tk.Canvas(root, bg=BG, highlightthickness=0, bd=0)
        self.bg_canvas.place(x=0, y=0, relwidth=1, relheight=1)
        self.bg_tile = load_background_tile()
        self.bg_tiles = {}
        self.bg_canvas.bind(
            "<Configure>",
            lambda e: self._tile_background(self.bg_canvas, self.bg_tiles),
        )

        root.grid_columnconfigure(0, weight=1, uniform="half")
        root.grid_columnconfigure(1, weight=1, uniform="half")
        root.grid_columnconfigure(2, weight=0, minsize=300)
        root.grid_rowconfigure(5, weight=1)

        self._build_toolbar()
        self._build_checker_header()
        self._build_create_header()
        self._build_counter()
        self._build_scan_card()
        self._build_banner()
        self._build_status_card()
        self._build_checker_bottom()
        self._build_create_bottom()
        self._build_footer()
        self._build_big_view()
        self._build_readme_panel()

        root.bind("<Configure>", self._on_resize, add="+")
        self.entry.focus_set()

    def _tile_background(self, canvas, tiles):
        """Repite el patrón en mosaico (nunca se estira)."""
        if self.bg_tile is None:
            return

        tile_w = self.bg_tile.width()
        tile_h = self.bg_tile.height()
        width = canvas.winfo_width()
        height = canvas.winfo_height()

        for col in range(width // tile_w + 1):
            for row in range(height // tile_h + 1):
                if (col, row) not in tiles:
                    tiles[(col, row)] = canvas.create_image(
                        col * tile_w, row * tile_h,
                        image=self.bg_tile, anchor="nw", tags="tile",
                    )

        canvas.tag_lower("tile")

    # ----- barra superior -----

    def _build_toolbar(self):
        bar = Card(self.root, padx=10, pady=8)
        self._put(bar, 0, 0, 3)
        body = bar.body

        segment = tk.Frame(
            body, bg=FIELD, padx=3, pady=3,
            highlightthickness=1, highlightbackground=BORDER,
        )
        segment.pack(side="left")
        self.mode_buttons = {}

        for mode in ("checker", "create"):
            button = RoundButton(
                segment, bg=FIELD, size=10, height=32,
                command=lambda m=mode: self.set_mode(m),
            )
            button.pack(side="left", padx=(0 if mode == "checker" else 3, 0))
            self.mode_buttons[mode] = button

        self.alt_matching_check = CheckBox(
            body, self.alt_matching_var, command=self.on_alt_toggle,
        )
        self.alt_matching_check.pack(side="left", padx=(18, 0))

        self.rack_check = CheckBox(
            body, self.rack_mode_var, command=self.on_rack_toggle,
        )
        self.validate_check = CheckBox(
            body, self.validate_var, command=self.on_validate_toggle,
        )
        Tooltip(
            self.rack_check,
            lambda: self.t(
                "tip_rack_locked" if self.validation_active() else "tip_rack"
            ),
        )
        Tooltip(
            self.validate_check,
            lambda: self.t(
                "tip_validate_locked" if self.rack_mode_var.get()
                else "tip_validate"
            ),
        )

        self.language_menu = ttk.Combobox(
            body,
            textvariable=self.language_var,
            values=list(LANGUAGES.keys()),
            state="readonly",
            width=11,
            justify="left",
            style="Dark.TCombobox",
            font=ui(10),
        )
        # README al extremo derecho, separado del resto.
        self.readme_button = RoundButton(
            body, icon="book", size=10, height=32, padx=12,
            command=self.toggle_readme,
        )
        self.readme_button.pack(side="right")

        self.language_menu.pack(side="right", padx=(0, 16))
        self.language_menu.bind("<<ComboboxSelected>>", self.change_language)

        self._icon(body, "globe", MUTED, size=16, width=1.6).pack(
            side="right", padx=(0, 8),
        )

        self.big_button = RoundButton(
            body, variant="ghost", icon="expand", size=10, bold=False,
            height=32, command=self.toggle_big_view,
        )
        self.big_button.pack(side="right", padx=(0, 14))

    # ----- cabecera de cada modo -----

    def _build_checker_header(self):
        card = Card(self.root, padx=20, pady=14)
        self.checker_header = card
        body = card.body

        self.title_label = self._label(
            body, 20, "bold", anchor="w",
        )
        self.title_label.config(text=APP_TITLE)
        self.title_label.pack(fill="x")

        info = tk.Frame(body, bg=PANEL)
        info.pack(fill="x", pady=(2, 12))

        self.file_icon = tk.Canvas(
            info, width=18, height=18, bg=PANEL, highlightthickness=0, bd=0,
        )
        self.file_icon.pack(side="left", padx=(0, 6))

        self.file_info = self._label(info, 10, fg=MUTED, anchor="w")
        self.file_info.pack(side="left", fill="x", expand=True)

        buttons = tk.Frame(body, bg=PANEL)
        buttons.pack(fill="x")

        self.load_button = RoundButton(
            buttons, icon="folder", height=40, command=self.load_file,
        )
        self.load_button.pack(side="left", padx=(0, 10))

        self.export_button = RoundButton(
            buttons, variant="primary", icon="download", height=40,
            command=self.export_report,
        )
        self.export_button.pack(side="left", padx=(0, 10))

        self.print_button = RoundButton(
            buttons, icon="printer", height=40, command=self.ask_print,
        )
        self.print_button.pack(side="left")

    def _build_create_header(self):
        card = Card(self.root, padx=20, pady=14)
        self.create_header = card
        body = card.body

        row_one = tk.Frame(body, bg=PANEL)
        row_one.pack(fill="x")

        dest = tk.Frame(row_one, bg=PANEL)
        dest.pack(side="left", padx=(0, 12))
        self.dest_label = self._label(dest, 9, fg=MUTED, anchor="w")
        self.dest_label.pack(fill="x", pady=(0, 4))
        self.dest_menu = self._combo(dest, 22)
        self.dest_menu.pack(anchor="w")
        self.dest_menu.bind("<<ComboboxSelected>>", self.on_dest_selected)

        name = tk.Frame(row_one, bg=PANEL)
        name.pack(side="left", padx=(0, 12))
        self.name_label = self._label(name, 9, fg=MUTED, anchor="w")
        self.name_label.pack(fill="x", pady=(0, 4))
        self.pallet_name_entry = tk.Entry(
            name,
            textvariable=self.pallet_name_var,
            width=15,
            font=ui(10),
            bg=FIELD,
            fg=WHITE,
            insertbackground=WHITE,
            relief="flat",
            bd=0,
            highlightthickness=1,
            highlightbackground=BORDER_STRONG,
            highlightcolor=BLUE_LIGHT,
        )
        self.pallet_name_entry.pack(anchor="w", ipady=6, ipadx=4)
        self.pallet_name_entry.bind("<Return>", self.on_next_pallet_key)

        self.next_pallet_button = RoundButton(
            row_one, variant="primary", icon="plus", height=34,
            command=self.next_pallet,
        )
        self.next_pallet_button.pack(side="left", anchor="s")
        self.pallet_row = row_one

        # Modo Rack: ubicación, tamaño y orden de llenado.
        rack_row = tk.Frame(body, bg=PANEL)
        self.rack_row = rack_row

        def small_field(parent, label_attr, variable, width):
            box = tk.Frame(parent, bg=PANEL)
            label = self._label(box, 9, fg=MUTED, anchor="w")
            label.pack(fill="x", pady=(0, 4))
            setattr(self, label_attr, label)
            entry = tk.Entry(
                box, textvariable=variable, width=width, font=ui(10),
                bg=FIELD, fg=WHITE, insertbackground=WHITE, relief="flat",
                bd=0, highlightthickness=1, highlightbackground=BORDER_STRONG,
                highlightcolor=BLUE_LIGHT,
            )
            entry.pack(anchor="w", ipady=6, ipadx=4)
            entry.bind("<Return>", lambda e: self.focus_entry())
            return box, entry

        box, self.rack_location_entry = small_field(
            rack_row, "rack_location_label", self.rack_location_var, 12,
        )
        box.pack(side="left", padx=(0, 10))
        box, _ = small_field(rack_row, "rack_rows_label", self.rack_rows_var, 4)
        box.pack(side="left", padx=(0, 10))
        box, _ = small_field(rack_row, "rack_cols_label", self.rack_cols_var, 4)
        box.pack(side="left", padx=(0, 12))

        order = tk.Frame(rack_row, bg=PANEL)
        order.pack(side="left", padx=(0, 12), anchor="s")
        self.rack_order_label = self._label(order, 9, fg=MUTED, anchor="w")
        self.rack_order_label.pack(fill="x", pady=(0, 4))
        segment = tk.Frame(
            order, bg=FIELD, padx=2, pady=2,
            highlightthickness=1, highlightbackground=BORDER,
        )
        segment.pack(anchor="w")
        self.order_buttons = {}

        for key in ("h", "v"):
            button = RoundButton(
                segment, bg=FIELD, size=9, height=28, padx=10,
                command=lambda k=key: self.set_rack_order(k),
            )
            button.pack(side="left", padx=(0 if key == "h" else 2, 0))
            self.order_buttons[key] = button

        self.add_rack_button = RoundButton(
            rack_row, variant="primary", icon="plus", height=34,
            command=self.add_rack,
        )
        self.add_rack_button.pack(side="left", anchor="s")

        row_two = tk.Frame(body, bg=PANEL)
        row_two.pack(fill="x", pady=(14, 0))
        self.create_row_two = row_two

        self.delete_button = RoundButton(
            row_two, icon="trash", size=9, bold=False, height=32, padx=12,
            command=self.delete_selected_item,
        )
        self.delete_button.pack(side="left", padx=(0, 8))

        self.undo_button = RoundButton(
            row_two, icon="undo", size=9, bold=False, height=32, padx=12,
            command=self.undo_last,
        )
        self.undo_button.pack(side="left", padx=(0, 8))

        self.export_list_button = RoundButton(
            row_two, icon="download", size=9, height=32, padx=12,
            command=self.export_create_list,
        )
        self.export_list_button.pack(side="left", padx=(0, 8))

        self.print_list_button = RoundButton(
            row_two, icon="printer", size=9, bold=False, height=32, padx=12,
            command=self.ask_print,
        )
        self.print_list_button.pack(side="left")

        # Validar lista: fila propia con el archivo y las alternativas A/B/U.
        self.validate_row = tk.Frame(body, bg=PANEL)
        self.load_valid_button = RoundButton(
            self.validate_row, icon="file", size=9, height=32, padx=12,
            command=self.load_validation_list,
        )
        self.load_valid_button.pack(side="left", padx=(0, 10))
        self.valid_file_label = self._label(self.validate_row, 9, fg=MUTED)
        self.valid_file_label.pack(side="left", padx=(0, 18))
        self.valid_alt_check = CheckBox(
            self.validate_row, self.valid_alt_var, command=self.on_alt_toggle,
        )
        self.valid_alt_check.pack(side="left")

    # ----- contador -----

    def _build_counter(self):
        card = Card(self.root, padx=20, pady=14)
        self.counter_card = card
        self._put(card, 1, 2)
        body = card.body

        self.counter_caption = self._label(body, 9, "bold", fg=MUTED, anchor="w")
        self.counter_caption.pack(fill="x")

        self.counter_row = tk.Frame(body, bg=PANEL)
        self.counter_row.pack(fill="x", pady=(2, 8))

        self.counter = self._label(self.counter_row, 34, "bold")
        self.counter.config(text="0")
        self.counter.pack(side="left")

        self.counter_total = self._label(
            self.counter_row, 16, "bold", fg="#71717a",
        )
        self.counter_total.pack(side="left", anchor="s", padx=(6, 0), pady=(0, 8))

        self.progress = ProgressBar(body)
        self.progress.pack(fill="x")

        self.counter_info = self._label(
            body, 9, fg=MUTED, anchor="w", justify="left", wraplength=250,
        )
        self.counter_info.pack(fill="x", pady=(8, 0))

        self.rate_label = self._label(body, 9, fg=MUTED, anchor="w")
        self.rate_label.pack(fill="x")

    # ----- campo de escaneo -----

    def _build_scan_card(self):
        card = Card(self.root, padx=20, pady=14)
        self._put(card, 2, 0, 3)
        body = card.body

        self.input_label = self._label(body, 10, "bold", anchor="w")
        self.input_label.pack(fill="x", pady=(0, 8))

        self.field = Card(
            body, bg=FIELD, border=BLUE_LIGHT, outer=PANEL, radius=11,
            padx=14, pady=8, border_width=2,
        )
        self.field.pack(fill="x")
        field_body = self.field.body

        self.scan_icon = tk.Canvas(
            field_body, width=28, height=28, bg=FIELD,
            highlightthickness=0, bd=0,
        )
        self.scan_icon.pack(side="left", padx=(0, 10))

        self.length_label = self._label(
            field_body, bg=FIELD, fg="#71717a", font=mono(10),
        )
        self.length_label.pack(side="right", padx=(10, 0))

        self.entry = tk.Entry(
            field_body,
            font=mono(19),
            bg=FIELD,
            fg=WHITE,
            insertbackground=WHITE,
            selectbackground=BUTTON_ACTIVE,
            selectforeground=WHITE,
            relief="flat",
            bd=0,
            highlightthickness=0,
        )
        self.entry.pack(side="left", fill="x", expand=True, ipady=4)

        self.legend_container = tk.Frame(body, bg=PANEL)
        self.legend_container.pack(fill="x", pady=(10, 0))

        self.checker_legend = self._build_legend(
            self.legend_container,
            [
                ("legend_green", GREEN),
                ("legend_red", RED),
                ("legend_blue", BLUE_LIGHT),
                ("legend_yellow", YELLOW),
            ],
        )
        self.create_legend = self._build_legend(
            self.legend_container,
            [
                ("legend_dup", DUP_BG),
                ("legend_warn", YELLOW),
            ],
        )

    def set_ring(self, color):
        """Color del borde del campo de escaneo (y de su icono)."""
        self.field.set_colors(border=color)
        self.scan_icon.delete("all")
        draw_icon(self.scan_icon, "scan", 14, 14, color, size=22, width=2)

    # ----- aviso integrado y mensajes -----

    def _build_banner(self):
        card = Card(self.root, bg=YELLOW_BG, border=YELLOW_BORDER, padx=14, pady=10)
        self.banner = card
        body = card.body

        self._icon(body, "warn", YELLOW_TEXT, bg=YELLOW_BG).pack(
            side="left", padx=(0, 10),
        )

        # Hasta tres botones (el primero resaltado) y casillas opcionales.
        buttons = tk.Frame(body, bg=YELLOW_BG)
        buttons.pack(side="right")
        self.banner_buttons = [
            RoundButton(
                buttons, variant="warn" if index == 0 else "warn_ghost",
                size=9, height=30, padx=12, bg=YELLOW_BG,
            )
            for index in range(3)
        ]
        self.banner_button_one, self.banner_button_two = self.banner_buttons[:2]

        self.banner_checks = tk.Frame(body, bg=YELLOW_BG)
        self.print_vars = {
            "summary": tk.BooleanVar(value=True),
            "pending": tk.BooleanVar(value=False),
            "extras": tk.BooleanVar(value=False),
        }
        self.print_checks = {}

        for key, variable in self.print_vars.items():
            check = CheckBox(self.banner_checks, variable, bg=YELLOW_BG,
                             command=self.focus_entry)
            check.label.config(fg=YELLOW_TEXT)
            check.pack(side="left", padx=(0, 16))
            self.print_checks[key] = check

        self.banner_label = self._label(
            body, 10, "bold", fg=YELLOW_TEXT, bg=YELLOW_BG,
            anchor="w", justify="left", wraplength=560,
        )
        self.banner_label.pack(side="left", fill="x", expand=True)

    def _build_status_card(self):
        card = Card(self.root, padx=16, pady=10)
        self.status_card = card
        self._put(card, 4, 0, 3)
        body = card.body

        self.status_icon = tk.Canvas(
            body, width=38, height=38, bg=PANEL, highlightthickness=0, bd=0,
        )
        self.status_icon.pack(side="left", padx=(0, 14))

        self.status_texts = tk.Frame(body, bg=PANEL)
        self.status_texts.pack(side="left", fill="x", expand=True)

        self.status = self._label(
            self.status_texts, 15, "bold", fg=MUTED,
            anchor="w", justify="left", wraplength=860,
        )
        self.status.pack(fill="x")

        self.details = self._label(
            self.status_texts, 10, fg=MUTED,
            anchor="w", justify="left", wraplength=860,
        )
        self.details.pack(fill="x")

    # ----- parte inferior: Checker -----

    def _build_checker_bottom(self):
        # Historial de escaneos.
        card = Card(self.root, fill=True, padx=18, pady=14, height=200)
        self.history_card = card
        body = card.body

        head = tk.Frame(body, bg=PANEL)
        head.pack(fill="x", pady=(0, 8))

        self.history_title = self._label(head, 9, "bold", fg=MUTED, anchor="w")
        self.history_title.pack(side="left")

        self.undo_scan_button = RoundButton(
            head, variant="ghost", icon="undo", size=9, bold=False,
            height=28, padx=10, command=self.undo_scan,
        )
        self.undo_scan_button.pack(side="right")

        self.history_rows = []

        for _ in range(HISTORY_SIZE):
            row = tk.Frame(body, bg=PANEL)
            row.pack(fill="x", pady=3)

            dot = tk.Canvas(
                row, width=10, height=10, bg=PANEL,
                highlightthickness=0, bd=0,
            )
            dot.pack(side="left", padx=(0, 10))

            serial = self._label(row, font=mono(11), anchor="w")
            serial.pack(side="left")

            tag = self._label(row, 9, fg=MUTED, anchor="e")
            tag.pack(side="right")

            self.history_rows.append((dot, serial, tag))

        # Pendientes y sobrantes (pestañas) con buscador.
        card = Card(self.root, fill=True, padx=16, pady=12, height=200)
        self.lists_card = card
        body = card.body

        head = tk.Frame(body, bg=PANEL)
        head.pack(fill="x", pady=(0, 8))
        self.tab_buttons = {}

        for tab in ("pending", "extras"):
            button = RoundButton(
                head, variant="tab", size=9, height=28, padx=12,
                command=lambda name=tab: self.set_list_tab(name),
            )
            button.pack(side="left", padx=(0, 6))
            self.tab_buttons[tab] = button

        search = tk.Frame(
            head, bg=FIELD, highlightthickness=1,
            highlightbackground=BORDER_STRONG, highlightcolor=BLUE_LIGHT,
        )
        search.pack(side="right")
        self._icon(search, "search", MUTED, size=14, box=24, bg=FIELD,
                   width=1.8).pack(side="left", padx=(4, 0))

        self.search_entry = tk.Entry(
            search,
            textvariable=self.search_var,
            width=18,
            font=mono(10),
            bg=FIELD,
            fg=WHITE,
            insertbackground=WHITE,
            relief="flat",
            bd=0,
            highlightthickness=0,
        )
        self.search_entry.pack(side="left", ipady=4, padx=(2, 6))
        self.search_entry.bind("<Return>", lambda e: self.focus_entry())
        self.search_entry.bind("<Escape>", self.on_search_escape)
        self.search_var.trace_add("write", lambda *a: self.rebuild_pending_tree())

        self.search_label = self._label(head, 9, fg=MUTED)
        self.search_label.pack(side="right", padx=(0, 8))

        tree_frame = tk.Frame(body, bg=PANEL)
        tree_frame.pack(fill="both", expand=True)

        self.pending_tree = ttk.Treeview(
            tree_frame,
            columns=("serial", "count", "time"),
            show="headings",
            style="Dark.Treeview",
            selectmode="browse",
        )
        self.pending_tree.column("serial", width=300, anchor="w", stretch=True)
        self.pending_tree.column("count", width=70, anchor="center", stretch=False)
        self.pending_tree.column("time", width=90, anchor="center", stretch=False)
        self.pending_tree.tag_configure("odd", background="#171a20")
        self.pending_tree.tag_configure("muted", foreground=MUTED)
        self.pending_tree.bind("<<TreeviewSelect>>", lambda e: self.focus_entry())
        # Copiar el S/N: clic derecho → "Copiar S/N", o Ctrl+C con la fila
        # seleccionada (el foco vuelve al escaneo, ver on_copy).
        self.pending_tree.bind("<Button-3>", self.show_copy_menu)
        self.pending_tree.bind("<Control-c>", self.copy_selected_serial)
        self.copy_menu = tk.Menu(
            self.root, tearoff=0, bg=CONTROL, fg=WHITE, bd=0,
            activebackground=BLUE, activeforeground="#ffffff", font=ui(10),
        )
        self.copy_menu.add_command(label="", command=self.copy_selected_serial)

        pending_scroll = ttk.Scrollbar(
            tree_frame, orient="vertical",
            command=self.pending_tree.yview,
            style="Thin.Vertical.TScrollbar",
        )
        self.pending_tree.configure(yscrollcommand=pending_scroll.set)
        self.pending_tree.pack(side="left", fill="both", expand=True)
        pending_scroll.pack(side="right", fill="y")

    # ----- parte inferior: Crear lista -----

    def _build_create_bottom(self):
        card = Card(self.root, fill=True, padx=14, pady=12, height=220)
        self.sheet_card = card
        body = card.body

        # Pestañas abajo; se empaquetan antes para que el visor no las tape.
        tabs_row = tk.Frame(body, bg=PANEL)
        tabs_row.pack(side="bottom", fill="x", pady=(8, 0))
        self.sheet_tabs = tk.Frame(tabs_row, bg=PANEL)
        self.sheet_tabs.pack(side="left")
        self.sheet_tab_buttons = []

        self.cell_info_label = self._label(tabs_row, fg=MUTED, font=mono(10))
        self.cell_info_label.pack(side="right")

        self.sheet_grid = SheetGrid(
            body,
            on_select=self.on_grid_select,
            on_edit=self.on_grid_edit,
            on_done=self.focus_entry,
            on_resize=self.on_columns_resized,
        )
        self.sheet_grid.pack(fill="both", expand=True)

    # ----- pie -----

    def _build_footer(self):
        card = Card(self.root, padx=14, pady=8)
        self._put(card, 6, 0, 3, last_row=True)
        body = card.body

        self.reset_button = RoundButton(
            body, icon="reset", size=9, height=32, padx=12,
            command=self.reset_app,
        )
        self.reset_button.pack(side="right")

        self.save_label = self._label(body, 9, fg=DIM)
        self.save_label.pack(side="right", padx=(0, 12))

        self.keys_frame = tk.Frame(body, bg=PANEL)
        self.keys_frame.pack(side="left")

        # Último en empaquetarse: si falta espacio es lo único que se recorta.
        self.footer = self._label(body, 9, fg=DIM, anchor="center")
        self.footer.pack(side="left", padx=(10, 0), fill="x", expand=True)

    def _build_keys(self):
        """Atajos del modo activo como teclas en el pie."""
        for widget in self.key_widgets:
            widget.destroy()

        self.key_widgets = []
        if self.mode == "checker":
            spec = self.t("keys_checker")
        elif self.rack_mode_var.get():
            spec = self.t("keys_rack")
        else:
            spec = self.t("keys_create")

        for pair in spec.split(";"):
            key, label = pair.split("|", 1)
            chip = tk.Label(
                self.keys_frame, text=key, font=mono(9), bg=PANEL, fg=WHITE,
                padx=5, highlightthickness=1,
                highlightbackground=BORDER_STRONG,
            )
            chip.pack(side="left", padx=(0, 5))
            text = self._label(self.keys_frame, 9, fg=DIM)
            text.config(text=label)
            text.pack(side="left", padx=(0, 12))
            self.key_widgets.extend((chip, text))

    # ----- vista grande -----

    def _build_big_view(self):
        """Solo contador y resultado en tamaño enorme (F11)."""
        canvas = tk.Canvas(self.root, bg=BG, highlightthickness=0, bd=0)
        self.big_canvas = canvas
        self.big_tiles = {}

        self.big_caption = canvas.create_text(
            0, 0, text="", fill=MUTED, font=ui(18, "bold"),
        )
        self.big_counter = canvas.create_text(
            0, 0, text="", fill=WHITE, font=ui(110, "bold"),
        )
        self.big_status = canvas.create_text(
            0, 0, text="", fill=MUTED, font=ui(30, "bold"),
            justify="center", width=1000,
        )
        self.big_detail = canvas.create_text(
            0, 0, text="", fill=MUTED, font=ui(14), justify="center", width=1000,
        )
        self.big_hint = canvas.create_text(
            0, 0, text="", fill=DIM, font=ui(10),
        )
        canvas.bind("<Configure>", self._layout_big_view)
        canvas.bind("<Button-1>", lambda e: self.focus_entry())

    def _layout_big_view(self, event=None):
        canvas = self.big_canvas
        self._tile_background(canvas, self.big_tiles)
        width = canvas.winfo_width()
        height = canvas.winfo_height()
        cx = width / 2

        canvas.coords(self.big_caption, cx, height * 0.18)
        canvas.coords(self.big_counter, cx, height * 0.36)
        canvas.coords(self.big_status, cx, height * 0.60)
        canvas.coords(self.big_detail, cx, height * 0.70)
        canvas.coords(self.big_hint, cx, height - 28)
        canvas.itemconfigure(self.big_status, width=max(width - 120, 200))
        canvas.itemconfigure(self.big_detail, width=max(width - 120, 200))

    def toggle_big_view(self, event=None):
        self.big_view = not self.big_view

        if self.big_view:
            self.big_canvas.place(x=0, y=0, relwidth=1, relheight=1)
            # En un Canvas, lift() mueve dibujos; esto sube la ventana.
            tk.Misc.tkraise(self.big_canvas)
            self.refresh_big_view()
        else:
            self.big_canvas.place_forget()

        self.focus_entry()
        return "break"

    def refresh_big_view(self):
        if not self.big_view:
            return

        canvas = self.big_canvas
        alert = self.alert_active()
        canvas.itemconfigure(
            self.big_caption, text=self.counter_caption.cget("text"),
            fill=YELLOW_TEXT if alert else MUTED,
        )
        total = self.counter_total.cget("text") or ""
        canvas.itemconfigure(
            self.big_counter,
            text=f"{self.counter.cget('text')} {total}".strip(),
            fill=YELLOW_TEXT if alert else WHITE,
        )
        canvas.itemconfigure(
            self.big_status, text=self.status.cget("text"),
            fill=self.status.cget("fg"),
        )
        canvas.itemconfigure(self.big_detail, text=self.details.cget("text"))
        canvas.itemconfigure(self.big_hint, text=self.t("big_hint"))

    # ----- panel README (guía dentro de la app) -----

    def _build_readme_panel(self):
        """Guía de uso sobre el fondo de líneas (sin ventanas aparte)."""
        canvas = tk.Canvas(self.root, bg=BG, highlightthickness=0, bd=0)
        self.readme_canvas = canvas
        self.readme_tiles = {}

        card = Card(canvas, fill=True, padx=22, pady=16, height=400)
        self.readme_window = canvas.create_window(0, 0, anchor="nw", window=card)
        body = card.body

        head = tk.Frame(body, bg=PANEL)
        head.pack(fill="x", pady=(0, 10))
        self._icon(head, "book", LIME, size=20, box=26, width=2).pack(
            side="left", padx=(0, 10),
        )
        title = self._label(head, 16, "bold", fg=LIME)
        title.config(text="README")
        title.pack(side="left")

        close = RoundButton(
            head, icon="x", size=10, height=32, padx=8,
            command=self.toggle_readme,
        )
        close.pack(side="right")

        frame = tk.Frame(body, bg=PANEL)
        frame.pack(fill="both", expand=True)
        text = tk.Text(
            frame, bg=PANEL, fg="#d4d4d8", font=ui(11), wrap="word",
            relief="flat", bd=0, highlightthickness=0, padx=6, pady=4,
            spacing1=2, spacing3=4, cursor="arrow",
            selectbackground=BUTTON_ACTIVE, insertwidth=0,
        )
        scroll = ttk.Scrollbar(
            frame, orient="vertical", command=text.yview,
            style="Thin.Vertical.TScrollbar",
        )
        text.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        text.pack(side="left", fill="both", expand=True)
        self.readme_text = text

        text.tag_configure("title", font=ui(22, "bold"), foreground=LIME,
                           spacing1=6, spacing3=10)
        text.tag_configure("h2", font=ui(14, "bold"), foreground=LIME,
                           spacing1=22, spacing3=8)
        text.tag_configure("h3", font=ui(11, "bold"), foreground=WHITE,
                           spacing1=12, spacing3=4)
        text.tag_configure("item", lmargin1=12, lmargin2=30)
        text.tag_configure("step", lmargin1=12, lmargin2=34)
        text.tag_configure("mono", font=mono(10), foreground=WHITE,
                           lmargin1=18, lmargin2=18, spacing1=0, spacing3=0)
        text.tag_configure("name", foreground=LIME, font=ui(11, "bold"))
        text.tag_configure("email", foreground=BLUE_TEXT)
        text.tag_configure("link", foreground=BLUE_TEXT, underline=True)
        text.tag_bind("link", "<Enter>", lambda e: text.config(cursor="hand2"))
        text.tag_bind("link", "<Leave>", lambda e: text.config(cursor="arrow"))
        text.tag_bind("link", "<Button-1>", self._open_readme_link)

        canvas.bind("<Configure>", self._layout_readme)

    def _open_readme_link(self, event):
        text = self.readme_text
        index = text.index(f"@{event.x},{event.y}")
        ranges = text.tag_prevrange("link", f"{index}+1c")

        if ranges:
            import webbrowser
            webbrowser.open(text.get(*ranges))

        return "break"

    def _layout_readme(self, event=None):
        canvas = self.readme_canvas
        self._tile_background(canvas, self.readme_tiles)
        width = canvas.winfo_width()
        height = canvas.winfo_height()
        card_w = min(980, max(width - 80, 300))
        canvas.coords(self.readme_window, (width - card_w) / 2, 30)
        canvas.itemconfigure(
            self.readme_window, width=card_w, height=max(height - 60, 200),
        )

    def toggle_readme(self, event=None):
        self.readme_open = not self.readme_open

        if self.readme_open:
            self.render_readme()
            self.readme_canvas.place(x=0, y=0, relwidth=1, relheight=1)
            tk.Misc.tkraise(self.readme_canvas)
            self.readme_text.focus_set()
        else:
            self.readme_canvas.place_forget()
            self.focus_entry()

        return "break"

    @staticmethod
    def readme_blocks(text):
        """Convierte el README de texto en bloques (estilo, texto).

        Une las líneas de un mismo párrafo o punto de lista y conserva tal
        cual las tablas y los bloques con sangría.
        """
        lines = text.splitlines()
        blocks = []
        index = 0
        last_length = 0

        def is_rule(line):
            return len(line.strip()) >= 3 and set(line.strip()) <= {"=", "-"}

        def is_table(line):
            return re.search(r"\S {2,}\S", line.strip()) is not None

        while index < len(lines):
            line = lines[index].rstrip()
            following = lines[index + 1] if index + 1 < len(lines) else ""

            if not line.strip():
                blocks.append(("gap", ""))
                index += 1
                continue

            if is_rule(following) and not is_rule(line):
                style = "title" if "=" in following else "h2"
                blocks.append((style, line.strip()))
                index += 2
                continue

            if is_rule(line):
                index += 1
                continue

            indent = len(line) - len(line.lstrip())
            previous = blocks[-1][0] if blocks else "gap"
            stripped = line.strip()

            if is_table(line) or (indent >= 4 and previous in ("gap", "mono")):
                blocks.append(("mono", line))
            elif re.match(r"^(\d+\.|-)\s", stripped):
                blocks.append(("step" if stripped[0].isdigit() else "item", stripped))
            elif previous in ("item", "step") and indent > 0 or (
                previous == "body" and last_length >= 55
            ):
                # Línea cortada al ancho del .txt: sigue el mismo párrafo.
                style, value = blocks[-1]
                blocks[-1] = (style, value + " " + stripped)
            elif re.match(r"^[A-ZÁÉÍÓÚÑÇ0-9 .]+ - ", stripped):
                blocks.append(("h3", stripped))
            else:
                blocks.append(("body", stripped))

            last_length = len(line)
            index += 1

        return [block for block in blocks if block[0] != "gap"]

    def render_readme(self):
        text = self.readme_text
        text.configure(state="normal")
        text.delete("1.0", "end")

        for style, value in self.readme_blocks(README_TEXTS.get(self.lang, "")):
            start = text.index("end-1c")
            text.insert("end", value + "\n", () if style == "body" else (style,))

            # Enlaces (GitHub): se abren en el navegador con un clic.
            for match in re.finditer(r"https?://\S+", value):
                text.tag_add(
                    "link", f"{start}+{match.start()}c",
                    f"{start}+{match.end()}c",
                )

            if style == "mono":
                continue

            # Nombres de personas y el correo de contacto en color.
            for pattern, tag in (
                (r"Desblack|\bDess\b", "name"),
                (r"[\w.+-]+@[\w-]+\.[\w.]+", "email"),
            ):
                for match in re.finditer(pattern, value):
                    text.tag_add(
                        tag, f"{start}+{match.start()}c",
                        f"{start}+{match.end()}c",
                    )

        text.configure(state="disabled")
        text.yview_moveto(0)

    def _on_resize(self, event):
        if event.widget is not self.root:
            return

        width = max(event.width - 160, 320)
        self.status.config(wraplength=width)
        self.details.config(wraplength=width)
        self.banner_label.config(wraplength=max(width - 340, 240))

    def focus_entry(self):
        self.entry.focus_set()

    def _bind_keys(self):
        self.entry.bind("<KeyRelease>", self.on_key_release)
        self.entry.bind("<Return>", self.on_enter)
        self.entry.bind("<Delete>", self.on_entry_delete)
        self.entry.bind("<BackSpace>", self.on_backspace)
        self.entry.bind("<<Paste>>", self.on_paste)
        self.entry.bind("<<Copy>>", self.on_copy)
        self.root.bind("<F2>", self.on_next_pallet_key)
        self.root.bind("<F11>", self.toggle_big_view)
        self.root.bind("<Control-z>", self.on_undo_key)
        self.root.bind("<Control-Z>", self.on_undo_key)
        self.root.bind("<Control-e>", self.on_export_key)
        self.root.bind("<Control-E>", self.on_export_key)
        self.root.bind("<Escape>", self.on_escape)
        self.root.bind("<Control-p>", self.ask_print)
        self.root.bind("<Control-P>", self.ask_print)

    # ========================================================
    # IDIOMAS Y MODOS
    # ========================================================

    def change_language(self, event=None):
        self.lang = LANGUAGES.get(self.language_var.get(), "es")
        self.update_language()
        self.schedule_save()
        self.focus_entry()

    def update_language(self):
        """Actualiza los textos visibles sin reiniciar la aplicación."""
        t = self.t
        checker = self.mode == "checker"

        self.mode_buttons["checker"].config(text=t("mode_checker"))
        self.mode_buttons["create"].config(text=t("mode_create"))

        for mode, button in self.mode_buttons.items():
            button.config(variant="primary" if mode == self.mode else "ghost")

        self.alt_matching_check.set_text(t("alt_match"))
        self.alt_matching_check.redraw()
        self.big_button.config(text=t("big_view"))
        self.readme_button.config(text=t("readme"))

        if self.readme_open:
            self.render_readme()

        self.dest_label.config(text=t("destination"))
        self.dest_menu.configure(
            values=[t("dest_column"), t("dest_sheet"), t("dest_cell")]
        )
        self.dest_menu.current(self.dest_index)
        self.name_label.config(text=t("pallet_name"))
        self.next_pallet_button.config(text=t("next_pallet"))
        self.delete_button.config(text=t("delete_selected"))
        self.undo_button.config(text=t("undo"))
        self.export_list_button.config(text=t("export_list"))
        self.rack_check.set_text(t("rack_mode"))
        self.rack_check.redraw()
        self.validate_check.set_text(t("validate_mode"))
        # Una bloquea a la otra: Modo Rack solo coloca máquinas y no valida.
        self.validate_check.set_enabled(not self.rack_mode_var.get())
        self.rack_check.set_enabled(not self.validation_active())
        self.valid_alt_check.set_text(t("alt_match"))
        self.valid_alt_check.redraw()
        self.load_valid_button.config(text=t("validate_load"))
        self.valid_file_label.config(
            text=t("validate_file", name=os.path.basename(self.valid_path))
            if self.valid_path else t("validate_no_list"),
            fg=WHITE if self.valid_path else YELLOW_TEXT,
        )
        self.rack_location_label.config(text=t("rack_location"))
        self.rack_rows_label.config(text=t("rack_rows"))
        self.rack_cols_label.config(text=t("rack_cols"))
        self.rack_order_label.config(text=t("rack_order"))
        self.order_buttons["h"].config(text=t("rack_order_h"))
        self.order_buttons["v"].config(text=t("rack_order_v"))
        self._refresh_order_buttons()
        self.add_rack_button.config(text=t("add_rack"))

        self.load_button.config(text=t("load"))
        self.print_button.config(text=t("print"))
        self.print_list_button.config(text=t("print"))
        self.export_button.config(text=t("export"))
        self.input_label.config(
            text=t("input_label" if checker else "input_label_create")
        )
        self.footer.config(text=t("footer"))
        self._build_keys()

        for label, key in self.legend_labels:
            label.config(text=t(key))

        self.history_title.config(text=t("history_title").upper())
        self.undo_scan_button.config(text=t("undo"))
        self.search_label.config(text=t("search"))
        self.refresh_file_info()

        if self.reset_armed:
            self.reset_button.config(text=t("reset_arm"))
        else:
            self.reset_button.config(text=t("reset_label"))

        self.update_counter()
        self.update_create_counts()
        self.refresh_history()
        self.rebuild_pending_tree()
        self.refresh_status()
        self.refresh_banner()

    def set_mode(self, new_mode):
        if new_mode == self.mode:
            self.focus_entry()
            return

        self.mode = new_mode
        self.standby = False
        self._cancel_flash()
        self.disarm_reset()
        self.set_status("")
        self.entry.delete(0, tk.END)
        self.update_mode_view()
        self.schedule_save()
        self.focus_entry()

    def update_mode_view(self):
        checker = self.mode == "checker"

        for card in (
            self.checker_header, self.create_header,
            self.history_card, self.lists_card, self.sheet_card,
        ):
            card.grid_remove()

        self.checker_legend.pack_forget()
        self.create_legend.pack_forget()
        self.alt_matching_check.pack_forget()
        self.rack_check.pack_forget()
        self.validate_check.pack_forget()
        self.validate_row.pack_forget()
        self.pallet_row.pack_forget()
        self.rack_row.pack_forget()

        if checker:
            self._put(self.checker_header, 1, 0, 2)
            self._put(self.history_card, 5, 0)
            self._put(self.lists_card, 5, 1, 2)
            self.checker_legend.pack(fill="x")
            self.alt_matching_check.pack(side="left", padx=(18, 0))
            self.progress.pack(fill="x", before=self.counter_info)
            self.rate_label.pack(fill="x")
            self.length_label.pack(side="right", padx=(10, 0))
        else:
            self._put(self.create_header, 1, 0, 2)
            self._put(self.sheet_card, 5, 0, 3)
            self.create_legend.pack(fill="x")
            self.rack_check.pack(side="left", padx=(18, 0))
            self.validate_check.pack(side="left", padx=(18, 0))

            if self.validation_active():
                self.validate_row.pack(fill="x", pady=(12, 0))
            row = self.rack_row if self.rack_mode_var.get() else self.pallet_row
            row.pack(fill="x", before=self.create_row_two)
            self.progress.pack_forget()
            self.rate_label.pack_forget()
            self.length_label.pack_forget()

        self.update_language()

        if checker:
            self.refresh_history()
        else:
            self.refresh_create_view()

        self.apply_counter_colors()
        self.refresh_banner()

    def on_alt_toggle(self):
        self.schedule_save()
        self.focus_entry()

    def get_file_info_text(self):
        if not self.csv_path:
            return self.t("no_file")

        return self.t(
            "file_summary",
            name=os.path.basename(self.csv_path),
            count=self.loaded_count,
            unique=len(self.serials),
        )

    def refresh_file_info(self):
        self.file_info.config(
            text=self.get_file_info_text(),
            fg=WHITE if self.csv_path else MUTED,
        )
        self.file_icon.delete("all")
        draw_icon(
            self.file_icon, "file", 9, 9, GREEN if self.csv_path else DIM,
            size=16, width=1.8,
        )

    # ========================================================
    # MENSAJES, COLORES Y AVISOS
    # ========================================================

    # tipo -> (texto, fondo, borde, icono, color del icono)
    STATUS_STYLES = {
        "ok": (GREEN, GREEN_BG, GREEN_BORDER, "check", "#052e16"),
        "error": (RED, RED_BG, RED_BORDER, "x", "#2a0606"),
        "warn": (YELLOW_TEXT, YELLOW_BG, YELLOW_BORDER, "warn", "#1c1300"),
        "info": (WHITE, PANEL, BORDER, "info", PANEL),
    }

    def set_status(self, key, _details=None, **params):
        self.status_key = key
        self.status_params = params
        self.status_details = _details or []
        self.refresh_status()

    def refresh_status(self):
        if not self.status_key:
            text = self.t("details" if self.mode == "checker" else "create_hint")
            fg, bg, border, icon, icon_fg = MUTED, PANEL, BORDER, None, None
            font = ui(11)
            detail = ""
        else:
            kind = STATUS_KINDS.get(self.status_key, "info")
            fg, bg, border, icon, icon_fg = self.STATUS_STYLES[kind]
            text = self.t(self.status_key, **self.status_params)
            font = ui(15, "bold")
            detail = "  ·  ".join(
                self.t(key, **params) for key, params in self.status_details
            )

            if self.status_key == "already_scanned":
                icon = "reset"

        self.status_card.set_colors(bg=bg, border=border)
        self.status_texts.config(bg=bg)
        self.status.config(text=text, fg=fg, bg=bg, font=font)
        self.details.config(
            text=detail, bg=bg, fg=MUTED if bg == PANEL else "#d4d4d8",
        )

        canvas = self.status_icon
        canvas.config(bg=bg)
        canvas.delete("all")

        if icon:
            circle = fg if kind != "info" else BORDER_STRONG
            canvas.create_oval(1, 1, 37, 37, fill=circle, outline="")
            draw_icon(
                canvas, icon, 19, 19, icon_fg if kind != "info" else WHITE,
                size=20, width=2.6,
            )
        else:
            draw_icon(canvas, "info", 19, 19, DIM, size=22, width=1.8)

        self.refresh_big_view()

    def alert_active(self):
        """True mientras la app debe permanecer en amarillo."""
        if self.mode == "checker":
            return self.standby

        return self.dup_count() > 0

    def apply_counter_colors(self):
        alert = self.alert_active()
        bg = YELLOW_BG if alert else PANEL

        self.counter_card.set_colors(
            bg=bg, border=YELLOW_BORDER if alert else BORDER,
        )

        for widget in (
            self.counter_row, self.counter_total,
            self.counter_info, self.rate_label,
        ):
            widget.config(bg=bg)

        self.counter_caption.config(bg=bg, fg=YELLOW_TEXT if alert else MUTED)
        self.counter.config(bg=bg, fg=YELLOW_TEXT if alert else WHITE)
        self.progress.config(bg=bg)
        self.progress.set(self.progress.value, YELLOW if alert else GREEN)

        if self.flash_job is None:
            self.set_ring(YELLOW if alert else BLUE_LIGHT)

        self.refresh_big_view()

    def update_counter(self):
        t = self.t

        if self.mode == "checker":
            found = len(self.found)
            total = len(self.serials)
            left = total - found
            self.counter.config(text=str(found))
            self.counter_total.config(text=f"/ {total}")
            self.counter_caption.config(text=t("caption_scanned").upper())
            self.progress.set(found / total if total else 0)
            self.counter_info.config(
                text=t(
                    "progress_info",
                    pct=int(found * 100 / total) if total else 0,
                    left=left,
                )
            )
            rate = self.scan_rate()

            if rate and left:
                self.rate_label.config(
                    text=t(
                        "rate_info",
                        rate=round(rate),
                        eta=max(1, round(left / rate)),
                    )
                )
            else:
                self.rate_label.config(text="")
        elif self.rack_mode_var.get():
            # Modo Rack: la siguiente ubicación en grande.
            pallet, slot = self.target_slot()
            self.counter_caption.config(text=t("caption_next").upper())
            self.counter_total.config(text="")

            if pallet is None:
                self.counter.config(text="—")
            elif slot is None:
                self.counter.config(text="✓")
            else:
                self.counter.config(text=f"({slot[0]},{slot[1]})")
        else:
            self.counter.config(text=str(len(self.create_seen)))
            self.counter_total.config(text="")
            self.counter_caption.config(text=t("caption_unique").upper())

        self.refresh_big_view()

    def scan_rate(self):
        """S/N encontrados por minuto en los últimos minutos (o None)."""
        now = time.time()
        recent = []

        for entry in reversed(self.scan_log):
            if now - entry[0] > RATE_WINDOW_S:
                break

            if entry[2] in ("exact", "variant"):
                recent.append(entry[0])

        # Hace falta un poco de historia para que la estimación tenga sentido.
        if len(recent) < 3 or now - recent[-1] < 30:
            return None

        minutes = max((now - recent[-1]) / 60, 1)
        return len(recent) / minutes

    def flash(self, color):
        """Destello breve del borde del campo de escaneo."""
        if self.batch:
            return

        self._cancel_flash()
        self.set_ring(color)
        self.flash_job = self.root.after(FLASH_MS, self.restore_colors)

    def restore_colors(self):
        self.flash_job = None
        self.apply_counter_colors()

    def _cancel_flash(self):
        if self.flash_job is not None:
            try:
                self.root.after_cancel(self.flash_job)
            except Exception:
                pass

            self.flash_job = None

    def show_banner(self, text, first, second=None, third=None, checks=False):
        """Aviso amarillo integrado (sin ventanas emergentes).

        first/second/third = (texto, función). checks=True muestra las
        casillas de qué imprimir.
        """
        self.banner_label.config(text=text)
        actions = [action for action in (first, second, third) if action]
        state = (text, tuple(label for label, _ in actions), checks)

        if state != getattr(self, "_banner_state", None):
            self._banner_state = state

            for button in self.banner_buttons:
                button.pack_forget()

            for button, (label, command) in zip(self.banner_buttons, actions):
                button.config(text=label, command=command)
                button.pack(side="left", padx=(8, 0))

            self.banner_checks.pack_forget()

            if checks:
                self.banner_checks.pack(side="right", padx=(12, 4))

        if not self.banner_visible:
            self._put(self.banner, 3, 0, 3)
            self.banner_visible = True

    def hide_banner(self):
        if self.banner_visible:
            self.banner.grid_remove()
            self.banner_visible = False

    def refresh_banner(self):
        if self.print_prompt == self.mode == "checker":
            for key, check in self.print_checks.items():
                check.set_text(self.t("print_" + key))
                check.redraw()

            self.show_banner(
                self.t("print_ask_checker"),
                (self.t("print"), self.print_checker),
                (self.t("cancel"), self.cancel_print),
                checks=True,
            )
        elif self.print_prompt == self.mode == "create":
            self.show_banner(
                self.t("print_ask"),
                (self.t("print_current"), lambda: self.print_create(True)),
                (self.t("print_all"), lambda: self.print_create(False)),
                (self.t("cancel"), self.cancel_print),
            )
        elif self.mode == "checker" and self.pending_resume:
            self.show_banner(
                self.t(
                    "resume_prompt",
                    done=len(self.pending_resume),
                    total=len(self.serials),
                ),
                (self.t("resume_continue"), self.resume_continue),
                (self.t("resume_fresh"), self.resume_fresh),
            )
        elif (
            self.mode == "create" and self.validation_active()
            and not self.valid_set
        ):
            self.show_banner(
                self.t("validate_banner"),
                (self.t("validate_load"), self.load_validation_list),
                (self.t("validate_off"), self.turn_off_validation),
            )
        elif self.mode == "create" and self.rack_full_pending():
            pallet = self.rack_full_pending()
            self.show_banner(
                self.t(
                    "rack_full_banner", rack=pallet["name"],
                    total=len(self.rack_slots(pallet)),
                ),
                (self.t("add_rack"), self.add_rack),
                (self.t("dup_hide"), self.hide_rack_banner),
            )
        elif (
            self.mode == "create"
            and self.dup_count() > 0
            and self.dup_count() > self.dup_dismissed_at
        ):
            names = self.dup_serials()
            shown = ", ".join(names[:5])

            if len(names) > 5:
                shown += "…"

            self.show_banner(
                self.t(
                    "dup_banner",
                    count=self.dup_count(),
                    list=shown,
                ),
                (self.t("dup_remove"), self.remove_duplicates),
                (self.t("dup_hide"), self.hide_dup_banner),
            )
        else:
            self.hide_banner()

    def on_escape(self, event=None):
        if self.sheet_grid.editor is not None:
            return None

        if self.readme_open:
            return self.toggle_readme()

        if self.big_view:
            return self.toggle_big_view()

        if self.print_prompt:
            self.cancel_print()
            return "break"

        if self.mode == "create":
            self.dup_dismissed_at = self.dup_count()

            if self.rack_full_pending():
                self.rack_dismissed.add(self.rack_full_pending()["number"])
        elif self.pending_resume:
            self.pending_resume = set()
        else:
            self.standby = False

            if self.status_key == "already_scanned":
                self.set_status("")

        self.refresh_banner()
        self.apply_counter_colors()
        self.focus_entry()
        return "break"

    def on_search_escape(self, event=None):
        self.search_var.set("")
        self.focus_entry()
        return "break"

    # ========================================================
    # LECTURA DE ARCHIVOS
    # ========================================================

    @staticmethod
    def normalize_serial(value):
        """Convierte una celda en un SN limpio o devuelve '' si no sirve."""
        if value is None:
            return ""

        # Números de Excel: 123456.0 -> "123456".
        if isinstance(value, float) and value.is_integer():
            value = int(value)

        text = str(value).strip().upper()

        if text in {"NONE", "NAN"}:
            return ""

        return text

    @staticmethod
    def is_green_cell(cell):
        """True si la celda tiene el verde que usa el reporte exportado."""
        fill = getattr(cell, "fill", None)

        if fill is None or fill.fill_type != "solid":
            return False

        rgb = getattr(fill.fgColor, "rgb", None)

        return isinstance(rgb, str) and rgb.upper().endswith("C6EFCE")

    def extract_done(self, rows):
        """SN en verde de un reporte exportado antes (celdas, no valores)."""
        done = set()
        active_sn_columns = set()

        for row in rows:
            header_indices = {
                i
                for i, cell in enumerate(row)
                if normalize_header(cell.value) in SN_HEADERS
            }

            if header_indices:
                active_sn_columns = header_indices
                continue

            for index in active_sn_columns:
                if index >= len(row):
                    continue

                cell = row[index]

                if not self.is_green_cell(cell):
                    continue

                serial = self.normalize_serial(cell.value)

                if serial:
                    done.add(serial)

        return done

    def extract_serials(self, rows, sheet=None, first_row=1, locations=None):
        """Extrae valores de todas las columnas SN reconocidas.

        Si se pasa locations, guarda la hoja y fila de la primera
        aparición de cada S/N (para mostrar dónde está al escanearlo).
        """
        values = []
        active_sn_columns = set()

        for number, row in enumerate(rows, start=first_row):
            row = list(row)

            header_indices = {
                i
                for i, cell in enumerate(row)
                if normalize_header(cell) in SN_HEADERS
            }

            if header_indices:
                active_sn_columns = header_indices
                continue

            for index in active_sn_columns:
                if index >= len(row):
                    continue

                value = self.normalize_serial(row[index])

                if value:
                    values.append(value)

                    if locations is not None and value not in locations:
                        locations[value] = (sheet, number)

        return values

    def _read_csv_rows(self, path):
        """Lee un CSV probando codificaciones y separador (, ; tab |)."""
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                with open(path, "r", encoding=encoding, newline="") as file:
                    sample = file.read(8192)
                    file.seek(0)

                    try:
                        dialect = csv.Sniffer().sniff(
                            sample, delimiters=",;\t|"
                        )
                    except csv.Error:
                        dialect = csv.excel

                    return list(csv.reader(file, dialect))
            except UnicodeDecodeError:
                continue

        raise ValueError(self.t("csv_encoding"))

    def read_csv(self, path):
        locations = {}
        values = self.extract_serials(
            self._read_csv_rows(path), locations=locations,
        )

        return values, set(), locations

    def read_excel(self, path):
        if load_workbook is None:
            raise ImportError(self.t("excel_dependency"))

        # Se lee el libro completo (no solo valores) para poder
        # reconocer las celdas verdes de un reporte exportado antes.
        workbook = load_workbook(path, data_only=True)
        values = []
        done = set()
        locations = {}

        try:
            for worksheet in workbook.worksheets:
                # Hojas añadidas por un reporte exportado antes.
                if worksheet.title.lower() in REPORT_SHEETS:
                    continue

                rows = [list(row) for row in worksheet.iter_rows()]

                values.extend(
                    self.extract_serials(
                        [[cell.value for cell in row] for row in rows],
                        sheet=worksheet.title,
                        locations=locations,
                    )
                )
                done.update(self.extract_done(rows))
        finally:
            workbook.close()

        return values, done, locations

    # ========================================================
    # CARGAR ARCHIVO (Checker)
    # ========================================================

    def load_file(self):
        path = filedialog.askopenfilename(
            title=self.t("select_file"),
            filetypes=[
                (self.t("file_types"), "*.csv *.xlsx"),
                (self.t("csv_files"), "*.csv"),
                (self.t("excel_files"), "*.xlsx"),
                (self.t("all_files"), "*.*"),
            ],
        )

        if not path:
            self.focus_entry()
            return

        extension = os.path.splitext(path)[1].lower()

        try:
            if extension == ".csv":
                values, done, locations = self.read_csv(path)
            elif extension == ".xlsx":
                values, done, locations = self.read_excel(path)
            else:
                self.set_status("unsupported")
                self.focus_entry()
                return
        except Exception as exc:
            self.set_status("read_error", error=str(exc)[:200])
            self.focus_entry()
            return

        if not values:
            self.set_status("no_serials")
            self.focus_entry()
            return

        self._cancel_flash()
        self._clear_checker_state()
        self.serial_order = list(OrderedDict.fromkeys(values))
        self.serials = set(self.serial_order)
        self.serial_lengths = {len(s) for s in self.serial_order}
        self.locations = locations
        self.csv_path = path
        self.loaded_count = len(values)
        self.pending_resume = done & self.serials

        self.refresh_file_info()
        self.update_counter()
        self.refresh_history()
        self.rebuild_pending_tree()
        self.set_status("file_loaded")
        self.refresh_banner()
        self.restore_colors()
        self.schedule_save()
        self.focus_entry()

    def resume_continue(self):
        """Recupera los S/N en verde del reporte exportado antes."""
        self.found |= self.pending_resume & self.serials
        self.pending_resume = set()
        self.resumed = True

        self.update_counter()
        self.rebuild_pending_tree()
        self.set_status(
            "resumed", done=len(self.found), total=len(self.serials),
        )
        self.refresh_banner()
        self.schedule_save()
        self.focus_entry()

    def resume_fresh(self):
        self.pending_resume = set()
        self.refresh_banner()
        self.focus_entry()

    # ========================================================
    # EXPORTAR REPORTE (Checker)
    # ========================================================

    def export_report(self):
        if not self.csv_path or not self.serials:
            self.set_status("load_first")
            self.focus_entry()
            return

        if Workbook is None or PatternFill is None:
            self.set_status("export_dependency")
            self.focus_entry()
            return

        source_name = os.path.splitext(
            os.path.basename(self.csv_path)
        )[0]

        # Si se está continuando un reporte ya exportado, se propone
        # el mismo nombre para poder guardar el avance sobre él.
        if self.resumed and source_name.lower().endswith("_resultado"):
            default_name = f"{source_name}.xlsx"
        else:
            default_name = f"{source_name}_resultado.xlsx"

        output_path = filedialog.asksaveasfilename(
            title=self.t("save_report"),
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[(self.t("output_file"), "*.xlsx")],
        )

        if not output_path:
            self.focus_entry()
            return

        if not self.resumed and os.path.abspath(
            output_path
        ) == os.path.abspath(self.csv_path):
            self.set_status("same_file")
            self.focus_entry()
            return

        green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
        red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
        blue_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")

        workbook = None
        values_workbook = None

        try:
            extension = os.path.splitext(self.csv_path)[1].lower()

            if extension == ".xlsx":
                workbook = load_workbook(self.csv_path)

                # Segundo libro con los valores calculados de las
                # fórmulas, igual que al leer el archivo.
                values_workbook = load_workbook(
                    self.csv_path, data_only=True,
                )

            elif extension == ".csv":
                rows = self._read_csv_rows(self.csv_path)
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.title = "Datos"

                for row in rows:
                    worksheet.append(row)

            else:
                raise ValueError(self.t("invalid_format"))

            colored_cells = 0
            duplicate_cells = 0
            seen_serials = set()

            def cell_value(sheet_values, cell):
                """Valor calculado de la celda (no la fórmula)."""
                if sheet_values is None:
                    return cell.value

                return sheet_values.cell(
                    row=cell.row, column=cell.column,
                ).value

            for worksheet in workbook.worksheets:
                if worksheet.title.lower() in REPORT_SHEETS:
                    continue

                active_sn_columns = set()
                sheet_values = None
                sn_columns = set()

                if (
                    values_workbook is not None
                    and worksheet.title in values_workbook.sheetnames
                ):
                    sheet_values = values_workbook[worksheet.title]

                for row in worksheet.iter_rows():
                    header_indices = {
                        cell.column - 1
                        for cell in row
                        if normalize_header(cell_value(sheet_values, cell))
                        in SN_HEADERS
                    }

                    if header_indices:
                        active_sn_columns = header_indices
                        sn_columns |= {index + 1 for index in header_indices}
                        continue

                    for index in active_sn_columns:
                        if index >= len(row):
                            continue

                        cell = row[index]
                        serial = self.normalize_serial(
                            cell_value(sheet_values, cell)
                        )

                        if not serial:
                            continue

                        if serial in seen_serials:
                            # Repetido: azul. La primera aparición
                            # conserva verde/rojo según el escaneo.
                            cell.fill = blue_fill
                            duplicate_cells += 1
                        else:
                            seen_serials.add(serial)
                            cell.fill = (
                                green_fill if serial in self.found
                                else red_fill
                            )

                        colored_cells += 1

                if sn_columns:
                    autofit_columns(worksheet, columns=sn_columns)

            self._add_report_sheets(workbook)
            workbook.save(output_path)

            self.set_status(
                "export_success",
                path=output_path,
                cells=colored_cells,
                duplicates=duplicate_cells,
                found=len(self.found),
                remaining=len(self.serials - self.found),
                extras=len(self.extras),
            )

        except Exception as exc:
            self.set_status("export_error", error=str(exc)[:200])

        finally:
            if workbook is not None:
                workbook.close()

            if values_workbook is not None:
                values_workbook.close()

        self.focus_entry()

    def _add_report_sheets(self, workbook):
        """Añade al reporte las hojas Resumen, Escaneos y Sobrantes."""
        t = self.t
        bold = Font(bold=True)

        def stamp(seconds):
            return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(seconds))

        def add_row(sheet, values, header=False):
            sheet.append(values)

            for cell in sheet[sheet.max_row]:
                # Texto escaneado: nunca como fórmula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"

                if header:
                    cell.font = bold

        # Un reporte reanudado ya trae estas hojas: se reemplazan.
        for name in list(workbook.sheetnames):
            if name.lower() in REPORT_SHEETS:
                del workbook[name]

        summary = workbook.create_sheet(t("sheet_summary"))

        for label, value in self.summary_rows():
            add_row(summary, [label, value])
            summary.cell(row=summary.max_row, column=1).font = bold

        autofit_columns(summary)

        log = workbook.create_sheet(t("sheet_log"))
        add_row(log, t("log_headers").split("|"), header=True)

        for seconds, read, kind, match in self.scan_log:
            add_row(log, [stamp(seconds), read, t("hist_" + kind), match])

        autofit_columns(log)

        if self.extras:
            extras = workbook.create_sheet(t("sheet_extras"))
            add_row(extras, t("extras_headers").split("|"), header=True)

            for serial, entry in self.extras.items():
                add_row(extras, [serial, entry["count"], stamp(entry["first"])])

            autofit_columns(extras)

    def summary_rows(self):
        """Filas (etiqueta, valor) del resumen del Checker."""
        t = self.t
        times = [entry[0] for entry in self.scan_log]

        def stamp(seconds):
            return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(seconds))

        return [
            (t("sum_file"), os.path.basename(self.csv_path or "")),
            (t("sum_total"), len(self.serials)),
            (t("sum_found"), len(self.found)),
            (t("sum_pending"), len(self.serials) - len(self.found)),
            (t("sum_extras"), len(self.extras)),
            (t("sum_scans"), len(self.scan_log)),
            (t("sum_first"), stamp(min(times)) if times else ""),
            (t("sum_last"), stamp(max(times)) if times else ""),
            (t("sum_exported"), stamp(time.time())),
        ]

    # ========================================================
    # IMPRIMIR (la pregunta va en el aviso integrado)
    # ========================================================

    def ask_print(self, event=None):
        if self.mode == "checker" and not self.serials:
            self.set_status("load_first")
        elif self.mode == "create" and not self.create_items:
            self.set_status("nothing_to_export")
        elif Workbook is None:
            self.set_status("export_dependency")
        else:
            self.print_prompt = self.mode
            self.refresh_banner()

        self.focus_entry()
        return "break"

    def cancel_print(self):
        self.print_prompt = None
        self.refresh_banner()
        self.focus_entry()

    def print_create(self, current_only):
        self.print_prompt = None
        self.refresh_banner()
        workbook = self.build_create_workbook(
            only_sheet=self.view_sheet if current_only else None,
        )
        self.send_to_printer(workbook)

    def print_checker(self):
        choices = [key for key, var in self.print_vars.items() if var.get()]

        if not choices:
            self.set_status("print_nothing")
            self.focus_entry()
            return

        self.print_prompt = None
        self.refresh_banner()
        self.send_to_printer(self.build_checker_print(choices))

    def build_checker_print(self, choices):
        """Libro con lo elegido: resumen, pendientes y/o sobrantes."""
        t = self.t
        workbook = Workbook()
        bold = Font(bold=True)
        first = True

        def new_sheet(title):
            nonlocal first
            sheet = workbook.active if first else workbook.create_sheet()
            sheet.title = title
            first = False
            return sheet

        def add(sheet, values, header=False):
            sheet.append(values)

            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"

                if header:
                    cell.font = bold

        if "summary" in choices:
            sheet = new_sheet(t("sheet_summary"))

            for label, value in self.summary_rows()[:-1]:
                add(sheet, [label, value])
                sheet.cell(row=sheet.max_row, column=1).font = bold

            autofit_columns(sheet)

        if "pending" in choices:
            sheet = new_sheet(t("print_pending"))
            add(sheet, ["#", t("pending_col")], header=True)

            for number, serial in enumerate(self.pending_list(), start=1):
                add(sheet, [number, serial])

            autofit_columns(sheet)

        if "extras" in choices:
            sheet = new_sheet(t("sheet_extras"))
            add(sheet, t("extras_headers").split("|"), header=True)

            for serial, entry in self.extras.items():
                add(sheet, [
                    serial, entry["count"],
                    time.strftime("%H:%M:%S", time.localtime(entry["first"])),
                ])

            autofit_columns(sheet)

        return workbook

    @staticmethod
    def workbook_text(workbook):
        """Versión en texto (para imprimir con el Bloc de notas)."""
        lines = []

        for sheet in workbook.worksheets:
            lines += [sheet.title.upper(), "=" * len(sheet.title), ""]

            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                lines.append("    ".join(values).rstrip())

            lines += ["", ""]

        return "\r\n".join(lines)

    def send_to_printer(self, workbook):
        """Imprime en la impresora predeterminada de Windows.

        Con Excel instalado se imprime el libro; si no, un texto con el
        Bloc de notas.
        """
        try:
            folder = os.path.join(tempfile.gettempdir(), "SerialNumberChecker")
            os.makedirs(folder, exist_ok=True)
            base = os.path.join(folder, "print_" + time.strftime("%Y%m%d_%H%M%S"))

            for sheet in workbook.worksheets:
                # Todas las hojas seleccionadas: Excel imprime todas.
                sheet.sheet_view.tabSelected = True
                sheet.sheet_properties.pageSetUpPr.fitToPage = True
                sheet.page_setup.fitToWidth = 1
                sheet.page_setup.fitToHeight = 0

            workbook.save(base + ".xlsx")
            text = self.workbook_text(workbook)
            workbook.close()
            startfile = getattr(os, "startfile", None)

            if startfile is None:
                self.set_status("print_unavailable", path=base + ".xlsx")
                self.focus_entry()
                return

            try:
                startfile(base + ".xlsx", "print")
            except OSError:
                with open(base + ".txt", "w", encoding="utf-8-sig") as file:
                    file.write(text)

                startfile(base + ".txt", "print")

            self.set_status("print_sent")
        except Exception as exc:
            self.set_status("print_error", error=str(exc)[:200])

        self.focus_entry()

    def export_current(self):
        if self.mode == "create":
            self.export_create_list()
        else:
            self.export_report()

    # ========================================================
    # COMPROBACIÓN DE NÚMEROS DE SERIE (Checker)
    # ========================================================

    def expected_length(self):
        """Largo más común de los S/N del archivo (o 17 si no hay)."""
        if not self.serial_order:
            return CHECKER_LENGTH

        # Se calcula una vez por lista cargada (se consulta en cada tecla).
        key = (id(self.serial_order), len(self.serial_order))

        if getattr(self, "_expected_key", None) != key:
            counts = {}

            for serial in self.serial_order:
                counts[len(serial)] = counts.get(len(serial), 0) + 1

            self._expected_key = key
            self._expected_value = max(counts, key=counts.get)

        return self._expected_value

    def update_length_label(self):
        if self.mode != "checker":
            return

        typed = len(self.entry.get().strip())
        self.length_label.config(
            text=f"{typed}/{self.expected_length()}" if typed else "",
        )

    def on_key_release(self, event=None):
        if self.mode != "checker":
            return

        raw = self.entry.get()
        value = raw.strip().upper()

        if value != raw:
            cursor = self.entry.index(tk.INSERT)
            self.entry.delete(0, tk.END)
            self.entry.insert(0, value)
            self.entry.icursor(min(cursor, len(value)))

        if len(value) == CHECKER_LENGTH and self.serials:
            self.check_serial()

        self.update_length_label()

    def on_enter(self, event=None):
        if self.mode == "create":
            self.add_create_item()
        elif self.serials and self.entry.get().strip():
            self.check_serial()
        elif not self.serials:
            self.set_status("load_first")

        return "break"

    def on_entry_delete(self, event=None):
        """Supr con el campo vacío elimina el registro seleccionado."""
        if self.mode == "create" and not self.entry.get():
            self.delete_selected_item()
            return "break"

        return None

    def on_backspace(self, event=None):
        """Retroceso con el campo vacío elimina el registro seleccionado.

        Solo con una pulsación nueva: si se mantiene pulsado para vaciar el
        campo, la repetición automática no debe borrar registros.
        """
        now = time.time()
        previous = self._last_backspace
        self._last_backspace = now

        if self.mode != "create" or self.entry.get():
            return None

        if now - previous < BACKSPACE_REPEAT_S:
            return "break"

        self.delete_selected_item()
        return "break"

    def location_details(self, serial):
        """Dónde está el S/N en el archivo original (hoja y fila)."""
        where = self.locations.get(serial)

        if not where:
            return []

        sheet, row = where

        if sheet is None:
            return [("detail_location_csv", {"row": row})]

        return [("detail_location", {"sheet": sheet, "row": row})]

    def check_serial(self):
        """Comprueba el S/N del campo. Devuelve el resultado o None."""
        original = self.entry.get().strip().upper()

        if not original:
            return None

        if not self.serials:
            self.set_status("load_first")
            return None

        candidates = [original]

        if (
            self.alt_matching_var.get()
            and original not in self.serials
            and len(original) >= 7
        ):
            for replacement in ("A", "B", "U"):
                candidate = original[:6] + replacement + original[7:]

                if candidate not in candidates:
                    candidates.append(candidate)

        match = next(
            (c for c in candidates if c in self.serials),
            None,
        )

        self.entry.delete(0, tk.END)
        self.update_length_label()
        self._cancel_flash()
        undo = {"serial": original, "match": match or "", "new": False}

        if match is None:
            self.standby = False

            if self.serial_lengths and len(original) not in self.serial_lengths:
                # Otro código de la caja (EAN, lote...): no cuenta como
                # sobrante.
                kind = "bad_format"
                expected = "/".join(
                    str(n) for n in sorted(self.serial_lengths)[:3]
                )
                self.set_status(
                    "bad_format", serial=original,
                    _details=[(
                        "detail_bad_format",
                        {"length": len(original), "expected": expected},
                    )],
                )
            else:
                kind = "not_found"
                entry = self.extras.get(original)

                if entry:
                    entry["count"] += 1
                else:
                    self.extras[original] = {"count": 1, "first": time.time()}

                self.set_status(
                    "not_found", serial=original,
                    _details=[("detail_extra", {})],
                )

            self.apply_counter_colors()
            self.flash(RED)
            self.play_sound(False)
        else:
            already_found = match in self.found
            self.found.add(match)

            if already_found:
                # Mismo S/N escaneado otra vez: la app queda en amarillo
                # hasta el siguiente escaneo (o Esc) para que se note.
                kind = "repeated"
                self.standby = True
                self.set_status(
                    "already_scanned", serial=match,
                    _details=[("detail_repeat", {})],
                )
                self.apply_counter_colors()
                self.play_duplicate_sound()
            else:
                undo["new"] = True
                self.standby = False
                details = self.location_details(match)

                if match == original:
                    kind = "exact"
                    self.set_status("found_exact", _details=details)
                else:
                    kind = "variant"
                    self.set_status(
                        "found_variant", serial=match,
                        _details=details + [
                            ("detail_variant", {"original": original}),
                        ],
                    )

                self.apply_counter_colors()
                self.flash(GREEN)
                self.play_sound(True)

        undo["kind"] = kind
        self.checker_undo.append(undo)
        del self.checker_undo[:-CHECKER_UNDO_LIMIT]

        self.scan_log.append([time.time(), original, kind, match or ""])
        del self.scan_log[:-SCAN_LOG_LIMIT]

        self.push_history(match or original, kind, original)
        self.update_counter()

        if kind == "exact" or kind == "variant":
            self.mark_pending_found(match)
        elif kind == "not_found" and self.list_tab == "extras":
            self.rebuild_pending_tree()
        else:
            self.update_list_tabs()

        self.schedule_save()
        self.focus_entry()
        return kind

    def undo_scan(self):
        """Deshace el último escaneo del Checker (Ctrl+Z)."""
        if not self.checker_undo:
            self.set_status("nothing_to_undo")
            self.focus_entry()
            return

        action = self.checker_undo.pop()
        kind = action["kind"]

        if kind in ("exact", "variant") and action["new"]:
            self.found.discard(action["match"])
        elif kind == "not_found":
            entry = self.extras.get(action["serial"])

            if entry:
                entry["count"] -= 1

                if entry["count"] <= 0:
                    del self.extras[action["serial"]]

        if self.scan_log and self.scan_log[-1][1] == action["serial"]:
            self.scan_log.pop()

        if self.history and self.history[0]["original"] == action["serial"]:
            self.history.pop(0)

        self.standby = False
        self._cancel_flash()
        self.set_status("scan_undone", serial=action["serial"])
        self.update_counter()
        self.refresh_history()
        self.rebuild_pending_tree()
        self.apply_counter_colors()
        self.schedule_save()
        self.focus_entry()

    # ----- pegar varios a la vez -----

    def on_paste(self, event=None):
        """Ctrl+V con varias líneas: procesa cada valor por separado."""
        try:
            text = self.root.clipboard_get()
        except Exception:
            return None

        if not isinstance(text, str):
            return None

        # En Crear lista se acepta cualquier texto (comas incluidas).
        pattern = r"[\r\n\t;,]+" if self.mode == "checker" else r"[\r\n\t]+"
        values = [part.strip() for part in re.split(pattern, text) if part.strip()]

        if len(values) < 2:
            return None

        self.paste_values(values)
        return "break"

    def paste_values(self, values):
        if self.mode == "checker" and not self.serials:
            self.set_status("load_first")
            self.focus_entry()
            return

        self.batch = True
        kinds = []
        duplicates_before = self.dup_count()
        items_before = len(self.create_items)

        try:
            for value in values:
                self.entry.delete(0, tk.END)
                self.entry.insert(0, value)

                if self.mode == "checker":
                    kinds.append(self.check_serial())
                else:
                    self.add_create_item()
        finally:
            self.batch = False

        if self.mode == "checker":
            self.standby = False
            self.set_status(
                "paste_summary",
                total=len(values),
                found=sum(k in ("exact", "variant") for k in kinds),
                missing=sum(k in ("not_found", "bad_format") for k in kinds),
                repeated=kinds.count("repeated"),
            )
        elif self.validation_active():
            added = len(self.create_items) - items_before
            self.set_status(
                "paste_validated",
                total=len(values),
                added=added,
                invalid=len(values) - added,
                duplicates=self.dup_count() - duplicates_before,
            )
        else:
            self.set_status(
                "paste_added",
                count=len(values),
                duplicates=self.dup_count() - duplicates_before,
            )

        self.apply_counter_colors()
        self.play_sound(True)
        self.focus_entry()

    # ----- historial de escaneos -----

    HISTORY_COLORS = {
        "exact": GREEN,
        "variant": GREEN,
        "not_found": RED,
        "bad_format": RED,
        "repeated": YELLOW,
    }

    def push_history(self, serial, kind, original):
        self.history.insert(
            0, {"serial": serial, "kind": kind, "original": original},
        )
        del self.history[HISTORY_SIZE:]
        self.refresh_history()

    def refresh_history(self):
        for index, (dot, serial, tag) in enumerate(self.history_rows):
            dot.delete("all")

            if index < len(self.history):
                entry = self.history[index]
                color = self.HISTORY_COLORS.get(entry["kind"], MUTED)
                dot.create_oval(1, 1, 9, 9, fill=color, outline="")
                label = self.t("hist_" + entry["kind"])

                if entry["kind"] == "variant":
                    label += f" ← {entry['original']}"

                serial.config(text=entry["serial"], fg=WHITE, font=mono(11))
                tag.config(text=label, fg=color)
            elif index == 0:
                serial.config(text=self.t("history_empty"), fg=DIM, font=ui(10))
                tag.config(text="")
            else:
                serial.config(text="")
                tag.config(text="")

    # ----- pendientes y sobrantes -----

    def pending_list(self):
        return [s for s in self.serial_order if s not in self.found]

    def set_list_tab(self, tab):
        self.list_tab = tab
        self.rebuild_pending_tree()
        self.focus_entry()

    def update_list_tabs(self):
        counts = {
            "pending": len(self.serials) - len(self.found),
            "extras": len(self.extras),
        }

        for tab, key in (("pending", "tab_pending"), ("extras", "tab_extras")):
            self.tab_buttons[tab].config(
                text=self.t(key, count=counts[tab]),
                variant="primary" if tab == self.list_tab else "tab",
            )

    def rebuild_pending_tree(self):
        self.update_list_tabs()
        tree = self.pending_tree

        for item_id in tree.get_children():
            tree.delete(item_id)

        self.pending_iids = {}
        query = self.search_var.get().strip().upper()

        def stripe(number):
            return ("odd",) if number % 2 else ()

        if self.list_tab == "extras":
            tree.configure(displaycolumns=("serial", "count", "time"))
            tree.heading("serial", text=self.t("extras_col"))
            tree.heading("count", text=self.t("count_col"))
            tree.heading("time", text=self.t("time_col"))
            rows = [
                (serial, entry) for serial, entry in self.extras.items()
                if query in serial
            ]

            for number, (serial, entry) in enumerate(rows[:PENDING_LIMIT]):
                tree.insert(
                    "", "end", iid=f"x{number}", tags=stripe(number),
                    values=(
                        serial, entry["count"],
                        time.strftime("%H:%M:%S", time.localtime(entry["first"])),
                    ),
                )

            extra_rows = len(rows) - PENDING_LIMIT
            empty_text = self.t("extras_none")
        else:
            tree.configure(displaycolumns=("serial",))
            tree.heading("serial", text=self.t("pending_col"))
            rows = [s for s in self.pending_list() if query in s]

            for number, serial in enumerate(rows[:PENDING_LIMIT]):
                iid = f"p{number}"
                tree.insert(
                    "", "end", iid=iid, tags=stripe(number),
                    values=(serial, "", ""),
                )
                self.pending_iids[serial] = iid

            extra_rows = len(rows) - PENDING_LIMIT
            empty_text = self.t("pending_none" if self.serials else "no_file")

        if extra_rows > 0:
            tree.insert(
                "", "end", iid="more", tags=("muted",),
                values=(self.t("pending_more", count=extra_rows), "", ""),
            )
        elif not rows:
            tree.insert(
                "", "end", iid="none", tags=("muted",),
                values=(empty_text, "", ""),
            )

    # ----- copiar S/N de las listas -----

    def selected_list_serial(self):
        """S/N de la fila seleccionada en Pendientes / Sobrantes (o None)."""
        try:
            selection = self.pending_tree.selection()
        except Exception:
            return None

        if not isinstance(selection, (list, tuple)) or not selection:
            return None

        iid = selection[0]

        if iid in ("none", "more"):
            return None

        values = self.pending_tree.item(iid, "values")
        return str(values[0]) if values else None

    def copy_selected_serial(self, event=None):
        serial = self.selected_list_serial()

        if serial:
            self.root.clipboard_clear()
            self.root.clipboard_append(serial)
            self.set_status("copied", serial=serial)

        self.focus_entry()
        return "break"

    def on_copy(self, event=None):
        """Ctrl+C en el campo de escaneo: si no hay texto seleccionado ahí,
        copia el S/N seleccionado en la lista."""
        try:
            if self.entry.selection_present():
                return None
        except Exception:
            pass

        if self.mode == "checker" and self.selected_list_serial():
            return self.copy_selected_serial()

        return None

    def show_copy_menu(self, event):
        row = self.pending_tree.identify_row(event.y)

        if not row or row in ("none", "more"):
            return "break"

        self.pending_tree.selection_set(row)
        self.copy_menu.entryconfigure(0, label=self.t("copy_sn"))

        try:
            self.copy_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.copy_menu.grab_release()

        return "break"

    def mark_pending_found(self, serial):
        self.update_list_tabs()

        if self.list_tab != "pending":
            return

        iid = self.pending_iids.pop(serial, None)

        if iid and self.pending_tree.exists(iid):
            self.pending_tree.delete(iid)

        if not self.pending_iids:
            self.rebuild_pending_tree()
    # ========================================================
    # MODO CREAR LISTA S/N
    # ========================================================

    def dup_count(self):
        return sum(1 for item in self.create_items if item["dup"])

    def dup_serials(self):
        names = []

        for item in self.create_items:
            if item["dup"] and item["serial"] not in names:
                names.append(item["serial"])

        return names

    def _recompute_dups(self):
        """Marca como duplicado todo S/N que ya apareció antes."""
        seen = set()

        for item in self.create_items:
            item["dup"] = item["serial"] in seen
            seen.add(item["serial"])

        self.create_seen = seen

    def _cleanup_pallets(self):
        """Quita pallets vacíos, salvo el último (el activo) y los racks."""
        used = {item["pallet"] for item in self.create_items}
        last = len(self.pallets) - 1

        self.pallets = [
            pallet for index, pallet in enumerate(self.pallets)
            if index == last or pallet["number"] in used
            or pallet["mode"] == "rack"
        ]

    def _restore_pallets(self, snapshot):
        existing = {pallet["number"] for pallet in self.pallets}

        for pallet in snapshot:
            if pallet["number"] not in existing:
                self.pallets.append(dict(pallet))

        self.pallets.sort(key=lambda pallet: pallet["number"])

    def _after_items_changed(self):
        self._recompute_dups()
        self._cleanup_pallets()

        if self.dup_count() < self.dup_dismissed_at:
            self.dup_dismissed_at = self.dup_count()

    def pallet_count(self, number):
        return sum(1 for item in self.create_items if item["pallet"] == number)

    def dest_mode(self):
        return ("column", "sheet", "cell")[self.dest_index % 3]

    def on_dest_selected(self, event=None):
        index = self.dest_menu.current()
        self.dest_index = index if index in (0, 1, 2) else 0
        self.schedule_save()
        self.focus_entry()

    def _create_pallet(self, mode, name, sheet, cell, **extra):
        # Si el pallet activo sigue vacío se reutiliza en vez de crear otro
        # (un rack vacío no: es una ubicación física).
        last = self.pallets[-1] if self.pallets else None

        if (
            last and last["mode"] != "rack" and mode != "rack"
            and self.pallet_count(last["number"]) == 0
        ):
            number = self.pallets.pop()["number"]
        else:
            self.pallet_counter += 1
            number = self.pallet_counter

        pallet = {
            "number": number,
            "mode": mode,
            "name": name,
            "sheet": sheet or "",
            "cell": cell,
        }
        pallet.update(extra)
        self.pallets.append(pallet)

        return pallet

    def _show_pallet_sheet(self, pallet):
        """Lleva el visor a la hoja del pallet (si tiene una propia)."""
        layout = compute_layout(self.create_items, self.pallets, self.default_sheet)
        sheet = layout["pallet_sheets"].get(pallet["number"])

        if sheet:
            self.view_sheet = sheet

    def next_pallet(self, event=None):
        """Crea el siguiente pallet según el destino elegido (sin diálogos)."""
        if self.mode != "create":
            return "break"

        if self.rack_mode_var.get():
            return self.add_rack()

        mode = self.dest_mode()
        name = self.pallet_name_var.get().strip()
        sheet = ""
        cell = "A1"

        if mode == "cell":
            selected = self.sheet_grid.selected

            if not selected:
                self.set_status("select_cell_first")
                self.focus_entry()
                return "break"

            sheet = self.view_sheet
            cell = f"{column_letter(selected[1])}{selected[0]}"

        pallet = self._create_pallet(mode, name, sheet, cell)
        self.pallet_name_var.set("")

        if mode == "sheet":
            self._show_pallet_sheet(pallet)

        self.set_status("pallet_ready", pallet=pallet["number"])
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()
        return "break"

    def new_sheet(self):
        """Botón "+" de las pestañas: hoja nueva (o rack nuevo) y salta a ella."""
        if self.rack_mode_var.get():
            return self.add_rack()

        name = self.pallet_name_var.get().strip()
        pallet = self._create_pallet("sheet", name, "", "A1")
        self.pallet_name_var.set("")
        self._show_pallet_sheet(pallet)
        self.set_status("pallet_ready", pallet=pallet["number"])
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()
        return "break"

    def _ensure_pallet(self):
        """Primer escaneo: crea el Pallet 1 sin preguntar nada."""
        if self.pallets and self.pallets[-1]["mode"] != "rack":
            return

        mode = self.dest_mode()
        sheet = ""
        cell = "A1"
        selected = self.sheet_grid.selected

        if mode == "cell" and selected and self.view_sheet not in self.rack_sheets():
            sheet = self.view_sheet
            cell = f"{column_letter(selected[1])}{selected[0]}"
        elif mode == "cell":
            mode = "column"

        self._create_pallet(mode, "", sheet, cell)

    def add_create_item(self):
        value = self.entry.get().strip()

        if not value:
            return

        if self.rack_mode_var.get():
            self.add_rack_item(value)
            return

        if self.validation_active():
            if not self.valid_set:
                # Aviso 1: sin lista cargada no se registra nada.
                self.entry.delete(0, tk.END)
                self.set_status("validate_need_list")
                self.play_sound(False)
                self.refresh_banner()
                self.focus_entry()
                return

            match = self.match_valid(value)

            if match is None:
                # Aviso 2: no está en la lista, no se registra.
                self.entry.delete(0, tk.END)
                self._cancel_flash()
                self.set_status(
                    "validate_invalid", serial=value,
                    _details=[("validate_invalid_detail", {})],
                )
                self.flash(RED)
                self.play_sound(False)
                self.focus_entry()
                return

            value = match

        self._ensure_pallet()
        pallet = self.pallets[-1]

        item = {"serial": value, "pallet": pallet["number"], "dup": False}
        self.create_items.append(item)
        self._recompute_dups()
        self._push_undo({"type": "add", "item": item})

        self.entry.delete(0, tk.END)
        self._cancel_flash()

        if item["dup"]:
            self.set_status("duplicate_warning", serial=value)
            self.play_duplicate_sound()
        else:
            self.set_status("create_added", serial=value)
            self.play_sound(True)

        self.refresh_create_view(focus_index=len(self.create_items) - 1)
        self.schedule_save()
        self.focus_entry()

    # ----- Validar lista -----

    def validation_active(self):
        """Validar lista marcada (nunca en Modo Rack: ahí solo se colocan
        máquinas y no se valida)."""
        return bool(self.validate_var.get()) and not self.rack_mode_var.get()

    def on_validate_toggle(self):
        self.update_mode_view()
        self.schedule_save()
        self.focus_entry()

    def turn_off_validation(self):
        self.validate_var.set(False)
        self.on_validate_toggle()

    def load_validation_list(self):
        """Carga el CSV o Excel con los S/N esperados en los pallets."""
        path = filedialog.askopenfilename(
            title=self.t("validate_select"),
            filetypes=[
                (self.t("file_types"), "*.csv *.xlsx"),
                (self.t("csv_files"), "*.csv"),
                (self.t("excel_files"), "*.xlsx"),
                (self.t("all_files"), "*.*"),
            ],
        )

        if not path:
            self.focus_entry()
            return

        extension = os.path.splitext(path)[1].lower()

        try:
            if extension == ".csv":
                values, _done, _where = self.read_csv(path)
            elif extension == ".xlsx":
                values, _done, _where = self.read_excel(path)
            else:
                self.set_status("unsupported")
                self.focus_entry()
                return
        except Exception as exc:
            self.set_status("read_error", error=str(exc)[:200])
            self.focus_entry()
            return

        if not values:
            self.set_status("no_serials")
            self.focus_entry()
            return

        self.valid_order = list(OrderedDict.fromkeys(values))
        self.valid_set = set(self.valid_order)
        self.valid_path = path
        self.set_status(
            "validate_loaded", name=os.path.basename(path),
            count=len(self.valid_set),
        )
        self.update_language()
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()

    def match_valid(self, value):
        """S/N de la lista que corresponde a lo leído, o None.

        Igual que en Checker: exacto y, si está marcado, con el 7.º
        carácter cambiado por A, B o U.
        """
        text = self.normalize_serial(value)

        if text in self.valid_set:
            return text

        if self.valid_alt_var.get() and len(text) >= 7:
            for replacement in ("A", "B", "U"):
                candidate = text[:6] + replacement + text[7:]

                if candidate in self.valid_set:
                    return candidate

        return None

    def validated_count(self):
        scanned = {self.normalize_serial(item["serial"]) for item in self.create_items}
        return len(self.valid_set & scanned)

    # ----- Modo Rack -----

    def rack_pallets(self):
        return [pallet for pallet in self.pallets if pallet["mode"] == "rack"]

    def rack_sheets(self):
        """{hoja: datos del rack} de las hojas que son racks."""
        return {
            name: sheet["rack"]
            for name, sheet in self.layout["sheets"].items()
            if "rack" in sheet
        }

    def rack_by_number(self, number):
        return next(
            (p for p in self.pallets if p["number"] == number
             and p["mode"] == "rack"),
            None,
        )

    @staticmethod
    def rack_slots(pallet):
        """Ubicaciones (x, y) en el orden de llenado elegido."""
        rows = pallet.get("rows") or RACK_ROWS
        cols = pallet.get("cols") or RACK_COLS

        if pallet.get("order") == "v":
            return [(x, y) for x in range(cols) for y in range(rows)]

        return [(x, y) for y in range(rows) for x in range(cols)]

    def rack_used(self, number):
        return {
            tuple(item["slot"]) for item in self.create_items
            if item["pallet"] == number and item.get("slot")
        }

    def next_free_slot(self, pallet, after=None):
        """Siguiente ubicación libre (después de "after"), o None si lleno."""
        slots = self.rack_slots(pallet)
        used = self.rack_used(pallet["number"])
        start = slots.index(tuple(after)) + 1 if after in slots else 0

        for slot in slots[start:] + slots[:start]:
            if slot not in used:
                return slot

        return None

    def target_rack(self):
        """Rack donde irá el siguiente escaneo."""
        cursor = self.rack_cursor

        if cursor:
            pallet = self.rack_by_number(cursor["pallet"])

            if pallet:
                return pallet

        racks = self.rack_pallets()
        return racks[-1] if racks else None

    def target_slot(self):
        """(rack, (x, y)) del siguiente escaneo; la ubicación es None si
        el rack está completo."""
        pallet = self.target_rack()

        if pallet is None:
            return None, None

        cursor = self.rack_cursor
        used = self.rack_used(pallet["number"])

        if (
            cursor and cursor["pallet"] == pallet["number"]
            and tuple(cursor["slot"]) in self.rack_slots(pallet)
            and tuple(cursor["slot"]) not in used
        ):
            return pallet, tuple(cursor["slot"])

        slot = self.next_free_slot(
            pallet, tuple(cursor["slot"]) if cursor else None,
        )
        return pallet, slot

    def _rack_settings(self):
        """Ubicación, filas y columnas escritas (o None si no son válidas)."""
        location = self.rack_location_var.get().strip()

        try:
            rows = int(self.rack_rows_var.get())
            cols = int(self.rack_cols_var.get())
        except (TypeError, ValueError):
            return location, None, None

        if not (1 <= rows <= RACK_MAX and 1 <= cols <= RACK_MAX):
            return location, None, None

        return location, rows, cols

    def add_rack(self, event=None):
        """Agrega el rack siguiente de la ubicación (DC05-1, DC05-2...)."""
        if self.mode != "create":
            return "break"

        location, rows, cols = self._rack_settings()

        if not location:
            self.set_status("rack_need_location")
            self.rack_location_entry.focus_set()
            return "break"

        if rows is None:
            self.set_status("rack_bad_size", maximum=RACK_MAX)
            self.focus_entry()
            return "break"

        numbers = [
            pallet.get("rack", 0) for pallet in self.rack_pallets()
            if pallet.get("location", "").lower() == location.lower()
        ]
        rack = max(numbers, default=0) + 1
        name = f"{location}-{rack}"

        pallet = self._create_pallet(
            "rack", name, "", "A1",
            location=location, rack=rack, rows=rows, cols=cols,
            order=self.rack_order,
        )
        self.rack_cursor = {"pallet": pallet["number"], "slot": [0, 0]}
        self._show_pallet_sheet(pallet)
        self.set_status("rack_ready", rack=name, rows=rows, cols=cols)
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()
        return "break"

    def add_rack_item(self, value):
        """Escaneo en Modo Rack: va a la ubicación guiada."""
        pallet, slot = self.target_slot()

        if pallet is None:
            location, rows, _cols = self._rack_settings()

            if not location or rows is None:
                self.add_rack()
                return

            self.add_rack()
            pallet, slot = self.target_slot()

            if pallet is None:
                return

        if slot is None:
            # Rack completo: no se registra para no perder la ubicación.
            self.entry.delete(0, tk.END)
            self.set_status("rack_full_scan", rack=pallet["name"])
            self.play_duplicate_sound()
            self.refresh_banner()
            self.focus_entry()
            return

        x, y = slot
        item = {
            "serial": value, "pallet": pallet["number"], "dup": False,
            "slot": [x, y],
        }
        self.create_items.append(item)
        self._recompute_dups()
        self._push_undo({"type": "add", "item": item})

        following = self.next_free_slot(pallet, slot)
        self.rack_cursor = {
            "pallet": pallet["number"],
            "slot": list(following) if following else [x, y],
        }

        self.entry.delete(0, tk.END)
        self._cancel_flash()
        where = [("rack_where", {"rack": pallet["name"], "x": x, "y": y})]

        if item["dup"]:
            self.set_status("duplicate_warning", serial=value, _details=where)
            self.play_duplicate_sound()
        elif following is None:
            self.set_status(
                "rack_completed", rack=pallet["name"],
                total=len(self.rack_slots(pallet)),
            )
            self.play_duplicate_sound()
        else:
            self.set_status(
                "rack_added", serial=value, rack=pallet["name"], x=x, y=y,
            )
            self.play_sound(True)

        self.refresh_create_view(focus_index=len(self.create_items) - 1)
        self.schedule_save()
        self.focus_entry()

    def set_rack_order(self, order):
        self.rack_order = order
        self._refresh_order_buttons()
        pallet = self.target_rack()

        # El orden se aplica al rack activo mientras siga vacío.
        if pallet and not self.rack_used(pallet["number"]):
            pallet["order"] = order

        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()

    def _refresh_order_buttons(self):
        for order, button in self.order_buttons.items():
            button.config(
                variant="primary" if order == self.rack_order else "ghost",
            )

    def on_rack_toggle(self):
        if self.rack_mode_var.get():
            self.validate_var.set(False)

        self.update_mode_view()
        self.schedule_save()
        self.focus_entry()

    def rack_full_pending(self):
        """Rack activo completo (y su aviso no se ha ocultado)."""
        if not self.rack_mode_var.get():
            return None

        pallet, slot = self.target_slot()

        if pallet and slot is None and pallet["number"] not in self.rack_dismissed:
            return pallet

        return None

    def hide_rack_banner(self):
        pallet = self.target_rack()

        if pallet:
            self.rack_dismissed.add(pallet["number"])

        self.refresh_banner()
        self.focus_entry()

    # ----- deshacer y eliminar -----

    def _push_undo(self, action):
        self.undo_stack.append(action)
        del self.undo_stack[:-50]

    def _index_of(self, item):
        for index, other in enumerate(self.create_items):
            if other is item:
                return index

        return -1

    def undo_last(self):
        if self.mode != "create":
            return

        if not self.undo_stack:
            self.set_status("nothing_to_undo")
            self.focus_entry()
            return

        action = self.undo_stack.pop()
        kind = action["type"]

        if kind == "add":
            index = self._index_of(action["item"])

            if index >= 0:
                del self.create_items[index]

            if action["item"].get("slot"):
                # En un rack, el siguiente escaneo vuelve a esa ubicación.
                self.rack_cursor = {
                    "pallet": action["item"]["pallet"],
                    "slot": list(action["item"]["slot"]),
                }

        elif kind == "delete":
            for index, item in sorted(action["entries"], key=lambda e: e[0]):
                self.create_items.insert(
                    min(index, len(self.create_items)), item,
                )

            self._restore_pallets(action["pallets"])

        elif kind == "edit":
            action["item"]["serial"] = action["old"]

        self._after_items_changed()
        self.set_status("undone")
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()

    def _delete_entries(self, entries, status_key, **params):
        snapshot = [dict(pallet) for pallet in self.pallets]

        for index, _item in sorted(entries, key=lambda e: e[0], reverse=True):
            del self.create_items[index]

        self._push_undo(
            {"type": "delete", "entries": entries, "pallets": snapshot}
        )
        self._after_items_changed()
        self.set_status(status_key, **params)
        self.refresh_create_view()
        self.schedule_save()
        self.focus_entry()

    def delete_selected_item(self):
        if self.mode != "create":
            return

        selected = self.sheet_grid.selected
        info = self.sheet_grid.info.get(selected) if selected else None

        if not info or info["kind"] not in ("item", "dup"):
            self.set_status("nothing_to_select")
            self.focus_entry()
            return

        index = info["index"]
        item = self.create_items[index]

        self._delete_entries(
            [(index, item)], "deleted_msg", serial=item["serial"],
        )

    def remove_duplicates(self):
        entries = [
            (index, item) for index, item in enumerate(self.create_items)
            if item["dup"]
        ]

        if not entries:
            self.refresh_banner()
            return

        self._delete_entries(entries, "dups_removed", count=len(entries))

    def hide_dup_banner(self):
        self.dup_dismissed_at = self.dup_count()
        self.refresh_banner()
        self.focus_entry()

    def on_grid_select(self, row, col, info):
        rack = self.rack_sheets().get(self.view_sheet)

        # En un rack, clic en una ubicación vacía: el siguiente escaneo va ahí.
        if (
            rack and info is None and self.rack_mode_var.get()
            and row <= rack["rows"] and col <= rack["cols"]
        ):
            pallet = self.rack_by_number(rack["pallet"])

            if pallet:
                self.rack_cursor = {"pallet": pallet["number"],
                                    "slot": [col - 1, row - 1]}
                self.set_status(
                    "rack_target", rack=pallet["name"], x=col - 1, y=row - 1,
                )
                self.refresh_create_view()
                self.schedule_save()

        self._update_cell_info()

    def on_grid_edit(self, row, col, text):
        """Edición directa de una celda del visor (doble clic)."""
        info = self.sheet_grid.info.get((row, col))

        if not info or info["index"] is None:
            return

        item = self.create_items[info["index"]]
        old = item["serial"]

        if text == old:
            return

        self._push_undo({"type": "edit", "item": item, "old": old})
        item["serial"] = text
        self._after_items_changed()

        if item["dup"]:
            self.set_status("duplicate_warning", serial=text)
            self.play_duplicate_sound()

        self.refresh_create_view()
        self.schedule_save()

    def cell_name(self, sheet, row, col):
        """"Lista!B3" o, en un rack, "DC05-1 (1,2)"."""
        if sheet in self.rack_sheets():
            return f"{sheet} ({col - 1},{row - 1})"

        return f"{sheet}!{column_letter(col)}{row}"

    def _update_cell_info(self):
        selected = self.sheet_grid.selected

        if not selected:
            self.cell_info_label.config(text="")
            return

        row, col = selected
        info = self.sheet_grid.info.get(selected)
        text = self.cell_name(self.view_sheet, row, col)

        if info:
            text += f":  {info['text']}"

        self.cell_info_label.config(text=text)

    def update_create_counts(self):
        if self.mode != "create":
            return

        if self.rack_mode_var.get():
            pallet, slot = self.target_slot()

            if pallet is None:
                self.counter_info.config(text=self.t("rack_none"))
            else:
                self.counter_info.config(
                    text=self.t(
                        "rack_count",
                        rack=pallet["name"],
                        used=len(self.rack_used(pallet["number"])),
                        total=len(self.rack_slots(pallet)),
                        duplicates=self.dup_count(),
                    )
                )

            return

        active = self.pallets[-1] if self.pallets else None
        text = self.t(
            "create_count",
            pallet=active["number"] if active else "-",
            pallet_count=self.pallet_count(active["number"]) if active else 0,
            count=len(self.create_seen),
            duplicates=self.dup_count(),
        )

        if self.validation_active():
            if self.valid_set:
                done = self.validated_count()
                text += "\n" + self.t(
                    "validate_count", done=done, total=len(self.valid_set),
                    left=len(self.valid_set) - done,
                )
            else:
                text += "\n" + self.t("validate_no_list")

        self.counter_info.config(text=text)

    def refresh_create_view(self, focus_index=None):
        """Recalcula el diseño y actualiza el visor de la hoja."""
        self.layout = compute_layout(self.create_items, self.pallets, self.default_sheet)
        sheets = self.layout["sheets"]
        positions = self.layout["positions"]

        names = [
            name for name, sheet in sheets.items()
            if sheet["cells"] or "rack" in sheet
        ]

        if not names:
            names = [self.default_sheet]

        if focus_index is not None and focus_index in positions:
            self.view_sheet = positions[focus_index][0]

        if self.view_sheet not in names:
            self.view_sheet = names[0]

        self._refresh_sheet_tabs(names)
        sheet = sheets.get(self.view_sheet) or {
            "cells": {}, "max_row": 0, "max_col": 0,
        }
        rack = sheet.get("rack")

        if self.view_sheet != self.drawn_sheet:
            self.sheet_grid.clear()
            self.drawn_sheet = self.view_sheet
            self.sheet_grid.configure_sheet(
                self.col_widths.setdefault(self.view_sheet, {}),
                zero_based=bool(rack),
            )

        if rack:
            self.sheet_grid.sync(
                sheet["cells"], rack["rows"], rack["cols"], exact=True,
            )
        else:
            self.sheet_grid.sync(
                sheet["cells"], sheet["max_row"], sheet["max_col"],
            )

        # Ubicación guiada del siguiente escaneo (Modo Rack).
        guide = None

        if self.rack_mode_var.get() and rack:
            pallet, slot = self.target_slot()

            if pallet and slot and pallet["number"] == rack["pallet"]:
                guide = (slot[1] + 1, slot[0] + 1)

        self.sheet_grid.set_guide(guide)

        if focus_index is not None and focus_index in positions:
            _name, row, col = positions[focus_index]
            self.sheet_grid.ensure_visible(row, col)
        elif guide:
            self.sheet_grid.ensure_visible(*guide)

        self._update_cell_info()
        self.update_counter()
        self.update_create_counts()
        self.refresh_banner()
        self.apply_counter_colors()

    def _refresh_sheet_tabs(self, names):
        """Pestañas de hojas bajo el visor (como en Excel) y botón "+"."""
        rack_mode = bool(self.rack_mode_var.get())
        state = (tuple(names), self.view_sheet, rack_mode, self.lang)

        if state == self.sheet_tab_names:
            return

        self.sheet_tab_names = state

        for button in self.sheet_tab_buttons:
            button.destroy()

        self.sheet_tab_buttons = []

        for name in names:
            button = RoundButton(
                self.sheet_tabs, text=name, size=9, height=26, padx=12,
                variant="primary" if name == self.view_sheet else "tab",
                command=lambda n=name: self.show_sheet(n),
            )
            button.pack(side="left", padx=(0, 4))
            button.bind(
                "<Double-Button-1>",
                lambda e, n=name, b=button: self.begin_rename_sheet(n, b),
            )
            Tooltip(button, lambda: self.t("tip_tab"))
            self.sheet_tab_buttons.append(button)

        add = RoundButton(
            self.sheet_tabs, icon="plus", size=9, height=26, padx=10,
            text=self.t("add_rack") if rack_mode else self.t("new_sheet"),
            variant="secondary", command=self.new_sheet,
        )
        add.pack(side="left", padx=(6, 0))
        self.sheet_tab_buttons.append(add)

    def begin_rename_sheet(self, name, button):
        """Doble clic en una pestaña: se escribe el nombre nuevo ahí mismo."""
        self._close_rename()
        entry = tk.Entry(
            self.sheet_tabs, font=ui(9), bg=FIELD, fg=WHITE,
            insertbackground=WHITE, relief="flat", bd=0,
            highlightthickness=1, highlightbackground=BLUE_LIGHT,
            highlightcolor=BLUE_LIGHT, width=max(len(name) + 4, 14),
        )
        entry.insert(0, name)
        entry.select_range(0, tk.END)
        entry.place(in_=button, x=0, y=0, relheight=1)
        entry.bind("<Return>", lambda e: self.finish_rename(name))
        entry.bind("<KP_Enter>", lambda e: self.finish_rename(name))
        entry.bind("<Escape>", lambda e: self._cancel_rename())
        entry.bind("<FocusOut>", lambda e: self.finish_rename(name))
        self._rename_entry = entry
        # El clic que soltó el doble clic devuelve el foco al escaneo: se
        # espera un momento antes de dárselo al nombre.
        self.root.after(
            80, lambda: entry.winfo_exists() and entry.focus_set(),
        )
        return "break"

    def _close_rename(self):
        entry = self._rename_entry
        self._rename_entry = None

        if entry is not None:
            try:
                entry.destroy()
            except tk.TclError:
                pass

    def _cancel_rename(self):
        self._close_rename()
        self.focus_entry()
        return "break"

    def finish_rename(self, old):
        if self._rename_entry is None:
            return "break"

        new = self._rename_entry.get()
        self._close_rename()
        self.rename_sheet(old, new)
        self.focus_entry()
        return "break"

    def rename_sheet(self, old, new):
        """Cambia el nombre de una hoja (se usa también al exportar)."""
        new = "".join(ch for ch in str(new) if ch not in "[]:*?/\\").strip()[:31]

        if not new or new == old:
            return False

        names = list(self.layout["sheets"]) or [self.default_sheet]

        if new.lower() in {name.lower() for name in names if name != old}:
            self.set_status("sheet_name_taken", name=new)
            return False

        if old == self.default_sheet:
            self.default_sheet = new
        else:
            for number, sheet in self.layout["pallet_sheets"].items():
                pallet = next(
                    (p for p in self.pallets if p["number"] == number), None,
                )

                if sheet == old and pallet and pallet["mode"] in ("sheet", "rack"):
                    pallet["name"] = new

        # Pallets que empiezan en una celda de esa hoja.
        for pallet in self.pallets:
            if pallet.get("sheet") == old:
                pallet["sheet"] = new

        if old in self.col_widths:
            self.col_widths[new] = self.col_widths.pop(old)

        if self.view_sheet == old:
            self.view_sheet = new
            self.drawn_sheet = new

        self.sheet_tab_names = None
        self.set_status("sheet_renamed", old=old, new=new)
        self.refresh_create_view()
        self.schedule_save()
        return True

    def show_sheet(self, name):
        self.view_sheet = name
        rack = self.rack_sheets().get(name)

        # Al abrir un rack, el siguiente escaneo continúa en ese rack.
        if rack and self.rack_mode_var.get():
            pallet = self.rack_by_number(rack["pallet"])
            slot = self.next_free_slot(pallet)
            self.rack_cursor = {
                "pallet": pallet["number"],
                "slot": list(slot) if slot else [0, 0],
            }

        self.refresh_create_view()

        if self._rename_entry is None:
            self.focus_entry()

    def on_columns_resized(self):
        self.schedule_save()

    # ----- exportar -----

    def build_create_workbook(self, only_sheet=None):
        """Libro de Excel de la lista (todas las hojas o solo una).

        Los racks se dibujan como cuadrícula con sus coordenadas y, si se
        exportan todos, se añade la hoja con la lista de ubicaciones.
        """
        layout = compute_layout(self.create_items, self.pallets, self.default_sheet)
        workbook = Workbook()
        blue_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")
        first = True

        for name, sheet in layout["sheets"].items():
            if only_sheet and name != only_sheet:
                continue

            if not sheet["cells"] and "rack" not in sheet:
                continue

            worksheet = workbook.active if first else workbook.create_sheet()
            worksheet.title = name
            first = False

            if "rack" in sheet:
                self._write_rack_sheet(worksheet, name, sheet, blue_fill)
                continue

            widths = {}

            for (row, col), info in sheet["cells"].items():
                cell = worksheet.cell(row=row, column=col, value=info["text"])
                # Siempre como texto: un valor que empiece con "=" no debe
                # convertirse en fórmula.
                cell.data_type = "s"

                if info["kind"] == "header":
                    cell.font = Font(bold=True)
                elif info["kind"] == "dup":
                    cell.fill = blue_fill

                widths[col] = max(widths.get(col, 0), len(info["text"]))

            custom = self.col_widths.get(name, {})

            for col, width in widths.items():
                if col in custom:
                    # Mismo ancho que en el visor (≈ 7 píxeles por carácter).
                    width = custom[col] / 7
                else:
                    width = min(max(width + 2, 6), 90)

                worksheet.column_dimensions[column_letter(col)].width = width

        racks = [p for p in self.pallets if p["mode"] == "rack"]

        if racks and not only_sheet:
            self._write_rack_list(workbook, layout, racks)

        return workbook

    def _write_rack_sheet(self, worksheet, name, sheet, blue_fill):
        """Rack como en la estantería: coordenadas x arriba, y a la izquierda."""
        from openpyxl.styles import Alignment, Border, Side

        rack = sheet["rack"]
        thin = Side(style="thin", color="7F8C9D")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill = PatternFill(fill_type="solid", fgColor="D9DEE7")
        center = Alignment(horizontal="center", vertical="center")

        corner = worksheet.cell(row=1, column=1, value=name)
        corner.font = Font(bold=True)

        for x in range(rack["cols"]):
            cell = worksheet.cell(row=1, column=x + 2, value=x)
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = center
            cell.border = border
            longest = max(
                [len(info["text"]) for (row, col), info in sheet["cells"].items()
                 if col == x + 1] or [6],
            )
            worksheet.column_dimensions[column_letter(x + 2)].width = longest + 2

        for y in range(rack["rows"]):
            cell = worksheet.cell(row=y + 2, column=1, value=y)
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = center
            cell.border = border

            for x in range(rack["cols"]):
                slot = worksheet.cell(row=y + 2, column=x + 2)
                slot.border = border
                slot.alignment = center
                info = sheet["cells"].get((y + 1, x + 1))

                if info:
                    slot.value = info["text"]
                    slot.data_type = "s"

                    if info["kind"] == "dup":
                        slot.fill = blue_fill

            worksheet.row_dimensions[y + 2].height = 28

        worksheet.column_dimensions["A"].width = max(6, len(name) + 2)

        # Sin las letras y números propios de Excel (ni en pantalla ni al
        # imprimir): solo se ven las coordenadas del rack, desde 0.
        worksheet.sheet_view.showRowColHeaders = False
        worksheet.print_options.headings = False
        worksheet.print_title_rows = "1:1"
        worksheet.print_title_cols = "A:A"
        worksheet.page_setup.orientation = "landscape"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    def _write_rack_list(self, workbook, layout, racks):
        """Hoja con todas las ubicaciones: Ubicación | Rack | (x,y) | S/N."""
        used = {name.lower() for name in workbook.sheetnames}
        title = safe_sheet_name(self.t("rack_list_sheet"), used)
        worksheet = workbook.create_sheet(title)
        headers = self.t("rack_list_headers").split("|")
        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        for pallet in racks:
            order = {slot: n for n, slot in enumerate(self.rack_slots(pallet))}
            entries = sorted(
                (item for item in self.create_items
                 if item["pallet"] == pallet["number"] and item.get("slot")),
                key=lambda item: order.get(tuple(item["slot"]), 0),
            )

            for item in entries:
                x, y = item["slot"]
                worksheet.append([
                    pallet.get("location", ""),
                    pallet["name"],
                    f"({x},{y})",
                    x,
                    y,
                    item["serial"],
                    self.t("yes") if item["dup"] else self.t("no"),
                ])

                for cell in worksheet[worksheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"

        autofit_columns(worksheet)

    def export_create_list(self):
        if not self.create_items:
            self.set_status("nothing_to_export")
            self.focus_entry()
            return

        path = filedialog.asksaveasfilename(
            title=self.t("export_list_title"),
            defaultextension=".xlsx",
            filetypes=[("Excel workbook", "*.xlsx"), ("CSV", "*.csv")],
        )

        if not path:
            self.focus_entry()
            return

        layout = compute_layout(self.create_items, self.pallets, self.default_sheet)
        duplicates = self.dup_count()

        try:
            if path.lower().endswith(".csv"):
                headers = self.t("csv_headers").split("|")
                racks = set(self.rack_sheets())

                with open(path, "w", newline="", encoding="utf-8-sig") as file:
                    writer = csv.writer(file)
                    writer.writerow(headers)

                    for index, item in enumerate(self.create_items):
                        position = layout["positions"].get(index)
                        sheet, row, col = position if position else ("", 0, 0)

                        if not position:
                            cell = ""
                        elif sheet in racks:
                            cell = f"({col - 1},{row - 1})"
                        else:
                            cell = f"{column_letter(col)}{row}"

                        writer.writerow([
                            item["serial"],
                            item["pallet"],
                            sheet,
                            cell,
                            self.t("yes") if item["dup"] else self.t("no"),
                        ])
            else:
                if Workbook is None or Font is None or PatternFill is None:
                    self.set_status("export_dependency")
                    self.focus_entry()
                    return

                workbook = self.build_create_workbook()

                try:
                    workbook.save(path)
                finally:
                    workbook.close()

            self.set_status(
                "list_exported",
                path=path,
                count=len(self.create_items),
                duplicates=duplicates,
            )
        except Exception as exc:
            self.set_status("export_error", error=str(exc)[:200])

        self.focus_entry()

    # ========================================================
    # ATAJOS DE TECLADO
    # ========================================================

    def on_next_pallet_key(self, event=None):
        if self.mode == "create":
            self.next_pallet()

        return "break"

    def on_undo_key(self, event=None):
        if self.mode == "checker":
            self.undo_scan()
            return "break"

        if self.sheet_grid.editor is None:
            self.undo_last()
            return "break"

        return None

    def on_export_key(self, event=None):
        self.export_current()
        return "break"

    # ========================================================
    # REINICIAR (flechas en círculo)
    # ========================================================

    def disarm_reset(self):
        self.reset_armed = False

        if self.reset_job is not None:
            try:
                self.root.after_cancel(self.reset_job)
            except Exception:
                pass

            self.reset_job = None

        self.reset_button.config(
            text=self.t("reset_label"), icon="reset", variant="secondary",
        )

    def _clear_checker_state(self):
        self._cancel_flash()
        self.serials = set()
        self.serial_order = []
        self.serial_lengths = set()
        self.found = set()
        self.locations = {}
        self.extras = OrderedDict()
        self.scan_log = []
        self.checker_undo = []
        self.csv_path = None
        self.loaded_count = 0
        self.resumed = False
        self.pending_resume = set()
        self.pending_iids = {}
        self.history = []
        self.standby = False

    def _clear_create_state(self):
        self.rack_cursor = None
        self.rack_dismissed = set()
        self.col_widths = {}
        self.create_items = []
        self.create_seen = set()
        self.pallets = []
        self.pallet_counter = 0
        self.undo_stack = []
        self.default_sheet = DEFAULT_SHEET
        self.view_sheet = DEFAULT_SHEET
        self.dup_dismissed_at = 0
        self.valid_order = []
        self.valid_set = set()
        self.valid_path = None

    def reset_app(self):
        """Reinicia solo la sesión del modo activo.

        Si hay datos, el primer clic pide confirmación en el mismo botón
        ("¿Seguro?") y el segundo clic, dentro de 3 segundos, reinicia.
        """
        if self.mode == "checker":
            has_data = bool(self.csv_path)
        else:
            has_data = bool(self.create_items)

        if has_data and not self.reset_armed:
            self.reset_armed = True
            self.reset_button.config(
                text=self.t("reset_arm"), icon=None, variant="danger",
            )
            self.reset_job = self.root.after(RESET_ARM_MS, self.disarm_reset)
            return

        self.disarm_reset()
        self.entry.delete(0, tk.END)

        if self.mode == "checker":
            self._clear_checker_state()
            self.refresh_file_info()
            self.update_counter()
            self.refresh_history()
            self.rebuild_pending_tree()
            self.restore_colors()
        else:
            self._clear_create_state()
            self.sheet_grid.clear()
            self.drawn_sheet = None
            self.pallet_name_var.set("")
            self.refresh_create_view()

        self.set_status("reset_done")
        self.refresh_banner()
        self.apply_counter_colors()
        self.schedule_save()
        self.focus_entry()

    # ========================================================
    # SESIONES (se guardan y se recuperan al abrir la app)
    # ========================================================

    def schedule_save(self):
        if self.save_job is not None:
            try:
                self.root.after_cancel(self.save_job)
            except Exception:
                pass

        self.save_job = self.root.after(SAVE_DELAY_MS, self.save_sessions)

    def _session_payload(self):
        return {
            "version": SESSION_VERSION,
            "active_mode": self.mode,
            "alt_matching": bool(self.alt_matching_var.get()),
            "language": self.lang,
            "dest_index": self.dest_index,
            "view_sheet": self.view_sheet,
            "checker": {
                "csv_path": self.csv_path,
                "loaded_count": self.loaded_count,
                "serials": self.serial_order,
                "found": sorted(self.found),
                "resumed": self.resumed,
                "locations": {
                    serial: list(where)
                    for serial, where in self.locations.items()
                },
                "extras": [
                    [serial, entry["count"], entry["first"]]
                    for serial, entry in self.extras.items()
                ],
                "scan_log": self.scan_log,
            },
            "create": {
                "items": self.create_items,
                "pallets": self.pallets,
                "pallet_counter": self.pallet_counter,
                "rack_mode": bool(self.rack_mode_var.get()),
                "rack_location": self.rack_location_var.get(),
                "rack_rows": self.rack_rows_var.get(),
                "rack_cols": self.rack_cols_var.get(),
                "rack_order": self.rack_order,
                "rack_cursor": self.rack_cursor,
                "default_sheet": self.default_sheet,
                "validate": bool(self.validate_var.get()),
                "valid_alt": bool(self.valid_alt_var.get()),
                "valid_path": self.valid_path,
                "valid_list": self.valid_order,
                "col_widths": {
                    sheet: {str(col): width for col, width in widths.items()}
                    for sheet, widths in self.col_widths.items() if widths
                },
            },
        }

    def save_sessions(self):
        self.save_job = None
        temp_path = self.session_file + ".tmp"

        try:
            with open(temp_path, "w", encoding="utf-8") as file:
                json.dump(self._session_payload(), file, ensure_ascii=False)

            os.replace(temp_path, self.session_file)
            self.show_saved()
        except Exception as exc:
            self.save_label.config(
                text=self.t("session_save_error", error=str(exc)[:80]),
                fg=YELLOW,
            )

    def show_saved(self):
        if self.save_label_job is not None:
            try:
                self.root.after_cancel(self.save_label_job)
            except Exception:
                pass

        self.save_label.config(text=self.t("session_saved"), fg="#737373")
        self.save_label_job = self.root.after(2000, self.clear_saved)

    def clear_saved(self):
        self.save_label_job = None
        self.save_label.config(text="")

    def load_sessions(self):
        if not os.path.exists(self.session_file):
            return

        try:
            with open(self.session_file, "r", encoding="utf-8") as file:
                payload = json.load(file)

            if (
                not isinstance(payload, dict)
                or payload.get("version") != SESSION_VERSION
            ):
                return

            language = payload.get("language")

            if language in LANGUAGES.values():
                self.lang = language
                self.language_var.set(
                    next(n for n, code in LANGUAGES.items() if code == language)
                )

            self.alt_matching_var.set(bool(payload.get("alt_matching", True)))

            if payload.get("active_mode") in ("checker", "create"):
                self.mode = payload["active_mode"]

            if payload.get("dest_index") in (0, 1, 2):
                self.dest_index = payload["dest_index"]

            view = payload.get("view_sheet")

            if isinstance(view, str) and view:
                self.view_sheet = view

            # --- Checker ---
            checker = payload.get("checker") or {}
            order = [
                s for s in checker.get("serials", [])
                if isinstance(s, str) and s
            ]
            self.serial_order = list(OrderedDict.fromkeys(order))
            self.serials = set(self.serial_order)
            self.found = {
                s for s in checker.get("found", []) if s in self.serials
            }
            path = checker.get("csv_path")
            self.csv_path = path if isinstance(path, str) and self.serials else None
            self.loaded_count = int(checker.get("loaded_count") or len(self.serials))
            self.resumed = bool(checker.get("resumed"))
            self.serial_lengths = {len(s) for s in self.serial_order}

            # Campos añadidos después (una sesión antigua no los tiene).
            self.locations = {}

            for serial, where in (checker.get("locations") or {}).items():
                if (
                    serial in self.serials
                    and isinstance(where, list) and len(where) == 2
                    and isinstance(where[1], int)
                ):
                    self.locations[serial] = (where[0], where[1])

            self.extras = OrderedDict()

            for entry in checker.get("extras") or []:
                if (
                    isinstance(entry, list) and len(entry) == 3
                    and isinstance(entry[0], str)
                    and isinstance(entry[1], int)
                    and isinstance(entry[2], (int, float))
                ):
                    self.extras[entry[0]] = {
                        "count": entry[1], "first": entry[2],
                    }

            self.scan_log = [
                entry for entry in checker.get("scan_log") or []
                if isinstance(entry, list) and len(entry) == 4
                and isinstance(entry[0], (int, float))
                and entry[2] in self.HISTORY_COLORS
            ][-SCAN_LOG_LIMIT:]

            # --- Crear lista ---
            create = payload.get("create") or {}
            modes = ("column", "sheet", "cell", "rack")
            pallets = []

            def size(value, default):
                return value if isinstance(value, int) and 1 <= value <= RACK_MAX else default

            for entry in create.get("pallets", []):
                if (
                    isinstance(entry, dict)
                    and isinstance(entry.get("number"), int)
                    and entry.get("mode") in modes
                ):
                    pallet = {
                        "number": entry["number"],
                        "mode": entry["mode"],
                        "name": str(entry.get("name") or ""),
                        "sheet": str(entry.get("sheet") or ""),
                        "cell": str(entry.get("cell") or "A1"),
                    }

                    if entry["mode"] == "rack":
                        pallet.update({
                            "location": str(entry.get("location") or ""),
                            "rack": entry.get("rack") if isinstance(entry.get("rack"), int) else 0,
                            "rows": size(entry.get("rows"), RACK_ROWS),
                            "cols": size(entry.get("cols"), RACK_COLS),
                            "order": "v" if entry.get("order") == "v" else "h",
                        })

                    pallets.append(pallet)

            items = []

            for entry in create.get("items", []):
                if (
                    isinstance(entry, dict)
                    and isinstance(entry.get("serial"), str)
                    and entry["serial"]
                    and isinstance(entry.get("pallet"), int)
                ):
                    item = {
                        "serial": entry["serial"],
                        "pallet": entry["pallet"],
                        "dup": False,
                    }
                    slot = entry.get("slot")

                    if (
                        isinstance(slot, list) and len(slot) == 2
                        and all(isinstance(v, int) and v >= 0 for v in slot)
                    ):
                        item["slot"] = slot

                    items.append(item)

            known = {pallet["number"] for pallet in pallets}

            for item in items:
                if item["pallet"] not in known:
                    known.add(item["pallet"])
                    pallets.append({
                        "number": item["pallet"], "mode": "column",
                        "name": "", "sheet": "", "cell": "A1",
                    })

            pallets.sort(key=lambda pallet: pallet["number"])
            self.pallets = pallets
            self.create_items = items
            self._recompute_dups()

            self.rack_mode_var.set(bool(create.get("rack_mode")))
            default = create.get("default_sheet")

            if isinstance(default, str) and default.strip():
                self.default_sheet = default.strip()[:31]
            else:
                # Sesión de la 2.0: la hoja principal se llamaba "Lista".
                for pallet in self.pallets:
                    if pallet.get("sheet") == OLD_DEFAULT_SHEET:
                        pallet["sheet"] = DEFAULT_SHEET

                if self.view_sheet == OLD_DEFAULT_SHEET:
                    self.view_sheet = DEFAULT_SHEET

            self.validate_var.set(bool(create.get("validate")))
            self.valid_alt_var.set(bool(create.get("valid_alt", True)))
            self.valid_order = [
                str(v) for v in create.get("valid_list") or []
                if isinstance(v, str) and v
            ]
            self.valid_set = set(self.valid_order)
            path = create.get("valid_path")
            self.valid_path = path if isinstance(path, str) and self.valid_set else None
            self.rack_location_var.set(str(create.get("rack_location") or ""))
            self.rack_rows_var.set(str(create.get("rack_rows") or RACK_ROWS))
            self.rack_cols_var.set(str(create.get("rack_cols") or RACK_COLS))
            self.rack_order = "v" if create.get("rack_order") == "v" else "h"
            cursor = create.get("rack_cursor")

            if (
                isinstance(cursor, dict) and isinstance(cursor.get("pallet"), int)
                and isinstance(cursor.get("slot"), list) and len(cursor["slot"]) == 2
            ):
                self.rack_cursor = {"pallet": cursor["pallet"], "slot": cursor["slot"]}

            self.col_widths = {}

            for sheet, widths in (create.get("col_widths") or {}).items():
                if isinstance(widths, dict):
                    if sheet == OLD_DEFAULT_SHEET and "default_sheet" not in create:
                        sheet = DEFAULT_SHEET

                    self.col_widths[str(sheet)] = {
                        int(col): int(width) for col, width in widths.items()
                        if str(col).isdigit() and isinstance(width, int)
                    }
            self.pallet_counter = max(
                int(create.get("pallet_counter") or 0),
                max(known) if known else 0,
            )

            if self.mode == "checker" and self.serials:
                self.set_status("checker_session_restored")
            elif self.mode == "create" and self.create_items:
                self.set_status("create_session_restored")

        except Exception:
            # Una sesión dañada no debe impedir iniciar la aplicación.
            self._clear_checker_state()
            self._clear_create_state()

    def close_app(self):
        for job in (
            self.save_job, self.save_label_job, self.reset_job, self.flash_job,
        ):
            if job is not None:
                try:
                    self.root.after_cancel(job)
                except Exception:
                    pass

        self.save_sessions()
        self.root.destroy()

    # ========================================================
    # SONIDOS
    # ========================================================

    def play_duplicate_sound(self):
        if winsound is None or self.batch:
            return

        try:
            winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
        except Exception:
            pass

    def play_sound(self, positive):
        if winsound is None or self.batch:
            return

        try:
            sound = (
                winsound.MB_ICONASTERISK if positive
                else winsound.MB_ICONHAND
            )
            winsound.MessageBeep(sound)
        except Exception:
            pass


# ============================================================
# WINDOWS: NITIDEZ EN PANTALLAS CON ESCALA Y BARRA DE TÍTULO OSCURA
# ============================================================

def prepare_windows_process():
    """Nitidez con escala 125 %/150 % y app propia en la barra de tareas.

    Sin un AppUserModelID propio, Windows puede agrupar la ventana como
    genérica y mostrar el símbolo de ventana en vez del icono.
    """
    if os.name != "nt":
        return

    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
            "Desblack.SerialNumberChecker.2",
        )
    except Exception:
        pass

    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass


def use_dark_title_bar(root):
    """Barra de título oscura en Windows 10/11 (si el sistema lo permite)."""
    if os.name != "nt":
        return

    try:
        import ctypes
        root.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(root.winfo_id())
        value = ctypes.c_int(1)

        # 20 = DWMWA_USE_IMMERSIVE_DARK_MODE (19 en Windows 10 antiguos).
        for attribute in (20, 19):
            result = ctypes.windll.dwmapi.DwmSetWindowAttribute(
                hwnd, attribute, ctypes.byref(value), ctypes.sizeof(value),
            )

            if result == 0:
                break
    except Exception:
        pass


# Patrón de fondo (PNG en escala de grises, 665×335, se repite en mosaico).
BACKGROUND_TILE = (
    "iVBORw0KGgoAAAANSUhEUgAAApkAAAFPAgMAAAB7Tj1uAAAADFBMVEUWFhYLCwsCAgIAAAD6Bmw8"
    "AACJm0lEQVR42qT9f5Tc1n0niH5uG0CP3Y4bHQFtkfV2+DyRLSeSLXve2lqzely3vM3MjunnAiNA"
    "iUznmDOm3ySSoiJn1553IodseeyZfaM5j+Us7WRiZqJZS04EQAIqT8yaVLNwy+ninJecM26NpWQc"
    "+2yYfdWiiQu60D5NR4UL4vv+QLVEKqbH2elzLutHN1G3gIt7v/fz+X4+XxAfUEBEPh9EKb+QJERR"
    "ygbF1Eh83/eLHSNIrxlBMV6gwdONSTpuUDBuUNQ7XMidxkSOGxM5XqBoszEhCigY+1HFqSAiqniV"
    "VCgSIqJzelFEbSIinwaMigEL0hNGkA6MrpzO+/503i+m8xRcakTBpUYSXboniTYbfkJE/maXhAbn"
    "37dXA3jey30Hzh/2EEC3W0P2dpMDQDa1ufVFm49eMOW7pqvlSy/wzAq5XIwbePGF1fLFF5x16wVT"
    "mmurJeDC+rUHTDhHHrcBDJ3jPXI+/VkA0FfOlj89AgBQUy+1e3XL/qf/L8u+03Q0abFOalln32cp"
    "K3bhftGVdsw5DjkeAnhe+4MAaDxPREQ0blBwaV8SEAW02Y16H6uIiKJeVya8KxMepXIzSgftyB+w"
    "KL3OgmLAozRpR0HCo/Q6SyQR0cDwib7UkERE0bhBctwoiIiS3j75SoOIyK8uNYJqs5HQ5r6ENrsF"
    "9bpR1esmUa+bVu1EVWyQVvqACt2nlCgq9AEJ0JRFEZHvTw2KduYTn8gPxssDHhW+7/uDdjKpWDK5"
    "zoiqNslzBvnTOSrGRhVNdZLX68d0bFBARDSdj4iu6AkREU2NKtgxqvo0GLR5gogoqHrdoOJdqnon"
    "qeIni6o9SAp+MiraAzk1Ktpt4/ki9YmCsVGQAA14RL5f+c/yaFKxhHy/8qdsqtdjbGok0c58Eo3n"
    "ic4tEI0bRXJpgajXqKLxXqJxI6XxXqLNvVVERMVml4gKfoKIqKjagyBlCRFROmVVu75CBRtQwQaU"
    "6gNK6z4lV/RE7hhVNH5rkYwbBY0PU7LZoIh8v+odliTmwJx1eB7zmBOaw1YOz2PeqCWXAQCQtjS/"
    "e7s0s9szyDtihG7O4zsyxEeBzIsRugKhJ+mYFwJALDqA0J0MAMA6fWvUigEAeuv5i1cBwJH6CqTO"
    "8fzyCp63y1jaDOvLdyG3mRm6QN1yOubmDjwP4qgANGTueycOhNd837+m3P19BTB323lYEzwA8DWX"
    "NX/+CGv+vAM645Sq/xtQcY/U8MQzbvvEVdX/jXep+CTU8ORdJkA4+wiQw/17HID9HvdeB+6/agLg"
    "zPm4BSAAUvuaGdnZS5PV7Rcnq7kZuphIF9/6TRfUP3F2sX/i7EfPnj374eFJoZSO5I+eDQHQy1N4"
    "gB0VU6NIrzEPWD4vd9D1yfM8jw1oyq/T1CjkdE7KK8a0mBpVOjWKoGDn5RWjoCmrii0joYKo8Bsp"
    "pQFNDd/3/a6czu8+p3SMRwZEvu873ahwul8Z9JJo0CuiileJ71A6QJVM9SqazhfRdL44/8q8nzzF"
    "AW0+LUhg4PscAB4ZJHzgFwDAfDnVkx9OAcCo5LWFZyfnGnJyrvFstbOPJl/fR+fH71bVuBHQzs+p"
    "ZLxXyXEjiTYB4BMU9XyqOABUScWrpOAAdElPw69OcgCsoKpNVBlEVxeIpgtE1xeIvt4gGjeIvv5u"
    "op13Ez3XfdbfBICH0mQT2se/1JQ2IK8id/qPCw9AOGrrHs+mzAV0bLwNP/rHxOprL7SJeQwAzFzF"
    "33E0VwA43hPO8V5gAlCf/6ztAb/+GAANkMMffUjbvfFViWMegFW0PgTtqml+ygb23/lk6N7by58O"
    "geY3wT5Fn3/bAw6go/78v/nDdBfAKgALgFO2hh4AZ71z/wL69w9hybCXu/f22AM6oOIV63eeH75t"
    "uQUYuSX1H31M3bnx1aHq8FMAc+KfH0GjXCttjqIQrUUJxjxg6zNPgMvO57oOAFRi4UceU3zQEfVh"
    "OQCyL+s+AMTSzkawfYh3vI2cQpLm3A9svV/LTW5/8YEeImbTZwHkDP+FH23L8AAE9l/a79D2r3di"
    "PYDVzM3n9gtQX6GjAHFXc5gggKPfadziMKzVB5YAmACA0AYAaYben8U9YoDdFCt6s/4u9mVTxo46"
    "lgBLOVRfr6//63+GrZvOxAsuAgBYb799zpmAOZ637YQ6c57gMD1v1Iolt6URk+fFIrNv9X1LdcOL"
    "VnwUAETrmN0LEQK51umDdfqQgBq1yDuIicEBlMh0/qOPSOKmV/GRzPOY14pdPueG6y4AuAGa7joA"
    "YNt5GEB4O0pA58c6tzomcwECkANoq2ErCIKAS/09lgSCQJVuOMrdkEQQxFedM9BiZSMIwEVkMUD8"
    "ly47miMOwIUaOdBsyqFiOITnOeVmqAdgRx+1AxYfMRHAU+j8+x99mFwjAASU4j0wYazoAMp078p/"
    "VPAACjtZfLedlTaHeM+1f94zka2eOQqhydL9t/gJfmgudAN4wf7l+AXNeNRC3wOF+zcOxO/7Yy83"
    "EW//kBw6fILDgyouvhcA8K33Aq/ejasvQ+MhZneDxnPES/chgBPj5cdvO9CSIwVA/8CpYesDkAE8"
    "/EllYvG9FQI4sH7x+79d330AIKy6QQO2D87eAwDx278C8qDcX3AV18j85VMMYN7Lfwk4/QweTLv1"
    "GH+7yQN4YLlp3uI28m6+/Z3cA7zkyw8xco7/ng5AOMd7wjl+0kYAa+XsSlMHPEjrLv1WZ9C88d5a"
    "jD/oEqBju3P3n82tH83sVQBwz4jSDe0DAHQnR7jKPA/4/zx84JbX5cYR1vm0KwBAjTQwvQ49cjdE"
    "7sZmHZIcg7QBYIi/vuWdyW+cq5pDjgyI1Yib0PgbFJdAgI6CW0iYQDN2//6pxx4CACz0n6Vb9JLf"
    "OANengMLwFw5ZwLusY0MgFdIcgqpBQD42z9jfdEFAhxA7P4koxNS50p4wO/YJaBp4oMPccCTo/0x"
    "H+3PswRuaU/U6CQCAB/WnfSW89KNB7VKS0Kj7HZObKnZ7wFQelOs6M0h4EHaio5fAABT1Y//pR9+"
    "xtte8AK4H3K3F4XGtns//yUA9rbzp9h2YHmAm6zIOU4egC0dt1iLoW46u0f7LgAVH0E54r0P92bx"
    "5wrrxL1Z/KmvAB7C/frKtVsd8uwNC2eciA48EBzRgibKOxAShOds/Us6mH5GAcPT7T/7uAsEAL58"
    "lG71hSc3/uZsN1cx7HeMBshzWFmwG3++x723jj97B+yzDgIgsre1Wxwyv+ESxYwjPe4guNjqAxq1"
    "jNQDdL1p5qFjfu1rgNF575FfOxG27geAE9B/gsuOuaaeAsHPzZUjyWCrAMDeu+0svtvKAMCWd8h9"
    "AIBAuRv8J7gzYUh7/WtfAxYyJwVO+L4DAN0q4oOIAICdU9NlKnbq+DPdaZCc7dqL8eFJcOkwRb3D"
    "FPUCCsbdNNrcfQ0AXaJe9yt1/BlUFa+CigNgVTTg1WAAACxIr88HdG6+K6fzXfnKfBSMF6gYNxI5"
    "bkTB5j1JutmlyWY3CTYBsIeSSQ/a755uKh1QYQ9O38k9APKP3msegJ7aHNBhf2k1s764mlkhz6zQ"
    "XLdiK1uM78kW44dghRhaa/tmrz0APPtm/J3H4dpAELvCOf5o/fz42jftEvcyAizL3rI03GU5mrQc"
    "LbNIy8xSz1r5i0sdZa9xjvvvgXmoxbNPeQBbzRZiYD4pqohoPE+T8UKSEhGdbqQUFUmvK4miatCO"
    "gopH/oBH/qCd+IN2kl5nSXCdRTRoR1H9PEmvM9/3ffKvMKOKUt/3vQZF40aRpL7v9xqUBCltLnve"
    "0x9LaPPBhHrd4rX2SErOwCc+UKQPqGAJTfUBTXWfiCi9pgN6VBERTfWEpkbdz+k8ERUXWERERNcN"
    "CqYs8aeM/HNGFUwNCsbzVTQ2ompqVNHUqKKxQcHYICKiYLzMBhER0dN6QlOjSogoumJUfhRUvS5R"
    "FVHFB1TwE0XVHhRVeyCn7UExZYN0ygY0na9oyyAaG9UMAknp9F4+d3AdgQBGzRij/bkNAKP9Iotx"
    "r44QULiskyZtrstlZd9ll660pZtZpRnaZint0rxsl2ZoS7eOP5UVPtDqOwTgUFMovRAcAJaVQKmV"
    "xzsAA7IRx2TEdVW3jXK0oku9bUudQVoMgQdk3rYIrRwgIUKPz61NoHOg72TYdsLZtjtM1NlMNzUg"
    "FtnRzMweCJ3wl3IeujmF7pCOHc1b8VGO0EUzPpq34qNDio8GQRDkzdjshQiDINQdCdbrQwaBg07s"
    "3edInQOixPPLJcLlu2Npl7G0S0zsHKGdI7TfIx52kcQuaM0Vk/j0OvCMR8Men7PCdQcA955XH3Gf"
    "r1cCN2476thRDkDnxxbN5poDFZtQfWdJHe9wNWwxNWwh73dyOs6ZGra4msXik4t33SHriWU3/gQA"
    "9xkAgc2CIF8Vk9VtUhxm6ORm6OqLD7ss6bsg4ZofvPYo1MghdfagqYYmBzTcpQOYoU1FYRTFVI+I"
    "iOSOQeSzAUVEVLALdEWvaMqqYmpUdM2Y0pZR0JZRUGpMg6lRvfqaiNR4n7zOaAYvdeUr82lEFNB0"
    "Pi2e7X0iIkpo0IuiivtFxQu/4gVd5xRd0FU1ZVRN59Pq3DwV4/m02pkPiMj3xw1/TdMffwnwoOPX"
    "NzK9uZgwDguKvndQL+EAsnjT98N/tMzCD+xl/Q/Y2+qP7B9s/JWnx3/15NmPFPYPN97nsfivPD3+"
    "q6MbbYT3/U9vW7mtf+rYe44A0Pev3IYPWCbgYjSFZt5XAcik+J+fcouL97Tf8h8+/aj6D2fLj/6H"
    "uLxzWZOFIeT7LTt+2ztFtvUrL2Vbv+ICcNkXj3qPzcU4U+9WzH9swekjB4DOcX2CAwDwUmhy7QyX"
    "VsilFTrrL8Uci7GVLcYwtdBZ58esrH7N20Cn7MROCS38CADAidnQOc4BgDtfyJo6AFj2P2UapOlo"
    "0uIjaZEmrYNa7kg7vI97sQtnjXMeL3IeLwIAU32eAMFTRlIQEY3nKRgv1OjgeCHpHSYiooRHcsCj"
    "dMCjdNCOgkE7krPZkgY8ms2cSXCdBURU+FODfNoypkRE9KUGReMGERHJcUNeahAR1bhn78GCet0i"
    "mc2gyaFBWvFBQe2ECjag6awRUbRlVMEa/CmrZ9DrLEl3Z9AdY8CigojoHKPJlFExZVRMjSo6ZxDV"
    "s2VEU72KrhkUjI0qGhsRERWDcaMKdj+BrtQz6O6Y73WJiKjqnaSCnywqfrIo6tkzLdggnbJE1bhn"
    "Vc+e9f9MegsUiDmXddZne+cco/25CQB/0nzZMHMAuFMXptwDXS5Dl3bmfMcG6tnSgVzOnO/Y0g3t"
    "0gxtE0DeDt3Mgt7qAwCWmkLpTQEAlqHE8Xrnmn26tTt7btezJx8qfduWe/INuVSKzGIi9LZF6G0D"
    "ABdu7mAO271JvRt+IsS2E+oAsN0JLT4EgHyPROhlCH8pQ+gKlG6IerYEMldAHB1SPJtNgSH1O4KX"
    "rBcCgNAdCdbpA4A42/L1FQDA88srCJd5LJf5xsRmmNhMm9ilCPUSoQs6dhqT+KiYxEcFAFT5EUHA"
    "QG4Z52dXplCvLn7o1osfPxFU/ERa8JNp0R7Igg2CKRvIKRsQHRrIgiVy93VRr49RWtCOQUQU0Xg5"
    "mdQDNIk2rcaEiCg5/QlZ9Q5T0usWA6dLA6dLldMNBu2uHLSTaKoTPc3qRkS0ZVC0tYZBms5W+6kx"
    "KKZ1n9Mddn2QEBEZUyrYuXTKqnRqTGlqFHRtd/bUp3TD7JkURNF4IU22iqSeQQOaGl05na+fX0M3"
    "IiIa9JKo4kFVcSoGvEqu8yq5DirO6ZIu6FX0ynxBO/Mp7cynNbC/UCRbHCf95NCA0oiI2gX5rKA0"
    "iSi9n6rIT4juj5KtRlFsNYpk63CVjBspPdUooqcakuhjVTRuBNTbN6HeQ2lKRM/2iArncFE5h4uI"
    "0mLQDiht+1RQSoN2fX6eNgK6YlT+uXmiqRFUTxkBPTkf0bgR0OATVIwPV/5Wl4qtwwlVAdGXugF9"
    "ic/B4zlgfu35eqfPNeDLCwLx89nBhyXgmtyUyE2pQTNxtwlkLQAtDhSsXAEQOxyZA9gBAB3QuZsD"
    "bt75mg2QFoLcmCkzzgGI3wJsx8WLYJ4GAIcILiPS6SPESiL2b5QG0eeU1YjQ5Q3AcTQ0MVdDL5qM"
    "ARLgJGCGb8hBOityDtg8I+iYgGFpFZQTM0uurWQApAMTHNAA0gAhgeseICDYkAsGKa7VUN6w1KRe"
    "AkArAvgN+ISuYQg4IRiFAk7Y/FbOMoKdIdQzABKAzk0IMXcG5PTBzG3AeQam84xA6xuAx7XMNgEO"
    "K7MwCl0B3YH1MCMtNtlta5ZGD1vQQpOboWma8SpDHyAboDjXgFxjiE3GRWvkxBx0rHXIA/5aAFfv"
    "AzIIvNMChi6HtGA3LVhNDfrLQ0A4wNptGsRtUMdMwCauYsCgajyf0JMLROP5iMbzUTVuBEVECe8S"
    "ERUD3k0HPJgkRNWgHaXVgYiu6xFVjIpBO0grFtF1PQquG0RFQjTFvEpOz6vk9FuL5HSjSk43Cho3"
    "Kvodokt7iYKKil4joV6XKKKCet2Uevtmj12q2CBK2SAp2CC5og+IqKArOuZUDJsy8hRgUwmbwMwM"
    "MGWLnQoBxDTsgYbuOgcYCceUQ8ThnElKB5SAqXSokpn4ng1oEpB2leBQleTe9wUOSUZLUgOTwCEg"
    "rq95HLtQsQmUpKu4w+j4722r450RHT+Iy3obz+ucS53n4fIKIIBQ53Otdej7qRztB/QmQW+W0Ato"
    "uVA1a+pd1qHfqdfo750GoLxFSIsjsOVZ+SYX0nQgPQ7pAowDodd8Dnuaz2GpGXO7uYbblSjtpsgA"
    "HHcAt08YOpiMOOAA8iKzLusrljS4eVnnQtr5orQZIi9TapUBOEvK43OncwwdxMwB+h1guwOMWmCm"
    "J+3hKgDcvScS2Z4AFAQiu11q2R0c5Sfr1SizgOyBEuHPQBOdGh+M7R7jrMcc3QH0xZyzVkx6RwgA"
    "dWRWQm+hXG4Cb0BM+t0gu4/g9j4i+1ou3HUu3IzWLQFx8DmAgPgJDpovaGu+oK2FlMYLaU1YNmRa"
    "OXXUQPxENKn5yiphA0p1PyJjkCbGIE1YkiZsEJHuF4Wxu1qc35lPq9lMXYwXiuhSw6dxg6rgUqMg"
    "iqqod5iiTSKiJO32ulHCu0HCu5TwoBi0qRgYQTqY8/3dFWJqkJhDsZFZxUZmKStjysqYtDLI24bl"
    "dgcAsmLYbr5p2AYAtl9DrphZ/tDk2G/y4X69Ndyvr5Q/ZOZ2YToAMJQ4cLYwh2cLGzpZMVNaqUk3"
    "1jOEVuwAAFuMrWyxZntEJ+blB4ccrSEPW0NozVGpvYNZ9vs15kgLADRIE5hDa0h6a0h6UyR2UyR2"
    "E4mthCZHHABI6RxqhvvL5Xetk7u6OlllkKtMk3apSRurE3dV5Ad0ANBCD3rrbK63RDzaP2R7mgJ/"
    "oYblS0XCqO/mgIPW0BGtYQ0dNi86KHWHK30RSge7bMfILOC7tqaFbk2nhS4wB+eMOXTOmMNOH3qn"
    "D33xMa63+mZk1/jREzbMMzWwKsJV3PbwEhYfPgTr4UMwQ1c3Q3d78eEl3PawOwOBXbXkhCtLzhnO"
    "nDN8uNgHa8WctZ7RmiM+BECX9ZZ2WQcAbkqbtO/YsL5ir2iB7SJ2AeFyOnbULGuEl6vYBkB1lFtF"
    "48Zuo+p0g3rdOj+Bd4OKdwMiokE7Si6wqLrOErrOEhogiQZIoussqq7vYhR6ktLYKGhcRyjR6QYF"
    "lxpFNF6mcWMQEVHRaxRJr0FEVFCv6xPvRsS7KTldStsDKvQBpWyQ7GYbXNETEqAZ7hJMDdpt6dSo"
    "eBIQEV1gFFxgs1wFo6ienk+K6TxVVwyqpgZFU6OKpvNJsbNAfh3dUEBTdp6mbECFPoi2DEp3jCq4"
    "ZlS9rh8RUcpPphU/SUSkivbAL9ggKNhAVozoKaOirYWCnlookqcW5O7uhcQchq2YD1sxH+2PMdof"
    "Q28Kc6SCiy0NAC7roO/poDpnAGXTyvPvWEDgAdIuO9LOnO9YeX7GBQOA8Gii4Qet5/CDloDeFKXd"
    "7MJQoqU31493dAAQs90/gJxGvC11Dqlz67IOMbGvIrK2xcTankiH6juEA5jDdi+ztnuZ1XeE1neE"
    "NlzMddYK9MwBgMwWjtxTk3+ZezbPXCXKo2FW3hFmmZWrzBpSeVSJfF+9EsSLMIVwmBDONuLedmfU"
    "Ka2zrdg6S5HeNJtAQvqKlRkrAJBIPbPIzlyyQ5Cd5aX7nFp3Nqh0RCk6NdMSOwRgQP68TLbmn/Wf"
    "WnjWf2r+2Wi8MKAee6hIfN/3+YNJyk8k5Pu+z5Iq+VAhLxhVMTVkUrGkOMeK4rqRFNeNoAqIKDXO"
    "R1tEW/NEr8wT7cwTPTdP9OQ80ZNoTKKCKHrxMEW9OiOHet3o13tEEa/bBUZ0QSe6YBC9olNa+L7/"
    "tJHIa2vaLK4j871FiQMF4LwMwDu+ypoGAHYyWxyezIv7Aehg+ze2Z2g3QemtkutAaQH/G9xw//2A"
    "xAFBgFUAwwIYKeAAAJMAwFm3viIAeU+2GLtEvwggfLb87z/fRet+4MO/CazoQMlqqFrYMOX9AOZ5"
    "/D5oAG+aWFqBqVoOUwDsEpkN87n3A4DOpdJ5LAHgEMTbGVDaDMoEAruEMvv0YQC5DcgAwGEg7MJo"
    "AfqsafuBAyuA3V4tW20A7KT4heFJqACA7oTNIbDAAIwAMEDaQMsCMh1aAACuMh9+jwaY/A3QCfli"
    "iXwRXOmUL+o6F4HNAUPhCRtm6AGweBJawESLD8a6Ce7q4u1vhJ6v0og7uO9D9UFbj3UheB2TsBYw"
    "nHVBZ52+YgRYrf5lneNlAFgmLdPrfYDSAbh12kFpAdxEGXsAXDTb79HK2IFg4Dx0gefdFZa/BCAB"
    "8S/e/78wiNgN3fC2OEEAHRRrgpcAYG0w24lNmJqJJeRYhRr6AMx8UQC5WdPyGZv1lwGMZc5f7ekB"
    "KM3QBuQDhwA99A5ZgGbWre8Czr/epWExGQ44QFffBMyhBEoO9uovBXgbGRjlqwzgejLsQV0MADjI"
    "RniDpmFR8FeDn0U+6xMZnud5fCgBrABospprzWYpD5wLxJ/yPM/LW/HRkNbe4XleEyomQFDNdehO"
    "nT4gWL3+2RwANiYO5kAwcwAZX8c2X0eJ1nUwF6W6qBAEmXOnAbrLRuB5QNPBm2ESNJPQQtlRZk2/"
    "qxzZGQtQItHC2WcxPkvJYHVfM2qrGhNjatiCusiDAKIt9e1ZbgO9xs3U/NJEPSCDIGSLD3PM5QAY"
    "BUOYq6w0VwEGmM8xQOqe8DxLZZZwMtdzgoDC7GdKAge2AQaRKw4CCcARSMJPBkGcwYx3MzFmn+sA"
    "qGdsi/RmEKjYyvQVLvW7PS9gFul4jeMuXz2XgGClcPqeZ8U8dwQSoscU0ee+T/ShyaxVREbSe6SO"
    "EPkJmuBkjW8Yf9B4dvL1fUQ771ZERLTzbpWMH1HJOJ3W8WcxNTaJ1KwljxHR5+r2/L8ILjXS3fhT"
    "Xapjy6T3zALRlQbR1/9uHe9+/d1Ez72b6LkHp0bdg3PzvbW5Wyap1NFKtjhsx2+axeE/hHkrEvEL"
    "0nQBgMlb/gmsmJu78efGEQAgmmXf/WjCuI6X9LssYO6WfzRaA4CW0jnkchMAVic/hpAODzAAeC70"
    "bplGQX1X340/6RsA4ITFxR9zSBcA1GZ434/r5+w3T9gww9UcABYfbt/6oLG1G3/eqpuiOXLim+NP"
    "2N+7dc4A4vp3zb7zY/r58C8xALBCN2zFzhAAqG/dikDFh4c1I928yG99HW0qAZShXZp3/QoA5PzH"
    "5AxQvaPoqxG/dT8rcQQAkAx7UMOOBgClwW+ZzCMNhPVO4tbnJ6xTM/L4aN6KOwDAqN+59WVf5nXm"
    "ns1u3c9ytiW60wDJvfUW5DtWzTL8yINaXKvjz1vnJR3vbL8afz7GAcC6PLrl6Req/pUZOvmt+0lG"
    "QABQz575QQAoj5a3Hp5HZ4+3Pj/Qm3B348831Rm3pGe3/vtZnByfEoBGUJt/nOni/xnYUAjoTw9g"
    "o9TBf87t/zD28IsDefYXBkL/GoD4Kfq//4//4K9Xyrj5pqs7CxqDXGg/13zT9cldbzKaqhMAuLiC"
    "+yPBxV1a9LXfsfh/ClwmfsVi4p3uduUZpYqhX/Z0BEdtYgG03zw6/Okydk5PoMBiJ24K/DQL8NMs"
    "dN9zTSvpl12kf6Y333ZME3eb9zA+asHD58w6UbKpAcMT2UEDAXRqv0Wj9lsOATBMNgGNPoBSzzjO"
    "vn8v2ATQM4BlOICX7wewjLM8dgAN7tcYy8kZvocslpMHwMro/hhAN7st7mZ//L/GcEW3XMFHiP87"
    "pa3gI3SwzzoFefJNpKFcMs++z/eBhrS/hrlMnHVi++DiFwD6AjgJAENAb6MAgKUVSy6vWBIA3FWr"
    "XMVG5vLN2IIIAGRY3IyBF2Ovfo0DOhCqGMfA+OdR4mxetoYoWwIJc8RiDACOoJGTHIyBUHe4XKTw"
    "XuQsB4VqQ9O+68D8rmMi/qVWfWu6oFjM9V0gBKwzICeG6TwDdELAxjCzPd//R9sItRyh96DvL0HE"
    "JpxjHBTzNzsxcNsaQPEqa8WuqlcWC4txoJkIAYQAD1dgngHMJ6DpLW2FAc79Le3yHJhiAJZhZRqs"
    "TAMeNmF9B2Z5jAPHOFfxUqmO/6rvf4xDCcCg4Mp8UpwzqHpyOaLxPNF4nqgbJTwioiiq2t2gaidU"
    "5zYk6VRPiqmRFNcNogusO5jqUTDVk3Sq+77vJ7Sj60TjeUq2jIrGRiHHBsnxfEWHi+TSXt/3kyLp"
    "faJINh/0fT+KyOkW5DxCFboFtU8kV5ikK0zSFT1J6wW+oLENsERWLCgqFlXXWUDTOaKpQZQEAxaR"
    "71fR1KiCqeGT7ydFzeRSVeedED09/2w1NpJgvEzBjCCcXNrLTtIVPYkKdj6Z6udpSx/Qll5RlFZr"
    "J4mSKq34ybRgFRGlQcFOyoKdpy19IAs2SJ6al/SleUlfWgjSnX0T36ci6e3lc6eeAGvF2F5bxxyn"
    "elfTBDjeoQOeVy5Iu7xN2gg8L8/f6qpRaAHhEYnQBe50TYR7OEIv0cI6q1occ099V750MIbeiml0"
    "8DkxasZi1FyjRaYeawGtkKnRSUzqvTykzi2pbws5t2JJm0+kq5FyNVIWXnzCKT0v0Knvrc2561rY"
    "I63s5MhaNQbabwFhp7QdAeQ8s8Azy/OCQAn8TI74kxLxfUKLOwB+BsiPQItvgxnPgs6+4whsr4XN"
    "sHcWbA2kO4x0Jy+5VefR2lY2p5aU3QcS3SX7LgQ2z4WVIziAUnzymxCf/CbiHvDMoxbAoDbWADKm"
    "NDWIthYkbc0XqsZAZTLhD9bxIJIqYSeJiCg1BkWhJz+cGkk6Nc5HpPvp1KAfTo3zuwxZMTWKc0YV"
    "vTKfVlMjrc7Np8XOvKyebARV1PtYSlT5Se8wdXt1RDvg3eBZ/hANeFDJx3x/MOf7NXKUVDX+GVXj"
    "eRJzUKOMKZFZP7wts4qNia6skMnb7l4YugAQFozYfq1ewpTJ++kSvyZNDDMcKH/ITKaAaxIHtmXN"
    "1zPJNE0xDAvGNhRjqKxYryymyw5bjJdsgHmt2Mpq1vMjxIcc5nA1aw6p/E9fYM47NI9LE6Y0eb/G"
    "PxULbwPmcKqf6B1B+klBemuIpaaAvV9QsYt52iuQS00AUMkq07ID+Hx4gG8EHlYn7upG5qnPhx42"
    "ZnHnc6HL5joPd/TW58Fan4d+cMjQjHNz/3NxqwbSQRcd0RwBAIVKd3jTcEjqDpq2pmV7YIYutNAl"
    "ZO4svnsSmIMbcqydMYdO3xw6Z8xRp6/pTt8+Y7Madz2Qi/BAXj8/hMVjLnViV3ViVy0+vITFIy46"
    "sasWY1fV8edRYC1E2wnBnBBtfoZfNEPTMC/gsm4CELjMWlqdswpL2rACC3ZmQ/vPnzTL+AgQWyhj"
    "u1SP1bFCMeQAaGpU0TWjiqbzs3ypBtF4fhfzrHg3GLTr51OW0PU5KqaM0qmRpNdZVF1nlE5vyDCY"
    "GhTQNVbR2KjU1gwDfdEokheXqdcgSZRQ7xNFUucMFMQficjpphVOqqo9SK6wAV3Rk9nsOVNYLaQk"
    "5jBq9su/o4blSA3LPQpkS2Q2JcNePSSHp7Tm6FQds+mclI5cGmDS5sPSMMvLBnRp8KHU+fDV+PON"
    "eKxcVuu0R62XNm3Apo3czik+KgQAFf9f81Ydd26roVOqfgdKP7mjLr4L68YKwuW7RLjMIWeIdmgN"
    "gTmwzjq+0Yk5+2DMh624Pdofi9H+GzBPSLs0ASCwgd+2Yju8XWihm2jfsfI8vB0ivD3RwiPJbvyp"
    "gXXy5/RWrLNW/JzeFAeXm+LgS83nhy1oACajR5mq405rhn1CLl8V8nZM5O05KY6J4tgI3byO6xwA"
    "c8iceDVey6y7n8is7Z6w+o5A35F2TU5kD2Axc1GjqkfDLP4kEB/JRdyBWR5VIv+9APERmPF99V40"
    "7uSOCHohst5ZhL2zkL1tCGcbrBPrK3AAlAZZmV3HnYGdubTn6lKmxWbofKsUvSFEB6VwyOw/Pqp3"
    "FI9aACNxV+M85Men4Ud/6pWr07cdIvVHO/rWf/OseVcggPDEyeGHBlw+BCA+f9fiT53LPvp3XKD/"
    "VIfefD776JvP37X4U69kiz+1A63Untn/tlcYgMH/UODaba+g+Kmx9fv/5KtQ/+T/J4qPH35SSDj5"
    "1/70yWNv+8uVzlkF63NdRzzWbS7+wwex/bVfxW3/6EksfvSrWPzoV7H/3q+yPxNgxcXv9Z0hKC30"
    "JPH1RE7ZIJgCAKuCLT2qn+sDmuqDyRYAGAO5Y1TpkwCwEBTTGqIvXnlrUew00uI0wHiNv2616yjh"
    "ApM0k1v5W+jKQZsD6MoBj4I6PTaaPK0XxRhAez5Kp/MBTeejdDofpeOFQZ362hgEJOYAAu5ymLJy"
    "imtsiwvYh1RLARzLK5DLK+sEAO7KRgYyc9iA4hvQFQJdodSuIQY7WwIkamRojtcYLaPP14dUQ6Z5"
    "jlCJAHRHNC86nGIARme9yaAJAAlMk0gDyDQB0zTLdunuSh6gARmrUYzn8cwX8eCndHytz3ULeOJj"
    "T3za+SbEC1p+28M+AIbFI2p49VPLAA5/C1AabKWB5GYTmQX30w7Un5oA8B+GgNKAOxH2hq4NqLin"
    "WKt/F7uPqT9t9aUNq6k7gFG6W8VZs+0BCHNNKQAq11SVa1SJ1qLuK/rlnAOgtHc/Rb2PUdLz5omo"
    "CGi8QERFxbt0knfloN0NLtS5L8V11usGRAVR78Rgyk8kFT8xqNqD6hxLgoINfF+viIjOzddc+6Zt"
    "VFFBlI7fWlFQJJsP+rMssIh6DxJVQUGO8wkZ+UTEH6oS5yEa8Ieqgj9UDXhUPM0TqnhEJOYghIfF"
    "+BBUvKcSEroLWyGDlgxX1P5hb6M5OoXMCIIgAC4bbqhJxEHAMzZ0roI5VxkoFgdaSSduxYfcOqvq"
    "zjtnXLun1tUIYPb3hXDz1prtIV7MW/1OqY67QL+VU99dBxxA9CbAmgRzJtCcCZgT6svmGTWss1Jp"
    "qifFlEXVFT1xumm9BkXBLLthoTg/bhSDccP3fT8oevumRuLXWUXJYDyfDMbzg8HpxqDqNQq/10gH"
    "vYeIiHoDIn6SqJ3wBwvf933iJxKSUz2JrugDOTWqYMuofD+K0i1jqg+iyveTLeN8tGUktGWcpy3j"
    "fDSus3+DJxtEAjSdJ7++v6JLDUl+5Ue9bpAmmw1KenvTYrNbVJtd3/d9P1kbDFhEfkJ0wUgG0/mB"
    "P52v/FfmB4Pxgk/j+ZSebBARfUgS6eeJDNrcp4ioiDa70YS2Fgoa75O0tSBp3AiIiNIn9qVsUBFR"
    "NdD9lPSkKHRfFrrvT43zwdQ4r64ZBQlN/fk/Fey3fsUcvv2o+Sc5Apt5sf3PH9fyxz4JeXxgDT80"
    "0OX7nvIAKPzDk9EpAQ8A+9c59A+01cYHwPQPtAP6oU6jwoT3y2pExbetYP/ygfB9Nt75PQ2A7pz/"
    "7UfNTDn6sbc51rG3Hd0wHxsDiJ2NJy7hVL8tbbD24y/B/tPfy9zHR5581Cr/cP+odXG//pyzP/nh"
    "0py2hjrb8hhx/sPbPA6YvNKufnjYDjsasf0a5YqZANBPgUMIa+xaO7YydI4A5vEWM4/xdktctVuf"
    "jkets9m77ubSDnlmxRysiGuEszLFi0esrBPzrNPnZc1rO5nJAC2c4apnWNMKNWghTPOXV+E8gdyJ"
    "zdz5IwCvZlvq56lYGxARJQU/GewYFOw0Chq/tQjGMy385uFX83vpqfk0GjeKaNwootONgnr7ZNLb"
    "J5PeJ1K5FqWzbHtZ7HLNThQN5pJgapA/Y3uJqPKn80RPG1QRET09T8XUSOhpg6qd+SSoM1KTdNxQ"
    "AuPdXHWjoGqzSxVRUvX2Jb0GRb19adI7nA6ch6qa446KqJjldqZsQFM9kdfqe4LGBtF4QdG4UVR1"
    "7OdPDQqqVxnCw2o6X0W1AKOKTte8tv9kg+j6LFP/AqPiOovoOovo+lxEu/Hi1CAxF3rYzbYU7Gd/"
    "GwAyZn2vFbvhYt8ZUt8dUr+Ow1vfcPJSO1Xndlonh0pvikW9EC1LDZWthLR/+IbMflkr9kCXtrKl"
    "Tdp1N6sBXE+84c9/pTRDV1mhGzo1VgzvjAvsZqC2OyFYZx2ss47hB4UatXJztD+3R/sF5mJ3N9ty"
    "I8dIxQAosKvnhz00xSk0Rx0tM8wSAJQOvcXq6Uww5wkwZ4JRr29ePBXnRkvAaIlFvSW+9zMZwgdy"
    "HrohaLkSAFDasow7eavfyZv9DtSwAwC5GnYAdrpOZL/ai3C1J4mtSWz3MvSdUN92QjAnFmCvZVsG"
    "lPCHKqLCH3B3nvzxcjEYN4pq3IhSIqJeg6JicKVmIWjHSCZPLkdRfRUV9fYq6jUq2nyodyEteJIW"
    "bJD4QeU8VO8s20+zSk6NKr1iVHKrjtXT+nHLSIiIgi1jENSzZxLsGGmNpCyTGu/lc0ZIBw2Jg4Yi"
    "orufyAFoLHNKSziZBZ6522XmKgEA8ScDaHhJAUAm/6SMS3tbNK1sCCvTyM005WYss+LP/Cxo2DJJ"
    "b3KuYa2+32VHvOmqlelXl5SdWWSHAGDSnquA0GtpmDUipo0ox0ihZUAs2Hh+0Z7EZEuBvUS9LtHa"
    "I0T3X6AteJ63TLTFIqJ2QtQeDKZ6UQWe53l68vL0niDqeZ7ndX9v0Iuo4t1kwLt+2n4kuc6i4Doj"
    "qngjSjYbRbHZkMVvFdUVg8iPUtqBT5MekewRTXg3KAZEPo9ocFJOew8SJQGlvUeINh8h6t1D5Jwg"
    "ch4icn6diEO7AceXLLMQAIyuWgwAI2D4GwDD9f0egGWeyWx1823HADRWmFbzrfw3obHW9zFs7eqY"
    "eQlwXWFVQ6Ers2CAnYK1AWwD2jYAxkv63IcAdlI2f/mzb+DHvgw0XBMAzBlwzr8P8O/sYvtzDsBm"
    "jfqkewDos6TDBEYc0LcBXMsJAA5gI0wXea3L7vAl4MIqUPKaHNyetaHpzghqU+iTY/GIA5Ai1hiD"
    "EHU/FBxeE69cyMtavz4kNI8BGgPcmWzZnPGOZs14tmYUY/j48GkPeMfBtaEGAPkSwL8JG9/6TR+A"
    "Touxemxtjw/gFzfaB0bAkUfBnFNLMAG+/c36KNyJASc+KJxcIHx89CkbkGdOj5gGzAHEAOiwvqQ7"
    "gAH3N0lg5AEIY+d2AC/O+mf9O2D1TwD+TcCJNQ7kpV7z2pmWM27jWqblaAEYOjPmlvoRAKjmcLPz"
    "kO0B+HK/PWM/cg0wp8DCNeAjWwB4oANUamg9A/6vwKHD+9I/+yqDo0ir0Wqw2AkfbMILS61/Ku7F"
    "v/M83F9Yd9gN7Af/18BrmmGt1LFLvrNiWAa6grV/WLLXCF/kyPQAgJcXhlXWesneQYejvG3G2dOr"
    "ivUu+AswAdAQICxVGyt9D/H9vyYYZqyOBlgo1fAkh6B8cfjVex834QH8Yn0iF4GNV7s3I5+5xgDk"
    "FlDqaHfiz77IAbEYf/Ym8eObpe0BiJ0v3m5LkzwAhQRMggDMK4AmgO1ZfgU2ftESGAIonVF8MFwh"
    "oDRb4tdaEKungXyGW+kSyIH0dltqXAKh9xlxoyJU7GrkS0BgTnv1pAFOZnEArdV6rWOdGW1PK7NZ"
    "IT7yhWadEb2hMvnaVx7WOi3kABQsVY/3Ppq0FvfdIHDhbM/SHlDWe0La0w4CR6fSu6iGsIPA3aOQ"
    "8RkHf2N/CTAB9Rr//rnUoGcjomRrvpiviOa/T+l8Uk0XCmNQRw3G+ahXxz9+756IiHYeUYPxI2ow"
    "7g7kOJhMNokm49f4d/XY1KgiIqKd+eRzRM9+iKj6h0SRMeC1JiVqE212g2mdzBRMb+Df3020c4Lo"
    "0kmiS3zuJjWqEksAoBUbN8t+ay+mfD8OmLzGt20tW/lJ1MwbqiZozxY3acZZHX+WeMsIsMLODBvs"
    "3NID42besCM4B7DUuln7PFllADAMPXBbLgIAP1Q//pd+WCtWtQz6xmtJdRYQJsOmDdhSO1tre+vH"
    "/3I/187M7EbO3EyyHgIA3Ba7Cnv2P1YvxmrjJ+qn8we1xMd54ibuuHbRMPGfPwmMmnWo6Ko+frJ+"
    "vl2KNwOAK298d7FvzeJPM5esV3+gfur4T9LPFW+i1RrGP7jxZqaP13dm63gbYJ14pb5v45WfrJ9/"
    "R+UawLHnJm27GnECgEDnQ8B6pV7kePYTeVbsqWMt2NVNF551AECXF+8GuLteK+FOry/+RP3Uv/Gq"
    "tuumHIBZysf33EQzyf4+AJjMlj9JP/WZTgo/WHndV6+nZwsA7LQeQyzd+MnOZ7yWA6Bgu3dzokyY"
    "1fEnTL38dLN+89vqJ3JZCNfq+F90brg3Ga/jzzi671sAXayBY6Hfmt+f+xcixESK5Kco1srezz7F"
    "GHuj49h0nQSZ9JJ291UYndYiY4xdLN+5/z39tsMYY59+YvvUsez5vIzehRW6WjaZKLEpEKt1ABNR"
    "54DlIO3baRB87VjWsd8QkqPFEBTjm3BZ/OX77y+e+Qc8WXtm/6lfvP/+r8UcnX8sUW6HBBMyB2QO"
    "ivM8ETkwINmWQcGKtGCJrH2iBnJqSJoaBU0bUu4Yhaw1+yTHjTR6EQAaMjq9T/rjRirHjSIYH1by"
    "0idktNlNgzERcaLJGlHCKaphVCquzBVFZST+YF5V11kUTDkAJotpo8Y/YVBxab5IXmmk9PW9VDy3"
    "l4pxg6JL+wbRFtcY3jKEJnUMP0xCeh4A2b9bd2bQPIjFXj56rwfAwIsvSMtJPQD9ReflfwnNBjOn"
    "ub56TRA0QbDiFV7Usds3hwCJh09tcAGomC/994g7rBU2QWFHdzIlAOy1lLys9ffPluDblZbr0orf"
    "kkE7kEFjGfi1rB2Ca++B1GFO7FybWNFD8ODhhc7jQ1dYL3CIF1xz40wXi0cSALGiuFj7NdsHMLfR"
    "ti5/C5Av8krB/KHUrJcvW2CXNbJrHRaNAIXw1PbTQ9Dvhw4OwQq0Uv9zq2aNvrh8GtChhdTHS/e7"
    "wG/Fzqjql6w5BFOxw1TsaE3B/6QQLfA5JUIX/GF3ZH7cu2zFAHDH97Scg/ouqM9BQxP1+mTmzeHm"
    "v08YAJzqQ28JU28KYL+A3RTQ1NDRlIkRIJeBOy1A2t/XsOx59x8IEHt5OzuaO8JF++NujP6vep7X"
    "BOufilfYv/U8p7OOuV7cutqJV97VWkd+ah1LVQy9mZtDzIlJbFGFDmjo3d0vgSDQi2HJeHOEa82R"
    "o102+PCyEQRBwIfSsBQQBEHAw5L1zkB0sow5MbRODNYS5VJLmACEC2QekD1QHm8GQRB7lJQYUuyA"
    "4s62GnZKtdEKgtiM1dBdX4QKgtyVaB1dN4UbLqInwVwJ1svAnFDXAP9pllRTY5BcmafeQ8FMgyCl"
    "HDeiaNygqtcYRL19RESDqPfQVK/VileMgsbGIHnyrT6NGwmdbiTRZqOoenuTgKgdEbUDovbA6RY+"
    "UULOI1JGV/RK1kpKCraWC6KIoq29U31QFERyaiTFVB8Mau1lMnjFSJLxckFjg6I1zX6fQeL9OkhZ"
    "GX/88wEAvfe4lSFzgcxVOv9+pue/FwCwzC+e+JdNm0IAloKoGHI7hyhkLswsh5WBuZ8BASMfGJ4H"
    "hud/utNzgGaw9u/MCZRRQu7J7MDO3C1dD8AFAu9iidwO4IYKo/0lo3+pNvgbFTBSSKocmY3nTSDp"
    "dSvqPUKR8wiRMfUjItoyiiRiya7e6OVd9en0tfjz2d6+yHeiwneoGPCgqhWjVF1gVAWXGuSPG5Pg"
    "UiM9Z1BUe1M8m3R73SjpdYOEdyfRWo34J3X8WRERbXYp3ewG1WaX0t6srQ0i4idoba4VLzLVN7PO"
    "tvPqcqapjezDGu3qjfQZtz6UOAA37NQ56qWj8YdzjR/JYcYM5nE+NI+V7daxq1a8mtmxs27F3NRU"
    "zettFEt3d2JetvocNOTri18AAKnt3wCssDWLP2FaYYtZIUxt1sxfBpx/T8CUDaItox431OtSQUQF"
    "PynHjYLGjSIYf6za1cJHvcMU1Pkn9Tl/ygiSpxoqGtc692jcKKi3Lxi0oyCptfKUOIeLGv/spjPc"
    "M5oaFLyyt7bc2GkQ7cb8NXFWRVODqqdnbWee0lfmCzEndY6JjRelzQC39tLSnDO3hVouwtXt2x5e"
    "Ar3KrbsKemsNAGJbbWCVNKymGytLqVhZkoQlqcG9rDVHDj446mjNkRNjbV0DgG0eLkmbNGnD/IoN"
    "9z8frePP71rAqFkv/ZqK2ajZL3UVM3fW3qhg6uos5kJv2/zs6nV83A0BWwoAb2auxDFnSH1ng/qr"
    "aA5rLXxzZOYt9P49AJTaqT6WmgJLTbGo121FU0LZypQ2aZftzJE2lfztE1bHoVIL3dAJ71M8dDOq"
    "M791HHNujD/XwTphi51avwEHjbXhBwXm4qPig/mRnJgDYNSsg8o96nl2Cs1RB5lhlq9y6wYfmsEd"
    "GQC0mBPmWi/LtV6OkSMx6vQXjVNxbrSS0M3M0BV1vucb6/iT9qik34HqO7EadhJ58V11fHfxJMDd"
    "cBZ/SsAVJlwJdrpu272QbfcyAYNoapA/naeAKOkdLoiKasDd+WIwbhTFuBGlvQdrLXyvO4gK/3qd"
    "pDWYGkExNoJi/NYoGi9H0bjxbDRuKOrtZYOiaCeyYIMoIKoHt5/0PmVUcssguWUUcrwvqFH+hqwF"
    "WEREdM1Ioql+PqhZpKTGQZdJjm0+xyQkCxxphS6QOaQBWN92wtvBM7dPmatE351x64swc+dzzXrX"
    "vqEo3lNSvCcXTTsXTTtjpZ1p5GZ6U5+MWhbpTbgicBIA0DM+0TOL7NgiO7fCzm0AEExcAdDF2V7r"
    "TQrqYqnRSEHodWsZE7G4PBE4QWntv0QFUU1Cs0TtsEGQAoAxSXyjmnmBnh+ce4hkDwAWJrT5EBU9"
    "z/N4lBavzqCVP8VylYxtz3MOB0SVP2We57kRPY0HC+p53tNt3x+0I4q+5D3lMSoGQUGO7XneYb/Y"
    "vCeZ9ADgMBVrD+3OoL7PgW4UcABsUBREqu5Pku6wge/Xfab0Y5Uc78af3ZkfVCOg3u7zh9IBP7zr"
    "X1A8yxtp/f7hlIio9oaKgimCoLZETYKBR7L+KFVtPVTQJQAw/GpzH9V2p4eJ1hpU9BpU9PZFaQ9a"
    "2NRtD5DxSuwhn+oOlMFlTuI9pVcr0E0Pn3+bB8gGkAHibR6EdEnG98yQrcnIhMD26Xt7/dV/dYo5"
    "p1i9JRGvpfzrB2NAg6rfWd28n5IPA+Dfie+Tb8jr7azVhoAyhc0BUpqQ0PZIaO4/N033GAw/8gui"
    "caOq7yP1bO8hKnpOg3a5g4powH1/0I6Kqh3RoO37T+tRcEWPqGr7vu8/FA34a0xSMG6o6infT0+/"
    "tSJKaGz4/rhRJGPz8KT6su/7pz8xSaLC7y2nfu+RoOqdoMr5Vb9Ad1LwblqxLpEfFVNtl0liUTUF"
    "QIFfEE0NiogSHtHgQEQVf3qeZi4tCRU0ZUSVkaQX5qPJK/O+vzP/bDSejyZbBlFEidwxas+SdNoe"
    "yCkbpERET+kJUUKbjYI2H0ypd9Qgiojo9AKRLIh3iQ4NqNDrxKgpO09TPfFroo5ovNd5iKoZHz/g"
    "czwmALjYjDlQDU2phg49P3RJ6ABgqxyxkAagGB/e+U4TmQXkVumElgmlBwHAKbfrGdS82OkvjlpC"
    "CBUEXhkDUMe+uq2OWYxij2IgCHBIEZJcjTph8RyHNIBMb4WBzrOQcSY9BUAi1D8vwXoS209MsM3n"
    "nIB7OpB3YgB36hB32SjvstmH6+25/sESOg9vD4LfchMt7HARup73cIdRv8dF7HieAwwfdt3wbO6G"
    "Wu6G2lXnd5HDZvqpHMBkyF+cXHwMf32x3RQxgHJP8zHlQi5lWrZ8FZFbisjOtNzTaeKVb3i4zlpK"
    "4iMvxAJ6LFascB1cO/qZp1QMMPfxEsj2tPqZizzbs9379bqfvb8gR8VHPLzwbpi/eQ7o/0aAfvmD"
    "UlwA4i4QAHo8sHPx6GI2hJ0N4T7qKaCjnMcBlPozUHvk0hX7Wu8TDg9AZecvc9tMdU8ELtT6syFK"
    "V0eJTJRWZOU0UvBsNRwIBegKpFN+4AWk7HrNtrEqifiJoo47E3nNoCQiSp6anxZT43xRGP7L04W0"
    "mM4X6poRyHEjKl6LS4uqnkGDihfJBebXPlEsiqLeYer2ulG3R/TkfB1njucjOskjdbKdVgOjSC6w"
    "ohgw37/AovTp2r29GM8Haa8b0GwGXZsTrc/NDGaH7/nwsM0+bLfYfr2VaEpMSiDnhX5N4gArTNuW"
    "JmPSwsZ3bP5iyEs2i0tH0tKGTsxgHm8J89O83RqFADBqlSuLsZV1Yl52YsCu3BnU5IQnhpicGDG8"
    "g4Hfqedo6sy5UzdHTatWbYSWZmqhBoQwzX+s8zm990TNLDlxqXSuSb2EtEsYnXiJAOitxzZDD3F2"
    "gG+EXHsudHQcc0Gx+6rWyPx4m+WuwPbp0Nw+HZbMiWeYYrzYGjrUvOigeRHQD8Z1PNwU+GvDEZdt"
    "C5nFkNlCk7amZT/bNGseeYL4q45uSxOWhOZ+z8Qc7vge12sjXeuMDRG62yJ0dS5qXAiW8wTFrrpt"
    "1jeKV0HC0ZojR9vN+ayGbu7YqVhZklppyjfA/euZH9rz2mWdD2dxJ9jskN9w4tUzFjmhBghXI+Ey"
    "hK5Zxk7e7N+Q8+mqHHtUzPYUOTCZspmpvV7xbjA4kASDdhJVwY6RKKKInpyfGkl63YhoNxYfzyfR"
    "eNmPXnWVmqeodn3CIJ22a788IpJTveo1Uup1/dotZOZGlIznCV1V4WTt+ocBFWyQXGGD4NXZs5EW"
    "UeU8RBV/iAb8ITH35lGr/pYXm8GwBzVsac0RB6M3qlwDAK+SOh+WzJTShpI2UXa7NMMHYO7Gpdnt"
    "KMEciVGvj4udPowa5F7Uyzg+OqR+B6rfAVZ0ldQDlAr9VK7YSWTGCjKjCamvYN1Y2U2zFaE2BNhp"
    "Sbsz6Bxt92oEHU6og+6yy0Vpl51j+EZnvQQOhqPmw0cS7bdcsIddhYfdIY4dYc14sc3Do3XOZ3xE"
    "tEZ5PYO+mLuhgPNdANDYqXDY4mrE23LEAWtY+1w19WYwB5J6SZG1HUf2NiL72kRa+Wx4guIe8hIs"
    "XocdixUrFHPayh0yB4C73cATTuZur2RurvJ2vDaJANC2E98HM/8n0LY/miPuAOyES0NuqbhTx6X9"
    "Ezni0s5F08oYkOlU77NY5jyvr1iZftWqdUb6DFbtd2K3RPhAvyydDZTaiEpXlOL0iOLHBQDIYcci"
    "IfSaiyedADWLOSuasqjWbn4oIpqPZgmRRTI2/MHU8CfnjEB+fW80+fregHb2+tWOURRTz/M8g+RO"
    "d1DxKCh4QIN24Rfc87yPFck5/rFKDYnUGlH0i2nxpO553uGUthjRhBE9oxNdAFGq189fMSjdAYAG"
    "+TvvJnK6Nf8u+Nz/FDIHAB2XjN8gvSyQMQBQG7B31bevunJzADo0pukyCILAquM7E9gwAU7HoSEI"
    "gq+dzduM1RAxA3hq44AKgsC3b5TpAnr9ezCA2dL+bm1yYZ9heFUuakLLaPSADqjw0Y0b2VoSqsZC"
    "h/eALc36qJVgZv2/dYC57NjbAA8fk+aXZ7psDjDn8R5sgItj/0bjP4d81ieTYsd2YkCJ/MMzLTn0"
    "2jOZdGBFA3QXdOxTOeDkrfY9AP+PAE8BCC04mnvMw9ad5vYN9A4QPgofwBcf/+pQ02oa0zQzgOc1"
    "122K5lG0/hdCAEdSfA9K55u17Rv/jAbHg/2p/0eLwQLEB2suGw87/d/5GLT/48xpvV2fv4+wWX8B"
    "mDYgONTw1w8g7GwrtAD3TwDndwANcwU48wC70EtEu4QRwGSNhbqy5mR2mXANAJs95+HBEeB5GGai"
    "5uFrQbMS4oDnwfuhlaBOWwADYmQE2/MOacHi6+14TWs2Dkxb6u8DXGzIm/9Gc+K71VkFtyn4b93o"
    "pU/9MgBw342sxg2MkZkvjtqQdgAPMSYz22P+Qm1Fz1kAxK21lVkf6zFaiU6sAO/XxC4hrM2Iam1X"
    "303ihT3rBA+7M9TuFZ5zQgCeFztnLCzdMLBPrcPzPG/Y6tcc7K5yM5l5ziswMXOgQes3nNfOjSXq"
    "WcnRnTPbkGdrdhPgeDyGAvhyS/AbyHQx0+nXA476DjwPOsXOEIi13XEx50oBMwjarnhVhi7qGGI1"
    "CIJg21mvcwXK2Te+4ZzmTtzxvCAwEbtAEu/68S9lEIAG96o1O3H1v85w6AGmcLbxuv69enSojVNu"
    "EEihhk7+mkWzCe1Pi58ecQDmy1TbKgBNAHiTfJOdOxDv/3v/by3n9UbR2kapAdCe/7nFQxrKaxdX"
    "AA/Jf9ybA+3pf1i9xr+DGHumnSEHnPOFIg/4FwD/Y+BP2Z7/fJfyAP6+205SfVIcNvNI1+v8hev8"
    "5TkGD8jetxeLgPsfX+Pj+B/W56j1OsGtGk1KAFpxS47suZxpAPCud66+Tsg6Q7laZ2/i0zeKevU+"
    "e8tDaiy8rXbjjZ3X84a9uD7jzutI5VP9JQKw1PrCLRXAoVv344kjN33H7c/XkU3bufmrs06sz7j4"
    "W+QDiLjW/+at+L7X9/OO781cb3+f3cRY1nfD3+Dib8wq6Nei4m997nVi5rfJety5l296e5s/POPi"
    "b8Wz8/3DNQAYqiF/fT/fWNSuyba82dXAljgEAO4t+evmxbpGwv430uu03nV2urQl3czFZ1THz0Hr"
    "FipxVeujNWn8Db74G62ayx697v+OmjUXb9+SF5azXJ3v/d/WXyewn3HZ+uvqceyZcfC2uoX4nEKr"
    "hpbDX7r6+n7mThQAAHMT7ebRtF7uzn23sBw4Uss54pMHbrpKWj4bM6x3U56JGD5eb5x+0LzF+NTi"
    "ngYAKna01/ezObuwdx+5+TeZM4kACp5xbtXP/i/kDgD8s5+9ecQMy1pd9JK62dRAz2r/BCF6tzDK"
    "T4YHaxXo7PHG9fGlZnE/wL7+g4+i1IYbZWxvQJ2lpQ597WuAsbM9fbQsAZpvIsMK/kG5sh0wXN0p"
    "p/r1g+nDAPbsCd+Hu6mkP9VfaJD5V4/H+gfuB/CUvgdKQ/IuRyTvcrTkD8ToPzHYRs8tVOBuMjXa"
    "X6rR28u4Q1mAn74e4C1v0u/a+r8AaOjBh3OlfXeYvSWfZG/JMdevMxfpj8A4ru43gf0mtverY/ks"
    "/hS2CSbBdnKA5cDGdx3dCp31F0McGAVBEAQ8s+KQMa1/IEfMcnyKwwyCIPh0Nmzl6wDCD46crEXA"
    "LxEgX8buasKVBg5NZe80mfM9jTkZZJ3owJV95hUMrczUrWdM3RJa2JOe7uHlUyu6yyAP5G15oDSv"
    "tEroPoAvD+/WOZ77K47RX7nsxRfcDfPYPSbF3ZJiDwg9AG5ixdLKEAPu44D7uCkOxNDK+LPbrnDV"
    "6FxcMkoOYsS/8SWBWCH8qIPh2y2GzJw1V9MyR9OyO0Dxr7YRunnrsUzLbanFH5ZaXECTsB8wXGz9"
    "cwuO9q3/ZOK2X9iHNzzs/qtTwgPw8uNjmBuLcbdu/UfwweEJNIeDvho+q6heoIPmEOJuuxBgqQCT"
    "2rZngxX/xMKhESSD/R2txHfMh/8yfvRB4Iuxo2vc/mJLo9jVaM3VzF980Cz7J1DGPajhgMHmmRry"
    "4yujpsCGErAA1F7fU5YEKQ0OROl1FsnrbGrMfN4MoiqaGs9GU4OiVwyajOcHyfitKY2XXxWoJ5cM"
    "fmJS8JNRwU9GRe2THxWsKopJ1eumtNYtyPlVo4jqQ1YUUMFPzFynE0rnZHLFoOSKoWbe+EU03usc"
    "TpNe7Qc84HBqTrByun4VTY0kmBo0edq4wCLf933/OkvSisYLRTJuEI0XiDb3VYMZchDUep+Uevds"
    "PkjUe1AVvQdV0XskIKKo4CeKhIp2Iqt2lRYs0c8T+b4/1QdJRFfYgGodKtFWQyZPNSh5akElvX2B"
    "71Oa9O55crmIxo0iGDcKKeZOP08hBQHjIYSSXkuTNoS0r3bCdp3JFIMhPBrz0C3x8FHQsSMoY7Ot"
    "YifRQtfzPA8qdqwziN0zBDckuE+36vjzdzXkk9Hd+OsRIHUcjKE8jxnNWCJW9koc2dfjyFNJ5N42"
    "kfclE3n0DdR3aiQiXvMmI8eWGpjUNMwt5aXGPK/08jar/TxdoDzaflJYnud5207ISoo7oLiTU/9U"
    "iOEJ0PCkRcMWzPijQBCYNGx9G4ANwM4BKzfjIFCBmw3f8/y6TuZEv26SfffaKNSCwA176wAyvY/S"
    "jVFaebbeEaUwUYpTQm08Tp5HUMP/7gelwB+poTZSQ1Mg4oMo8v3KZ9eJdJ9IP+8HxmCXi0+2DEqn"
    "hl9MjWpwrhG9vLNAsubWJy/v1DyaHC8XvhNVvkPFzL05qJ3N/aLGPrtRwomm80VBVLvZnWyTf51F"
    "1WDO968bdZvOUzpuUEpU+ZcaRJvdgjYfDKa9rhRzcJ4xS89jaP2zrGBm/MOlA/yDJgfUyAOQ60oM"
    "peluS5u96y5efvu7ZvZiuFpa8Wo5msWfL4YHdI2HTONHcjjHcpjHW3VO5Sivsc9ab7RR1FtLvRDx"
    "iRHcv9ZV9g6NOZcZcy6DOdKSduYoAZR27ABWCGH1O0wLTcyV7hNoAXB6oZDuajlZha1WGXCqjxJQ"
    "difeCA+wjZDXQ+KYC4odQf0j0GbxJ8UmlnSBJXudMT1k26dDc4Z/shr77EDpAOvE2wDgtWL8tR6z"
    "79oM2R5Ny27XtOyopoUu6Nh9cb0VvA+wX9YyW2pnbakJ0NRIlO8TjY2k91BVOQ8WlRMQ0biR+L6f"
    "JpuNAe/SgPt+0qboOksmU1Y7Ww74DJFkSURX9CSZKcLltddqbu66dZ5uEEW9umYo9Rq0uS8i3o2K"
    "tZNJwbtJwU8kBT+ZTlnlk++rKauIKt6NrrMous6Stbr+JhHRFIN2VFUHoqJiSV15UxFVwY4xNRKa"
    "GhRcWPCjnXmiS8uUjBtVNJ3VBh0bVUSzT6OixkF3a27uYp+8SyR3az08OV/wE1HBZj6fxiDZ/Zaz"
    "6qc0nq+Ikl6DKucTsnK6Yg7bTiiBGIYzGZmlGjr68zonYNSEBjB6o5I6QmnDvfMOmN+9HYgfUDx0"
    "w1e9lkIbJbRekGtOJnQnx6g3iz9bz9fYZ0bDzmu5oBO9ouGaGegrUMYK1vW7sW7cjfXlu0S4mkMC"
    "cfhLV4HczcBOry+y0yE0WnnfvxQAIPgf2Hle2AFSu44/vyIB5nyj8/BRHj/ciq5vddpW+whUfDI3"
    "j1+42/zy0Vn86Q5XUVpZWbp/7wDceznce+s6Tc7Ht/lVOeJXL+scEHYnPgDgpXbzob0l/5K9/a0/"
    "uz2fSBcT6dBEuqD+I8P7BED9RzQAdhWgmW60mcRcuGhPgiCAyNzAVSKjICjdMAMyJwiCINbitbgD"
    "5Au/H/QfttTwBCbDFiZ600bcqePPYQcm8j2KYhsl7Ew0rd34M59hn4kdAno20yKhl7t9Kt04W+8N"
    "S/GVYSk6Vil6UBsHrVIGwfrooAk4P1QiOFZC6EpgMEg4AAzoCr8Q1XmV5yn9GMlrtSCctgxKztW8"
    "9uDrC2lwbUFGOwspkUEnn6rfT3e6M846qCpeJT4HoBfVFB+rqD58VPlp8TQANNJizKioAID5/nTe"
    "9y/MR9WF+ah6ci/5YwBoJP6T9wTFaQDoBlWPzz2D3PM8j/ezJXxjBZ4HtDaEUEmiccvz1DCzTNx9"
    "l+d5nod3vVNZ0NRti7qyABMPNwHAzuwnYu7UOQ91xQ4TgPp0eZbrV2twbRultGFzAC/bsLVSe7lO"
    "C3Yy5QHKeQbKeUZvKzu0PQ+ce9mfu/GSDuD3W0wT2pnTYV1/c8V1H3809uDh5b+UuopXDOcvPChx"
    "l/4lhcwDoGe41BQdKPDnFYAvKcQeEDiJG2aLuqdDxQK5e2/Pq0s/PcqcUyhtDwjCr4Yqdvhv6g5U"
    "7NCh+L5MdyA+rGkBCcZaAuyDAraVtx578r2bwSdBsQLTnwKS37ZiG9CTtEhq9f/O/CyW21tFLx5O"
    "k3EjSTcbRUL+rGZAlLR7j1S+80g14CeJyC/aRE+zRE51NohSIt+fVewoiJItIw3GRlT5vu+ffms9"
    "hZ5+a0XjBlEQVL0Hfd+PkoI7XTlwusGzTjcYJHLKKjnVK7qm7ypUeRJdx1wzN+M8CEaFCN9YiAMA"
    "YF9m5dsDK7MlcIfUctLpIhAECaqL7jPQTz8D5p4BSK/+BGga3JJLrWFpA55XV+yIdQC6eonbapHp"
    "gO19f4hDGUkvRcYyil23PO4CHjAZHg0X2zx0Sx66jDbkMtuQy0x8xzZDFwpoO/EK43NOqDPH87Z7"
    "MVgrRhAEwagl+IZitKFy+41qWGa4851B4HkZfuadtmSwpcCnJJAh+2+D4GFLIvR6ZwhBQEFdsaME"
    "0HypFS+OWnEAIF8qY+ypBGw6TnolyrqkACSwrh+QG3Al2pYEYIb/Q2aGWo41+0kJ5EGQueEiW5tz"
    "A7gzmGbo5CXgecw5Y811jmGus45hJ+ZC1bl/Nn+iozfjckMJaAoQKjzpebkLdfy+O6QnPC+c1Ty6"
    "LwR07oTaVSeE8jwVn8rBHo9N1oNtPx4TEO3JPE+YE2npzTgzVGL9vEoIFB8hih1U/TVdJVJ4HtWJ"
    "+lPjPBW+n+7MT4KdRjqrWUM0/oSMxh/zo83DRGm7IPJ9KliR9KLgWaebDDgRpXMv+z4bqKlRDOZo"
    "slvzaDpf+b4f+GMjCXbmk5Soii41fP/SPVRsNqLg0j0pFU5X+j6R75xIel0/4Y8Q8UcGyTlWFedY"
    "QeeMy4o/4idE6YANEoGUDygiCqh9IUqN6bOR79PTbJBOjWk6NabF1KDoqdm6vNMgNW5ENG6ksrdP"
    "ptcaJMcGBZcactDzKyIi3wmqitc5MOdYVFxgqnZsYAVN56vonEHRdD6Z1I4PBQ2MJBov+MVmI6p6"
    "jagYL5B/qUHFpQZFmw2a1tXTAzFnOc/AAShe/EMT+79wiHueaLcYWHMEpkY5U8KUdu0X9F0ucyad"
    "jElLuJet0YSrF0M7s2JOTOuXAKDxmA2dY7PlXeTt1nNXgyAI9NZZbBQMmhIf2ShysVLHBvQOgNvK"
    "iy3pPGPJzih0lB07mRbzpLSyQAcotsIW5uA+UcvFTod82OtjHchsJzb1Tp/pp/rQO4Li0zU+tOYi"
    "sZtJYitBthJm6OYUOwn1XcaX6iVxSRfI3Tr+bDvxkDnHasLDETnrxJ25TvgG1hG4bM8UgjY39f2i"
    "/RdKYI8C4qN5K76PymPWVfp2oT0DQNnSEqCpEfm+T/7OPEXjvZHv+xGN99bT3LhRJJuNKaNJvfek"
    "QlHvEZ9qs/frjPwpS+TUIKL0Bhf2Vx2JjIjGBvm+70fjBiUvHq6SzbdWUa8h66IbsuqdIEqcT6iE"
    "d4OKP3SFVXLKKnWNVUFR8bpng/Ys/pS1Sj6aXGdJSkTFlFE0NaiYGuRfM3bmo1rjtkBEk3Gjil5s"
    "0OTJxni+qOX5jYKoqj+5rsDEduPPpJjR+3LK6urz1xjJLWNX7V+wQTJJXmxUk0sNKnqNcYOi8XxF"
    "47dWqkpPfyKtKzXV8aewgBBzndD8Riu3ARGPWrEaNeN81Izxd9R3b68h0EtWhvh5W7Ly7RLPe1lm"
    "5QhtidDNIVivvvCaI1FX/gCMlsiNVlIbU7aE1NUQb6RjtIeCWS2kyZtWsJ7bEqUlkR+VoZaZ4S9d"
    "RbgKAo78vg2E4G4df/79JgBx9+n3E5xA1aMp7L3YCZusE/a+0Tnm6M0AwJd/j1TY05sxNBVjoxAu"
    "U7ELdfyCqQCrBrlL9+9773Hv/VIdv334s3D+JADgDJ1nMvtUzFnvL4V9MLK3LQSANLcnpJaba58l"
    "lfyDqjgWaRR3tfP9R5j89tWPqgCA4LaEJlZIBgB7QNlZedfbP60A3bv77Y9bmfu4nTmPW/Hae89d"
    "RABg2NwmoQ07gn+jJVa2W/31qhqeWKr0JmQmZxqw3M7ju2f4J78j08j9jANAl+69rYw/7kinybM1"
    "/okNBwCedzdKcNH53/W4Fbelc7bD1PC8no8uxO8Pve1WALAH6KISGARpnVeZFO3rSVoHhel0/vtE"
    "81VazCfnpwsyHe+S4o1nJ19/t6rG71ZyvJDSzt4gGe9VybgbkQMAfuA7QVXHn1UyaBNNeU3jX3kD"
    "0W47x/xqwAHovn9hgWi8l+jJv0t0aaH2Ud7ZR/RcI6DTs3h1E9o/vkSeDci+k3f+GZ/Vf+evw6VH"
    "76t1PfZN3IfJAPAXYQIIT2E1BvDpP93gset8DYA63hN07FHLBFCcLZdu4BRWAdwpACyxTkozNhqA"
    "Zd1IYLnhoSGQZObZ90H7HvApGyj+qTN07+W6B6D5l+z1brgPfwyAndGNrsxs1aqPvgpA0nDZA/BH"
    "WLosoHt16dG8d/BRdkAH1LF/o3FAX6293pmO4dt1B9acpmW1lTvAas/012gKAdsD3F/WYEMrBIcH"
    "FP8EminN3AOwdW/vdd18S60rDz8qGjeo5py41p0vAVDDFd0DaH7k/NwxkR/4GJD+NwRLmsJ7ENi6"
    "t8UAvEC1Flu0uf3FB3ugvlkeM2vf91UAxo0fTPHKWZeBfqtf6nyulSEHMFK58/9V4EEA2C/jdVxM"
    "0wgCQGnyBvzfYp3yNS6+09dEEKCsM+pLOwhiNEXTVqPcBpSdWrtUDJtRM5/uAazEZGP7tVLwN+oc"
    "uRm6yAA4cYvxuSciAEHAToX4w078Zs8LglErfp3/cGZ5HpAvhks3vJmXr5ZKALSwjih7ITRHEoBy"
    "qdM39Vad83lxRfAbfNGRYMIoCEBYZzM+WeDmusA4IMkB8Iy7boLPWetkE5C7z4OfDmk2e/5Ip22s"
    "7LoPYybEfy0bxZaMex50S6J0Q094nqu7Yc7cELnn6bzzu7W8tP4fnMfKLj1PYCJ/acZrayWwsn0D"
    "22Q3RQ1VpqMMmp6LFXiQb8oDR0ktgIfsfY+/znX72PUAgFcMv/6/3sgHOa+x/RebsRNAb+5RyO0c"
    "7wngxJ1s6H7qfzRLIC7dR3e/0Uq9BGSn+y7Ahfi1Z4Af5R+d2y/Hma5g/0AJ+WfaD3+d9NiBDYWz"
    "HRVaVwELssDQqTlOrCr1loXvA4CSe2eDbP3nAObeWIG59ffEigcAp0RT2x/DAzz6QEtURYsBDgZq"
    "lybdICAW/90vVAC4uq3/DJvVbtDq+/49PwXsvBlwHl3r/MZtDEDvLxmfK50/FAgCoLXB0dporwYB"
    "ULtevjoBSROe53lnf5yimO9OWTGEKerV2/xC1m71cwBgNx3SZCd1AFDavfqtD2lliwgCQAs1zJEV"
    "1ukNvf5C1hvWzzv9m2W7fJYKdGtzc2bXUmFoob79+To3m63FYM6xEgC2F2+mXusa8DnL7Fv38yWl"
    "PVOPfRNzti3r29C8rA+tyyKpUwJvYnFj1wyCIKD+rT3bYae1GfohiXLp8ixzP2vBnfV/LbhJ//7d"
    "2Vc+5v44P/BdXXwODaP9ueUB2KPE6ptkzm0Ae+SNC6eUCTwA09Hg3K2P+njCAZR2U7TMmubP7B9a"
    "mf0yMQDOy/J7N/xx61irpjX6P+aIrCMetWvm2psDc3b59xijFuyZs/mNC+1czaN/78dcI9JrPfyi"
    "3onN0clac3SxVXvJAlB7yhtH4m6+wuVbs+bA3bP9CzstMSdyVwZBADAnLJlT++nhZnd2N1NBEARH"
    "7rul/TzKcjdGscJRfuQrAPBi7nx8iDpvQxt2blrj9rYAQEj7B7fup2mncRAAYLGYk027RtTuPgq6"
    "+2gQ19z7zTee4EEQBMNT4scM0P+8O6pkvnJ7BgCsdK9apRsDAMvXbjr9qxkA5ML9MS4PZJQ1gqwT"
    "MBgkM83RlA3S6QzPnBpTKoyKCoOiqVHMasH7L+8sSLmzj9T4Y5KIKB0fDuhSN5XjfV+pHACoa3+m"
    "HIAXFRdYEZ0DYC0XydNGEaUfKuj7P6/oOqsGFQcsllYX5n26MJ/Sc/tqf7RxQ1XPvZuC8c+m9SG7"
    "AW3yuT5ylwEUT1hrKJjnAWoNTI0CrTkMdCUWwbTtHABcPsrVbWCXCYzRzDsWIsbwxbB02CoAHMm3"
    "nePcMgEEIm+3zpYHOJClZzNLjQDSYSoNbLXEXwsg0xjuJDcGWbE2s0EOXY19t62s7H/H8AADeKjB"
    "Etpv90JvjqDCR3Xn9+7SPQDqEtNP9ZvDTtyMO/Gp5132x/+HB8BjWqjECimxYrmzAbGPfyvugmJ5"
    "Gw4NAawvwb23xzwdkPGK6xx7HLYHyGP/Zu5UfOpqK3ZyOtYTS3l7S3cgft6iTA3L9v5hyZzZJPhB"
    "kW9/Nm99iF7c9iCAiyZswEjSKiB60iC6ZNTFWE7Xe91o3CiizUZ1PqFzta9mNGg7jxQF7xZVnZFe"
    "V1muoiljg7pas57IGf5JW0ZE44WU/DrPnDYbldxsUNBr0CAIqt7HiCiiynG6snK6stYKplM2oOJD"
    "Uk513pVUET2DiK4DvK6nPGWJnNa7dzrHqJhpV/0dgwYUvGLUeu2pMV4uktONIqkRpFcMSl+ZL5Lx"
    "stONyPeJ+Mmiag8iIkoKlqRTViVEpKZ6RVtGJbfmq2g8T0lU+/VVPk3ZuFHIcaNIapxivExqZ1lG"
    "4329T6QJESXOQ1TxOSfUoQC9EyZ6KxYAoPM4HzUFRk04b1SijZl7SgZp2znwqUybZcieYTAzC83w"
    "AR6W8Lwg6OUYdvoKQZChdrSTCILJaEVktnpswy6EaVcCjiP1lSAIgcwwpSZsqeHuOhfuH2XI/84f"
    "m2HHWjdzAnhP0tU1jA2KyPfpdENSr5H6vu/T6cNp0uvKpHe4SJwuJZHT9X3fjwKne50FNGBRkNTV"
    "jR9KyemmUz4YGzP911spGDf8hIhoc69PvcNpVdsWJuR0B4nTHQycLhH1ln0iSdXpwwUbyKI9qFGZ"
    "QTui6tAjg0H7/NP6+Srx/WjLOB8I7BjnEyKSY4PU2Khnm50GBZcaMri0XHU39w0iqlGWJGXVwIno"
    "WR75fkREhEGatK8HU3b+OqOCyPevzyfpdL6IiCjaMaJqvFAURMlkvJDIzUbV3byn6m42KH2292BE"
    "pIqi92zEHynOO90gICI6pxfVwEiLqf7nKetWCVGU6n60BuKDQUSU+vy6mrLzfkJExfyfp1NjultR"
    "Wc149vSaQTRuSBrvmxARFTvG+RPnGlEwNmjQ636lxj/rGbQe80F1gaVJ/fk004ek1Svz5A9m1Z4H"
    "BtF4wadxYzYhG0Wy81YKLjWINvcWBZEf9bpSzOXmM+06+f1zE9bacHkQhNoHLoOpEXS1kVlqQ8tt"
    "DQBGEw9g8raQXa71wDkO/NpdvLRCD0wLnRr/rGdQANA7GxlrfSEDAH3lWC6U+IhQjA0Lwepa8Erc"
    "CUBXbmxLS8xmzzx36/xPq8QI8BwtNDGnTj8hSkBH7wnozu8i8DxiTzzG9FN9qZ3qk94a7uZ5mqGb"
    "ZXuU2E0N3Ai9mpPvuxlfqv2Yluz1pd2UTMbjYXtXi9TJ1/VO+Aa983norXwmOOlzaQMvqWGpN0UO"
    "gOI7BH7tvjr/0861fObLJJDO8E8azxNdMhKfKCp2GkTjRppsNoqk1xi0k0k9xilK6yqUJ2dIZJG0"
    "KbquU0TpbAxfmc2g9RwT0HhhVia7MaHeWyvqNYh6jRtrwVPCuzLhXVnPnkGh1ygkDdpJ4Pvk1/z7"
    "szP+/QISOWU1X1DNUbFbtXJsTGf813R+EBXJkw2iugbG2EjSC/N+NG4Mol38MynqGfQG/LN2ppzq"
    "RFuMaMuoaDz/qtupPiBKNhvF5FJj9s2CaUNGM/69q2Y5S2IOzBEAAszxMNFbAp4HXD0R56MmnFFz"
    "WO5R0qoXc3l7BkhPIqv54FDnwzvfCTN0M4hZzfdcq2dQABBGK8l3Dbf0FYE9dAx71Hq5THRDLXjc"
    "/XaplZYkQIRvBFeuMGf8+29rdfwZYg4H3d8PggAxnLjNnN8KguChEY6GYJ3DQUgxH7Y+7g5n+Z42"
    "j11dxWJU1LWEEi081eZ9x+bAHbKOQ60/0LN6gOqi97smnI/XA/TUGQx7wLAVP6fvf8g+CwA00a4B"
    "pq4EGUoA1Hdw2xMc1Hc0oLSLIAgCASahxQeXZAB45V3v/4yVuccJYPjWfmFnzqcVzMzado6fqHMX"
    "t0/MKcaSllj5wxYADE/A/M0HdbXRYiqTt6k6jLgvF63FDADMu6eZLt1/Xt8Ld6yLzGk6Wa9zIOvx"
    "WjoRrveEAy1uxbzf6rehNn7VpWcGpDbO6QBdbAUAbNxV458AwM4lKRvUunI2oOL+mRDIoHTHKJ6N"
    "AGA+pXS+qCsc75PJjkHFc40gmDk27wrXJQ+qijsMehJdYFFR+xoVyeDng2pG4xdXGFWV53mY98/7"
    "/zAtxg2fNhsRfX3eT87NSxrPR0REQa0xT4JNPve7ruUwm7X+KBet/oRxXWfoh5qtWM48zzOh52qj"
    "hpUqCyj0kEkrZJctAIQDOd91IhR1vZ7MQszgxAQVr7RbIu9zAOpshiWUnAFQGRZIMCWD4EDhbHOy"
    "Ykg3tqSDNrk5J0uwGvR1j9meB5PzY0ILn8Tp/+1L8f5/xfXewUd15y9aSy+HTebixRc8F7hfbb7Q"
    "Gq7lPoAvi7Z58tiKrQRsJVoK0JRa5G9RizV25OP87zrip0UDTDeeDIv4UeacWtGdv1ghdezf8Kst"
    "SOtLAEu8FFm+6D0IqUaaVQzZt4th+RdK5B/dPwQKkb/YFLkDpQ0feXQxbuEtGwBL6JqRRE/PJ7Xx"
    "9zxF4/lnIypm+UL+gPcaNXfZ6ybFZiMi3o3I6RYVG1RTNihmc8yWkdDOfBU9Z+sJXTGqWV2hiMZ7"
    "ia410oJS6j1Yu5v0nLq4T9o7TNQ7LGczaFT1urJwujJxuhGR2jKIdoyCxnsAHtCARzRgyWSKRJ7T"
    "KZqyJIiqKasNEKfGeDcTqxFV43miFxtE48Yg6e2rql63mM3Zlxq+v3m4SHrv5iepwK9HKUvSab3v"
    "KqZ6VVBAfOD7ka8q9pxB5Pt+MjYoGRsz55cGBXUsSnV546TXmNDm4SLp7VubcyJsOyFYK37zRSdM"
    "9Gas9FbM1VDqkHXMuUflmucFgS3BbIXs7RKZnQFuKYQT6jP0JHc899hi3hKHHCn03jU1quNPYi2R"
    "D1dETlBDeJ7jaZf1AyQE8zyyKYFNG5kttZYpAVtC2JJqoz+K3edpzcQk7vE5d52grSvuhgQet5kT"
    "OkMnBPK6Kp1QoTvsrCOIgdH+EsYHY6GpWOgf3CZbgu1Jke9JAeD4avD7Q5epPnfDHNpveY5zRqD3"
    "lbPbvcw+0PndlzLX39MOgiBAZOsHo9wNAoyaQxo1hcmUgKESR1cbnClxl6ESgIY9ps62zqrhRzCn"
    "5+LgUg75j6UmbWmtHF3XMzfXXORWZgeBx+PO9loO5nhe3wHinsA3WgLxgsCoiXxUinJUAsDoIzb0"
    "tpUZd9t5iU8xIV2p0x2ZgHUmDt0YQOaFnufZUeleXVt3As/D0MnBettvONuKrVErfgNr9QmtPl5q"
    "xcCKTpaac3DQIAHiJ88HGEQpv05X+HU1ZRUVzC8Kw/d93y+mRkH+Qur7RGq8kM487iWN933l2d6+"
    "3yucbvRsrxvJSw2Keocp6h2uKh5UFQ+qmqeNinPMp3N6FZ1oExGpasCIdubrOP+VBd+fLvjFpYZP"
    "lxYCGjd8utTwabMxGGzupWjzHop6XRJzcJ5YhdOHwOeyJXxuwlrDEK3RWWkyz/O8bWnp4D980fOA"
    "ia6szFK3ZUzeFrO/PnpIK9+sr4b8EELHit1sMbayxdjE0InZ0InZsHWM6x2R661RqK8cLx/dqNHn"
    "pg6cLcwhADBlMxRerEs31pUbQrqxLt34pYxjzYGzxuHEi8Ac3O8NcTo00XtC6L0noDsxjRyIYJXX"
    "MeYhADWuhaWm2K0TX9pS4JDs4JBcgCeJ+hytoUOt4SpyVyB3Bbad0NyNQVknhlyuNWiXbUCf4bej"
    "D2N11BTtPWpYflsJfLsYlt9Ww3JPAfQdqH5nhpnRlEU0NZJotyLbeNmPxo0B75Lv+/6gTRMqNhu+"
    "7/tFUte0KMjp+uR0VdUeUMEGUYHB1CAaGxSMDapoqic01ZNovDuDLgc0fqvarJGZotel2cxDlPQ+"
    "MUl6n1CzGTTYjUUT3t0yChrPVzQ2KhKgAe8Wz/KILiAJ6hg0mUzZVI/UzGPKp6pOaq8d7scNSl5s"
    "UDRupNTrBlXdLtXbVeo1qCoqfrKo+Ek5rWfQdMoG6ZRd5yd2OfdZYj4R0dYC0bhR0bhOCLkB4Nhs"
    "FMlmg5LNwxMSc2BOBuaEmOMx5niYjFp5OeJyrvaGkjYgdr0pyZbD0pbI75AobSnghoATtpgTxq6g"
    "+OiQ4qNgYI4EcyTpp2Ix86WHztfrSq6Q+goAWz1f+9Irgi2RmRIrd0jKbKmtLElN3SHF6W06dhp0"
    "7KgAQNHYSNV4mZJeI0h6DZn09k2SntNIfN/3/V438qNiwH3f9yM54HUNm+ssklVbFlNdVVOjGEz1"
    "9iCYsoGcsgFRMhkvJ5PxckCnG1GyuTel3sdS6rn1lrI63ZgQkazrqRVR+yRR+4QqWKKK9kDVFesH"
    "acHaVEx1oqlORGIO99oQZOfqrk9lq3d9SlrSLc3Mjd0cQRAEcQe8zK8+EQRBIKy+k1l9R2DYEtqw"
    "Y2GktHKkNIzUqMkzfcXK9JVAxKWdx6Wda7Ay8+47spfghrZ06zkUYekKQIiaynwOnRxhJ9dGnVg7"
    "uxi/4Wwrtvqt2Oq3LirQnALNSQCg2j/5Q9+fad8rovkJkfFs8RoWWtDWrhZ+H02+vu/ZavxuVfhE"
    "xZPvVmnvEXX+ib2JHHdpsnmYnn+h1rurNaJpmyj9EJH/IaK0HVQJB8CoqO5Pq62aW6crC0TTv1u3"
    "cYPoub9LtHMP0ZPvJnquS3TpIaLNW9YvBvqSeZ7n0kiqjczyPM9TGzfi3tqrjr7A6qsE4s0yLLZr"
    "zz3jQd8hAGildjlgzOLALMX4NaRbu6Wy85Zkw8bWr3JAL82Pt4Z367UW/i775j5o/F/DvAZg9Zuv"
    "ZgDwP7vROfq13usMUPpTCIz4Pkmf77FlH0Eg8v2vYyJXl37UN75VP5nZ+dxJXruDnHr80WGtRfrq"
    "TarXWV+1VyXZyXthlhDvfu1zZ9p9HUAJzuJlDxorIZywpz/gwPvimdPFj2BMcSO3/CN8dG+komrz"
    "OxL3Gq6scQlX8ht5Z/3GL23Ousp3aPIjrhgHBGx1/MkgWIJbHv9VScIJArg3a+EF/a2ve/R+I5M2"
    "Qju7fY8alrWr3fDtf/MPMwDi7a/qzvGjJI+iZsAnaDGIElJvq/Vr8ABdibt+ItvgW99H4e11Fff4"
    "yLAT8zrbIHudWjG5wbF71iOBpZuv3e45agPP6wgCcET2dis2gyCQdutvKE7XJ3+r84n4iEMBdOqv"
    "/+Ha41YAQHfif3Rz7kAGQHwAkK+e5wm7cWDRDXp8IMw+te1BaGrdFb2OBw+KOX/5E51OaGXdU86w"
    "C2/PEj5+8SQYcG8xumrawe8BwMrPf+bfojbVNOuKEt61mjO3wUxg5b3AIYglzIrTA9dnEnznW4CG"
    "2+73zQDOBvUrFJebWgDg6G2/bgIbM18OPgT+258CdsrZ6/M1N2/e+nxeXuKAhyyz7FxXGx6AoEMb"
    "+K/4aeouPKj2yzqwoc4e9AAUlXnDpJxlf+vxGR5ArYE6BGXP/OiN5vC/pp815pmTsgE2U1/rzRvP"
    "1LD1t73fKe5SCOC+Dz0LLpzH7wLAhTP6P9/LEqELBIATuwBzHj8dALBnMNmsN+bftp/NjUR5AKaG"
    "E96XuZ/xAFBmRY3/8x2l4xegPEAdvwCsvOMzwkNgw/20fdPS8be87mTXN2l2uyBGdS4iowfK/5rr"
    "rq8gBkB6E6A9Kgdsnt+UH0Libz3P67ETAPiaA4TOsBUAGDWZ883/ivNpn+0gAFL7BxZ0nHpeAWiK"
    "5l/8JPfRLfvZqeue9L8Ofvy+bScA8KY2/mv6CQ0lAkC6Gw4Q8lEAoKM7p25Ic4hvmS5y0vd7J/yE"
    "DwazFhEfRJXu156fyxVtGRVdA4D5utrZ1nyhxkZBO/sCOW5MgvG+KHnqMEW97iTajCjp+RX1/IJ6"
    "QVU5VVA5lA545QftqKo1R4Vf/XxRPQ0ARlqN//9lfX+MG9d95+dt+Iax1WapmLOxxKJ7TuMeWju1"
    "XcDt1VyUj66EAnZQjusZ56pNoRTd3sU/EK4PdymKOLurnHM9QAU0LpQ013NgFZWd83uMZgic7ipb"
    "mhkGSwWogev6Kh/QxqjtK1eyZpguN1k15hvqfe+P4Tq7OgLzz/xBviEf3/u+7+dXOTMXylJfKGdm"
    "XNY0PpDSuCzN5mKs/Has/FUKNjDzIhf4Fqs7XYyKqzKq9AQO2juZ7Xne0UmS5Umn5Hmep9eH1Xx9"
    "WDWlLU4lxSei8OyZLIxEdTgbinx2uaIo7A6B7ggIGcSzDWAZqDwrxNEJ8iKi0rv6OmeOAECVnm2q"
    "TBgPwlRR0hzQNkAe406Ra9dtTB5dA8qxfrVMenCA1OAAqUEtTv15nS5S4LczokDHvmhLI6WUYjU1"
    "00s3Izm2jB5bJhtzHaSxuslI3mRGvooDWXz2QEZna0Snazo6XdP6dE1v/Vlq3p3zvD84lpOPgMzX"
    "Pc8TQXBTBFtGBJkRQRY5AUUiDsYiJhNPc+1MNuaAFZhLM6SLnLbiGlikgl29pByLQU1OZZc69ms6"
    "9ud17LSNQaR1M9JarATGBAOL5MDSwcDm8dYmj2mTx/HLlgletozetAxl2vhflNIEihynRqSkTP3F"
    "LPYXyfiLufEXKfZrFPuLW8Y/vqtAH1gmGBwWcJ42kVA6agY6EoGOREBjxFukx1agpNTBwBqXp9C4"
    "ZYKBZYKzZVKDsok2ajn585r8+wKaPkEtJf8psZJrsZppsRqniIKURXrMoiCbntyDQLMLZR0Xb6mL"
    "/qdlaFCmfLM8HVnZBDr2F7difzGPnWBtRmxh2w3AWh2+Pc1r4yIEwoxXSp4XVoaH1vVoOK0/J4fy"
    "3sQ2IJuAT+UluEMOZwiMGuFSWA+XQGHVHxKcEZWcDLbfReVUZ5u1ulhHxgWm9adJRrtOIjathzYh"
    "tInI1klxz2AyKnLt3BF1HYGB9Vo2sKLdpgudniM6XctI+m0KCn3BOacA+YrTezuLRJBFIojGPKab"
    "jLbGllbZmEXpdRZlYxYPLB0NaikN5mIq5DUBna5FEZ2uaaJAkV8j56kgIMqN+KI2YiXWzSjWYsUE"
    "TptIrBglVlI1zbXTY4uSmTkaVWxiuY3X6xW8nj+2FSaPgcDDFqAU0FtlpztQSqltf1jddobVoZNU"
    "t52kebmO7LuNZNJ/CFQd8gVMrAVkt4tDAOwhEnskUB0C9hChPRRQE3fKL3TDtfOYANRzGLrOtjjf"
    "6onzzrA3WUvQcbZRckbVapFrp5DZKmH0wtY9+D+/PsQ3VsBOrDQ7X1/5Qav8wfbsT48BoPNrnxwD"
    "5R9ZysPw9o//iF/95I/41U/+cOb6z/3QeuFvPlN67+RpvHfKx3tvnA3fe6eOq4ft7JXjE3pl8XE8"
    "/PTj9JF49Eg57T3+sKwmvzEOPWTJ4yUD/dM7JQbFj97xI+g7/tHWH3+nYmb/0s1/etA/esc/2ps/"
    "9z/xv9852T8YfJk/Nu/3ZoAXuaj8HhfOdzByvoOS022g8aVeZk+5lx6AvA8PiEt6fVjV68Oq4Vvc"
    "8KH7bddxs1bFnTxa9D7RrIYQJbHcKImQwQkZc8KFvvO1I0nla8M65x5w0PsxB0r5+QngOec1Bzc2"
    "kFdDZjy2ntsLN0w1tA0QzqK+nN3It1vADNxhpeK+X6m4Z5oT90yCUgdwwvWOM/Q8YDZ0h8CprlIA"
    "rEaPeKNHvH4CB+snYp4TeA5hZ2j0nKTeF0ne5zhov185yBMwnqB7uoPt5zsHtk932DU7AVDC2/ZP"
    "MiJ5I+SoJ+H5R3qsW0/AGgnZj/Ym/Tqr9wTP+m6S9QUwA1uPSof0qGTnScPOR8LOKtx9vxW6VQCo"
    "95x1wO2AkxKJ82Kl57xY6TnvVvrOmabd6GW9xtqEP7R+jTdKmZUjm3GAI3pduGkCN00mB7Mrk0qG"
    "SSWrhm6mlFK07AJMhFwpqIPOM7hceVEw8aLgzhkw53MJWzsjmHMm4wLEmpXAngCgyHnaRNO84EjE"
    "FDUDiti4AHjkTjkKChUqkVG7K+igTGpQJtqYT2njPkkbbb+m4415HfvzRLFeWyUtVkmL1VQjzor+"
    "Z8wiKYkyzSKibJoRSQPLBFcsEwwsnQ3KJt+0AjU4oIMr5Y3Frdg/JmO/nVEyg21/q/Aodjp82xmB"
    "tTpgjYxPyPM87+07hwD6dQA5o90V9JAB2RqdT71/EJ8aIqmocGnUWP48b4SfRz4qOWdQcrKEO1ml"
    "L8Kq1erickPyBc8Dqlt8AaC38tGHXkx2nizM0fqsZUCHCO5cXsLdZm0pycM/4FnXSYASwDufX7it"
    "8/mFB790duHBB17h//wG4D+zhLwL4BuOfT/AWk941Dlc7zXC53qN8DlWD5v9RojbdOCcz9DIr11b"
    "YXnvFxT1I4RsUr2GifvLR+H+sj9y/lOdodPa9o/YNyoKQGpvV2EfPfV6Dtjih43uSd5ImKgnol9f"
    "O4FHZQ69doLq5wOkyVN4vd/oAiU82qURcRrNEl4XlCmRWIrdHUalq08AYH8kcmDoLCf18Hhz3T9Z"
    "3fZPVuEkzshJ1r4Klf9RLib9xuUFfrU0VDQzKYHC6gQjezSBPTo/sTOFUsZw97Cd5ApA7oZLwEj0"
    "FWBj6B9H4ryD0H+JwfkHdNbEtyG2877zvczO8Y0q+NABQOlrPCDJY/2qpYuak7X1Drs5xeJfI02k"
    "BgDQJjUoa9osa9o8oNOz7VS9CwCHg8CfVwVZc/6/EgUUNMmssIAuzaQ6Kjon+hLk1GpUx0SKCp/P"
    "drZTTs2gnJrxgVSePUy6eMtYbZQzU+iPFFFSQu/XDjqoM2zXWWn7Nlt46DjbY/p3X66ItwSydX5/"
    "BdWOBw98WO3o9U9X9fqnq3kJR9/oFV6Rzckd4UT0joQAJi0And9+YvKR1f+Yh4+UErfOXAAjvoC+"
    "WBcJeH5+ApQwYi5w2CmNNGO2qTLkVTAxvCO0BTAnhh5V6aDnQXXdjosSSpOjQOaI8AGXrd/15w7g"
    "Ifi/TufL3G/9zsE338kSBxQ+7QLo3NFt9O7ljd69/KHePbYAdSUAPmn0+B3DOQmFbgV4NMztrbfu"
    "5kBmo3fX4SJLOuOdk8z5++fWfyU8CVQw/EITTr/O/2L2eb9f74nzj/SY+9c0efZJAaePj/Ua4cKx"
    "Y0B1WJCyYxEU/ulOTJGISUqSF9jZMtFgjmhwOAqIpnyhYGx9WIMe12SC8TTTdmA1o6JGZRGRjge/"
    "pGN/MY39p+UU9Zbke5ZRU1EoEelmRGSUNmypRrG/SIHfJhPnNzgRkaGB5S8WDEXRVpSAbs7E6U0W"
    "0E2L1JgHW0RSfmBFLM4uMArGjAKiwZwkIk2D2p4VlHQ8zTEnv+a3dSxlkTWbUvNDP+XcTL1EDLvA"
    "4mzMosLDhK5bVPA/rU3LBFcsCgafMLGhwZyRMqDYXxyUKUgLVxRKZpBzoMOBawzIZipJYbjATnVi"
    "qxHm/Ua4AHRcD8AIHXfPCjpCveMCngcKlwr9TqFkYNT7DM96gudMlHJrQSmlDub8aCuc7Tc61Gsk"
    "ADr2SCnmhZlt5xdxN8ULc/+UUIKON/E8YCtcsokcAJmbETCDb1bBM3uCzp1JaVgVUIqJZ9x73bAp"
    "nI7HnE4VCKuAUizvtnqNUPQaoWD1ZYzo2ZbnKSXyfsO+2neUUupQBuDHvIQJrw+35lgydW7eyuzJ"
    "6U7pB+52/4b/HYDyI5OCs+eUGqHD6+EVXk+A17pHtgEHeY/6D63BViqZqycAaG0lIL8t5VoUkL9C"
    "ikhKFl29YVG6Y1E2sIiIXaCUiMbWONhYpGBjkWJ/kdTYGktJRIOyjqfst4gR5YP5ODg9n6p357PA"
    "LxJFUv++bGzFalCOs0GNyIg4IKLUiFRvzOvYb8fBRtvEr1pBTGTkTjmK/Kd0XGQ8UQKyZCC5TLVF"
    "esyDQude1pQ2b9J1FqXXmcl2alN4c5Hoxu4Kamk1qKVEUgb+ogn840GRi9QOYiegVdGm1abSq01F"
    "RDqKGKlLTNIlJmnMtRxbioja6bhMZmylZqcs4w/KRO/WFBEZ5bdNvFMuMs39ml7DppWqTes1tWNR"
    "tjlX0LoHNS2lv5ppJ061WCW/yAaK/fbWLjqYjVkU+4uplESRUFoNPkFSSkl+TZtmbm6yIL1Z2EtJ"
    "KYnGVqxMIfNWRgRqMEckJWWDGplLLKOxRcGrfDfCSBuxkqoxL1i7L1t5MqO8isidI6W3vfx5dRRK"
    "AUnH4bCdMwlfCyvcCePQHU11PRd/UoM2etR1C1+UnhtiLiXP8zx4GfLL1Pmx1Up+zMGL6TnBNXsL"
    "zAmn1/KRjovc8+hKx1WMNZ4HdAJWZ2K5igzo5j3Ra6DOAXRwZ76DGVRxx7LNZpddzE67O63uETio"
    "vu+i+m2eONfqvcJwtN5vlVx8WIO+WO87NqDUNU7M2ZOLlFVRffvOift2FeiUGDxvRKHLSnB/DLjX"
    "SnDfr4S2CoFOo9sqhczp4LJzRnDnc1fZqTwB3JyLUuVyJakDlLDGWoLmbi48pbt+8kV2zyUWqzEj"
    "NWaDT+hdU/g9p/iBNajlJCWRX0tjrTfmpZRSGaftrGoSq7kRKzmJFVU4UETx1hSHj7Mxu8BjGUup"
    "bliGgnzTMsGVMqlBbfAJowMi8he3iIIrUzbW6Z8VsGI1tmLaKdPWoFyMc3DASKVj0VaRCFQk/PmC"
    "Euu3JQXZTRar4vLbxX2xKknLYtanNKg1TaYtk42Zoet8rGiaHKRSEispiZXUiKUaFdT1WrEnvRaM"
    "eZxpy180JiAaizgwcVBEuFO6cwilg43vZPYIL1aRvOeCPwEg/3wPqjU6+6sLzAn/5Iqz3Gb5IoDh"
    "V+7VnYXdGrTXePYS8hBAb2U76zkfMQqAE/9qmqF6lU+QHZ7w1+dYq6g5b8xS4jl/nXjOG5Wes/Yo"
    "GgrAKy4QuGiEX6d6KPDQs+eYfuUVoC+QDVtc/5sR2O/U+TWU3IsYuhOM/lLx8NdHFQWAXQRC1GeH"
    "RyYP/stS5v4HUTEKgFWfSRNsT2vQbefyAksVAL7Qz1E6saAAuPho/Z2hraojdNxuaXjkfCtLgMxN"
    "ZrPD5H7JI/cPz2buskNDBSSZSWiCjt/y+847yci53Ar/WQEoZ3YethJ/GWAlR4kE0TS/eIvo4bTQ"
    "Gc3F2bhsiMqGdDk24wPnCv378Swe1M5t/dU8mcEvEQ1q2TQ/iIKN+a3YB4B2QP7e/GJpivrTRK89"
    "TGQeJvqn3yQ6Z7WLXCHeTq+XiS7NE/3VPNHOgWmu0LFYbZTlFKOXdEbs7yezkS089ET44C0KxTcl"
    "gH5s79EhTioCs6EHoDqk5Q9mC/1Ox///EAnmAphgL6Yj4JQy7gCW6N+9ZwhJFU7zySaUJ4YeebCf"
    "cKF/3w1v7Xuvv/nnDuBR5XO3tNa7EQBP7dO/cxeN3qsA7ORjvdIEHgfod+kWHLaoP/mo+d6+L6TO"
    "/2KpCdj1/l4P3ttayHurAh7wsV4jEaUnBeh3e5Nbx0nhKhTseygp3+JNrQC4rL5H/07MwWaRa9+9"
    "hz/0/Al4AH32xNp+eNx+wfWQvdV7fPnWR/9KE4nI6yf2oi8OUl5X8Ihl3DkjJkLBboTPnb9lnPXL"
    "o5EHlTx0i33xPbYHIKteq+7H3zvFblPp2M5LJ0IHgH9q/zjtfDlSAC9Nus09/Apw3QeAjN3e3wsN"
    "seSgF3pQ4q03PfcPXxr1PIRn6idvxY8yVkkA7t2qpi105YkIvT2eDejgQ5dr1x1SDnSG7q0I05a1"
    "4HlOxcn2C4nDjH9aKbD1rb0f1RjRsusB3N4Kl2wDuwXUD2WlW8fZ+YKAUg7Wlm7htLSUUgr62T26"
    "zlEGXeCRLO+13sqYU0gub0HHc7vgX6i94xmh0lkiQJRmn2nQXhVove9AKSfPe9SvryVcqYqVJ7fO"
    "z3AFEw/U2V75z7d4U1sAhhO+LyzZ/We+oAB4V28XJ+qho2Dz9q3J08NSl+B2Ce6VvbcbJ1bhgZLZ"
    "0X/Zo7m+IWgmDz1g9KgFOP+QwUNYb4S/vG+cIr88hKvgPPK5c/sA/l+xLQUPz//eoVLBVmAVTCZ2"
    "+J5dAsDD9xx278E1eADEJ//Fbu52eB8Is/fLkoO8FT974W+BbQBJDiTIv5e9JdDh9e9V9/Ez7qqG"
    "HKD+e19o4ud/7k8SgTw5U98/zt5vVxjgIXzg1ukphh4AZ83ZM/tczIZi4gGIw/uG7stFHvxjL0zE"
    "fsy9mNILt+jcs6oHwAuz6r51xV5zPCg4TvM+wKaQA9y5mokzM3s+uXTxqAsohdDd91mzXTdWCvj1"
    "rhPu++16RZZWo+ckODrNg/eyfdh0NocMwMVr+x69KYp/JmF/XINTP+8i97yu7rfCvPQQnwAT/F0O"
    "zOz5QvW0/pzt7ncBz/oF3vzBLeLqa7yVKEXqGm+UcPChfz+tP9f2crE6rkqA0Wi/zj2hsKIU0NnN"
    "O/pQmzwjKARExh4AZipJA3hcscYySnu2D7q8mtnwoC+PNvb9GrXcFgC+b2Ofr0vHntgCcL5hTyqZ"
    "7ZxaAMBKp9Ze2jOeZy+MbAXP+eylvQsWUX916Crw3cj23ZeqDhNPwf1TD5VSIk6eApjon/2Z/fu7"
    "DQAKD1Xz/WQst1BwrC3t02uH4VJoAyiFS6NGx3VOegAy5izvRYh5EVZFfGEftP567d4qPOBN+5b9"
    "0E1ceEjD1eTRLrwvMQD2wu1AaS+1zxWqKAr3JxAwB/kigG5777MTeis/EArIeyssBw6lRf1Z378f"
    "7fY8R/vn531EigGvuJOP7PN1byMPAdUX4dbFtTv0EwD7nYXbnFFpz4h6K2g+AYBdeuAWnwSQAmCJ"
    "eLKX78EXzmdF/cmvMvTrTwAw6O/7LdzkUaXAGm5vH0ENfp5/O4QTms7ewIYFy4HsKKAM5+polvnL"
    "ALOdUISItLbidMylHlv66pgJYC5ObyCK4mYUp+wmjWvT/uciBRvWmMYW0aalaedYlr1c3I83jqui"
    "Fp0PtB/r4uSem4iZaLf+vATSRpjYiJRWLEpvAgDP9asWBeNyShfKqd5cJHrVAdgxHW+UU3oZAA5I"
    "GogSur9WEcOcVVha5TzzGk+hJ/CC+Ivjo1bXSRonnoMYVjsSQH9YWc77n57Nk3ur+TrqjHpit/68"
    "fCbkHpC/gVIJbKX0JFZu26rW+QR17gCY4F9h+WTPedaPK1/7svg7VK5xB+DEszz5zHe1DZhqRziA"
    "NzmG0EHjYW3jyDEH+e+54a8AcbBxzATvHjOx/zTFTtF9icXpMqlBLZZ+LdXB1FczVmPmtyU5q6kR"
    "q9IEqdJSShmrMWcmSKWU0tJ0XQU7ZUk7B1Ia1FJ6tyalpJQ2qj9JOTaKjP+UlDo+FziVp03ktCly"
    "YlXoaImIgutcxIEhIiNipdcQFHjM0yZq0tYlFpCURhpWYPGkBpYmkmNLSilNMCgXQnYd+7UoitMi"
    "B9AEg9rUSplYtGVS2qhJ2pjXU846UUDKiCenGZ07ljEx6WZMRCa7zgafMIFfI/IXMyKi4hwdXKn5"
    "i1uapDSinekE6uYM6ZssMGOL5E5ZSSIpd8qRCHTUDGgMSgtPdyIdb9TGlgluWCYYWFFEqZqGZPvt"
    "jXkT7J6dc4MoJ7FSaABYJKVUgWbRbkYni1Oi6zySMkjzQblwtDO0aemg8JskqWK/PbCMlkSxX9PZ"
    "tP95zapMMgaMqiUGcO/tO5nT4duti7CcEKPpLjdqLLv9enfy0bw3OZSzIXoNKKXUqNFtfep95Erl"
    "yu0AyJkopb01ZHzByXjd80BOxpvOCP1WF5cbYZIknbl7PQ+gi+7ttEyH8mRi0/pkVPALmdvoOv08"
    "4QLAk1kpxoz7zWrCO3eORp2lEM+4TkmpSr78eeEqLLgBEvEMipQ6pVjeO8VaF/HdVih6jbCalCqF"
    "j03eX7hNdwCUYGciHmUW8PrhSbhln4eytwHPgbLhdvrC7VwZuf+NZ1u5C8Cx1888HvpIeq3uf+f1"
    "pNEDep5SLvI+8UaCqlKC54nADMLfzzB6KU+6j0Ns/xYmnoctrOQH8fqsPcyGXpZUqF8HPK86tCZD"
    "JzwSrg2r207YzICO53ledXho+6tgnucxr59XJgn+9XY4dLu4WEIAtwfkYQB3aA8xsocMS1sVNkl8"
    "G8gSfPexutPhQx/I/O1qFcjgeaQyLg462x14Xn6xEQJIrTi9wIN0bL2W7ZQ1EZHetIhIRLFuvkrj"
    "5s1sZ7HwcvSPZzS2xtlUF28CI4mI2j5RPO3HBv78VuYrEzVTippaR1Ofz5uMyIiYIhEENxHLcVnG"
    "RIEalIkuWKkZlzV9cMAYIh0QaXLaRB+UiaSM1bsHgjVszlE8sAI1sIiu1LSURHR6sfhHSL+ttFg5"
    "57ezKe9CZppFhUlloTKUkihqEgWbZSNloWmOuNTXmaaxlb42tiiWMqcdiyj221nktLeM094a1Iik"
    "pC3/GNElltElJumCZQp1NBVK6gucJFEQbFqUzKgjhI6br3c84BmHMw/AxB0C7ouA26G++61K10mU"
    "Aho9H8RPdWH9xF/H8wpXqc/YOTwPEHZWIsvBgdtH+TUb/JqthOcRvW0DI7eD7SOdSU90KgUCkVfC"
    "JWCm9TxYqzfkj4QAkAPdvC+AG/WEARPn73IAItDnRDuNRKANpg7vggKiD6xC2Tkojy0TSCnlwDJK"
    "06CWTjlARFNoukaUU6ElJyOiwX3abMzL1G+n5LcDIkpT0S6s4tUmlzSwtFiVkkgaFhPFGzWKr9R0"
    "vHE4IlJaygJjCvzjJCVRLNoJLNJjXvCrrlukpCIaWFFAFImYTDPeuskHtYJs5NfSmIKxtYvDBxRT"
    "oakhMrTRnmZktNfi3KBNmkVkWBRP+59E2ohYGxHlY3aBx4qkVIPDARHdsAyNy4YGVkSFpX3BRRsc"
    "ICJtgo3DosSAfAbIuOhs2kQegI495ABzVMZb4b29R59xRw0FIFxKPEz69V0cfgLkHoCHvwJMwke+"
    "6QPI4NSvpju3W6cALnCNL/DcA7Z4HQBzgnpP9BZYQ80JwANe4BMAt+UX76+apDmXh008HuagcDVx"
    "4WgdDzG5c+G27GrJTcQ3zDeQL8UluD1kr4fIXPt+QNx9V4bs+8+xtQfOsVwB6K0M9f8QV1rL9e82"
    "wud6jaSFrR5QVHLOR7QC4NChqxaSB5H/bsPerpB9Y+MXe4C0f1hF/373gWvI37hfOMvtvKQYKHEE"
    "ANYIR/aj4bbT6ryA/BWA+o2ugmfV/wwYfqVC9QSRnBqvR7EfxdIFwC6QJpp6GamtQdns6nqyG21S"
    "7x7Pgo1FCvxFHe+Kc0iZQg1COhI/a2IfYGL+pdhvB8HLgED7JaJ2NgYHrPbWAKTHDgArI6JYbsxL"
    "tXE4UBuHiQbTt5RyGoe+SoGPmW+JPAEgus17RG8EJoBGP+TANpgDzJYuwmWh4i7g8Svf7wyrYWU9"
    "SaqdVgI+UgDg9jvvlVgDALyQb4vHvlUPAUoKp/dKBiQQPwUs0DpyQCO3y8ujPATAK3TjZ0YM8FBl"
    "DsDycAiHYamCn38LeQLAu58qy2DlHyUjAOkfD/K//82zmdsBrr6xdH8FyZuXH2eAl7z5znN/+9XD"
    "Dn3zl56rPjb/5epvfHFyx1fbk98qt5/726/aQOfcpHoiHn7soy48Wwz/1x9/f2N2EQB978rGmycH"
    "ydUQyP/mDRuJ2LzLAXAMD/7qadx1VCC7z8dbf3bEeeGdk/033/HRzP6m9Bu/sJZ89tICPvvFe2Y/"
    "+qQAPOifymFRMNUVmutWnE0zylP6MIciasbqJiO6acXqJiuqvbFF2YDH6iaTwatlUpFl5E7ByZE7"
    "VW6kJMrGlqHNso61lOblsiaKaePp3fpKx+Kpc0EzigMR5Qbt1DRXtWHngrFF+djSwZhH6c6H2eUC"
    "rJ1JKaURAd1kQS6llEXhNWYBSUlp4c2hg7OWCQrjGBMUMLoJBlZMg5oJrpe19MtEUmo5+EWxoqb6"
    "x3zXf+M6i4jIFL4jBWuUrZJmEekPuagm37RMMKhpdbaWBYM5k56ukZRGxv78Wqk1mvUAUOWiw1qJ"
    "4wGIGyMbyOxJ7kHhIXtS6dghht6k8n0b6Dze8V5wQ+eZpUmlY4+o43a8N78Ar/lvAQ+596dfu6vl"
    "A8gqre4Cb3WboYPs4GqvCQz7qwCwfXtf/OB2fv8ktYcTbf8gvPvjk/UHP471N+1hpeOG9pvueqXj"
    "bV/pnFYeFKjrH8Q02UMNnqBs6k7W9heJaKMgU5LfTo1YJRKrqXFWlWlGUjcjpVmUGrFKWqxILaK0"
    "+OZ0qh/Oiiwbbd6tadqoGUmU0sY8Eb18QE/xUwoG81k+OB7kg3l1zmnr2G+bSDytIhEHUTMIjEPB"
    "JVY4pI+tLMGYR1JKqcZWnI3nlJRSZoM5IhIrRZXRLFSZmkXZmBs1tijbtIzatHQ2LtJMsk1m0k3L"
    "EEmZblrpmBc10phrc8MiTaTNwCKKp/7SxWmonZnHAjKPtdWrLKaoGdMlpmhsEb1qUXB9juhlq8D7"
    "BrU8+X89Xmr0Gn65MQAAAABJRU5ErkJggg=="
)


# Guía de uso que muestra el botón README (copia de README.txt,
# README_EN.txt y README_PT.txt; tests/test_sim.py comprueba que coinciden).
README_TEXTS = {
    "es": """\
Creador: Desblack "Dess"
Contacto: desblackve@gmail.com


SERIAL NUMBER CHECKER
=====================

Versión 2.1A


1. QUÉ HACE LA APLICACIÓN
-------------------------

Serial Number Checker tiene dos modos (botones arriba a la izquierda):

- CHECKER: carga una lista de números de serie (S/N) desde un archivo CSV
  o Excel y comprueba, escaneando o escribiendo, cuáles están, cuáles
  faltan y cuáles sobran. Al terminar exporta un reporte Excel coloreado.
- CREAR LISTA S/N: arma una lista nueva con lo que escaneas, organizada
  por pallets (en columnas o en hojas) o, con el Modo Rack, por
  ubicaciones de un rack. La exporta a Excel o CSV.

La interfaz está en español, inglés y portugués. El botón README (arriba
a la derecha) abre esta guía dentro de la aplicación.


2. EJEMPLOS RÁPIDOS
-------------------

EJEMPLO 1 - Revisar un inventario contra un archivo

  1. Pulsa "Cargar CSV / Excel" y elige el archivo (por ejemplo,
     inventario_octubre.xlsx).
  2. Escanea los equipos uno a uno. Cada lectura muestra:
       Verde      ENCONTRADO        (Hoja Lote1 · fila 7)
       Amarillo   YA ESCANEADO      (el contador no sube)
       Rojo       NO ENCONTRADO     (se guarda en Sobrantes)
  3. Mira el contador: "128 / 340" y "38% completado · 212 pendientes".
  4. Pulsa "Exportar reporte": se crea inventario_octubre_resultado.xlsx
     con los S/N en verde (escaneado), rojo (pendiente) y azul
     (duplicado), más las hojas Resumen, Escaneos y Sobrantes.

EJEMPLO 2 - Me equivoqué de código

  Si escaneas el código equivocado, pulsa Ctrl+Z (o el botón "Deshacer"
  de Últimos escaneos): se anula el último escaneo y el contador vuelve
  a como estaba.

EJEMPLO 3 - Tengo los S/N en un correo o en otra hoja de Excel

  Copia la lista (uno por línea) y pégala con Ctrl+V en el campo de
  escaneo. La aplicación procesa todos de una vez y muestra un resumen:
  "Pegados 25: 22 encontrados · 2 no encontrados · 1 repetidos".

EJEMPLO 4 - Armar una lista de pallets

  1. Pasa al modo "Crear lista S/N".
  2. Escanea los S/N del primer pallet: se guardan en la columna A
     (Pallet 1).
  3. Pulsa "Siguiente pallet" (o F2): los siguientes van a la columna B
     (Pallet 2), y así sucesivamente.
  4. Si quieres un pallet en una hoja aparte, pulsa "Nueva hoja" junto a
     las pestañas de hojas.
  5. Pulsa "Exportar lista" (o Ctrl+E) para guardarla en Excel o CSV.

EJEMPLO 5 - Llenar un rack ubicación por ubicación (Modo Rack)

  1. En "Crear lista S/N", marca la casilla "Modo Rack".
  2. Escribe la ubicación (por ejemplo, DC05), las filas y columnas del
     rack (por ejemplo, 6 y 8) y el orden: Horizontal o Vertical.
  3. Pulsa "Agregar nuevo rack": se crea DC05-1.
  4. Escanea: el contador muestra la siguiente ubicación, por ejemplo
     (0,0), y cada S/N se guarda ahí. Luego (1,0), (2,0)...
  5. Cuando se llena el rack, la aplicación avisa "Se completó el rack
     DC05-1". Pulsa "Agregar nuevo rack" para seguir con DC05-2.

EJEMPLO 6 - Comprobar que cada S/N va en el pallet correcto

  1. En "Crear lista S/N", marca "Validar lista" y pulsa "Cargar lista"
     para elegir el CSV o Excel con los S/N que deben ir en los pallets.
  2. Escanea normalmente. Si un S/N no está en esa lista, no se registra
     y aparece en rojo "S/N NO VÁLIDO PARA EL PALLET".
  3. El contador indica cuántos llevas y cuántos faltan, por ejemplo
     "Validados 45 de 60 · faltan 15".

EJEMPLO 7 - Ver el avance de lejos

  Pulsa F11 (o "Vista grande"): la pantalla muestra solo el contador y el
  resultado en tamaño enorme. F11 o Esc para volver.


3. NOVEDADES
------------

VERSIÓN 2.1A - Validar lista y nombres de hoja

- Validar lista (Crear lista): solo se registran los S/N de una lista
  cargada; los demás se avisan en rojo y no entran.
- Con Validar lista se puede usar también "Alternativas A/B/U".
- La primera hoja se llama "Pallet 1" y cualquier hoja se puede
  renombrar con doble clic en su pestaña.
- Leyenda al pasar el ratón sobre "Modo Rack", "Validar lista" y las
  pestañas de hojas.
- Al exportar a Excel, las columnas salen con el ancho justo para el S/N
  más largo.
- Checker: copiar el S/N pendiente con Ctrl+C o con clic derecho →
  "Copiar S/N".

VERSIÓN 2.0 - Diseño nuevo y Modo Rack

- Diseño nuevo: tarjetas oscuras sobre fondo de líneas, botones con
  iconos y barra de título oscura en Windows.
- Ejecutable para Windows (SerialNumberChecker.exe): ya no hace falta
  instalar Python.
- Botón README: esta guía se puede leer dentro de la aplicación.
- Modo Rack: lista guiada por ubicaciones (x,y) de cada rack.
- Sobrantes: los S/N que no están en el archivo se guardan en su propia
  pestaña y se exportan en una hoja aparte.
- Deshacer el último escaneo (Ctrl+Z o botón "Deshacer").
- Hoja y fila: al encontrar un S/N se indica dónde está en el archivo.
- Filtro de formato: si se escanea otro código de la caja (por ejemplo,
  uno de 13 caracteres), avisa "NO PARECE UN S/N" y no lo cuenta.
- Pegar varios S/N a la vez con Ctrl+V.
- Contador con barra de progreso, ritmo (S/N por minuto) y tiempo
  estimado para terminar.
- Pestañas Pendientes / Sobrantes con buscador.
- Vista grande (F11).
- El reporte incluye las hojas Resumen, Escaneos (hora de cada lectura)
  y Sobrantes.
- Imprimir (Ctrl+P): resumen, pendientes y sobrantes (Checker) o la hoja
  actual / todas las hojas (Crear lista).
- Crear lista: botón "Nueva hoja", columnas que se pueden estirar y
  tecla Retroceso para eliminar el registro seleccionado.
- Se reconocen más encabezados de S/N (ver sección 15).
- La sesión se guarda sola: al volver a abrir la aplicación sigues donde
  lo dejaste.


4. MODO CHECKER, PASO A PASO
----------------------------

  1. Abre Serial Number Checker.
  2. Si lo necesitas, cambia el idioma en el selector de arriba a la
     derecha (Español, English o Português). El cambio es inmediato.
  3. Pulsa "Cargar CSV / Excel" y selecciona el archivo con los S/N.
  4. Se muestra el nombre del archivo, los números leídos y los S/N
     únicos.
  5. Si el archivo es un reporte exportado antes, la aplicación pregunta
     si quieres continuar donde te quedaste (sección 9).
  6. Escanea o escribe un S/N en el campo de escaneo.
  7. Al llegar a 17 caracteres la comprobación es automática. También
     puedes pulsar ENTER.
  8. Repite con todos los S/N que quieras comprobar.
  9. Al terminar, pulsa "Exportar reporte".
  10. Para empezar de cero, usa el botón "Reiniciar" (sección 16).

La casilla "Alternativas A/B/U" de la barra superior activa o desactiva
la búsqueda alternativa de la sección 5.


5. REGLAS DE COINCIDENCIA
-------------------------

- Primero se busca una coincidencia exacta con un S/N cargado.
- Si no la hay y "Alternativas A/B/U" está marcada, se prueba sustituyendo
  el séptimo carácter por A, B y U. Si una alternativa coincide, se marca
  como encontrado el S/N que realmente aparece en la lista (el mensaje
  dice "variante" y qué se leyó).
- Un S/N ya encontrado no vuelve a sumar en el contador.
- Si el código leído no tiene el largo de ningún S/N del archivo, se
  avisa "NO PARECE UN S/N" (probablemente se escaneó otro código de la
  caja) y no se guarda como sobrante.

IMPORTANTE: la sustitución del séptimo carácter es una búsqueda
alternativa. No garantiza que el número de serie sea auténtico ni es una
validación completa del formato.


6. COLORES, SONIDOS Y MENSAJES
------------------------------

En la aplicación:

- VERDE: S/N encontrado. Sonido de confirmación. El mensaje indica si fue
  exacto o una variante, y en qué hoja y fila del archivo está.
- AMARILLO: S/N ya escaneado antes. Otro sonido. El contador no sube y la
  aplicación queda en amarillo hasta el siguiente escaneo o hasta pulsar
  Esc.
- ROJO: S/N no encontrado (se guarda en Sobrantes) o código que no
  parece un S/N. Sonido de error.
- El borde del campo de escaneo destella en el color del resultado.
- Bajo el campo de escaneo hay una leyenda con los colores del reporte.

En el reporte Excel:

- VERDE: S/N encontrado durante la revisión.
- ROJO: S/N que todavía no se ha encontrado.
- AZUL: S/N duplicado en el archivo. La primera aparición conserva su
  verde o rojo y las repeticiones (en cualquier hoja) salen en azul.

Si se encontró por una alternativa A/B/U, se colorea el S/N real de la
lista.


7. CONTADOR, PROGRESO Y RITMO
-----------------------------

El contador muestra los S/N encontrados respecto al total de S/N únicos:

    128 / 340
    38% completado · 212 pendientes
    12/min · faltan ~18 min

- Un S/N repetido en el archivo cuenta una sola vez.
- Al cargar un archivo se indica cuántos números se leyeron y cuántos son
  únicos. Ejemplo: "12 números leídos · 10 S/N únicos" significa que hay
  2 valores repetidos en el archivo.
- El ritmo y el tiempo estimado aparecen tras unos escaneos y se calculan
  con los últimos 10 minutos.


8. PENDIENTES, SOBRANTES E HISTORIAL
------------------------------------

- ÚLTIMOS ESCANEOS: los 5 últimos S/N leídos y su resultado. El botón
  "Deshacer" (o Ctrl+Z) anula el último escaneo; se puede repetir.
- PENDIENTES: los S/N del archivo que faltan por escanear.
- SOBRANTES: los S/N escaneados que no están en el archivo, con cuántas
  veces se leyeron y la hora de la primera lectura.
- COPIAR: selecciona un S/N de Pendientes o Sobrantes y pulsa Ctrl+C, o
  haz clic derecho sobre él y elige "Copiar S/N".
- BUSCAR: filtra la pestaña abierta mientras escribes. Enter o Esc
  devuelven el cursor al campo de escaneo.


9. CONTINUAR UN REPORTE EXPORTADO
---------------------------------

Si cargas un reporte XLSX exportado antes por la aplicación, puedes
seguir donde lo dejaste.

- La aplicación detecta los S/N en verde del reporte y pregunta en el
  aviso amarillo si quieres continuar.
- "Continuar": esos S/N se marcan como encontrados. "Empezar de cero": el
  archivo se carga sin progreso.
- Las hojas Resumen, Escaneos y Sobrantes del reporte se ignoran al leer
  y se reemplazan al volver a exportar.
- Los CSV no se pueden continuar porque no guardan colores.

NOTA: se reconoce el verde exacto del reporte. Si un archivo tiene celdas
con ese mismo verde por otra razón (por ejemplo, el estilo "Bueno" de
Excel), también preguntará; en ese caso elige "Empezar de cero".


10. EXPORTAR REPORTE
--------------------

"Exportar reporte" (o Ctrl+E) guarda un Excel (.xlsx) con el resultado.

- Si el archivo cargado es XLSX, el reporte conserva sus hojas.
- Si es CSV, se crea un XLSX para poder colorear las celdas.
- Las columnas de S/N se colorean en verde, rojo o azul (sección 6).
- Las columnas de S/N salen con el ancho justo para el S/N más largo.
- Se añaden tres hojas:
    Resumen      archivo, totales, encontrados, pendientes, sobrantes y
                 horas de la primera y última lectura.
    Escaneos     cada lectura con su hora, lo leído y el resultado.
    Sobrantes    los S/N que no están en el archivo (si hay).
- El archivo original no se modifica: el reporte se guarda aparte y no
  se permite guardarlo encima del original.
- Excepción: si estás continuando un reporte (sección 9), puedes guardar
  el avance sobre ese mismo reporte. Windows pedirá confirmar que deseas
  reemplazarlo.
- El nombre sugerido es nombre_del_archivo_resultado.xlsx.


11. MODO CREAR LISTA S/N
------------------------

Sirve para armar una lista nueva con lo que escaneas. Acepta cualquier
texto (sin cambiar mayúsculas ni limitar el largo).

- DESTINO del siguiente pallet:
    Columna nueva      cada pallet en la columna siguiente (por defecto).
    Hoja nueva         cada pallet en su propia hoja.
    Desde la celda     empieza en la celda seleccionada del visor.
- NOMBRE (opcional): título del pallet o de la hoja. Si se deja vacío se
  usa "Pallet 1", "Pallet 2"...
- "Siguiente pallet" (F2): empieza el pallet siguiente.
- "Nueva hoja" (junto a las pestañas de hojas): crea una hoja nueva y
  salta a ella.
- NOMBRE DE LAS HOJAS: la primera se llama "Pallet 1". Doble clic en la
  pestaña de cualquier hoja para cambiarle el nombre (Enter para
  guardar, Esc para cancelar). El nombre nuevo se usa al exportar e
  imprimir.
- VALIDAR LISTA: comprueba que cada S/N está en una lista cargada
  (sección 12).
- VISOR: muestra la hoja tal como se exportará. Clic para seleccionar una
  celda, doble clic para editarla. Las columnas se estiran arrastrando su
  borde en la fila de letras; doble clic en el borde las ajusta al texto.
- DUPLICADOS: entran en azul y la aplicación queda en amarillo con un
  aviso: "Eliminar duplicados" u "Ocultar aviso".
- ELIMINAR: Retroceso o Supr con el campo vacío, o el botón "Eliminar",
  borra el registro seleccionado. "Deshacer" (Ctrl+Z) lo recupera.
- EXPORTAR LISTA (Ctrl+E): a Excel (con las mismas hojas y colores del
  visor) o a CSV (S/N, pallet, hoja, celda y si es duplicado). En Excel
  las columnas salen con el ancho justo para el S/N más largo; un CSV no
  guarda anchos de columna.
- Un valor que empiece con "=" se guarda como texto, no como fórmula.


12. VALIDAR LISTA
-----------------

Para asegurarte de que en los pallets solo entran los S/N correctos. Se
activa con la casilla "Validar lista" en la barra superior del modo
Crear lista.

- CARGAR LISTA: elige un CSV o Excel con los S/N esperados (se reconocen
  los mismos encabezados que en Checker, sección 15).
- SIN LISTA: mientras no cargues una, la aplicación avisa en amarillo y
  no registra nada de lo que escanees.
- S/N VÁLIDO: entra en el pallet de forma normal.
- S/N DUPLICADO: entra en azul, como siempre.
- S/N QUE NO ESTÁ EN LA LISTA: no se registra y aparece en rojo "S/N NO
  VÁLIDO PARA EL PALLET", con sonido de error.
- FALTAN: el contador muestra "Validados X de Y · faltan Z".
- ALTERNATIVAS A/B/U: con Validar lista activo aparece también esta
  casilla, con la misma regla que en Checker (sección 5).
- MODO RACK: el Modo Rack solo coloca máquinas y no valida, así que una
  opción bloquea a la otra (en gris): con Modo Rack activo no se puede
  marcar "Validar lista", y con Validar lista activo no se puede marcar
  "Modo Rack".


13. MODO RACK
-------------

Para guardar equipos en racks, ubicación por ubicación. Se activa con la
casilla "Modo Rack" en la barra superior del modo Crear lista.

- UBICACIÓN: el lugar físico, por ejemplo DC05. Los racks se llaman
  DC05-1, DC05-2, DC05-3...
- FILAS y COLUMNAS: tamaño del rack (por defecto 6 filas y 8 columnas).
- COORDENADAS (x,y): x es la columna e y la fila, empezando en (0,0)
  arriba a la izquierda. El visor numera desde 0.
- ORDEN: Horizontal llena (0,0), (1,0), (2,0)... fila por fila. Vertical
  llena (0,0), (0,1), (0,2)... columna por columna.
- "AGREGAR NUEVO RACK" (o F2, o junto a las pestañas): crea el rack
  siguiente de la ubicación y empieza en (0,0).
- GUÍA: el contador muestra en grande la siguiente ubicación y el visor
  la marca con un recuadro verde. Cada escaneo se guarda ahí y avanza.
- SALTAR: haz clic en una ubicación vacía del visor y el siguiente
  escaneo irá ahí.
- RACK COMPLETO: al llenar la última ubicación la aplicación avisa "Se
  completó el rack DC05-1" y no registra más en ese rack hasta que
  agregues uno nuevo.
- EXPORTAR: una hoja por rack dibujada como la estantería (coordenadas x
  arriba e y a la izquierda) y una hoja "Lista" con Ubicación, Rack,
  Posición (x,y), X, Y, S/N y si es duplicado.


14. IMPRIMIR Y VISTA GRANDE
---------------------------

IMPRIMIR (botón "Imprimir" o Ctrl+P): la aplicación pregunta en el aviso
amarillo qué imprimir.

- Checker: marca Resumen, Pendientes y/o Sobrantes y pulsa "Imprimir".
- Crear lista: "Hoja actual" o "Todas las hojas".
- Se imprime en la impresora predeterminada de Windows. Si el equipo
  tiene Excel se imprime con Excel; si no, con el Bloc de notas.

VISTA GRANDE (F11): solo el contador y el último resultado en tamaño
grande, para verlos de lejos. Se puede seguir escaneando. F11 o Esc para
volver.


15. ARCHIVOS COMPATIBLES Y ENCABEZADOS
--------------------------------------

Formatos:

- XLSX: archivos de Excel. Se leen todas las hojas del libro.
- CSV: una sola tabla. El separador (coma, punto y coma, tabulación o |)
  se detecta solo.

Encabezados reconocidos como columna de S/N (sin importar mayúsculas,
acentos, espacios ni signos):

    SN, S/N, NS, N/S
    SERIAL, SERIAL NO, SERIAL NR, SERIAL NBR, SERIAL NUM, SERIAL NUMBER
    SERIE, N SERIE, Nº DE SERIE, NO SERIE, NO. DE SERIE
    NUM SERIE, NUM. DE SERIE, NÚMERO SERIE, NÚMERO DE SERIE
    N SERIAL, NÚMERO SERIAL

Reglas de lectura:

- Si hay varios encabezados de S/N (en columnas u hojas distintas), se
  leen todos.
- Las columnas que no son de S/N no se usan.
- Las celdas vacías y los valores NONE o NAN se ignoran.
- Si una celda tiene una fórmula, se usa el valor calculado que guardó
  Excel.
- Los números se leen sin decimales (123456.0 se lee como 123456).


16. ATAJOS, REINICIO Y SESIÓN
-----------------------------

Atajos de teclado:

    Enter        comprobar / registrar
    Ctrl+Z       deshacer
    Ctrl+E       exportar
    Ctrl+P       imprimir
    Ctrl+V       pegar varios S/N
    Ctrl+C       copiar el S/N seleccionado en Pendientes o Sobrantes
    F11          vista grande
    Esc          quitar el aviso / salir de la vista grande
    F2           siguiente pallet o nuevo rack (Crear lista)
    Retroceso    eliminar el registro seleccionado (Crear lista, con el
                 campo vacío). Supr hace lo mismo.

REINICIAR: el botón "Reiniciar" (abajo a la derecha) borra la sesión del
modo activo. Si hay datos, el primer clic muestra "¿Seguro?" y hay que
pulsar otra vez en 3 segundos. El archivo original no se borra de tu
equipo.

SESIÓN: el avance se guarda solo. Si cierras la aplicación, al abrirla
recuperas el archivo cargado, lo escaneado, los sobrantes y la lista o
los racks en creación. Aun así, exporta el reporte antes de reiniciar.


17. REQUISITOS E INSTALACIÓN
----------------------------

EJECUTABLE (recomendado):

- Descomprime la carpeta completa y abre SerialNumberChecker.exe. No lo
  separes de la carpeta _internal.
- No necesita Python ni otras instalaciones.
- Si Windows muestra "Windows protegió su PC", pulsa "Más información" y
  "Ejecutar de todas formas". Ocurre porque el programa no está firmado
  digitalmente.

DESDE EL CÓDIGO (opcional):

- Python 3 para Windows (incluye Tkinter) y la biblioteca openpyxl:

      pip install openpyxl

- Después ejecuta serial_number_checker.py.
- Para crear el ejecutable, usa build_exe.bat.

Los sonidos funcionan en Windows. El icono de la aplicación es
arthur.ico.


18. RECOMENDACIONES
-------------------

- Verifica que el archivo cargado sea el correcto.
- Comprueba que las columnas de S/N tengan un encabezado reconocido
  (sección 15).
- Recuerda que se leen todas las hojas de los archivos XLSX.
- Si un S/N no aparece como encontrado, revisa que esté bien escrito y
  que exista en una columna de S/N.
- Revisa la pestaña Sobrantes antes de exportar.
- Exporta el reporte antes de reiniciar.
- Conserva una copia del archivo original sin modificar.


19. GitHub
----------

Proyecto, novedades y descargas:

    https://github.com/Desblackve/serial-number-checker

------------------------------------------------------------
""",
    "en": """\
Creator: Desblack "Dess"
Contact: desblackve@gmail.com


SERIAL NUMBER CHECKER
=====================

Version 2.1A


1. WHAT THE APP DOES
--------------------

Serial Number Checker has two modes (buttons at the top left):

- CHECKER: loads a list of serial numbers (S/N) from a CSV or Excel file
  and checks, by scanning or typing, which ones are there, which are
  missing and which are extra. When you finish, it exports a colored
  Excel report.
- CREATE S/N LIST: builds a new list from what you scan, organized by
  pallets (in columns or sheets) or, with Rack mode, by the locations of
  a rack. It exports it to Excel or CSV.

The interface is available in Spanish, English and Portuguese. The
README button (top right) opens this guide inside the app.


2. QUICK EXAMPLES
-----------------

EXAMPLE 1 - Check an inventory against a file

  1. Press "Load CSV / Excel" and choose the file (for example,
     inventory_october.xlsx).
  2. Scan the items one by one. Each read shows:
       Green      FOUND             (Sheet Lot1 · row 7)
       Yellow     ALREADY SCANNED   (the counter does not go up)
       Red        NOT FOUND         (saved in Extras)
  3. Look at the counter: "128 / 340" and "38% done · 212 pending".
  4. Press "Export report": inventory_october_resultado.xlsx is created
     with the S/Ns in green (scanned), red (pending) and blue
     (duplicate), plus the Summary, Scans and Extras sheets.

EXAMPLE 2 - I scanned the wrong code

  If you scan the wrong code, press Ctrl+Z (or the "Undo" button in
  Latest scans): the last scan is cancelled and the counter goes back to
  what it was.

EXAMPLE 3 - I have the S/Ns in an email or another Excel sheet

  Copy the list (one per line) and paste it with Ctrl+V into the scan
  field. The app processes all of them at once and shows a summary:
  "Pasted 25: 22 found · 2 not found · 1 repeated".

EXAMPLE 4 - Build a list of pallets

  1. Switch to "Create S/N List" mode.
  2. Scan the S/Ns of the first pallet: they are saved in column A
     (Pallet 1).
  3. Press "Next pallet" (or F2): the next ones go to column B
     (Pallet 2), and so on.
  4. If you want a pallet on a separate sheet, press "New sheet" next to
     the sheet tabs.
  5. Press "Export list" (or Ctrl+E) to save it to Excel or CSV.

EXAMPLE 5 - Fill a rack location by location (Rack mode)

  1. In "Create S/N List", tick "Rack mode".
  2. Type the location (for example, DC05), the rows and columns of the
     rack (for example, 6 and 8) and the order: Horizontal or Vertical.
  3. Press "Add new rack": DC05-1 is created.
  4. Scan: the counter shows the next location, for example (0,0), and
     each S/N is saved there. Then (1,0), (2,0)...
  5. When the rack is full, the app warns "Rack DC05-1 is complete".
     Press "Add new rack" to continue with DC05-2.

EXAMPLE 6 - Check that each S/N goes on the right pallet

  1. In "Create S/N List", tick "Validate list" and press "Load list" to
     choose the CSV or Excel file with the S/Ns that belong on the
     pallets.
  2. Scan as usual. If an S/N is not in that list, it is not recorded and
     "S/N NOT VALID FOR THE PALLET" appears in red.
  3. The counter shows how many you have and how many are left, for
     example "Validated 45 of 60 · 15 left".

EXAMPLE 7 - See the progress from a distance

  Press F11 (or "Big view"): the screen shows only the counter and the
  result in a very large size. F11 or Esc to go back.


3. WHAT'S NEW
-------------

VERSION 2.1A - Validate list and sheet names

- Validate list (Create list): only S/Ns from a loaded list are
  recorded; the others get a red alert and are not added.
- With Validate list you can also use "A/B/U alternatives".
- The first sheet is called "Pallet 1" and any sheet can be renamed by
  double-clicking its tab.
- A short description appears when you hover over "Rack mode",
  "Validate list" and the sheet tabs.
- When exporting to Excel, columns have the exact width for the longest
  S/N.
- Checker: copy the pending S/N with Ctrl+C or with right-click → "Copy
  S/N".

VERSION 2.0 - New design and Rack mode

- New design: dark cards over a line-pattern background, buttons with
  icons and a dark title bar on Windows.
- Windows executable (SerialNumberChecker.exe): Python no longer needs
  to be installed.
- README button: this guide can be read inside the app.
- Rack mode: a list guided by the (x,y) locations of each rack.
- Extras: S/Ns that are not in the file are kept in their own tab and
  exported on a separate sheet.
- Undo the last scan (Ctrl+Z or the "Undo" button).
- Sheet and row: when an S/N is found, the app shows where it is in the
  file.
- Format filter: if another code on the box is scanned (for example, one
  with 13 characters), it warns "DOES NOT LOOK LIKE AN S/N" and does not
  count it.
- Paste several S/Ns at once with Ctrl+V.
- Counter with a progress bar, pace (S/Ns per minute) and estimated time
  to finish.
- Pending / Extras tabs with a search box.
- Big view (F11).
- The report includes the Summary, Scans (time of each read) and Extras
  sheets.
- Print (Ctrl+P): summary, pending and extras (Checker) or the current
  sheet / all sheets (Create list).
- Create list: "New sheet" button, columns that can be stretched and the
  Backspace key to delete the selected record.
- More S/N headers are recognized (see section 15).
- The session is saved automatically: when you open the app again you
  continue where you left off.


4. CHECKER MODE, STEP BY STEP
-----------------------------

  1. Open Serial Number Checker.
  2. If needed, change the language in the selector at the top right
     (Español, English or Português). The change is immediate.
  3. Press "Load CSV / Excel" and select the file with the S/Ns.
  4. The file name, the values read and the unique S/Ns are shown.
  5. If the file is a report exported before, the app asks whether you
     want to resume where you left off (section 9).
  6. Scan or type an S/N in the scan field.
  7. At 17 characters the check is automatic. You can also press ENTER.
  8. Repeat with all the S/Ns you want to check.
  9. When you finish, press "Export report".
  10. To start from scratch, use the "Reset" button (section 16).

The "A/B/U alternatives" box in the top bar turns the alternative search
of section 5 on or off.


5. MATCHING RULES
-----------------

- First, an exact match with a loaded S/N is searched for.
- If there is none and "A/B/U alternatives" is ticked, the seventh
  character is replaced with A, B and U. If an alternative matches, the
  S/N that actually appears in the list is marked as found (the message
  says "variant" and what was read).
- An S/N already found does not add to the counter again.
- If the code read does not have the length of any S/N in the file, the
  app warns "DOES NOT LOOK LIKE AN S/N" (probably another code on the box
  was scanned) and it is not saved as an extra.

IMPORTANT: replacing the seventh character is an alternative search. It
does not guarantee that the serial number is genuine and it is not a
complete format validation.


6. COLORS, SOUNDS AND MESSAGES
------------------------------

In the app:

- GREEN: S/N found. Confirmation sound. The message says whether it was
  exact or a variant, and on which sheet and row of the file it is.
- YELLOW: S/N already scanned. Another sound. The counter does not go up
  and the app stays yellow until the next scan or until you press Esc.
- RED: S/N not found (saved in Extras) or a code that does not look like
  an S/N. Error sound.
- The border of the scan field flashes in the color of the result.
- Below the scan field there is a legend with the report colors.

In the Excel report:

- GREEN: S/N found during the check.
- RED: S/N not found yet.
- BLUE: S/N duplicated in the file. The first occurrence keeps its green
  or red and the repetitions (on any sheet) are blue.

If it was found through an A/B/U alternative, the real S/N in the list
is colored.


7. COUNTER, PROGRESS AND PACE
-----------------------------

The counter shows the S/Ns found out of the total unique S/Ns:

    128 / 340
    38% done · 212 pending
    12/min · ~18 min left

- An S/N repeated in the file counts only once.
- When a file is loaded, the app shows how many values were read and how
  many are unique. Example: "12 values read · 10 unique S/Ns" means there
  are 2 repeated values in the file.
- The pace and estimated time appear after a few scans and are
  calculated with the last 10 minutes.


8. PENDING, EXTRAS AND HISTORY
------------------------------

- LATEST SCANS: the last 5 S/Ns read and their result. The "Undo" button
  (or Ctrl+Z) cancels the last scan; it can be repeated.
- PENDING: the S/Ns in the file that have not been scanned yet.
- EXTRAS: the scanned S/Ns that are not in the file, with how many times
  they were read and the time of the first read.
- COPY: select an S/N in Pending or Extras and press Ctrl+C, or
  right-click it and choose "Copy S/N".
- SEARCH: filters the open tab as you type. Enter or Esc return the
  cursor to the scan field.


9. RESUME AN EXPORTED REPORT
----------------------------

If you load an XLSX report exported before by the app, you can continue
where you left off.

- The app detects the green S/Ns in the report and asks in the yellow
  alert whether you want to resume.
- "Resume": those S/Ns are marked as found. "Start from scratch": the
  file is loaded without progress.
- The Summary, Scans and Extras sheets of the report are ignored when
  reading and replaced when exporting again.
- CSV files cannot be resumed because they do not keep colors.

NOTE: the exact green of the report is recognized. If a file has cells
with that same green for another reason (for example, Excel's "Good"
style), it will ask as well; in that case choose "Start from scratch".


10. EXPORT REPORT
-----------------

"Export report" (or Ctrl+E) saves an Excel file (.xlsx) with the result.

- If the loaded file is XLSX, the report keeps its sheets.
- If it is CSV, an XLSX is created so the cells can be colored.
- The S/N columns are colored green, red or blue (section 6).
- The S/N columns have the exact width for the longest S/N.
- Three sheets are added:
    Summary      file, totals, found, pending, extras and the times of the
                 first and last read.
    Scans        each read with its time, what was read and the result.
    Extras       the S/Ns that are not in the file (if any).
- The original file is not modified: the report is saved separately and
  cannot be saved over the original.
- Exception: if you are resuming a report (section 9), you can save the
  progress over that same report. Windows will ask you to confirm that
  you want to replace it.
- The suggested name is file_name_resultado.xlsx.


11. CREATE S/N LIST MODE
------------------------

Used to build a new list with what you scan. It accepts any text (case
is kept and there is no length limit).

- DESTINATION of the next pallet:
    New column         each pallet in the next column (default).
    New sheet          each pallet on its own sheet.
    From the cell      starts at the cell selected in the viewer.
- NAME (optional): title of the pallet or sheet. If left empty, "Pallet
  1", "Pallet 2"... are used.
- "Next pallet" (F2): starts the next pallet.
- "New sheet" (next to the sheet tabs): creates a new sheet and jumps to
  it.
- SHEET NAMES: the first one is called "Pallet 1". Double-click the tab
  of any sheet to rename it (Enter to save, Esc to cancel). The new name
  is used when exporting and printing.
- VALIDATE LIST: checks that each S/N is in a loaded list (section 12).
- VIEWER: shows the sheet exactly as it will be exported. Click to select
  a cell, double-click to edit it. Columns are stretched by dragging
  their border in the letter row; double-click the border to fit the
  text.
- DUPLICATES: they come in blue and the app stays yellow with an alert:
  "Delete duplicates" or "Hide alert".
- DELETE: Backspace or Del with the field empty, or the "Delete" button,
  deletes the selected record. "Undo" (Ctrl+Z) brings it back.
- EXPORT LIST (Ctrl+E): to Excel (with the same sheets and colors as the
  viewer) or to CSV (S/N, pallet, sheet, cell and whether it is a
  duplicate). In Excel the columns have the exact width for the longest
  S/N; a CSV file does not keep column widths.
- A value that starts with "=" is saved as text, not as a formula.


12. VALIDATE LIST
-----------------

To make sure only the right S/Ns go on the pallets. It is turned on with
the "Validate list" box in the top bar of Create list mode.

- LOAD LIST: choose a CSV or Excel file with the expected S/Ns (the same
  headers as in Checker are recognized, section 15).
- NO LIST: until you load one, the app shows a yellow alert and does not
  record anything you scan.
- VALID S/N: it goes on the pallet as usual.
- DUPLICATE S/N: it comes in blue, as always.
- S/N NOT IN THE LIST: it is not recorded and "S/N NOT VALID FOR THE
  PALLET" appears in red, with the error sound.
- LEFT: the counter shows "Validated X of Y · Z left".
- A/B/U ALTERNATIVES: with Validate list on, this box also appears, with
  the same rule as in Checker (section 5).
- RACK MODE: Rack mode only places machines and does not validate, so
  each option locks the other (grey): while Rack mode is on, "Validate
  list" cannot be ticked, and while Validate list is on, "Rack mode"
  cannot be ticked.


13. RACK MODE
-------------

To store items in racks, location by location. It is turned on with the
"Rack mode" box in the top bar of Create list mode.

- LOCATION: the physical place, for example DC05. The racks are named
  DC05-1, DC05-2, DC05-3...
- ROWS and COLUMNS: size of the rack (6 rows and 8 columns by default).
- COORDINATES (x,y): x is the column and y the row, starting at (0,0) at
  the top left. The viewer counts from 0.
- ORDER: Horizontal fills (0,0), (1,0), (2,0)... row by row. Vertical
  fills (0,0), (0,1), (0,2)... column by column.
- "ADD NEW RACK" (or F2, or next to the tabs): creates the next rack of
  the location and starts at (0,0).
- GUIDE: the counter shows the next location in large size and the
  viewer marks it with a green box. Each scan is saved there and moves
  on.
- JUMP: click an empty location in the viewer and the next scan goes
  there.
- RACK COMPLETE: when the last location is filled the app warns "Rack
  DC05-1 is complete" and does not record anything else in that rack
  until you add a new one.
- EXPORT: one sheet per rack drawn like the shelf (x coordinates at the
  top and y on the left) and a "List" sheet with Location, Rack, Position
  (x,y), X, Y, S/N and whether it is a duplicate.


14. PRINT AND BIG VIEW
----------------------

PRINT ("Print" button or Ctrl+P): the app asks in the yellow alert what
to print.

- Checker: tick Summary, Pending and/or Extras and press "Print".
- Create list: "Current sheet" or "All sheets".
- It prints on the Windows default printer. If the computer has Excel,
  it prints with Excel; otherwise, with Notepad.

BIG VIEW (F11): only the counter and the last result in large size, to
see them from a distance. You can keep scanning. F11 or Esc to go back.


15. SUPPORTED FILES AND HEADERS
-------------------------------

Formats:

- XLSX: Excel files. All the sheets of the workbook are read.
- CSV: a single table. The separator (comma, semicolon, tab or |) is
  detected automatically.

Headers recognized as an S/N column (regardless of case, accents, spaces
or punctuation):

    SN, S/N, NS, N/S
    SERIAL, SERIAL NO, SERIAL NR, SERIAL NBR, SERIAL NUM, SERIAL NUMBER
    SERIE, N SERIE, Nº DE SERIE, NO SERIE, NO. DE SERIE
    NUM SERIE, NUM. DE SERIE, NÚMERO SERIE, NÚMERO DE SERIE
    N SERIAL, NÚMERO SERIAL

Reading rules:

- If there are several S/N headers (in different columns or sheets), all
  of them are read.
- Columns that are not S/N columns are not used.
- Empty cells and NONE or NAN values are ignored.
- If a cell has a formula, the calculated value saved by Excel is used.
- Numbers are read without decimals (123456.0 is read as 123456).


16. SHORTCUTS, RESET AND SESSION
--------------------------------

Keyboard shortcuts:

    Enter        check / record
    Ctrl+Z       undo
    Ctrl+E       export
    Ctrl+P       print
    Ctrl+V       paste several S/Ns
    Ctrl+C       copy the S/N selected in Pending or Extras
    F11          big view
    Esc          dismiss the alert / leave the big view
    F2           next pallet or new rack (Create list)
    Backspace    delete the selected record (Create list, with the field
                 empty). Del does the same.

RESET: the "Reset" button (bottom right) clears the session of the
active mode. If there is data, the first click shows "Sure?" and you
must click again within 3 seconds. The original file is not deleted from
your computer.

SESSION: progress is saved automatically. If you close the app, when you
open it again you get back the loaded file, what was scanned, the extras
and the list or racks being created. Even so, export the report before
resetting.


17. REQUIREMENTS AND INSTALLATION
---------------------------------

EXECUTABLE (recommended):

- Unzip the whole folder and open SerialNumberChecker.exe. Do not
  separate it from the _internal folder.
- It does not need Python or any other installation.
- If Windows shows "Windows protected your PC", press "More info" and
  "Run anyway". It happens because the program is not digitally signed.

FROM THE CODE (optional):

- Python 3 for Windows (includes Tkinter) and the openpyxl library:

      pip install openpyxl

- Then run serial_number_checker.py.
- To build the executable, use build_exe.bat.

Sounds work on Windows. The app icon is arthur.ico.


18. RECOMMENDATIONS
-------------------

- Make sure the loaded file is the right one.
- Check that the S/N columns have a recognized header (section 15).
- Remember that all the sheets of XLSX files are read.
- If an S/N does not appear as found, check that it is written correctly
  and that it exists in an S/N column.
- Review the Extras tab before exporting.
- Export the report before resetting.
- Keep an unmodified copy of the original file.


19. GitHub
----------

Project, updates and downloads:

    https://github.com/Desblackve/serial-number-checker

------------------------------------------------------------
""",
    "pt": """\
Criador: Desblack "Dess"
Contato: desblackve@gmail.com


SERIAL NUMBER CHECKER
=====================

Versão 2.1A


1. O QUE O APLICATIVO FAZ
-------------------------

O Serial Number Checker tem dois modos (botões no canto superior
esquerdo):

- CHECKER: carrega uma lista de números de série (S/N) de um arquivo CSV
  ou Excel e verifica, escaneando ou digitando, quais estão presentes,
  quais faltam e quais sobram. No final, exporta um relatório Excel
  colorido.
- CRIAR LISTA S/N: monta uma lista nova com o que você escaneia,
  organizada por pallets (em colunas ou planilhas) ou, com o Modo Rack,
  pelas posições de um rack. Exporta para Excel ou CSV.

A interface está em espanhol, inglês e português. O botão README (no
canto superior direito) abre este guia dentro do aplicativo.


2. EXEMPLOS RÁPIDOS
-------------------

EXEMPLO 1 - Conferir um inventário com um arquivo

  1. Clique em "Carregar CSV / Excel" e escolha o arquivo (por exemplo,
     inventario_outubro.xlsx).
  2. Escaneie os equipamentos um por um. Cada leitura mostra:
       Verde      ENCONTRADO        (Planilha Lote1 · linha 7)
       Amarelo    JÁ ESCANEADO      (o contador não sobe)
       Vermelho   NÃO ENCONTRADO    (salvo em Sobras)
  3. Veja o contador: "128 / 340" e "38% concluído · 212 pendentes".
  4. Clique em "Exportar relatório": é criado
     inventario_outubro_resultado.xlsx com os S/N em verde (escaneado),
     vermelho (pendente) e azul (duplicado), mais as planilhas Resumo,
     Leituras e Sobras.

EXEMPLO 2 - Escaneei o código errado

  Se escanear o código errado, pressione Ctrl+Z (ou o botão "Desfazer"
  de Últimas leituras): a última leitura é anulada e o contador volta ao
  que era.

EXEMPLO 3 - Tenho os S/N em um e-mail ou em outra planilha do Excel

  Copie a lista (um por linha) e cole com Ctrl+V no campo de leitura. O
  aplicativo processa todos de uma vez e mostra um resumo:
  "Colados 25: 22 encontrados · 2 não encontrados · 1 repetidos".

EXEMPLO 4 - Montar uma lista de pallets

  1. Mude para o modo "Criar lista S/N".
  2. Escaneie os S/N do primeiro pallet: eles são salvos na coluna A
     (Pallet 1).
  3. Clique em "Próximo pallet" (ou F2): os seguintes vão para a coluna
     B (Pallet 2), e assim por diante.
  4. Se quiser um pallet em uma planilha separada, clique em "Nova
     planilha" ao lado das abas de planilhas.
  5. Clique em "Exportar lista" (ou Ctrl+E) para salvá-la em Excel ou
     CSV.

EXEMPLO 5 - Preencher um rack posição por posição (Modo Rack)

  1. Em "Criar lista S/N", marque "Modo Rack".
  2. Digite a localização (por exemplo, DC05), as linhas e colunas do
     rack (por exemplo, 6 e 8) e a ordem: Horizontal ou Vertical.
  3. Clique em "Adicionar novo rack": é criado o DC05-1.
  4. Escaneie: o contador mostra a próxima posição, por exemplo (0,0), e
     cada S/N é salvo ali. Depois (1,0), (2,0)...
  5. Quando o rack ficar cheio, o aplicativo avisa "O rack DC05-1 foi
     completado". Clique em "Adicionar novo rack" para seguir com o
     DC05-2.

EXEMPLO 6 - Conferir se cada S/N vai no pallet certo

  1. Em "Criar lista S/N", marque "Validar lista" e clique em "Carregar
     lista" para escolher o CSV ou Excel com os S/N que devem ir nos
     pallets.
  2. Escaneie normalmente. Se um S/N não estiver nessa lista, ele não é
     registrado e aparece em vermelho "S/N NÃO VÁLIDO PARA O PALLET".
  3. O contador mostra quantos já foram e quantos faltam, por exemplo
     "Validados 45 de 60 · faltam 15".

EXEMPLO 7 - Ver o avanço de longe

  Pressione F11 (ou "Visão ampliada"): a tela mostra apenas o contador e
  o resultado em tamanho grande. F11 ou Esc para voltar.


3. NOVIDADES
------------

VERSÃO 2.1A - Validar lista e nomes de planilha

- Validar lista (Criar lista): só são registrados os S/N de uma lista
  carregada; os outros recebem um aviso em vermelho e não entram.
- Com Validar lista também é possível usar "Alternativas A/B/U".
- A primeira planilha se chama "Pallet 1" e qualquer planilha pode ser
  renomeada com um clique duplo na aba.
- Uma descrição curta aparece ao passar o mouse sobre "Modo Rack",
  "Validar lista" e as abas de planilhas.
- Ao exportar para Excel, as colunas saem com a largura exata para o S/N
  mais longo.
- Checker: copiar o S/N pendente com Ctrl+C ou com o botão direito →
  "Copiar S/N".

VERSÃO 2.0 - Design novo e Modo Rack

- Design novo: cartões escuros sobre um fundo de linhas, botões com
  ícones e barra de título escura no Windows.
- Executável para Windows (SerialNumberChecker.exe): não é mais preciso
  instalar o Python.
- Botão README: este guia pode ser lido dentro do aplicativo.
- Modo Rack: lista guiada pelas posições (x,y) de cada rack.
- Sobras: os S/N que não estão no arquivo ficam em uma aba própria e são
  exportados em uma planilha separada.
- Desfazer a última leitura (Ctrl+Z ou botão "Desfazer").
- Planilha e linha: ao encontrar um S/N, o aplicativo indica onde ele
  está no arquivo.
- Filtro de formato: se for escaneado outro código da caixa (por
  exemplo, um de 13 caracteres), avisa "NÃO PARECE UM S/N" e não o
  conta.
- Colar vários S/N de uma vez com Ctrl+V.
- Contador com barra de progresso, ritmo (S/N por minuto) e tempo
  estimado para terminar.
- Abas Pendentes / Sobras com busca.
- Visão ampliada (F11).
- O relatório inclui as planilhas Resumo, Leituras (hora de cada
  leitura) e Sobras.
- Imprimir (Ctrl+P): resumo, pendentes e sobras (Checker) ou a planilha
  atual / todas as planilhas (Criar lista).
- Criar lista: botão "Nova planilha", colunas que podem ser esticadas e
  tecla Backspace para excluir o registro selecionado.
- Mais cabeçalhos de S/N são reconhecidos (ver seção 15).
- A sessão é salva sozinha: ao abrir o aplicativo de novo, você continua
  de onde parou.


4. MODO CHECKER, PASSO A PASSO
------------------------------

  1. Abra o Serial Number Checker.
  2. Se precisar, mude o idioma no seletor do canto superior direito
     (Español, English ou Português). A mudança é imediata.
  3. Clique em "Carregar CSV / Excel" e selecione o arquivo com os S/N.
  4. São mostrados o nome do arquivo, os valores lidos e os S/N únicos.
  5. Se o arquivo for um relatório exportado antes, o aplicativo pergunta
     se você quer continuar de onde parou (seção 9).
  6. Escaneie ou digite um S/N no campo de leitura.
  7. Ao chegar a 17 caracteres, a verificação é automática. Também é
     possível pressionar ENTER.
  8. Repita com todos os S/N que quiser verificar.
  9. No final, clique em "Exportar relatório".
  10. Para começar do zero, use o botão "Reiniciar" (seção 16).

A caixa "Alternativas A/B/U" da barra superior liga ou desliga a busca
alternativa da seção 5.


5. REGRAS DE CORRESPONDÊNCIA
----------------------------

- Primeiro é procurada uma correspondência exata com um S/N carregado.
- Se não houver e "Alternativas A/B/U" estiver marcada, o sétimo
  caractere é substituído por A, B e U. Se uma alternativa corresponder,
  é marcado como encontrado o S/N que realmente aparece na lista (a
  mensagem diz "variante" e o que foi lido).
- Um S/N já encontrado não soma de novo no contador.
- Se o código lido não tiver o tamanho de nenhum S/N do arquivo, o
  aplicativo avisa "NÃO PARECE UM S/N" (provavelmente foi escaneado
  outro código da caixa) e ele não é salvo como sobra.

IMPORTANTE: a substituição do sétimo caractere é uma busca alternativa.
Não garante que o número de série seja autêntico nem é uma validação
completa do formato.


6. CORES, SONS E MENSAGENS
--------------------------

No aplicativo:

- VERDE: S/N encontrado. Som de confirmação. A mensagem indica se foi
  exato ou uma variante, e em qual planilha e linha do arquivo ele está.
- AMARELO: S/N já escaneado. Outro som. O contador não sobe e o
  aplicativo fica amarelo até a próxima leitura ou até pressionar Esc.
- VERMELHO: S/N não encontrado (salvo em Sobras) ou código que não
  parece um S/N. Som de erro.
- A borda do campo de leitura pisca na cor do resultado.
- Abaixo do campo de leitura há uma legenda com as cores do relatório.

No relatório Excel:

- VERDE: S/N encontrado durante a conferência.
- VERMELHO: S/N ainda não encontrado.
- AZUL: S/N duplicado no arquivo. A primeira ocorrência mantém o verde
  ou vermelho e as repetições (em qualquer planilha) ficam em azul.

Se foi encontrado por uma alternativa A/B/U, é colorido o S/N real da
lista.


7. CONTADOR, PROGRESSO E RITMO
------------------------------

O contador mostra os S/N encontrados em relação ao total de S/N únicos:

    128 / 340
    38% concluído · 212 pendentes
    12/min · faltam ~18 min

- Um S/N repetido no arquivo conta uma única vez.
- Ao carregar um arquivo, o aplicativo mostra quantos valores foram
  lidos e quantos são únicos. Exemplo: "12 valores lidos · 10 S/N
  únicos" significa que há 2 valores repetidos no arquivo.
- O ritmo e o tempo estimado aparecem após algumas leituras e são
  calculados com os últimos 10 minutos.


8. PENDENTES, SOBRAS E HISTÓRICO
--------------------------------

- ÚLTIMAS LEITURAS: os 5 últimos S/N lidos e seu resultado. O botão
  "Desfazer" (ou Ctrl+Z) anula a última leitura; pode ser repetido.
- PENDENTES: os S/N do arquivo que ainda faltam escanear.
- SOBRAS: os S/N escaneados que não estão no arquivo, com quantas vezes
  foram lidos e a hora da primeira leitura.
- COPIAR: selecione um S/N em Pendentes ou Sobras e pressione Ctrl+C, ou
  clique nele com o botão direito e escolha "Copiar S/N".
- BUSCAR: filtra a aba aberta enquanto você digita. Enter ou Esc
  devolvem o cursor ao campo de leitura.


9. CONTINUAR UM RELATÓRIO EXPORTADO
-----------------------------------

Se você carregar um relatório XLSX exportado antes pelo aplicativo, pode
continuar de onde parou.

- O aplicativo detecta os S/N em verde do relatório e pergunta no aviso
  amarelo se você quer continuar.
- "Continuar": esses S/N são marcados como encontrados. "Começar do
  zero": o arquivo é carregado sem progresso.
- As planilhas Resumo, Leituras e Sobras do relatório são ignoradas na
  leitura e substituídas ao exportar de novo.
- Arquivos CSV não podem ser continuados porque não guardam cores.

NOTA: é reconhecido o verde exato do relatório. Se um arquivo tiver
células com esse mesmo verde por outro motivo (por exemplo, o estilo
"Bom" do Excel), também vai perguntar; nesse caso escolha "Começar do
zero".


10. EXPORTAR RELATÓRIO
----------------------

"Exportar relatório" (ou Ctrl+E) salva um Excel (.xlsx) com o
resultado.

- Se o arquivo carregado for XLSX, o relatório mantém suas planilhas.
- Se for CSV, é criado um XLSX para poder colorir as células.
- As colunas de S/N são coloridas de verde, vermelho ou azul (seção 6).
- As colunas de S/N saem com a largura exata para o S/N mais longo.
- São adicionadas três planilhas:
    Resumo       arquivo, totais, encontrados, pendentes, sobras e as
                 horas da primeira e da última leitura.
    Leituras     cada leitura com sua hora, o que foi lido e o resultado.
    Sobras       os S/N que não estão no arquivo (se houver).
- O arquivo original não é modificado: o relatório é salvo à parte e não
  é permitido salvá-lo por cima do original.
- Exceção: se você estiver continuando um relatório (seção 9), pode
  salvar o progresso sobre esse mesmo relatório. O Windows pedirá para
  confirmar que deseja substituí-lo.
- O nome sugerido é nome_do_arquivo_resultado.xlsx.


11. MODO CRIAR LISTA S/N
------------------------

Serve para montar uma lista nova com o que você escaneia. Aceita
qualquer texto (sem mudar maiúsculas nem limitar o tamanho).

- DESTINO do próximo pallet:
    Nova coluna        cada pallet na coluna seguinte (padrão).
    Nova planilha      cada pallet em sua própria planilha.
    Da célula          começa na célula selecionada no visor.
- NOME (opcional): título do pallet ou da planilha. Se ficar vazio, é
  usado "Pallet 1", "Pallet 2"...
- "Próximo pallet" (F2): começa o pallet seguinte.
- "Nova planilha" (ao lado das abas de planilhas): cria uma planilha
  nova e vai para ela.
- NOME DAS PLANILHAS: a primeira se chama "Pallet 1". Clique duas vezes
  na aba de qualquer planilha para renomeá-la (Enter para salvar, Esc
  para cancelar). O nome novo é usado ao exportar e imprimir.
- VALIDAR LISTA: confere se cada S/N está em uma lista carregada (seção
  12).
- VISOR: mostra a planilha exatamente como será exportada. Clique para
  selecionar uma célula, clique duas vezes para editá-la. As colunas são
  esticadas arrastando sua borda na linha de letras; clique duas vezes na
  borda para ajustá-las ao texto.
- DUPLICADOS: entram em azul e o aplicativo fica amarelo com um aviso:
  "Excluir duplicados" ou "Ocultar aviso".
- EXCLUIR: Backspace ou Del com o campo vazio, ou o botão "Excluir",
  apaga o registro selecionado. "Desfazer" (Ctrl+Z) o recupera.
- EXPORTAR LISTA (Ctrl+E): para Excel (com as mesmas planilhas e cores
  do visor) ou para CSV (S/N, pallet, planilha, célula e se é
  duplicado). No Excel as colunas saem com a largura exata para o S/N
  mais longo; um CSV não guarda larguras de coluna.
- Um valor que começa com "=" é salvo como texto, não como fórmula.


12. VALIDAR LISTA
-----------------

Para garantir que nos pallets só entrem os S/N corretos. É ativado com a
caixa "Validar lista" na barra superior do modo Criar lista.

- CARREGAR LISTA: escolha um CSV ou Excel com os S/N esperados (são
  reconhecidos os mesmos cabeçalhos do Checker, seção 15).
- SEM LISTA: enquanto não carregar uma, o aplicativo mostra um aviso
  amarelo e não registra nada do que for escaneado.
- S/N VÁLIDO: entra no pallet normalmente.
- S/N DUPLICADO: entra em azul, como sempre.
- S/N QUE NÃO ESTÁ NA LISTA: não é registrado e aparece em vermelho "S/N
  NÃO VÁLIDO PARA O PALLET", com o som de erro.
- FALTAM: o contador mostra "Validados X de Y · faltam Z".
- ALTERNATIVAS A/B/U: com Validar lista ativo, esta caixa também aparece,
  com a mesma regra do Checker (seção 5).
- MODO RACK: o Modo Rack só coloca máquinas e não valida, então uma
  opção bloqueia a outra (cinza): com o Modo Rack ativo não é possível
  marcar "Validar lista", e com Validar lista ativo não é possível marcar
  "Modo Rack".


13. MODO RACK
-------------

Para guardar equipamentos em racks, posição por posição. É ativado com a
caixa "Modo Rack" na barra superior do modo Criar lista.

- LOCALIZAÇÃO: o lugar físico, por exemplo DC05. Os racks se chamam
  DC05-1, DC05-2, DC05-3...
- LINHAS e COLUNAS: tamanho do rack (6 linhas e 8 colunas por padrão).
- COORDENADAS (x,y): x é a coluna e y a linha, começando em (0,0) no
  canto superior esquerdo. O visor numera a partir de 0.
- ORDEM: Horizontal preenche (0,0), (1,0), (2,0)... linha por linha.
  Vertical preenche (0,0), (0,1), (0,2)... coluna por coluna.
- "ADICIONAR NOVO RACK" (ou F2, ou ao lado das abas): cria o rack
  seguinte da localização e começa em (0,0).
- GUIA: o contador mostra a próxima posição em tamanho grande e o visor
  a marca com um quadro verde. Cada leitura é salva ali e avança.
- PULAR: clique em uma posição vazia do visor e a próxima leitura irá
  para lá.
- RACK COMPLETO: ao preencher a última posição, o aplicativo avisa "O
  rack DC05-1 foi completado" e não registra mais nada nesse rack até
  você adicionar um novo.
- EXPORTAR: uma planilha por rack desenhada como a estante (coordenadas
  x em cima e y à esquerda) e uma planilha "Lista" com Localização,
  Rack, Posição (x,y), X, Y, S/N e se é duplicado.


14. IMPRIMIR E VISÃO AMPLIADA
-----------------------------

IMPRIMIR (botão "Imprimir" ou Ctrl+P): o aplicativo pergunta no aviso
amarelo o que imprimir.

- Checker: marque Resumo, Pendentes e/ou Sobras e clique em "Imprimir".
- Criar lista: "Planilha atual" ou "Todas as planilhas".
- A impressão sai na impressora padrão do Windows. Se o computador tiver
  Excel, imprime com o Excel; se não, com o Bloco de Notas.

VISÃO AMPLIADA (F11): apenas o contador e o último resultado em tamanho
grande, para vê-los de longe. É possível continuar escaneando. F11 ou
Esc para voltar.


15. ARQUIVOS COMPATÍVEIS E CABEÇALHOS
-------------------------------------

Formatos:

- XLSX: arquivos do Excel. Todas as planilhas da pasta de trabalho são
  lidas.
- CSV: uma única tabela. O separador (vírgula, ponto e vírgula,
  tabulação ou |) é detectado sozinho.

Cabeçalhos reconhecidos como coluna de S/N (sem importar maiúsculas,
acentos, espaços ou sinais):

    SN, S/N, NS, N/S
    SERIAL, SERIAL NO, SERIAL NR, SERIAL NBR, SERIAL NUM, SERIAL NUMBER
    SERIE, N SERIE, Nº DE SERIE, NO SERIE, NO. DE SERIE
    NUM SERIE, NUM. DE SERIE, NÚMERO SERIE, NÚMERO DE SERIE
    N SERIAL, NÚMERO SERIAL

Regras de leitura:

- Se houver vários cabeçalhos de S/N (em colunas ou planilhas
  diferentes), todos são lidos.
- As colunas que não são de S/N não são usadas.
- Células vazias e valores NONE ou NAN são ignorados.
- Se uma célula tiver uma fórmula, é usado o valor calculado salvo pelo
  Excel.
- Os números são lidos sem decimais (123456.0 é lido como 123456).


16. ATALHOS, REINÍCIO E SESSÃO
------------------------------

Atalhos de teclado:

    Enter        verificar / registrar
    Ctrl+Z       desfazer
    Ctrl+E       exportar
    Ctrl+P       imprimir
    Ctrl+V       colar vários S/N
    Ctrl+C       copiar o S/N selecionado em Pendentes ou Sobras
    F11          visão ampliada
    Esc          ocultar o aviso / sair da visão ampliada
    F2           próximo pallet ou novo rack (Criar lista)
    Backspace    excluir o registro selecionado (Criar lista, com o campo
                 vazio). Del faz o mesmo.

REINICIAR: o botão "Reiniciar" (canto inferior direito) apaga a sessão
do modo ativo. Se houver dados, o primeiro clique mostra "Certeza?" e é
preciso clicar de novo em até 3 segundos. O arquivo original não é
apagado do seu computador.

SESSÃO: o progresso é salvo sozinho. Se você fechar o aplicativo, ao
abri-lo recupera o arquivo carregado, o que foi escaneado, as sobras e a
lista ou os racks em criação. Mesmo assim, exporte o relatório antes de
reiniciar.


17. REQUISITOS E INSTALAÇÃO
---------------------------

EXECUTÁVEL (recomendado):

- Descompacte a pasta inteira e abra o SerialNumberChecker.exe. Não o
  separe da pasta _internal.
- Não precisa de Python nem de outras instalações.
- Se o Windows mostrar "O Windows protegeu o computador", clique em
  "Mais informações" e "Executar assim mesmo". Isso acontece porque o
  programa não tem assinatura digital.

A PARTIR DO CÓDIGO (opcional):

- Python 3 para Windows (inclui Tkinter) e a biblioteca openpyxl:

      pip install openpyxl

- Depois execute serial_number_checker.py.
- Para criar o executável, use build_exe.bat.

Os sons funcionam no Windows. O ícone do aplicativo é arthur.ico.


18. RECOMENDAÇÕES
-----------------

- Verifique se o arquivo carregado é o correto.
- Confira se as colunas de S/N têm um cabeçalho reconhecido (seção 15).
- Lembre-se de que todas as planilhas dos arquivos XLSX são lidas.
- Se um S/N não aparecer como encontrado, verifique se está escrito
  corretamente e se existe em uma coluna de S/N.
- Revise a aba Sobras antes de exportar.
- Exporte o relatório antes de reiniciar.
- Guarde uma cópia do arquivo original sem modificações.


19. GitHub
----------

Projeto, novidades e downloads:

    https://github.com/Desblackve/serial-number-checker

------------------------------------------------------------
""",
}


def main():
    prepare_windows_process()
    root = tk.Tk()
    SerialNumberChecker(root)
    root.mainloop()


if __name__ == "__main__":
    main()


# Rework de un proyecto viejo.
# Aviso de autoría y mantenimiento:
# El código se proporciona tal como está, sin garantías de
# funcionamiento ni soporte técnico futuro.
# Consultas sobre funciones: Desblackve@gmail.com
